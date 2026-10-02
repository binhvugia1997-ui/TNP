"""Regression tests for real-file structures: multi-line / merged headers and
PPTX content that python-pptx does not expose (OLE objects, mc:AlternateContent)."""
import copy
from io import BytesIO
from pathlib import Path

from lxml import etree
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.classifier import find_qpn_slide, heuristic_classify
from app.excel_writer import ExcelWriter
from app.extractor import extract_record
from app.pptx_parser import parse_pptx
from app.qpn_renderer import SlideRenderer, render_qpn

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}


# --------------------------------------------------------------------------- Excel headers
def _template_multiline(path: Path, weeks_style: str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kiểm chứng"
    ws["A1"] = "CÔNG TY ABC"
    ws["A2"] = "BẢNG KIỂM CHỨNG ĐỐI SÁCH"
    hdr = ["STT", "Management\nnumber", "Tên\nvendor", "Ngày\nphát sinh", "Model", "Item",
           "Nội dung\nlỗi", "QPN", "Nguyên\nnhân", "Nội dung\nđối sách cải tiến", "Hình ảnh\ncải tiến"]
    for i, h in enumerate(hdr, start=1):
        ws.cell(row=4, column=i, value=h)
        ws.merge_cells(start_row=4, start_column=i, end_row=5, end_column=i)
    ws.merge_cells(start_row=4, start_column=12, end_row=4, end_column=19)
    ws.cell(row=4, column=12, value="Kiểm chứng\n(sau đối sách)")
    for k in range(1, 9):
        label = {"plus": f"+{k}", "wn": f"WEEK\n+{k}", "w": f"W{k}"}[weeks_style]
        ws.cell(row=5, column=11 + k, value=label)
    wb.save(path)
    return path


def test_headers_with_line_breaks_and_merged_week_parent(tmp_path):
    for style in ("plus", "wn", "w"):
        tpl = _template_multiline(tmp_path / f"t_{style}.xlsx", style)
        w = ExcelWriter(tpl, tmp_path / f"o_{style}.xlsx")
        assert w.header_row == 4 and w.data_start == 6, style
        assert w.columns["management_number"] == 2, style
        assert w.columns["vendor"] == 3 and w.columns["occurrence_date"] == 4
        assert w.columns["defect_content"] == 7 and w.columns["root_cause"] == 9
        assert w.columns["improvement"] == 10 and w.columns["improvement_image"] == 11
        for k in range(1, 9):
            assert w.columns[f"week_{k}"] == 11 + k, (style, k)
        row = w.append_record(extract_record_stub())
        assert row == 6
        w.save()
        ws = load_workbook(tmp_path / f"o_{style}.xlsx")["Kiểm chứng"]
        assert ws.cell(row=6, column=2).value == "260918080-VOC"
        assert all(ws.cell(row=6, column=11 + k).value is None for k in range(1, 9))


def extract_record_stub():
    from app.extractor import ExtractedRecord
    return ExtractedRecord(management_number="260918080-VOC", model="A185", item="Rear",
                           root_cause="x", improvement="y", qpn_slide=2)


def test_plain_number_subheaders_not_taken_as_weeks_without_parent(tmp_path):
    """Sub-cells '1'..'8' under an unrelated parent must NOT become WEEK columns."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Kiểm chứng"
    for i, h in enumerate(["STT", "Management number", "Model", "QPN", "Nguyên nhân", "Nội dung đối sách cải tiến"], start=1):
        ws.cell(row=1, column=i, value=h)
        ws.merge_cells(start_row=1, start_column=i, end_row=2, end_column=i)
    ws.merge_cells("G1:I1")
    ws["G1"] = "Số lượng lỗi"
    for k in range(1, 4):
        ws.cell(row=2, column=6 + k, value=str(k))
    p = tmp_path / "t.xlsx"
    wb.save(p)
    w = ExcelWriter(p, tmp_path / "o.xlsx")
    assert not any(f.startswith("week_") for f in w.columns)


# --------------------------------------------------------------------------- PPTX: OLE / AlternateContent
def _png_bytes(color="#ddddff"):
    im = Image.new("RGB", (400, 300), color)
    bio = BytesIO()
    im.save(bio, "PNG")
    return bio.getvalue()


def _add_ole_qpn(slide, prs, wrap_in_alternate_content: bool, alt_text: str = "Quality Problem Notice"):
    """Emulate an embedded Excel object (p:graphicFrame/p:oleObj with preview picture)."""
    # add a real picture first to obtain an image relationship id, then remove the p:pic
    pic = slide.shapes.add_picture(BytesIO(_png_bytes()), Inches(1), Inches(1), Inches(8), Inches(5))
    blip = pic._element.xpath(".//a:blip")[0]
    rid = blip.get("{%s}embed" % NS["r"])
    pic._element.getparent().remove(pic._element)
    xml = f"""
<p:graphicFrame xmlns:p="{NS['p']}" xmlns:a="{NS['a']}" xmlns:r="{NS['r']}">
  <p:nvGraphicFramePr>
    <p:cNvPr id="77" name="Object 7" descr="{alt_text}"/>
    <p:cNvGraphicFramePr/><p:nvPr/>
  </p:nvGraphicFramePr>
  <p:xfrm><a:off x="914400" y="914400"/><a:ext cx="7315200" cy="4572000"/></p:xfrm>
  <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/presentationml/2006/ole">
    <p:oleObj name="Worksheet" r:id="{rid}" imgW="7315200" imgH="4572000" progId="Excel.Sheet.12">
      <p:embed/>
      <p:pic>
        <p:nvPicPr><p:cNvPr id="78" name="Preview"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>
        <p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
        <p:spPr><a:xfrm><a:off x="914400" y="914400"/><a:ext cx="7315200" cy="4572000"/></a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
      </p:pic>
    </p:oleObj>
  </a:graphicData></a:graphic>
</p:graphicFrame>"""
    gf = etree.fromstring(xml)
    if wrap_in_alternate_content:
        ac = etree.Element("{%s}AlternateContent" % NS["mc"], nsmap={"mc": NS["mc"]})
        choice = etree.SubElement(ac, "{%s}Choice" % NS["mc"], Requires="v")
        choice.append(gf)
        fb = etree.SubElement(ac, "{%s}Fallback" % NS["mc"])
        fb.append(copy.deepcopy(gf))
        slide.shapes._spTree.append(ac)
    else:
        slide.shapes._spTree.append(gf)


def _add_alternate_content_textbox(slide, text: str):
    tb = slide.shapes.add_textbox(Inches(1), Inches(0.2), Inches(6), Inches(0.6))
    tb.text_frame.text = text
    sp = tb._element
    ac = etree.Element("{%s}AlternateContent" % NS["mc"], nsmap={"mc": NS["mc"]})
    choice = etree.SubElement(ac, "{%s}Choice" % NS["mc"], Requires="a14")
    sp.getparent().remove(sp)
    choice.append(sp)
    slide.shapes._spTree.append(ac)


def _deck_with_ole_qpn(path: Path, wrap: bool, alt: str = "Quality Problem Notice", header_text: str = ""):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "BÁO CÁO LỖI XƯỚC REAR SM-A185"
    s2 = prs.slides.add_slide(blank)
    if header_text:
        _add_alternate_content_textbox(s2, header_text)
    _add_ole_qpn(s2, prs, wrap, alt)
    s3 = prs.slides.add_slide(blank)
    s3.shapes.add_textbox(Inches(1), Inches(0.3), Inches(6), Inches(0.6)).text_frame.text = "NGUYÊN NHÂN"
    s3.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "- Khay không có lót"
    s4 = prs.slides.add_slide(blank)
    s4.shapes.add_textbox(Inches(1), Inches(0.3), Inches(6), Inches(0.6)).text_frame.text = "CẢI TIẾN"
    s4.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "- Thay khay có lót EVA"
    prs.save(str(path))
    return path


def test_python_pptx_skips_alternate_content_but_parser_recovers(tmp_path):
    p = _deck_with_ole_qpn(tmp_path / "ole_ac.pptx", wrap=True)
    prs = Presentation(str(p))
    assert len(list(prs.slides[1].shapes)) == 0                # python-pptx sees nothing on the QPN slide
    r = parse_pptx(p)
    s2 = r.slides[1]
    assert s2.xml_stats["AlternateContent"] == 1 and s2.xml_stats["oleObj"] == 2
    assert len(s2.pictures) == 1                                # Choice/Fallback duplicates de-duplicated
    pic = s2.pictures[0]
    assert pic.origin == "ole" and pic.image_blob and pic.image_ext == "png"
    assert "Quality Problem Notice" in pic.alt_text and "Excel.Sheet.12" in pic.alt_text
    assert find_qpn_slide(r) == 2


def test_ole_qpn_without_alternate_content(tmp_path):
    p = _deck_with_ole_qpn(tmp_path / "ole.pptx", wrap=False, alt="")
    r = parse_pptx(p)
    assert len(r.slides[1].pictures) == 1 and r.slides[1].pictures[0].origin == "ole"
    # alt text empty -> QPN only recognisable by OLE name / progId? 'Worksheet' is not QPN -> not found
    assert find_qpn_slide(r) is None
    # but a header text box (hidden in AlternateContent) saying QPN makes it detectable
    p2 = _deck_with_ole_qpn(tmp_path / "ole_hdr.pptx", wrap=True, alt="", header_text="QUALITY PROBLEM NOTICE")
    r2 = parse_pptx(p2)
    assert r2.slides[1].text_blocks and r2.slides[1].text_blocks[0].origin == "alternate_content"
    assert find_qpn_slide(r2) == 2


def test_ole_qpn_renders_and_flows_to_record(tmp_path):
    p = _deck_with_ole_qpn(tmp_path / "ole_ac.pptx", wrap=True)
    r = parse_pptx(p)
    cls = heuristic_classify(r)
    assert cls.qpn_slide == 2 and cls.cause_slides == [3] and cls.improvement_slides == [4]
    assert 2 not in cls.improvement_image_slides
    rec = extract_record(r, cls)
    assert rec.qpn_slide == 2 and cls.qpn_source == "metadata"
    assert not any("Không tìm thấy QPN" in x or "chỉ do AI" in x for x in rec.review_reasons)
    png, backend = render_qpn(r, 2, tmp_path / "qpn.png", SlideRenderer(prefer=("builtin",)))
    with Image.open(png) as im:
        # the OLE preview picture (light blue) must be drawn on the rendered slide
        px = im.getpixel((im.width // 2, im.height // 2))
        assert px[2] > 200 and px[0] < 240
