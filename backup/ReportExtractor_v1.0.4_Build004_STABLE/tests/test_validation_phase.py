"""Real-report validation phase: tolerant Qwen JSON, ambiguity handling,
'Xử lý tạm thời' column, per-record diagnostics."""
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

from app.batch_processor import BatchOptions, BatchProcessor, format_file_diagnostics
from app.classifier import classify, heuristic_classify, merge_llm_with_heuristic, normalize_llm_response
from app.excel_writer import ExcelWriter
from app.extractor import ExtractedRecord, extract_record
from app.ollama_client import OllamaError, parse_json_response
from app.pptx_parser import parse_pptx


# ------------------------------------------------------------------ Qwen JSON tolerance
def test_normalize_llm_response_key_variants():
    raw = {"Management Number": "260918080-VOC", "modelName": "SM-A185", "Item": "Rear",
           "QPN": "Slide 2", "root_cause_slides": "3, 4", "temporary_handling": [4],
           "corrective_action_slides": ["5", "6"], "long_term_action_slides": [7],
           "before_after_slides": [5, 6], "verification": {"slides": [8]}, "confidence": 85}
    n = normalize_llm_response(raw)
    assert n["management_number"] == "260918080-VOC" and n["model"] == "SM-A185"
    assert n["qpn_slide"] == "Slide 2"
    assert n["confidence"] == 0.85


def test_normalize_llm_response_nested_and_ranges(a185_report):
    r = parse_pptx(a185_report)
    raw = {"result": {"slides": {"qpn": 2, "cause": {"slide": 3}, "countermeasures": "5-7",
                                 "temporary": "4", "images": [{"slide": 5}, {"slide": 6}]}}}
    c = merge_llm_with_heuristic(raw, heuristic_classify(r), r)
    assert c.qpn_slide == 2 and c.cause_slides == [3]
    assert c.improvement_slides == [5, 6, 7] and c.temporary_slides == [4]
    assert c.improvement_image_slides == [5, 6]
    assert c.source == "qwen"


def test_llm_garbage_values_are_ignored(a185_report):
    r = parse_pptx(a185_report)
    raw = {"qpn_slide": "null", "cause_slides": None, "improvement_slides": "n/a", "model": None, "item": ["x"]}
    c = merge_llm_with_heuristic(raw, heuristic_classify(r), r)
    assert c.qpn_slide == 2 and c.cause_slides == [3] and c.improvement_slides == [5, 6, 7]
    assert c.model == "" and c.item == ""
    assert c.source == "qwen+heuristic"       # heuristics filled the gaps -> not reported as pure qwen


def test_think_block_is_stripped():
    from app.ollama_client import OllamaClient
    import re
    text = "<think>reasoning…</think>\n{\"qpn_slide\": 2}"
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.DOTALL)
    assert parse_json_response(text) == {"qpn_slide": 2}


# ------------------------------------------------------------------ ambiguity -> Cần kiểm tra
def test_llm_only_qpn_and_disagreement_flag_review(report_factory):
    p = report_factory("noqpn.pptx", with_qpn=False)
    r = parse_pptx(p)
    heur = heuristic_classify(r)
    assert heur.qpn_slide is None
    raw = {"qpn_slide": 1, "cause_slides": [7], "improvement_slides": [2], "temporary_slides": [3]}
    c = merge_llm_with_heuristic(raw, heur, r)
    # the cover slide is never accepted as QPN on the LLM's word alone
    assert c.qpn_slide is None and "cover slide without QPN evidence" in c.qpn_override
    assert not any("chỉ do AI" in a for a in c.ambiguities)
    assert any("đối sách không thống nhất" in a for a in c.ambiguities)
    assert any("nguyên nhân không thống nhất" in a for a in c.ambiguities)
    rec = extract_record(r, c)
    assert "Không tìm thấy QPN trong báo cáo" in rec.review_reasons   # never guessed silently
    # a non-cover slide proposed by the LLM without evidence is accepted but flagged
    raw2 = {"qpn_slide": 3, "cause_slides": [7], "improvement_slides": [2], "temporary_slides": [3]}
    c2 = merge_llm_with_heuristic(raw2, heur, r)
    assert c2.qpn_slide == 3 and c2.qpn_source == "llm" and any("chỉ do AI" in a for a in c2.ambiguities)


def test_temporary_heading_slide_removed_from_llm_improvement(a185_report):
    r = parse_pptx(a185_report)
    raw = {"qpn_slide": 2, "cause_slides": [3], "improvement_slides": [4, 5, 6, 7], "temporary_slides": []}
    c = merge_llm_with_heuristic(raw, heuristic_classify(r), r)
    assert 4 not in c.improvement_slides and 4 in c.temporary_slides
    rec = extract_record(r, c)
    assert "Sorting 100%" not in rec.improvement


def test_low_confidence_flags_review(a185_report):
    r = parse_pptx(a185_report)
    raw = {"qpn_slide": 2, "cause_slides": [3], "improvement_slides": [5, 6, 7], "confidence": 0.3}
    c = merge_llm_with_heuristic(raw, heuristic_classify(r), r)
    assert c.confidence == 0.3 and any("tin cậy thấp" in a for a in c.ambiguities)


# ------------------------------------------------------------------ Excel: Xử lý tạm thời column
def _template_with_temporary(template):
    wb = load_workbook(template)
    ws = wb["Kiểm chứng"]
    # insert a header "Nội dung đối sách tạm thời" (tricky: contains 'đối sách') at column T
    ws.cell(row=2, column=20, value="Nội dung đối sách tạm thời").font = Font(bold=True)
    ws.merge_cells(start_row=2, start_column=20, end_row=3, end_column=20)
    wb.save(template)
    return template


def test_temporary_column_detected_and_left_blank(template, tmp_path):
    _template_with_temporary(template)
    out = tmp_path / "o.xlsx"
    w = ExcelWriter(template, out)
    assert w.columns["temporary"] == 20
    assert w.columns["improvement"] == 10                      # not hijacked by the temporary header
    rec = ExtractedRecord(model="A185", improvement="đối sách", temporary_excluded="Sorting 100%")
    row = w.append_record(rec)
    w.save()
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=row, column=20).value is None
    assert ws.cell(row=row, column=10).value == "đối sách"
    # opt-in fill keeps it separate from improvement
    w2 = ExcelWriter(template, out)
    row2 = w2.append_record(rec, fill_temporary=True)
    w2.save()
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=row2, column=20).value == "Sorting 100%"
    assert "Sorting" not in ws.cell(row=row2, column=10).value


def test_mapping_sheet_without_header_row(template, tmp_path):
    wb = load_workbook(template)
    ws = wb["Phân loại"]
    wb.remove(ws)
    m = wb.create_sheet("Phân loại")
    m.append(["Nắp lưng", "Rear"])
    m.append(["Bo mạch", "PBA"])
    m.append(["Main"])
    wb.save(template)
    w = ExcelWriter(template, tmp_path / "o.xlsx")
    assert w.item_mapping["Nắp lưng"] == "Rear" and w.item_mapping["Main"] == "Main"


# ------------------------------------------------------------------ diagnostics
def test_batch_result_has_diagnostics_and_review_report(sample_tree, tmp_path):
    out = tmp_path / "Output" / "r.xlsx"
    files = [sample_tree["files"][0], sample_tree["files"][3]]
    opts = BatchOptions(files=files, template=sample_tree["template"], output_file=out, use_ollama=False, row_mode="append")
    proc = BatchProcessor(opts)
    proc.run()
    res = json.loads((out.parent / "logs" / "batch_result.json").read_text(encoding="utf-8"))["results"]
    r0, r1 = res
    assert r0["classifier"] == "heuristic"
    assert r0["temporary_slides"] == [4] and r0["verify_slides"] == [8] and r0["defect_slide"] == 2
    assert r0["blank_fields"] == [] and r0["review_reasons"] == []
    assert r1["status"] == "needs_review" and r1["qpn_slide"] is None
    assert any("QPN" in x for x in r1["review_reasons"])
    txt = (out.parent / "logs" / "review_report.txt").read_text(encoding="utf-8")
    assert "Bộ phân loại       : heuristic" in txt and "Slide xử lý tạm thời: [4]" in txt
    diag = format_file_diagnostics(proc.results[0])
    assert "Management number  : 260918080-VOC" in diag and "WEEK +1..+8 luôn để trống" in diag


def test_classifier_name_is_qwen_when_llm_used(a185_report):
    class Fake:
        def generate_json(self, model, prompt, system="", **kw):
            return {"qpn_slide": 2, "cause_slides": [3], "improvement_slides": [5, 6, 7],
                    "improvement_image_slides": [5, 6], "temporary_slides": [4], "confidence": 0.9}
    c = classify(parse_pptx(a185_report), Fake(), "qwen3:4b")
    assert c.source == "qwen" and c.confidence == 0.9 and c.ambiguities == []


# ---------------------------------------------------------------- compact classifier prompt
def test_prompt_is_compact_and_contains_only_slide_text(sample_tree):
    from app.classifier import SYSTEM_PROMPT, build_prompt, compact_slide_line
    from app.pptx_parser import parse_pptx
    report = parse_pptx(sample_tree["files"][0])
    prompt = build_prompt(report)
    import re as _re
    lines = [ln for ln in prompt.splitlines() if _re.match(r"^S\d+ \| ", ln)]
    assert len(lines) == len(report.slides)
    # one line per slide, hard cap per slide; no XML / image / shape data
    for s in report.slides:
        assert len(compact_slide_line(s)) <= 420 + len("S99 |  | pics=99 | ") + 90
    slide_part = "\n".join(lines)
    assert "<" not in slide_part and "xml" not in slide_part.lower() and "image/" not in slide_part
    assert "shape" not in slide_part.lower() and "picture(s)" not in slide_part
    assert "management_number" not in prompt and '"model"' not in prompt   # text values are not asked from the LLM
    assert "think" not in SYSTEM_PROMPT.lower()
    assert len(prompt) < 3000 and len(SYSTEM_PROMPT) < 300
    # duplicated shape text is collapsed
    from app.pptx_parser import Block, SlideData
    sd = SlideData(number=1, blocks=[Block(kind="paragraph", text="Tiêu đề\nDòng A"), Block(kind="paragraph", text="Dòng A\nDòng A")])
    assert compact_slide_line(sd) == "S1 | Tiêu đề | pics=0 | Dòng A"


def test_ollama_failure_never_stops_batch_and_is_logged_with_timing(sample_tree, tmp_path, monkeypatch):
    from app.batch_processor import BatchOptions, BatchProcessor
    out = tmp_path / "o" / "r.xlsx"
    # unreachable server, 1s timeout -> heuristic fallback, batch completes
    opts = BatchOptions(files=[sample_tree["files"][0]], template=sample_tree["template"], output_file=out,
                        ollama_server="127.0.0.1:1", model="qwen3:4b", use_ollama=True, request_timeout=1,
                        row_mode="append")
    proc = BatchProcessor(opts, on_file=lambda *a: None)
    s = proc.run()
    assert s.failed == 0 and s.completed + s.needs_review == 1
    fr = proc.results[0]
    assert fr.classifier == "heuristic"
    assert any("heuristic fallback" in n and "s," in n for n in fr.classifier_notes)
