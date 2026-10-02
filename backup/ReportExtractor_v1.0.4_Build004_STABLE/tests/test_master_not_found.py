"""Management Number absent from the master – PROMPT-004D authoritative rule.

History: PROMPT-004 reported such a report as MASTER_NOT_FOUND and never wrote a row.  PROMPT-004D superseded that:
a VALID Management Number parsed from the filename that is absent from the master now gets exactly ONE new production
row (PROCESS_NEW_ROW), visible in prescan/GUI, processed normally.  A filename WITHOUT a valid Management Number still
creates nothing ("Cần kiểm tra — Không xác định được Management Number từ tên file").
"""
import os
import shutil
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

import app.batch_processor as bp
from app.batch_processor import BatchOptions, BatchProcessor
from app.config import AppConfig
from app.gui_controller import GuiController
from app.prescan import (ACTION_INVALID_MGMT, ACTION_OUTSIDE_PERIOD, ACTION_PROCESS, ACTION_PROCESS_NEW_ROW,
                         ACTION_SOURCE_DUPLICATE, CACHE_FILE_NAME, PROCESS_ACTIONS, FastScanCache, MasterLookup,
                         month_period, prescan)

SHEET = "Kiểm chứng"
SEP = month_period(2026, 9)
RUN_TODAY = date(2026, 9, 20)


def _fake(folder: Path, name: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    p = folder / name
    p.write_bytes(b"not a pptx")
    return p


def _run(files, template, out, **kw):
    events = []
    opts = BatchOptions(files=[Path(f) for f in files], template=template, output_file=out, use_ollama=False,
                        **{"today": RUN_TODAY, "period": SEP, **kw})
    proc = BatchProcessor(opts, on_file=lambda i, s, d: events.append((i, s, d)))
    return proc.run(), events, proc


def _keys(out):
    ws = load_workbook(out)[SHEET]
    return [ws.cell(row=r, column=2).value for r in range(4, 10)]


# ---------------------------------------------------------------- absent VALID key -> PROCESS_NEW_ROW -> one row
def test_absent_key_is_planned_as_new_row_created_once_and_processed(sample_tree, template, tmp_path, monkeypatch):
    calls = {"parse": 0}
    real = bp.parse_pptx

    def spy(p):
        calls["parse"] += 1
        return real(p)
    monkeypatch.setattr(bp, "parse_pptx", spy)
    out = tmp_path / "Output" / "k.xlsx"
    s, ev, proc = _run([sample_tree["files"][0]], template, out)
    it = proc.prescan_result.items[0]
    assert it.action == ACTION_PROCESS_NEW_ROW and it.action in PROCESS_ACTIONS and s.candidates == 1
    assert it.reason == "Sẽ thêm dòng mới cho Management Number 260918080-VOC"
    assert s.new_rows_planned == 1 and s.new_rows == 1 and s.completed == 1 and s.not_written == 0
    assert calls["parse"] == 1                                                   # processed normally, once
    fr = proc.results[0]
    assert fr.status == "completed" and fr.new_row and fr.excel_row == 4
    assert _keys(out) == ["260918080-VOC", None, None, None, None, None]         # exactly one row, key written
    assert load_workbook(out)[SHEET].cell(row=4, column=5).value == "A185"
    assert ev[-1][1] == "completed_new"
    assert bp.STAGE_LABELS_VI["completed_new"] == "Hoàn thành — đã thêm Management Number mới"
    assert s.backup_file and Path(s.backup_file).exists()                       # backup taken (before the row)
    assert _keys(s.backup_file) == [None] * 6                                    # … and it holds the pre-change master


def test_bare_and_suffixed_forms_match_existing_rows(sample_tree, template, tmp_path):
    wb = load_workbook(template)
    ws = wb[SHEET]
    ws.cell(row=4, column=2, value=260918080)             # numeric bare form stored by Excel
    ws.cell(row=5, column=2, value=" 260918081-voc ")     # spacing / case
    wb.save(template)
    a = _fake(tmp_path / "in", "(CTMS)_11108_260918080_ Đối sách LỖI X 18.9.2026.pptx")
    b = _fake(tmp_path / "in", "260918081-VOC_b.pptx")
    c = _fake(tmp_path / "in", "260918080-VOC_c.pptx")   # suffixed form of a bare key: NOT the same key -> new row
    s, ev, proc = _run([a, b, c], template, tmp_path / "Output" / "k.xlsx")
    acts = [it.action for it in proc.prescan_result.items]
    assert acts[0] == ACTION_PROCESS and acts[1] == ACTION_PROCESS and acts[2] == ACTION_PROCESS_NEW_ROW
    assert proc.prescan_result.items[0].excel_row == 4 and proc.prescan_result.items[1].excel_row == 5


# ---------------------------------------------------------------- no valid key / other rejections -> nothing written
def test_invalid_key_outside_period_and_duplicates_leave_master_untouched_only_absent_key_gets_a_row(template, tmp_path):
    nokey = _fake(tmp_path / "in", "bao_cao_khong_ma.pptx")
    baddate = _fake(tmp_path / "in", "(CTMS)_11107_261345001_ Đối sách.pptx")
    octo = _fake(tmp_path / "in", "261001002-VOC_october.pptx")
    d1 = _fake(tmp_path / "in" / "A", "260915003-VOC_a.pptx")
    d2 = _fake(tmp_path / "in" / "B", "260915003-VOC_b.pptx")
    os.utime(d1, (1_700_000_000, 1_700_000_000))
    os.utime(d2, (1_700_000_900, 1_700_000_900))
    out = tmp_path / "Output" / "k.xlsx"
    s, ev, proc = _run([nokey, baddate, octo, d1, d2], template, out)
    acts = [it.action for it in proc.prescan_result.items]
    assert acts[0] == ACTION_INVALID_MGMT and acts[1] == ACTION_INVALID_MGMT     # no valid key -> never a row
    assert acts[2] == ACTION_OUTSIDE_PERIOD and acts[3] == ACTION_SOURCE_DUPLICATE
    assert acts[4] == ACTION_PROCESS_NEW_ROW                                     # the only one that may write
    assert s.not_written == 2 and s.new_rows_planned == 1
    nk = next(r for r in proc.results if r.source_file.endswith("bao_cao_khong_ma.pptx"))
    assert nk.status == "not_written" and nk.excel_row is None and not nk.new_row
    assert "Không xác định được Management Number từ tên file" in nk.error
    assert bp.STAGE_LABELS_VI["not_written"] == "Cần kiểm tra — Không xác định được Management Number từ tên file"
    # d2 is a fake file: the row + key are created and saved, parsing then fails -> error, still only ONE row
    assert s.new_rows == 1 and _keys(out)[0] == "260915003-VOC" and _keys(out)[1:] == [None] * 5


def test_existing_partial_rows_still_filled_and_duplicates_marked(sample_tree, template, tmp_path):
    wb = load_workbook(template)
    ws = wb[SHEET]
    ws.cell(row=4, column=2, value="260918080-VOC")
    ws.cell(row=5, column=2, value="260918081-VOC")
    ws.cell(row=6, column=2, value="260918081-VOC")
    wb.save(template)
    out = tmp_path / "Output" / "k.xlsx"
    s, ev, proc = _run([sample_tree["files"][0], sample_tree["files"][1]], template, out)
    assert [it.action for it in proc.prescan_result.items] == [ACTION_PROCESS, ACTION_PROCESS]
    assert s.new_rows == 0 and all(not r.new_row for r in proc.results)
    ws2 = load_workbook(out)[SHEET]
    assert ws2.cell(row=4, column=5).value == "A185" and ws2.cell(row=5, column=5).value == "A185"
    assert ws2.cell(row=6, column=5).value is None and ws2.cell(row=7, column=2).value is None
    assert ws2.cell(row=6, column=2).fill.fgColor.rgb.endswith("FFC7CE")
    s2, ev2, proc2 = _run([sample_tree["files"][0]], template, out)
    assert s2.skipped == 1 and s2.new_rows == 0 and load_workbook(out)[SHEET].cell(row=7, column=2).value is None


# ---------------------------------------------------------------- cache hit never hides a missing row -> new row
def test_cache_hit_with_missing_row_is_a_new_row_not_a_fast_skip(tmp_path):
    f = _fake(tmp_path / "in", "260915001-VOC_r.pptx")
    cache = FastScanCache(tmp_path / CACHE_FILE_NAME)
    cache.record("260915001-VOC", f, "completed")
    res = prescan([f], SEP, cache, MasterLookup(lambda m: [], lambda r: []), today=RUN_TODAY)
    it = res.items[0]
    assert it.action == ACTION_PROCESS_NEW_ROW and it.cache_decision == "miss:master_row_missing"
    assert res.candidates == [it] and res.counts()["fast_skipped"] == 0 and res.counts()["new_rows"] == 1
    assert "Management Number mới (sẽ thêm dòng): 1" in res.summary_lines_vi()


# ---------------------------------------------------------------- new monthly workbook without the key -> one new row
def test_new_monthly_workbook_without_the_key_gets_one_new_row(sample_tree, template, tmp_path):
    wb = load_workbook(template)
    wb[SHEET].cell(row=4, column=2, value="260918080-VOC")
    wb.save(template)
    out1 = tmp_path / "Output" / "Kiem_chung_08_2026.xlsx"
    s1, _, _ = _run([sample_tree["files"][0]], template, out1)
    assert s1.completed == 1 and s1.new_rows == 0
    # a different workbook created from a template WITHOUT the key -> the cache must not hide it; one row is created
    tpl2 = template.with_name("tpl2.xlsx")
    wb = load_workbook(template)
    wb[SHEET].cell(row=4, column=2).value = None
    wb.save(tpl2)
    out2 = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    s2, ev2, proc2 = _run([sample_tree["files"][0]], tpl2, out2)
    assert s2.fast_skipped == 0 and s2.new_rows == 1 and s2.not_written == 0 and s2.completed == 1
    assert _keys(out2) == ["260918080-VOC", None, None, None, None, None]
    # rerun on the new workbook: row found, complete -> skipped, no second row
    s3, ev3, proc3 = _run([sample_tree["files"][0]], tpl2, out2)
    assert s3.skipped == 1 and s3.new_rows == 0 and not s3.backup_file
    assert proc3.results[0].error.startswith("Bỏ qua")          # fast-cache skip or "Bỏ qua — đã cập nhật": no Qwen, no row
    assert _keys(out2) == ["260918080-VOC", None, None, None, None, None]


# ---------------------------------------------------------------- GUI: new row visible, counted, in the denominator
def test_gui_prescan_counts_new_row_inside_the_denominator(sample_tree, template, tmp_path, monkeypatch):
    class _D(date):
        @classmethod
        def today(cls):
            return RUN_TODAY
    monkeypatch.setattr(bp, "date", _D)
    reports = tmp_path / "Bao_cao" / "2026-09"
    reports.mkdir(parents=True)
    shutil.copy(sample_tree["files"][0], reports / "260918080-VOC_a.pptx")
    shutil.copy(sample_tree["files"][1], reports / "260918081-VOC_b.pptx")
    wb = load_workbook(template)
    wb[SHEET].cell(row=4, column=2, value="260918080-VOC")
    wb.save(template)
    ctl = GuiController(AppConfig(), config_path=tmp_path / "c.json")
    ctl.set_report_folder(str(reports))
    ctl.set_template(str(template))
    ctl.set_output(str(tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"))
    assert ctl.start(use_ollama=False)
    ctl.processor._thread.join(120)
    ctl.pump()
    lines = ctl.summary_lines()
    assert "Management Number mới (sẽ thêm dòng): 1" in lines and "Cần xử lý thực tế: 2" in lines
    assert "Management Number mới: 1" in lines
    assert ctl.progress.total == 2 and len(ctl.files) == 2                        # new-row report IS in the denominator
    scan = {r.path.name: r for r in ctl.scan_rows()}
    row = scan["260918081-VOC_b.pptx"]
    assert row.status_vi == "Sẽ xử lý — Management Number mới" and row.will_process and row.is_new_row
    new = next(ctl.result_for(i) for i in range(2) if ctl.result_for(i).management_number == "260918081-VOC")
    assert new.new_row and new.excel_row == 5 and new.status in ("completed", "needs_review")
    gui_row = next(r for r in ctl.rows if r.management_number == "260918081-VOC")
    if new.status == "completed":
        assert gui_row.stage == "completed_new" and gui_row.status_vi == "Hoàn thành — đã thêm Management Number mới"
    assert ctl.diagnostics_for(ctl.files.index(Path(new.source_file)))["Dòng sử dụng"] == "5"


# ---------------------------------------------------------------- the batch DOES use ExcelWriter.create_row (once per key)
def test_excel_writer_create_row_is_called_exactly_once_per_new_key(sample_tree, template, tmp_path, monkeypatch):
    import app.excel_writer as xw
    calls = []
    orig = xw.ExcelWriter.create_row

    def spy(self, mgmt):
        calls.append((mgmt, self.backup_path is not None))
        return orig(self, mgmt)
    monkeypatch.setattr(xw.ExcelWriter, "create_row", spy)
    out = tmp_path / "Output" / "k.xlsx"
    s, _, proc = _run([sample_tree["files"][0]], template, out)
    assert calls == [("260918080-VOC", False)] and proc.results[0].new_row       # called once, backup not yet taken …
    assert s.backup_file and Path(s.backup_file).exists()                       # … create_row takes it before changing
    s2, _, proc2 = _run([sample_tree["files"][0]], template, out)
    assert len(calls) == 1 and s2.skipped == 1                                  # rerun: existing row found, no call
