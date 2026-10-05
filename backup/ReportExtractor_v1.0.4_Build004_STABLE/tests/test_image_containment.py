"""Rules #31-#43: every QPN / "Sau cải tiến" picture stays completely inside its Excel destination rectangle.

Real root causes fixed (both produced pictures wider than the real column):
1. ``col_width_to_px`` added 5 px padding to a width that openpyxl already returns padded (stored XML width).
2. ``column_dimensions[letter]`` returns a DEFAULT dimension for columns inside a ``<col min max>`` range, so a
   narrow grouped column was treated as 64 px wide.
Vertical: layouts are now fitted against the FINAL row height (not only the 409 pt cap) and validated before anchors.
"""
import itertools
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

from app.excel_writer import (EMU_PER_PX, IMAGE_MARGIN_PX, MAX_ROW_HEIGHT_PT, ExcelWriter, TemplateError,
                              assert_image_inside_area, col_width_to_px, emu_to_px, images_inside_area,
                              plan_image_grid, pt_to_px, px_to_emu)
from app.extractor import ExtractedRecord

MGMT = "260918080-VOC"
IMG_COL, QPN_COL = 11, 8


def _png(path: Path, w: int, h: int) -> Path:
    Image.new("RGB", (w, h), (20, 120, 200)).save(path)
    return path


def _record():
    return ExtractedRecord(management_number=MGMT, model="A185", item="Rear", defect_content="Xước",
                           root_cause="NGUYÊN NHÂN\n- a", improvement="CẢI TIẾN\n- b", qpn_slide=2)


def _prepare(template: Path, row=4, widths=None):
    wb = load_workbook(template)
    ws = wb["Kiểm chứng"]
    ws.cell(row=row, column=2, value=MGMT)
    for col, w in (widths or {}).items():
        ws.column_dimensions[get_column_letter(col)].width = w
    wb.save(template)


def _geom(img):
    a = img.anchor
    return (emu_to_px(a._from.colOff), emu_to_px(a._from.rowOff), emu_to_px(a.ext.cx), emu_to_px(a.ext.cy))


def _images_in(ws, row, col):
    return [i for i in ws._images if i.anchor._from.row == row - 1 and i.anchor._from.col == col - 1]


def _check_row(w: ExcelWriter, ws, row: int, col: int, field: str, n_expected: int):
    """#41/#42 on the SAVED workbook: inside area, no overlap, bottom-most image + margin <= row height."""
    imgs = _images_in(ws, row, col)
    assert len(imgs) == n_expected
    items = [_geom(i) for i in imgs]
    area_w, area_h = w.image_area_px(row, field), pt_to_px(ws.row_dimensions[row].height)
    assert_image_inside_area(items, area_w, area_h)
    assert max(y + h for _, y, _, h in items) + IMAGE_MARGIN_PX <= area_h + 1
    assert ws.row_dimensions[row].height <= MAX_ROW_HEIGHT_PT
    return items


# ---------------------------------------------------------------- helpers / unit layer
def test_unit_helpers_are_consistent():
    assert px_to_emu(10) == 10 * EMU_PER_PX and emu_to_px(px_to_emu(37)) == 37
    assert pt_to_px(15) == 20 and pt_to_px(409) == 545
    assert col_width_to_px(9.140625) == 64 and col_width_to_px(0) == 0


def test_assert_image_inside_area_detects_each_edge_and_overlap():
    assert images_inside_area([(4, 4, 92, 50)], 100, 60)
    with pytest.raises(ValueError, match="right edge"):
        assert_image_inside_area([(4, 4, 94, 50)], 100, 60)
    with pytest.raises(ValueError, match="bottom edge"):
        assert_image_inside_area([(4, 4, 92, 54)], 100, 60)
    with pytest.raises(ValueError, match="margin"):
        assert_image_inside_area([(1, 4, 50, 50)], 100, 60)
    with pytest.raises(ValueError, match="overlaps"):
        assert_image_inside_area([(4, 4, 40, 40), (20, 20, 40, 30)], 200, 200)
    assert images_inside_area([(4, 4, 93, 50)], 100, 60)          # 1 px rounding tolerance only


SIZES = [(4000, 1000), (300, 2400), (640, 480), (1920, 1080), (100, 100), (2, 3000), (3000, 2)]


@pytest.mark.parametrize("n", [1, 2, 3, 5, 7])
@pytest.mark.parametrize("area", [(120, None), (300, None), (640, None), (300, 200), (900, 120), (64, 545)])
def test_every_generated_layout_is_inside_its_area(n, area):
    area_w, area_h = area
    for combo in itertools.islice(itertools.permutations(SIZES, n), 6):
        plan = plan_image_grid(list(combo), area_w, area_h_px=area_h)
        limit_h = pt_to_px(MAX_ROW_HEIGHT_PT) if area_h is None else min(area_h, pt_to_px(MAX_ROW_HEIGHT_PT))
        assert_image_inside_area(plan["items"], area_w, limit_h)
        assert plan["height_pt"] <= MAX_ROW_HEIGHT_PT
        assert len(plan["items"]) == n
        for (w, h), (_, _, pw, ph) in zip(combo, plan["items"]):
            if pw >= 12 and ph >= 12:
                assert abs(pw / ph - w / h) / (w / h) < 0.08       # aspect preserved (floor rounding only)


def test_fit_respects_height_as_well_as_width_but_still_fills_width_when_possible():
    wide = plan_image_grid([(400, 100)], 300, area_h_px=400)
    assert wide["items"][0][2] == 300 - 2 * IMAGE_MARGIN_PX            # #43: narrow area, fill width
    tall = plan_image_grid([(100, 400)], 300, area_h_px=120)           # #34: height is the binding constraint
    (x, y, w, h), = tall["items"]
    assert y + h <= 120 - IMAGE_MARGIN_PX and w < 300 - 2 * IMAGE_MARGIN_PX and abs(w / h - 0.25) < 0.05
    assert x > IMAGE_MARGIN_PX                                          # centred in the slot, no stretch


# ---------------------------------------------------------------- real workbook geometry
def test_column_width_resolution_uses_col_ranges_hidden_and_sheet_default(template, tmp_path):
    _prepare(template)
    w = ExcelWriter(template, tmp_path / "o.xlsx")
    ws = w.ws
    dim = ws.column_dimensions[get_column_letter(IMG_COL)]
    dim.width, dim.min, dim.max = 12, IMG_COL, IMG_COL + 2               # <col min=11 max=13 width=12>
    assert w._col_width_chars(IMG_COL + 2) == 12                        # inside the range, not a 64 px default
    assert w.image_area_px(4, "improvement_image") == col_width_to_px(12)
    dim.hidden = True
    assert w.image_area_px(4, "improvement_image") == 0
    dim.hidden = False
    ws.sheet_format.defaultColWidth = 20                                # column never defined -> sheet default
    assert w._col_width_chars(40) == 20


def test_merged_destination_uses_sum_of_all_columns_and_rows(template, tmp_path):
    _prepare(template, widths={IMG_COL: 10, IMG_COL + 1: 14})
    wb = load_workbook(template)
    ws = wb["Kiểm chứng"]
    ws.merge_cells(start_row=4, start_column=IMG_COL, end_row=5, end_column=IMG_COL + 1)
    ws.row_dimensions[5].height = 30
    wb.save(template)
    w = ExcelWriter(template, tmp_path / "o.xlsx")
    assert w.image_area_px(4, "improvement_image") == col_width_to_px(10) + col_width_to_px(14)
    w.update_record(4, _record(), improvement_jpg=[_png(tmp_path / "a.png", 800, 300)])
    out = w.save()
    ws2 = load_workbook(out)["Kiểm chứng"]
    (img,) = _images_in(ws2, 4, IMG_COL)
    x, y, iw, ih = _geom(img)
    assert x + iw <= w.image_area_px(4, "improvement_image") - IMAGE_MARGIN_PX
    area_h = pt_to_px(ws2.row_dimensions[4].height) + pt_to_px(ws2.row_dimensions[5].height)
    assert y + ih + IMAGE_MARGIN_PX <= area_h + 1


@pytest.mark.parametrize("qpn_size", [(3000, 400), (400, 3000), (640, 640)])
def test_qpn_and_five_after_images_never_leave_their_cells(template, tmp_path, qpn_size):
    _prepare(template, widths={QPN_COL: 18, IMG_COL: 28})
    qpn = _png(tmp_path / "qpn.png", *qpn_size)
    pics = [_png(tmp_path / f"p{i}.png", *s) for i, s in enumerate([(1600, 900), (300, 1200), (1000, 1000),
                                                                      (2400, 600), (900, 1600)])]
    w = ExcelWriter(template, tmp_path / "o.xlsx")
    w.update_record(4, _record(), qpn_png=qpn, improvement_jpg=pics)
    assert w.image_bounds_report() == []
    out = w.save()
    ws = load_workbook(out)["Kiểm chứng"]
    (qx, qy, qw, qh), = _check_row(w, ws, 4, QPN_COL, "qpn", 1)
    assert abs(qw / qh - qpn_size[0] / qpn_size[1]) / (qpn_size[0] / qpn_size[1]) < 0.05
    items = _check_row(w, ws, 4, IMG_COL, "improvement_image", 5)
    assert len({(x, y) for x, y, _, _ in items}) == 5                    # five independent objects


def test_fill_missing_fields_path_fits_images_to_final_row_height(template, tmp_path):
    _prepare(template, widths={IMG_COL: 20})
    wb = load_workbook(template)
    wb["Kiểm chứng"].row_dimensions[4].height = 18                       # short existing row
    wb.save(template)
    w = ExcelWriter(template, tmp_path / "o.xlsx")
    w.fill_missing_fields(4, _record(), ["qpn", "improvement_image"], qpn_png=_png(tmp_path / "q.png", 500, 2500),
                          improvement_jpg=[_png(tmp_path / "a.png", 600, 600), _png(tmp_path / "b.png", 600, 2400)])
    out = w.save()
    ws = load_workbook(out)["Kiểm chứng"]
    _check_row(w, ws, 4, QPN_COL, "qpn", 1)
    _check_row(w, ws, 4, IMG_COL, "improvement_image", 2)


def test_row_height_is_final_before_anchors_and_capped_at_excel_limit(template, tmp_path):
    _prepare(template, widths={IMG_COL: 40})
    pics = [_png(tmp_path / f"t{i}.png", 400, 2600) for i in range(4)]   # ideal stack >> 409 pt
    w = ExcelWriter(template, tmp_path / "o.xlsx")
    w.update_record(4, _record(), improvement_jpg=pics)
    out = w.save()
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.row_dimensions[4].height <= MAX_ROW_HEIGHT_PT
    items = _check_row(w, ws, 4, IMG_COL, "improvement_image", 4)
    assert all(h > 60 for _, _, _, h in items)                            # shrunk/gridded, not thumbnails
    for (x1, y1, w1, h1), (x2, y2, w2, h2) in itertools.combinations(items, 2):
        assert x1 + w1 <= x2 + 1 or x2 + w2 <= x1 + 1 or y1 + h1 <= y2 + 1 or y2 + h2 <= y1 + 1


def test_save_refuses_tampered_overflowing_geometry(template, tmp_path):
    _prepare(template)
    w = ExcelWriter(template, tmp_path / "o.xlsx")
    w.update_record(4, _record(), improvement_jpg=[_png(tmp_path / "a.png", 400, 200)])
    img, _, _ = w._placed[0]
    img.anchor.ext.cx = px_to_emu(5000)                                     # simulate a layout bug
    assert w.image_bounds_report()
    with pytest.raises(TemplateError):
        w.save()
