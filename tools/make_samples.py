"""Generate sample PPTX reports and an Excel verification template for tests/demo.

Usage:  python tools/make_samples.py [target_dir]
"""
from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt

IMPROVEMENT_TEXT = (
    "Lỗi xước rear do va chạm với khay chứa ( Áp dụng cải tiến 17/9 – Công đoạn Assy Daoltech):\n"
    "- Tại công đoạn lắp ráp Rear, thao tác đặt sản phẩm vào khay\n"
    "+ Trước: Khay nhựa cứng không có lót xốp, sản phẩm tiếp xúc trực tiếp gây xước\n"
    "+ Sau: Thay khay mới có lót mút EVA 3mm, bổ sung vách ngăn từng ô\n"
    "\n"
    "Cải tiến lỗi mẻ rear khi tháo jig\n"
    "- Tại công đoạn ép Rear vào Frame\n"
    "+ Trước: Jig ép kim loại, cạnh sắc, lực ép 25kgf\n"
    "+ Sau: Thay jig bằng POM, bo tròn cạnh R0.5, giảm lực ép còn 18kgf"
)
IMPROVEMENT_TEXT_2 = (
    "Cải tiến trong kiểm tra\n"
    "- Bổ sung đèn kiểm tra 3 hướng tại OQC, tăng độ sáng từ 800 lux lên 1500 lux\n"
    "- Đào tạo lại 12 nhân viên kiểm tra về tiêu chuẩn lỗi xước/mẻ (ngày 20/09)"
)
LONG_TERM_TEXT = (
    "- Đưa hạng mục kiểm tra khay/jig vào checklist đầu ca (daily check sheet)\n"
    "- Cập nhật SOP lắp ráp Rear phiên bản Rev.03, triển khai ngang cho dòng A175, A165"
)
TEMP_TEXT = (
    "- Sorting 100% hàng tồn kho tại kho thành phẩm: 2.350 pcs\n"
    "- Dán tem đánh dấu tạm thời lô hàng đã kiểm tra lại\n"
    "- Tăng cường kiểm tra ngoại quan 100% tại OQC trong 2 tuần"
)
CAUSE_TEXT_1 = (
    "Nguyên nhân trong sản xuất\n"
    "- Khay chứa Rear không có lớp lót, sản phẩm va chạm khi vận chuyển giữa các công đoạn gây xước\n"
    "- Jig ép có cạnh sắc, lực ép cao làm mẻ cạnh Rear"
)
CAUSE_TEXT_2 = (
    "Nguyên nhân trong kiểm tra\n"
    "- Đèn kiểm tra tại OQC không đủ sáng (800 lux), góc nhìn 1 hướng nên bỏ sót lỗi xước nhỏ"
)


def _pic(color: str, label: str, size=(640, 360)) -> BytesIO:
    im = Image.new("RGB", size, color)
    d = ImageDraw.Draw(im)
    d.rectangle([10, 10, size[0] - 10, size[1] - 10], outline="black", width=4)
    d.text((30, 30), label, fill="black")
    bio = BytesIO()
    im.save(bio, "PNG")
    bio.seek(0)
    return bio


def _textbox(slide, text, left, top, width, height, size=14, bold=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
    return tb


def make_report(path: Path, model="SM-A185", item="Rear", mgmt="260918080-VOC",
                with_qpn=True, with_improvement=True) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    W = prs.slide_width

    # slide 1 – cover
    s = prs.slides.add_slide(blank)
    _textbox(s, f"BÁO CÁO ĐỐI SÁCH LỖI XƯỚC / MẺ {item.upper()} {model}", Inches(0.5), Inches(1), W - Inches(1), Inches(1), 28, True)
    _textbox(s, f"Model: {model}\nItem: {item}\nManagement No: {mgmt}\nVendor: CÔNG TY TNHH ABC VINA\nNgày phát sinh: 18/09/2026",
             Inches(0.5), Inches(2.5), Inches(6), Inches(2.5), 16)

    # slide 2 – QPN
    if with_qpn:
        s = prs.slides.add_slide(blank)
        _textbox(s, "Quality Problem Notice", Inches(0.5), Inches(0.3), Inches(8), Inches(0.7), 24, True)
        rows, cols = 4, 4
        tbl = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.2), W - Inches(1), Inches(3.2)).table
        headers = ["Model", "Defect Info", "Cause & Action", "PIC"]
        for c, h in enumerate(headers):
            tbl.cell(0, c).text = h
        tbl.cell(1, 0).text = model
        tbl.cell(1, 1).text = "Xước: 15ea\nLệch ANT: 5ea\nMẻ: 13ea"
        tbl.cell(1, 2).text = "Xem báo cáo đối sách"
        tbl.cell(1, 3).text = "Nguyễn Văn A"
        tbl.cell(2, 0).text = "Line"
        tbl.cell(2, 1).text = "OQC-02"
        tbl.cell(3, 0).text = "Date"
        tbl.cell(3, 1).text = "2026-09-18"
        s.shapes.add_picture(_pic("#ffd6d6", "NG image"), Inches(0.5), Inches(4.6), Inches(3), Inches(2))
        s.shapes.add_picture(_pic("#e0e0ff", "Confirmation sheet"), Inches(4), Inches(4.6), Inches(3), Inches(2))

    # slide 3 – root cause
    s = prs.slides.add_slide(blank)
    _textbox(s, "2. NGUYÊN NHÂN", Inches(0.5), Inches(0.3), Inches(8), Inches(0.7), 24, True)
    _textbox(s, CAUSE_TEXT_1, Inches(0.5), Inches(1.2), W - Inches(1), Inches(2.4), 14)
    _textbox(s, CAUSE_TEXT_2, Inches(0.5), Inches(3.8), W - Inches(1), Inches(1.5), 14)

    # slide 4 – temporary handling (must be excluded)
    s = prs.slides.add_slide(blank)
    _textbox(s, "XỬ LÝ TẠM THỜI", Inches(0.5), Inches(0.3), Inches(8), Inches(0.7), 24, True)
    _textbox(s, TEMP_TEXT, Inches(0.5), Inches(1.2), W - Inches(1), Inches(3), 14)
    s.shapes.add_picture(_pic("#fff2cc", "Sorting 100%"), Inches(0.5), Inches(4.5), Inches(4), Inches(2.3))

    if with_improvement:
        # slide 5 – improvement in production (with before/after pictures)
        s = prs.slides.add_slide(blank)
        _textbox(s, "3. CẢI TIẾN TRONG SẢN XUẤT", Inches(0.5), Inches(0.3), Inches(8), Inches(0.7), 24, True)
        _textbox(s, IMPROVEMENT_TEXT, Inches(0.5), Inches(1.1), Inches(7.5), Inches(5.5), 12)
        s.shapes.add_picture(_pic("#ffe0b2", "Before: khay cũ"), Inches(8.3), Inches(1.2), Inches(4.5), Inches(2.5))
        s.shapes.add_picture(_pic("#c8e6c9", "After: khay lót EVA"), Inches(8.3), Inches(4.0), Inches(4.5), Inches(2.5))

        # slide 6 – improvement in inspection
        s = prs.slides.add_slide(blank)
        _textbox(s, "4. CẢI TIẾN TRONG KIỂM TRA", Inches(0.5), Inches(0.3), Inches(8), Inches(0.7), 24, True)
        _textbox(s, IMPROVEMENT_TEXT_2, Inches(0.5), Inches(1.1), Inches(7.5), Inches(3), 14)
        s.shapes.add_picture(_pic("#b3e5fc", "Đèn kiểm tra 3 hướng"), Inches(8.3), Inches(1.2), Inches(4.5), Inches(2.5))
        s.shapes.add_picture(_pic("#d1c4e9", "Đào tạo nhân viên"), Inches(8.3), Inches(4.0), Inches(4.5), Inches(2.5))

        # slide 7 – long-term countermeasure (no picture)
        s = prs.slides.add_slide(blank)
        _textbox(s, "5. ĐỐI SÁCH LÂU DÀI", Inches(0.5), Inches(0.3), Inches(8), Inches(0.7), 24, True)
        _textbox(s, LONG_TERM_TEXT, Inches(0.5), Inches(1.1), W - Inches(1), Inches(3), 14)

    # slide 8 – verification (no data yet)
    s = prs.slides.add_slide(blank)
    _textbox(s, "6. KIỂM CHỨNG HIỆU QUẢ", Inches(0.5), Inches(0.3), Inches(8), Inches(0.7), 24, True)
    _textbox(s, "Theo dõi 8 tuần sau cải tiến (cập nhật sau)", Inches(0.5), Inches(1.2), Inches(8), Inches(1), 14)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def make_template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kiểm chứng"
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.merge_cells("A1:S1")
    ws["A1"] = "BẢNG KIỂM CHỨNG ĐỐI SÁCH CẢI TIẾN"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    headers = ["STT", "Management number", "Tên vendor", "Ngày phát sinh", "Model", "Item",
               "Nội dung lỗi", "QPN", "Nguyên nhân", "Nội dung đối sách cải tiến", "Hình ảnh cải tiến"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=2, start_column=i, end_row=3, end_column=i)
    ws.merge_cells(start_row=2, start_column=12, end_row=2, end_column=19)
    ws.cell(row=2, column=12, value="Kiểm chứng").font = Font(bold=True)
    ws.cell(row=2, column=12).alignment = Alignment(horizontal="center")
    for k in range(1, 9):
        c = ws.cell(row=3, column=11 + k, value=f"WEEK +{k}")
        c.font = Font(bold=True)
        c.border = border
        c.alignment = Alignment(horizontal="center")
    widths = {"A": 5, "B": 18, "C": 16, "D": 13, "E": 9, "F": 8, "G": 22, "H": 40, "I": 45, "J": 70, "K": 40}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for k in range(12, 20):
        ws.column_dimensions[ws.cell(row=3, column=k).column_letter].width = 9
    # an empty pre-formatted data row (borders) like real templates
    for c in range(1, 20):
        ws.cell(row=4, column=c).border = border
        ws.cell(row=4, column=c).alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A4"

    m = wb.create_sheet("Phân loại")
    m.append(["Model", "Item", "Từ khóa"])
    m.append(["A185", "Rear", "rear; nắp lưng; back cover"])
    m.append(["A175", "Main", "main; chính"])
    m.append(["A165", "Sub", "sub; phụ"])
    m.append(["A166", "PBA", "pba; bo mạch"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def make_sample_tree(root: Path) -> dict:
    reports = root / "Reports"
    files = [
        make_report(reports / "September" / "A185" / "(CTMS)_20506_260918080-VOC_A185_Rear.pptx"),
        make_report(reports / "September" / "A185" / "260918081-VOC_report2.pptx", model="SM-A185", item="Rear", mgmt="260918081-VOC"),
        make_report(reports / "September" / "A175" / "260918082-VOC_report3.pptx", model="SM-A175", item="Main", mgmt="260918082-VOC"),
        make_report(reports / "October" / "(CTMS)_20509_260918083-VOC_ Đối sách A166 PBA 21.09.2026.pptx", model="SM-A166", item="PBA", mgmt="260918083-VOC", with_qpn=False),
    ]
    template = make_template(root / "Templates" / "Verification.xlsx")
    return {"reports": reports, "files": files, "template": template}


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("sample_data")
    info = make_sample_tree(target)
    print("Created:")
    for f in info["files"]:
        print("  ", f)
    print("  ", info["template"])
