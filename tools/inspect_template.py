"""Dump the structure of a real Excel template so column mapping can be verified.

    python tools/inspect_template.py "D:\\Templates\\Kiem_chung.xlsx"
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from app.excel_writer import ExcelWriter, validate_template  # noqa: E402
from app.pptx_parser import norm_key  # noqa: E402


def inspect_template(path: Path, max_rows: int = 12, max_cols: int = 40) -> str:
    out = [f"TEMPLATE: {path}"]
    wb = load_workbook(path)
    out.append(f"Sheets: {wb.sheetnames}")
    for ws in wb.worksheets:
        out.append(f"\n--- Sheet '{ws.title}'  ({ws.max_row} rows x {ws.max_column} cols)  key={norm_key(ws.title)!r}")
        merged = [str(r) for r in ws.merged_cells.ranges]
        out.append(f"Merged ranges ({len(merged)}): {merged[:40]}{' …' if len(merged) > 40 else ''}")
        widths = {k: v.width for k, v in ws.column_dimensions.items() if v.width}
        out.append(f"Column widths: {widths}")
        def merged_of(r, c):
            for rng in ws.merged_cells.ranges:
                if rng.min_row <= r <= rng.max_row and rng.min_col <= c <= rng.max_col:
                    return str(rng)
            return ""
        from app.excel_writer import FIELD_ALIASES, _alias_score
        for r in range(1, min(ws.max_row, max_rows) + 1):
            cells = []
            scores = {}
            for c in range(1, min(ws.max_column, max_cols) + 1):
                v = ws.cell(row=r, column=c).value
                if v in (None, ""):
                    continue
                coord = ws.cell(row=r, column=c).coordinate
                mg = merged_of(r, c)
                cells.append(f"{coord}{'[' + mg + ']' if mg else ''}={str(v)[:50]!r} key={norm_key(str(v))!r}")
                if isinstance(v, str):
                    for f in FIELD_ALIASES:
                        sc = _alias_score(norm_key(v), f)
                        if sc:
                            scores[f"{f}@{coord}"] = sc
            out.append(f"  R{r} (h={ws.row_dimensions[r].height}): " + ("  ".join(cells) if cells else "(empty)"))
            if scores:
                out.append(f"      alias hits: {scores}")
    ok, msg = validate_template(path)
    out.append(f"\nvalidate_template: {'OK' if ok else 'FAIL'} – {msg}")
    with tempfile.TemporaryDirectory() as tmp:
        try:
            w = ExcelWriter(path, Path(tmp) / "probe.xlsx")
            out.append(f"Target sheet     : {w.ws.title}")
            out.append(f"Header row       : {w.header_row}   data starts at row {w.data_start}   next free row {w.next_row()}")
            out.append("Detected columns :")
            from openpyxl.utils import get_column_letter
            for f, c in sorted(w.columns.items(), key=lambda kv: kv[1]):
                hdr = next((w.ws.cell(row=rr, column=c).value for rr in range(w.header_row, w.data_start)
                            if w.ws.cell(row=rr, column=c).value not in (None, "")), None)
                out.append(f"   {f:20s} -> col {c:3d} ({get_column_letter(c)})  header={hdr!r}")
            expected = ["management_number", "vendor", "occurrence_date", "model", "item", "defect_content", "qpn",
                        "root_cause", "improvement", "temporary", "improvement_image"] + [f"week_{i}" for i in range(1, 9)]
            missing = [f for f in expected if f not in w.columns]
            out.append(f"Missing (optional ok: temporary/status/note): {missing or 'none'}")
            out.append(f"Item mapping ({len(w.item_mapping)}): {dict(list(w.item_mapping.items())[:30])}")
            out.append(f"Known models ({len(w.known_models)}): {w.known_models[:30]}")
            w.close()
        except Exception as e:  # noqa: BLE001
            out.append(f"ExcelWriter FAILED: {type(e).__name__}: {e}")
    return "\n".join(out)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    print(inspect_template(Path(sys.argv[1])))
