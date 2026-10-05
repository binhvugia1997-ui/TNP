"""Write extracted records into the user's Excel verification template.

* Never overwrites the template: output workbook = copy of template (or the
  existing output workbook is re-opened so batches can be resumed / appended).
* The destination sheet is located by its STRUCTURE (header aliases of the
  verification table), never by its name; header cells are matched by field
  aliases so the exact column layout of the template does not matter.  A sheet
  named "Data" (source data) and the "Phân loại" mapping sheet are never chosen.
* Formatting of the previous data row (borders, fonts, alignment, single-row
  merges) is copied to every new row.
* Images are embedded (not linked) and scaled to fit the target column.
"""
from __future__ import annotations

import logging
import math
import os
import re
import shutil
from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

from .extractor import canonicalize_vendor_value

from .pptx_parser import norm_key

LOG = logging.getLogger("report_extractor.excel")

MAPPING_SHEET_KEYS = ("phan loai", "classification", "mapping", "danh muc")
# a sheet qualifies as the destination table only when ALL of these columns are present
REQUIRED_FIELDS = ("management_number", "qpn", "root_cause", "improvement")
# Fields the extractor owns in an existing (matched) row.  Everything else – Management Number, WEEK +1..+8,
# manual columns – is never touched.  Text fields are "complete" when non-blank, image fields when a picture is
# anchored in the cell.
MANAGED_TEXT_FIELDS = ("vendor", "occurrence_date", "model", "item", "defect_content", "root_cause", "improvement")
MANAGED_IMAGE_FIELDS = ("qpn", "improvement_image")
MANAGED_FIELDS = MANAGED_TEXT_FIELDS + MANAGED_IMAGE_FIELDS
# the full expected set – used for the diagnostic message and for ranking
EXPECTED_FIELDS = ("management_number", "vendor", "occurrence_date", "model", "item", "defect_content",
                   "qpn", "root_cause", "improvement", "improvement_image")
FIELD_LABELS = {"management_number": "Management number", "vendor": "Tên vendor", "occurrence_date": "Ngày phát sinh",
                "model": "Model", "item": "Item", "defect_content": "Nội dung lỗi", "qpn": "QPN",
                "root_cause": "Nguyên nhân", "improvement": "Nội dung đối sách cải tiến",
                "improvement_image": "Hình ảnh cải tiến"}
# sheets that are never a destination: source data / mapping
EXCLUDED_SHEET_KEYS = ("data",)

# field -> list of accent-insensitive aliases (norm_key form)
FIELD_ALIASES: Dict[str, List[str]] = {
    "stt": ["stt", "no", "so tt", "tt"],
    "management_number": ["management number", "management no", "mgmt no", "ma quan ly", "so quan ly",
                          "management"],
    "vendor": ["ten vendor", "vendor", "nha cung cap", "ncc"],
    "occurrence_date": ["ngay phat sinh", "occurrence date", "ngay xay ra", "ngay"],
    "model": ["model", "model name"],
    "item": ["item", "hang muc", "linh kien", "part"],
    "defect_content": ["noi dung loi", "defect content", "noi dung khong phu hop", "hien tuong", "defect"],
    "qpn": ["qpn", "quality problem notice", "hinh anh qpn"],
    "root_cause": ["nguyen nhan", "root cause", "cause"],
    "temporary": ["xu ly tam thoi", "phuong phap xu ly tam thoi", "doi sach tam thoi", "bien phap tam thoi",
                  "temporary action", "temporary handling", "containment"],
    "improvement": ["noi dung doi sach cai tien", "noi dung doi sach", "doi sach cai tien", "doi sach",
                    "countermeasure", "corrective action", "cai tien"],
    "improvement_image": ["hinh anh cai tien", "hinh anh doi sach", "hinh anh", "image", "improvement image"],
    "status": ["trang thai", "status", "ket qua"],
    "note": ["ghi chu", "note", "remark", "remarks"],
}
for _i in range(1, 9):
    FIELD_ALIASES[f"week_{_i}"] = [f"week +{_i}", f"week+{_i}", f"week {_i}", f"w +{_i}", f"w+{_i}",
                                   f"tuan +{_i}", f"tuan {_i}", f"w{_i}"]

# "WEEK +1", "WEEK + 1", "Week+1", "W1", "+1", "Tuần 1", "T+1", "1W" ...
WEEK_RE = re.compile(r"^(?:week|wk|w|tuan|t)?\s*\+?\s*([1-8])\s*(?:w|wk|week|tuan)?$")
# merged parent headers that group the WEEK sub-columns
WEEK_PARENT_KEYS = ("kiem chung", "theo doi", "week", "tuan", "verification", "xac nhan", "follow up",
                    "followup", "hieu qua", "monitoring", "giam sat")

MAX_ROW_HEIGHT_PT = 409.0          # hard Excel limit for a row
EXCEL_DEFAULT_COL_WIDTH = 9.140625  # stored (XML) width of Excel's default column = 64 px with Calibri 11
EXCEL_DEFAULT_ROW_HEIGHT_PT = 15.0
MAX_DIGIT_WIDTH_PX = 7             # Calibri 11 @ 96 dpi (Excel "MDW")
CONTAINMENT_TOL_PX = 1             # only EMU/pixel rounding may differ by this much
# ---- image layout (all derived from the real column widths of the workbook, never from a screenshot) ----
IMAGE_MARGIN_PX = 4                # left/right/top/bottom margin inside the destination area
IMAGE_GAP_PX = 6                   # gap between two independent images
MAX_IMAGE_HEIGHT_PT = 300.0        # one width-fitted picture is never taller than this (readable, < Excel row cap)
MAX_IMAGE_COLUMNS = 3              # grid fallback when stacking would exceed the Excel row limit
PX_PER_PT = 4.0 / 3.0
EMU_PER_PX = 9525


def col_width_to_px(width_chars: float, mdw: int = MAX_DIGIT_WIDTH_PX) -> int:
    """Stored Excel column width -> pixels, ECMA-376 §18.3.1.13.

    openpyxl exposes the width *as stored in the XML*, which ALREADY contains the 5 px cell padding
    (default column: stored 9.140625 -> 64 px).  ``trunc(((256*W + trunc(128/MDW)) / 256) * MDW)``.
    Adding another "+5 px" (the old formula) over-estimated every column by 5 px and let images spill
    into the next column.
    """
    if width_chars is None:
        width_chars = EXCEL_DEFAULT_COL_WIDTH
    w = float(width_chars)
    if w <= 0:                      # hidden column
        return 0
    return int(((256 * w + int(128 / mdw)) / 256) * mdw)


def pt_to_px(pt: float) -> int:
    """Points -> whole pixels at 96 dpi (Excel row heights)."""
    return int(math.floor(float(pt) * PX_PER_PT + 1e-6))


def px_to_pt(px: float) -> float:
    return float(px) / PX_PER_PT


def px_to_emu(px: float) -> int:
    return int(round(float(px) * EMU_PER_PX))


def emu_to_px(emu: float) -> int:
    return int(round(float(emu) / EMU_PER_PX))


def assert_image_inside_area(items: List[Tuple[int, int, int, int]], area_w_px: int, area_h_px: int,
                             margin_px: int = IMAGE_MARGIN_PX, tol_px: int = CONTAINMENT_TOL_PX) -> None:
    """Hard containment rule (#32/#41): every (x, y, w, h) stays inside the margins of the destination area
    and no two pictures overlap.  Raises ``ValueError`` describing the first violation."""
    for i, (x, y, w, h) in enumerate(items):
        if w <= 0 or h <= 0:
            raise ValueError(f"image {i}: empty size {w}x{h}")
        if x + tol_px < margin_px or y + tol_px < margin_px:
            raise ValueError(f"image {i}: top-left ({x},{y}) inside the {margin_px}px margin")
        if x + w - tol_px > area_w_px - margin_px:
            raise ValueError(f"image {i}: right edge {x + w} > {area_w_px - margin_px} (area {area_w_px}px)")
        if y + h - tol_px > area_h_px - margin_px:
            raise ValueError(f"image {i}: bottom edge {y + h} > {area_h_px - margin_px} (area {area_h_px}px)")
        for j, (x2, y2, w2, h2) in enumerate(items[:i]):
            if x < x2 + w2 - tol_px and x2 < x + w - tol_px and y < y2 + h2 - tol_px and y2 < y + h - tol_px:
                raise ValueError(f"image {i} overlaps image {j}")


def images_inside_area(items, area_w_px: int, area_h_px: int, margin_px: int = IMAGE_MARGIN_PX) -> bool:
    try:
        assert_image_inside_area(items, area_w_px, area_h_px, margin_px)
        return True
    except ValueError:
        return False


def plan_image_grid(sizes: List[Tuple[int, int]], area_px: int, max_row_pt: float = MAX_ROW_HEIGHT_PT,
                    max_image_pt: float = MAX_IMAGE_HEIGHT_PT, margin_px: int = IMAGE_MARGIN_PX,
                    gap_px: int = IMAGE_GAP_PX, max_columns: int = MAX_IMAGE_COLUMNS,
                    area_h_px: Optional[int] = None) -> Dict[str, Any]:
    """Deterministic layout of independent pictures inside one destination area.

    Every picture is scaled to the FULL usable slot width (aspect ratio preserved, height follows).
    Preferred = one column (pictures stacked, source order).  The vertical budget is the smaller of the
    Excel row cap (``max_row_pt`` = 409 pt) and, when given, the REAL area height ``area_h_px`` (final row
    height); a layout that would not fit is scaled down uniformly (never cropped / distorted).  Among
    1..max_columns columns the plan giving the LARGEST displayed picture width wins (ties -> fewer columns),
    so five pictures are never squeezed into thumbnails when a 2-column grid shows them bigger.
    Sizes are floored and positions clamped so ``assert_image_inside_area`` holds for every returned plan.
    Returns {"columns", "scale", "slot_px", "items": [(x_px, y_px, w_px, h_px)], "height_pt"}.
    """
    n = len(sizes)
    if n == 0 or area_px <= 0:
        return {"columns": 1, "scale": 1.0, "slot_px": 0, "items": [], "height_pt": 0.0}
    usable = max(20, area_px - 2 * margin_px)
    budget_px = (max_row_pt * PX_PER_PT) - 2 * margin_px
    if area_h_px is not None:
        budget_px = min(budget_px, area_h_px - 2 * margin_px)
    budget_px = max(1.0, budget_px)
    best = None
    for cols in range(1, min(max_columns, n) + 1):
        slot = (usable - gap_px * (cols - 1)) / cols
        # per picture: width = slot, height from aspect ratio, capped at MAX_IMAGE_HEIGHT_PT (aspect kept)
        dims = []
        for (w, h) in sizes:
            ww = slot
            hh = slot * h / max(1, w)
            cap = max_image_pt * PX_PER_PT
            if hh > cap:
                ww, hh = cap * w / max(1, h), cap
            dims.append((ww, hh))
        rows = [dims[i:i + cols] for i in range(0, n, cols)]
        row_h = [max(d[1] for d in r) for r in rows]
        # gaps are NOT scaled when placing, so only the picture heights share the remaining budget
        pic_budget = budget_px - gap_px * (len(rows) - 1)
        total = sum(row_h)
        scale = min(1.0, max(0.01, pic_budget) / total) if total > 0 else 1.0
        eff_w = min(d[0] for d in dims) * scale
        key = (round(eff_w, 1), -cols)
        if best is None or key > best[0]:
            best = (key, cols, scale, slot, dims, rows, row_h)
    _, cols, scale, slot, dims, rows, row_h = best
    items: List[Tuple[int, int, int, int]] = []
    y = float(margin_px)
    for r_idx, r in enumerate(rows):
        rh = row_h[r_idx] * scale
        for c_idx, (ww, hh) in enumerate(r):
            w_px, h_px = max(1, int(math.floor(ww * scale))), max(1, int(math.floor(hh * scale)))
            slot_w = slot * scale if scale < 1 else slot
            slot_x = margin_px + c_idx * slot_w + c_idx * gap_px
            x_px = int(round(slot_x + (slot_w - w_px) / 2))          # centre inside the slot
            x_px = max(margin_px, min(x_px, area_px - margin_px - w_px))
            y_px = int(math.floor(y))
            items.append((x_px, y_px, w_px, h_px))
        y += rh + gap_px
    bottom = max((it[1] + it[3] for it in items), default=margin_px)
    height_px = bottom + margin_px
    return {"columns": cols, "scale": scale, "slot_px": int(slot * scale), "items": items,
            "height_pt": min(max_row_pt, px_to_pt(height_px))}


def _as_paths(value) -> List[Path]:
    """One path or a list of paths -> list (independent pictures stay independent)."""
    if not value:
        return []
    if isinstance(value, (str, Path)):
        return [Path(value)]
    return [Path(v) for v in value if v]


def week_index(cell_key: str) -> Optional[int]:
    m = WEEK_RE.match(cell_key or "")
    return int(m.group(1)) if m else None


def _alias_score(cell_key: str, field: str) -> int:
    if not cell_key:
        return 0
    if field.startswith("week_"):
        idx = week_index(cell_key)
        if idx is None:
            return 0
        # bare numbers ("1") are only accepted through the merged-parent rule
        return 3 if (idx == int(field[5:]) and re.search(r"[a-z+]", cell_key)) else 0
    best = 0
    for alias in FIELD_ALIASES[field]:
        if cell_key == alias:
            best = max(best, 3)
        elif field.startswith("week_"):
            # week aliases must match exactly-ish ("week +1" != "week +10")
            if re.fullmatch(re.escape(alias) + r"(?:\D.*)?", cell_key):
                best = max(best, 2)
        elif len(alias) >= 4 and re.search(r"(?<![a-z0-9])" + re.escape(alias) + r"(?![a-z0-9])", cell_key):
            best = max(best, 2 if len(alias) > 5 else 1)
    return best


class TemplateError(RuntimeError):
    pass


class ExcelLockedError(TemplateError):
    """The destination workbook could not be replaced because another process (typically Microsoft Excel) holds
    it open.  ``path`` is the locked file; ``original`` the technical exception (logged, never shown as the primary
    GUI message)."""

    def __init__(self, path: Path, original: BaseException):
        self.path = Path(path)
        self.original = original
        super().__init__(locked_file_message(self.path))


def locked_file_message(path: Path) -> str:
    return ("Không thể cập nhật file Excel vì file đang được sử dụng.\n\nFile:\n" + str(path)
            + "\n\nHãy đóng file Excel hoặc chương trình đang sử dụng file, sau đó thử lại.")


_LOCK_WINERRORS = (32, 33)        # ERROR_SHARING_VIOLATION, ERROR_LOCK_VIOLATION
_LOCK_TEXT_RE = re.compile(r"WinError (32|33)\b|being used by another process|sharing violation|lock violation",
                           re.IGNORECASE)


def is_sharing_violation(err: BaseException) -> bool:
    """True only for Windows sharing/lock violations – a plain ``PermissionError`` (read-only folder, ACL) is NOT
    a lock and must not get the 'file is in use' message."""
    if not isinstance(err, OSError):
        return False
    if getattr(err, "winerror", None) in _LOCK_WINERRORS:
        return True
    return bool(_LOCK_TEXT_RE.search(str(err)))


def probe_writable(path: Path) -> Optional[BaseException]:
    """Best-effort preflight: can ``path`` be opened for writing right now?  Returns the exception when it cannot
    (never raises).  A file may still become locked afterwards – the commit boundary stays the real guard."""
    try:
        if Path(path).exists():
            with open(path, "r+b"):
                pass
        return None
    except OSError as e:
        return e


BACKUP_DIR_NAME = "backup"


class ExcelWriter:
    def __init__(self, template: Path, output: Path, probe: bool = False):
        """``probe=True`` opens the current master state READ-ONLY for the pre-scan (no copy of the template,
        nothing is ever written); otherwise the output workbook is created from the template when missing."""
        self.template = Path(template)
        self.output = Path(output)
        self.probe = probe
        if not self.template.exists():
            raise TemplateError(f"Không tìm thấy form Excel: {self.template}")
        if self.template.resolve() == self.output.resolve():
            raise TemplateError("File kết quả không được trùng với form Excel gốc")
        if probe:
            self.wb = load_workbook(self.output if self.output.exists() else self.template)
        else:
            self.output.parent.mkdir(parents=True, exist_ok=True)
            if not self.output.exists():
                shutil.copyfile(self.template, self.output)
            self.wb = load_workbook(self.output)
        self.columns: Dict[str, int] = {}
        self.header_row = 0
        self.data_start = 0
        self.ws = self._find_target_sheet()          # also sets columns / header_row / data_start
        self.item_mapping, self.known_models = self._read_mapping_sheet()
        self._dirty = False
        self._placed: List[Tuple[Any, int, str]] = []   # (image, row, field) written by THIS writer
        # PROMPT-004 §17: one safety copy of the master is taken right before the FIRST real modification
        self.backup_dir = self.output.parent / BACKUP_DIR_NAME
        self.backup_path: Optional[Path] = None
        self.backup_enabled = True

    # ------------------------------------------------------------------
    # Safety backup (before the first modification only)
    # ------------------------------------------------------------------
    def _ensure_backup(self) -> None:
        """Called by every mutating primitive.  First call of the batch: copy ``output`` to
        ``backup/<name>_backup_<timestamp>.xlsx``; failure raises *before* anything is modified.  A batch that only
        skips complete rows never calls this, so no unnecessary backup is created."""
        if self.probe:
            raise TemplateError("Workbook mở ở chế độ quét (probe) – không ghi")
        self._dirty = True
        if self.backup_path is not None or not self.backup_enabled:
            return
        try:
            self.backup_dir.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            target = self.backup_dir / f"{self.output.stem}_backup_{stamp}{self.output.suffix}"
            n = 1
            while target.exists():
                target = self.backup_dir / f"{self.output.stem}_backup_{stamp}_{n}{self.output.suffix}"
                n += 1
            shutil.copyfile(self.output, target)
        except OSError as e:
            self._dirty = False
            raise TemplateError(f"Không tạo được bản sao lưu Excel trước khi ghi ({e}) – không sửa file gốc") from e
        self.backup_path = target
        LOG.info("MASTER_BACKUP file=%s", target)

    # ------------------------------------------------------------------
    # Template analysis
    # ------------------------------------------------------------------
    def _find_target_sheet(self):
        """Pick the destination sheet by table structure, regardless of its name.

        * exactly one sheet with all REQUIRED_FIELDS -> use it;
        * several -> TemplateError('Cần kiểm tra: ...' + candidate names), no guessing;
        * none -> TemplateError listing, per sheet, which required headers are missing.
        'Data' (source data) and 'Phân loại' (mapping) sheets are never candidates.
        """
        candidates, report = [], []
        for ws in self.wb.worksheets:
            key = norm_key(ws.title)
            if key in EXCLUDED_SHEET_KEYS or any(k in key for k in MAPPING_SHEET_KEYS):
                report.append(f"'{ws.title}': bỏ qua (sheet dữ liệu nguồn / phân loại)")
                continue
            try:
                header, cols, data_start = self._analyse_sheet(ws)
            except TemplateError:
                report.append(f"'{ws.title}': không có dòng tiêu đề")
                continue
            missing = [f for f in REQUIRED_FIELDS if f not in cols]
            if missing:
                report.append(f"'{ws.title}': thiếu cột " + ", ".join(FIELD_LABELS[f] for f in missing))
                continue
            candidates.append((ws, header, cols, data_start))
        if len(candidates) == 1:
            ws, self.header_row, self.columns, self.data_start = candidates[0]
            LOG.info("Destination sheet: '%s' (header row %s, columns %s)", ws.title, self.header_row, sorted(self.columns))
            return ws
        if len(candidates) > 1:
            names = ", ".join(f"'{c[0].title}'" for c in candidates)
            raise TemplateError("Cần kiểm tra: nhiều sheet có cấu trúc bảng kiểm chứng, không thể tự chọn: " + names)
        raise TemplateError("Không tìm thấy sheet chứa bảng kiểm chứng (cần các cột: "
                            + ", ".join(FIELD_LABELS[f] for f in REQUIRED_FIELDS) + "). "
                            + "; ".join(report))

    def _detect_headers(self) -> None:
        self.header_row, self.columns, self.data_start = self._analyse_sheet(self.ws)

    def _analyse_sheet(self, ws) -> Tuple[int, Dict[str, int], int]:
        """Locate the header band of ``ws``; returns (header_row, field->column, first data row)."""
        max_scan = min(ws.max_row, 30)
        row_matches: Dict[int, Dict[str, Tuple[int, int]]] = {}
        for r in range(1, max_scan + 1):
            matches: Dict[str, Tuple[int, int]] = {}
            for c in range(1, min(ws.max_column, 80) + 1):
                v = ws.cell(row=r, column=c).value
                if not isinstance(v, str) or not v.strip():
                    continue
                if len(v.strip()) > 60 or v.count("\n") > 3:  # long data text, not a header label
                    continue
                key = norm_key(v)                            # collapses line breaks / spaces / accents
                # merged parents of the WEEK group ("Theo dõi cải tiến", "Kiểm chứng") are not data columns
                if any(k in key for k in ("theo doi", "kiem chung")) and week_index(key) is None:
                    continue
                # "Nội dung đối sách tạm thời" must never be taken for the improvement column
                if "tam thoi" in key or "temporary" in key or "containment" in key:
                    s = _alias_score(key, "temporary") or 1
                    if "temporary" not in matches or s > matches["temporary"][1]:
                        matches["temporary"] = (c, s)
                    continue
                for field in FIELD_ALIASES:
                    s = _alias_score(key, field)
                    if s and (field not in matches or s > matches[field][1]):
                        matches[field] = (c, s)
            if matches:
                row_matches[r] = matches
        if not row_matches:
            raise TemplateError(f"Không nhận diện được dòng tiêu đề trong sheet '{ws.title}'")
        # header row = row with most strong matches
        def strength(r: int) -> Tuple[int, int]:
            m = row_matches[r]
            core = sum(1 for f in ("improvement", "root_cause", "qpn", "model", "item", "defect_content") if f in m)
            return (core, len(m))
        header = max(row_matches, key=strength)
        self.header_row = header
        cols: Dict[str, int] = {}
        # header band: the header row, up to 2 rows above (merged parent titles) and 3 rows below
        band = [r for r in range(max(1, header - 2), header + 4) if r <= max_scan]
        for r in [header] + [r for r in band if r != header]:
            if r in row_matches:
                for f, (c, s) in row_matches[r].items():
                    if f not in cols:
                        cols[f] = c
        # WEEK columns inherited from a merged parent header ("Kiểm chứng" / "WEEK" / "Tuần")
        # whose sub-cells are only "+1", "1", "W1", "1W", "T1" ...
        for r in band:
            for c in range(1, min(ws.max_column, 80) + 1):
                v = ws.cell(row=r, column=c).value
                if v in (None, ""):
                    continue
                k = norm_key(str(v))
                idx = week_index(k)
                if idx is None:
                    continue
                field = f"week_{idx}"
                if field in cols:
                    continue
                parent = self._parent_header_key(ws, r, c, header)
                if parent and any(x in parent for x in WEEK_PARENT_KEYS):
                    cols[field] = c
        last_header = header
        for r in range(header + 1, header + 4):
            if r in row_matches and any(f.startswith("week_") or f in ("model", "item") for f in row_matches[r]):
                last_header = r
            elif any(cols.get(f"week_{i}") and self._cell_row_of_week(ws, r, cols[f"week_{i}"]) for i in range(1, 9)):
                last_header = r
        return header, cols, last_header + 1

    @staticmethod
    def _cell_row_of_week(ws, row: int, col: int) -> bool:
        v = ws.cell(row=row, column=col).value
        return v not in (None, "") and week_index(norm_key(str(v))) is not None

    def _parent_header_key(self, ws, row: int, col: int, header_row: int) -> str:
        """Text of the merged/plain header cell(s) directly above (row-1, row-2) covering ``col``."""
        texts = []
        for r in (row - 1, row - 2, header_row):
            if r < 1 or r == row:
                continue
            cell, rng = self._anchor_in(ws, r, col)
            if cell.value not in (None, ""):
                texts.append(norm_key(str(cell.value)))
        return " ".join(texts)

    def _read_mapping_sheet(self) -> Tuple[Dict[str, str], List[str]]:
        mapping: Dict[str, str] = {}
        models: List[str] = []
        for ws in self.wb.worksheets:
            if not any(k in norm_key(ws.title) for k in MAPPING_SHEET_KEYS):
                continue
            rows = [[c for c in row] for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 2000), values_only=True)]
            # detect header
            hdr_idx, item_col, key_col, model_col = None, None, None, None
            for i, row in enumerate(rows[:10]):
                keys = [norm_key(str(v)) if v is not None else "" for v in row]
                for ci, k in enumerate(keys):
                    if k in ("item", "hang muc", "linh kien", "ten item", "phan loai item") and item_col is None:
                        item_col = ci
                    elif k in ("model", "ten model", "model name") and model_col is None:
                        model_col = ci
                    elif k in ("tu khoa", "keyword", "keywords", "ten", "ten linh kien", "mo ta", "noi dung",
                               "ten goi", "alias") and key_col is None:
                        key_col = ci
                if item_col is not None or model_col is not None:
                    hdr_idx = i
                    break
            body = rows[hdr_idx + 1:] if hdr_idx is not None else rows
            for row in body:
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if hdr_idx is not None:
                    if item_col is not None and item_col < len(row) and row[item_col]:
                        item = str(row[item_col]).strip()
                        mapping.setdefault(item, item)
                        if key_col is not None and key_col < len(row) and row[key_col]:
                            for kw in re.split(r"[;,/\n]+", str(row[key_col])):
                                kw = kw.strip()
                                if kw:
                                    mapping.setdefault(kw, item)
                    if model_col is not None and model_col < len(row) and row[model_col]:
                        for mv in re.split(r"[;,/\n]+", str(row[model_col])):
                            mv = mv.strip()
                            if mv and mv not in models:
                                models.append(mv)
                else:
                    if len(vals) >= 2:
                        mapping.setdefault(vals[0], vals[1])
                        mapping.setdefault(vals[1], vals[1])
                    elif len(vals) == 1:
                        mapping.setdefault(vals[0], vals[0])
            break
        if mapping:
            LOG.info("Item mapping loaded: %d keywords, %d models", len(mapping), len(models))
        return mapping, models

    # ------------------------------------------------------------------
    # Row helpers
    # ------------------------------------------------------------------
    def _anchor(self, row: int, col: int):
        """Top-left cell of the merged range containing (row, col), or the cell itself."""
        return self._anchor_in(self.ws, row, col)

    @staticmethod
    def _anchor_in(ws, row: int, col: int):
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
                return ws.cell(row=rng.min_row, column=rng.min_col), rng
        return ws.cell(row=row, column=col), None

    def _row_used(self, row: int) -> bool:
        check = [f for f in self.columns if f not in ("stt", "vendor", "occurrence_date", "status", "note", "temporary")
                 and not f.startswith("week_")]
        for f in check:
            cell, _ = self._anchor(row, self.columns[f])
            if cell.value not in (None, ""):
                return True
        # rows holding only an image count as used too
        for img in getattr(self.ws, "_images", []):
            try:
                if img.anchor._from.row + 1 == row:
                    return True
            except Exception:
                pass
        return False

    def next_row(self) -> int:
        r = self.data_start
        limit = max(self.ws.max_row, self.data_start) + 1
        while r <= limit and self._row_used(r):
            r += 1
        return r

    def used_count(self) -> int:
        return sum(1 for r in range(self.data_start, self.next_row()) if self._row_used(r))

    def _copy_row_style(self, src_row: int, dst_row: int) -> None:
        ws = self.ws
        if src_row < self.data_start or src_row == dst_row:
            return
        self._ensure_backup()
        max_col = max(list(self.columns.values()) + [ws.max_column])
        for c in range(1, max_col + 1):
            s = ws.cell(row=src_row, column=c)
            d = ws.cell(row=dst_row, column=c)
            if isinstance(d, MergedCell):
                continue
            if s.has_style:
                d._style = copy(s._style)
        # replicate single-row merges of the source row
        for rng in list(ws.merged_cells.ranges):
            if rng.min_row == src_row and rng.max_row == src_row:
                try:
                    ws.merge_cells(start_row=dst_row, start_column=rng.min_col,
                                   end_row=dst_row, end_column=rng.max_col)
                except Exception:
                    pass
        if ws.row_dimensions[src_row].height:
            ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height

    def _col_width_chars(self, col: int) -> float:
        """Stored width of ONE column, honouring <col min=.. max=..> ranges, hidden columns and the sheet default
        (``column_dimensions[letter]`` alone silently returns a *default* dimension for columns inside a range)."""
        for dim in list(self.ws.column_dimensions.values()):
            lo, hi = dim.min or 0, dim.max or 0
            if lo and hi and lo <= col <= hi:
                if dim.hidden:
                    return 0.0
                if dim.width:
                    return float(dim.width)
        dim = self.ws.column_dimensions.get(get_column_letter(col))
        if dim is not None:
            if dim.hidden:
                return 0.0
            if dim.width:
                return float(dim.width)
        fmt = self.ws.sheet_format
        if fmt is not None and fmt.defaultColWidth:
            return float(fmt.defaultColWidth)
        if fmt is not None and fmt.baseColWidth:
            return float(fmt.baseColWidth) + 5.0 / MAX_DIGIT_WIDTH_PX
        return EXCEL_DEFAULT_COL_WIDTH

    def _col_width_px(self, col: int, rng=None) -> int:
        cols = range(rng.min_col, rng.max_col + 1) if rng is not None else [col]
        return int(sum(col_width_to_px(self._col_width_chars(c)) for c in cols))

    def _row_height_pt(self, row: int) -> float:
        h = self.ws.row_dimensions[row].height
        if h:
            return float(h)
        fmt = self.ws.sheet_format
        if fmt is not None and fmt.defaultRowHeight:
            return float(fmt.defaultRowHeight)
        return EXCEL_DEFAULT_ROW_HEIGHT_PT

    def _area_height_px(self, row: int, rng=None) -> int:
        rows = range(rng.min_row, rng.max_row + 1) if rng is not None else [row]
        return int(sum(pt_to_px(self._row_height_pt(r)) for r in rows))

    def image_area_height_px(self, row: int, field: str) -> int:
        """Pixel height of the destination image area as it is NOW (merged rows included)."""
        if field not in self.columns:
            return 0
        col = self.columns[field]
        _, rng = self._anchor(row, col)
        return self._area_height_px(row, rng)

    def _set_cell(self, row: int, field: str, value: Any, wrap: bool = True) -> Optional[Any]:
        if field not in self.columns:
            return None
        cell, _ = self._anchor(row, self.columns[field])
        if cell.value == (value if value not in ("",) else None) and not (wrap and isinstance(value, str) and value):
            return cell                              # nothing changes -> no backup, no dirty flag
        self._ensure_backup()
        cell.value = value if value not in ("",) else None
        if wrap and isinstance(value, str):
            al = copy(cell.alignment) if cell.alignment else Alignment()
            cell.alignment = Alignment(horizontal=al.horizontal or "left", vertical="top", wrap_text=True,
                                       indent=al.indent, text_rotation=al.text_rotation)
        return cell

    def image_area_px(self, row: int, field: str) -> int:
        """Pixel width of the destination image area (merged span included)."""
        if field not in self.columns:
            return 0
        col = self.columns[field]
        _, rng = self._anchor(row, col)
        return self._col_width_px(col, rng)

    def plan_field_images(self, row: int, field: str, paths: List[Path]) -> Optional[Dict[str, Any]]:
        """Step 1 (#35): read the real destination width and compute the candidate layout + required height.
        Returns {"paths", "sizes", "plan", "height_pt"} or None when there is nothing to place."""
        paths = [Path(p) for p in (paths or []) if p and Path(p).exists()]
        if field not in self.columns or not paths:
            return None
        sizes = []
        for p in paths:
            with PILImage.open(p) as im:
                sizes.append(im.size)
        plan = plan_image_grid(sizes, self.image_area_px(row, field))
        return {"paths": paths, "sizes": sizes, "plan": plan, "height_pt": float(plan["height_pt"]) + 2}

    def insert_planned_images(self, row: int, field: str, planned: Optional[Dict[str, Any]]) -> None:
        """Steps 2-4 (#35): with the FINAL row height applied, re-fit the pictures to the final rectangle,
        validate containment (#39) and only then create the anchors.  Never writes an overflowing geometry."""
        if not planned:
            return
        self._ensure_backup()
        col = self.columns[field]
        cell, rng = self._anchor(row, col)
        area_w = self._col_width_px(col, rng)
        area_h = self._area_height_px(row, rng)
        plan = plan_image_grid(planned["sizes"], area_w, area_h_px=area_h)
        items = plan["items"]
        shrink = 1.0
        while not images_inside_area(items, area_w, area_h) and shrink > 0.05:
            shrink -= 0.02                                    # defensive: uniform proportional shrink only
            items = [(x, y, max(1, int(w * shrink)), max(1, int(h * shrink))) for (x, y, w, h) in plan["items"]]
        assert_image_inside_area(items, area_w, area_h)      # raises on a genuine layout bug instead of writing it
        for p, (x, y, w, h) in zip(planned["paths"], items):
            img = XLImage(str(p))
            img.width, img.height = max(1, w), max(1, h)
            marker = AnchorMarker(col=cell.column - 1, colOff=px_to_emu(x), row=cell.row - 1, rowOff=px_to_emu(y))
            img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(cx=px_to_emu(img.width), cy=px_to_emu(img.height)))
            self.ws.add_image(img)
            self._placed.append((img, row, field))
        self._dirty = True

    def place_images(self, row: int, field: str, paths: List[Path]) -> float:
        """Convenience for a single field: plan -> raise the row height if needed -> fit -> insert.
        Returns the row height (pt) the pictures required."""
        planned = self.plan_field_images(row, field, paths)
        if not planned:
            return 0.0
        self._remove_images_in_cell(row, field)
        need = min(MAX_ROW_HEIGHT_PT, max(self._row_height_pt(row), planned["height_pt"]))
        self.ws.row_dimensions[row].height = need
        self.insert_planned_images(row, field, planned)
        return planned["height_pt"]

    def _write_images_with_row_height(self, row: int, heights: List[float], qpn_png, improvement_jpg,
                                      want_qpn: bool, want_imp: bool) -> None:
        """Shared #35 sequence: candidate plans -> final row height -> fit to final rectangle -> anchors."""
        plans = {}
        if (want_qpn and qpn_png) or (want_imp and improvement_jpg):
            self._ensure_backup()
        if want_qpn and qpn_png:
            self._remove_images_in_cell(row, "qpn")
            plans["qpn"] = self.plan_field_images(row, "qpn", [Path(qpn_png)])
        if want_imp and improvement_jpg:
            self._remove_images_in_cell(row, "improvement_image")
            plans["improvement_image"] = self.plan_field_images(row, "improvement_image", _as_paths(improvement_jpg))
        heights.extend(pl["height_pt"] for pl in plans.values() if pl)
        self.ws.row_dimensions[row].height = min(MAX_ROW_HEIGHT_PT, max(heights))
        for field, pl in plans.items():
            self.insert_planned_images(row, field, pl)
        self._dirty = True

    def replace_improvement_images(self, row: int, paths: List[Path]) -> int:
        """PROMPT-006 re-apply after user confirmation: swap ONLY the pictures of the ``improvement_image`` cell of
        ``row`` (text fields, QPN, WEEK, vendor/date, other rows and sheets are untouched).  Returns the number of
        pictures placed.  An empty list clears the cell (nothing confirmed as After)."""
        if "improvement_image" not in self.columns:
            raise TemplateError("Không tìm thấy cột Hình ảnh cải tiến trong form Excel")
        self._ensure_backup()
        self._remove_images_in_cell(row, "improvement_image")
        self._placed = [t for t in self._placed if not (t[1] == row and t[2] == "improvement_image")]
        if not paths:
            return 0
        self.place_images(row, "improvement_image", [Path(x) for x in paths])
        return len(paths)

    def replace_improvement_text(self, row: int, text: str) -> bool:
        """PROMPT-006B re-apply after content confirmation: rewrite ONLY the ``improvement`` text cell of ``row``
        (verbatim text supplied by the extractor).  Every other cell, picture, WEEK/manual column, formula and
        sheet is untouched.  Returns True when the cell changed (backup taken before the first modification)."""
        if "improvement" not in self.columns:
            raise TemplateError("Không tìm thấy cột Nội dung đối sách cải tiến trong form Excel")
        cell, _ = self._anchor(row, self.columns["improvement"])
        if (cell.value or "") == (text or ""):
            return False
        self._ensure_backup()
        self._set_cell(row, "improvement", text or "")
        need = self._text_height_pt(row, "improvement", text or "")
        cur = self.ws.row_dimensions[row].height or 15.0
        if need > cur:
            self.ws.row_dimensions[row].height = min(MAX_ROW_HEIGHT_PT, need)
        self._dirty = True
        return True

    def image_bounds_report(self, row: Optional[int] = None) -> List[str]:
        """#39/#42: compare every picture placed by this writer with its destination rectangle as the workbook
        stands now (row heights converted back to pixels).  Returns human-readable violations (empty = OK)."""
        problems: List[str] = []
        by_cell: Dict[Tuple[int, str], List[Tuple[int, int, int, int]]] = {}
        for img, r, field in self._placed:
            if row is not None and r != row:
                continue
            if img not in getattr(self.ws, "_images", []):
                continue
            a = img.anchor
            by_cell.setdefault((r, field), []).append((emu_to_px(a._from.colOff), emu_to_px(a._from.rowOff),
                                                       emu_to_px(a.ext.cx), emu_to_px(a.ext.cy)))
        for (r, field), items in by_cell.items():
            col = self.columns[field]
            _, rng = self._anchor(r, col)
            try:
                assert_image_inside_area(items, self._col_width_px(col, rng), self._area_height_px(r, rng))
            except ValueError as e:
                problems.append(f"row {r} {field}: {e}")
        return problems

    def _text_height_pt(self, row: int, field: str, text: str) -> float:
        if field not in self.columns or not text:
            return 0.0
        col = self.columns[field]
        _, rng = self._anchor(row, col)
        width_chars = max(8.0, self._col_width_px(col, rng) / 7.0)
        lines = 0
        for ln in text.split("\n"):
            lines += max(1, math.ceil(len(ln) / max(1.0, width_chars * 1.05)))
        return lines * 15.0 + 4

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------
    # ------------------------------------------------------------------
    # Management-Number keyed rows
    # ------------------------------------------------------------------
    @staticmethod
    def _norm_mgmt(v: Any) -> str:
        """Exact-key normalisation: trim + upper-case.  Bare numeric keys (260601017) may be
        stored by Excel as numbers -> compare their integer text; suffixes are never stripped."""
        if v in (None, ""):
            return ""
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v).strip().upper()

    def find_rows_by_management_number(self, mgmt: str) -> List[int]:
        """Exact (trim/case-insensitive) matches in the Management Number column, data rows only."""
        if "management_number" not in self.columns or not mgmt:
            return []
        key = self._norm_mgmt(mgmt)
        col = self.columns["management_number"]
        rows: List[int] = []
        for r in range(self.data_start, self.ws.max_row + 1):
            cell, rng = self._anchor(r, col)
            if rng is not None and cell.row != r:
                continue                         # covered by a vertical merge – counted once at its anchor
            if self._norm_mgmt(cell.value) == key:
                rows.append(r)
        return rows

    @staticmethod
    def _as_date(value: Any) -> Optional[date]:
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        txt = str(value).strip()
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d", "%d/%m/%y", "%d.%m.%y"):
            try:
                return datetime.strptime(txt, fmt).date()
            except ValueError:
                continue
        return None

    def _write_date(self, row: int, value: Optional[date]) -> None:
        if "occurrence_date" not in self.columns or value is None:
            return
        cell, _ = self._anchor(row, self.columns["occurrence_date"])
        self._ensure_backup()
        template_fmt = cell.number_format
        cell.value = value                      # openpyxl switches 'General' to 'yyyy-mm-dd' here
        if not template_fmt or template_fmt == "General":
            cell.number_format = "DD/MM/YYYY"   # template had no date format -> DD/MM/YYYY
        else:
            cell.number_format = template_fmt   # keep the template's own date format

    def _has_image_in_cell(self, row: int, field: str) -> bool:
        if field not in self.columns:
            return False
        cell, _ = self._anchor(row, self.columns[field])
        for img in getattr(self.ws, "_images", []):
            try:
                fr = img.anchor._from
                if fr.row + 1 == cell.row and fr.col + 1 == cell.column:
                    return True
            except Exception:
                pass
        return False

    def missing_managed_fields(self, row: int) -> List[str]:
        """Extractor-managed fields of ``row`` that are still blank (only columns present in the sheet)."""
        missing: List[str] = []
        for f in MANAGED_TEXT_FIELDS:
            if f not in self.columns:
                continue
            cell, _ = self._anchor(row, self.columns[f])
            if cell.value in (None, "") or (isinstance(cell.value, str) and not cell.value.strip()):
                missing.append(f)
        for f in MANAGED_IMAGE_FIELDS:
            if f in self.columns and not self._has_image_in_cell(row, f):
                missing.append(f)
        return missing

    def _remove_images_in_cell(self, row: int, field: str) -> None:
        if field not in self.columns:
            return
        if self._has_image_in_cell(row, field):
            self._ensure_backup()
        cell, _ = self._anchor(row, self.columns[field])
        keep = []
        for img in getattr(self.ws, "_images", []):
            try:
                fr = img.anchor._from
                if fr.row + 1 == cell.row and fr.col + 1 == cell.column:
                    continue
            except Exception:
                pass
            keep.append(img)
        self.ws._images = keep

    # ------------------------------------------------------------------
    # Writing
    # ------------------------------------------------------------------
    def _write_content(self, row: int, rec, qpn_png: Optional[Path], improvement_jpg,
                       fill_temporary: bool) -> None:
        """Report content shared by append/update: model, item, defect, cause, improvement, images."""
        self._set_cell(row, "model", rec.model)
        self._set_cell(row, "item", rec.item)
        self._set_cell(row, "defect_content", rec.defect_content)
        self._set_cell(row, "root_cause", rec.root_cause)
        self._set_cell(row, "improvement", rec.improvement)
        # "Xử lý tạm thời" column: kept separate from improvement; blank unless explicitly enabled
        if fill_temporary:
            self._set_cell(row, "temporary", rec.temporary_excluded)
        heights = [self.ws.row_dimensions[row].height or 15.0]
        heights.append(self._text_height_pt(row, "improvement", rec.improvement))
        heights.append(self._text_height_pt(row, "root_cause", rec.root_cause))
        heights.append(self._text_height_pt(row, "defect_content", rec.defect_content))
        self._write_images_with_row_height(row, heights, qpn_png, improvement_jpg, want_qpn=True, want_imp=True)

    def _write_vendor_and_date(self, row: int, rec, overwrite_blank_only: bool) -> List[str]:
        """Vendor / Ngày phát sinh: fill when blank, keep when equal, never overwrite a different value."""
        notes: List[str] = []
        if "vendor" in self.columns:
            cell, _ = self._anchor(row, self.columns["vendor"])
            existing = str(cell.value).strip() if cell.value not in (None, "") else ""
            ppt_set = list(getattr(rec, "vendors", None) or [v for v in (rec.vendor or "").split("\n") if v])
            if ppt_set:
                if not existing:
                    cell.value = "\n".join(ppt_set)                      # one vendor per line
                    al = copy(cell.alignment) if cell.alignment else Alignment()
                    cell.alignment = Alignment(horizontal=al.horizontal, vertical=al.vertical or "top",
                                               wrap_text=True, indent=al.indent, text_rotation=al.text_rotation)
                else:
                    xl_set, unknown = canonicalize_vendor_value(existing)
                    if set(xl_set) != set(ppt_set) or unknown:
                        shown = " / ".join(xl_set + unknown) if (xl_set or unknown) else existing
                        notes.append(f"Vendor trong Excel: {shown}; Vendor trong báo cáo: {' / '.join(ppt_set)} "
                                     f"(giữ giá trị Excel)")
        if "occurrence_date" in self.columns:
            cell, _ = self._anchor(row, self.columns["occurrence_date"])
            existing = self._as_date(cell.value)
            if rec.occurrence_date:
                if cell.value in (None, ""):
                    self._write_date(row, rec.occurrence_date)
                elif existing is None:
                    notes.append(f"Ngày phát sinh trong Excel '{cell.value}' không đọc được, ngày theo Management Number là "
                                 f"{rec.occurrence_date_text} (giữ giá trị Excel)")
                elif existing != rec.occurrence_date:
                    notes.append(f"Ngày phát sinh trong Excel {existing.strftime('%d/%m/%Y')} khác ngày theo Management Number "
                                 f"{rec.occurrence_date_text} (giữ giá trị Excel)")
        return notes

    def create_row(self, mgmt: str) -> int:
        """Create ONE blank report row for a Management Number that is absent from the master.

        * destination = ``next_row()``: the first unused report row of the table (a blank form row the template
          intentionally contains is reused; otherwise the row directly after the last report row) – never header /
          summary rows, never another sheet;
        * FORMAT only is cloned from the nearest report row above (styles, borders, fills, fonts, number formats,
          alignment/wrap, single-row merges, row height); values and pictures are never copied;
        * the Management Number (and STT) is written immediately so a later lookup finds the row.
        Returns the row number.
        """
        mgmt = (mgmt or "").strip()
        if not mgmt or "management_number" not in self.columns:
            raise ValueError("Management Number trống – không tạo dòng mới")
        existing = self.find_rows_by_management_number(mgmt)
        if existing:
            return existing[0]                                   # never a second row for the same key
        self._ensure_backup()                                    # one backup per batch, BEFORE the first change
        row = self.next_row()
        src = row - 1
        if src >= self.data_start and self._row_used(src):
            self._copy_row_style(src, row)
        if "stt" in self.columns:
            self._set_cell(row, "stt", self.used_count() + 1, wrap=False)
        self._set_cell(row, "management_number", mgmt)
        self._dirty = True
        return row

    def append_record(self, rec, qpn_png: Optional[Path] = None, improvement_jpg: Optional[Path] = None,
                      status_text: str = "", note_text: str = "", fill_temporary: bool = False) -> int:
        """Write one report as one NEW row (append mode). Returns the Excel row number."""
        row = self.next_row()
        src = row - 1 if row - 1 >= self.data_start else self.data_start
        if src != row:
            self._copy_row_style(src, row)
        if "stt" in self.columns:
            self._set_cell(row, "stt", self.used_count() + 1, wrap=False)
        self._set_cell(row, "management_number", rec.management_number)
        self._set_cell(row, "vendor", rec.vendor or "")
        self._write_date(row, rec.occurrence_date)
        for i in range(1, 9):   # WEEK +1..+8: blank unless real source data (none parsed) -> blank
            self._set_cell(row, f"week_{i}", rec.weeks.get(i, "") or "", wrap=False)
        self._write_content(row, rec, qpn_png, improvement_jpg, fill_temporary)
        if status_text:
            self._set_cell(row, "status", status_text)
        if note_text:
            self._set_cell(row, "note", note_text)
        return row

    def update_record(self, row: int, rec, qpn_png: Optional[Path] = None, improvement_jpg: Optional[Path] = None,
                      status_text: str = "", note_text: str = "", fill_temporary: bool = False) -> List[str]:
        """Write the report into an EXISTING row located by Management Number.

        * the Management Number cell is preserved untouched;
        * Vendor / Ngày phát sinh: fill if blank, keep if equal, conflict -> keep + note;
        * WEEK +1..+8 and any other columns are left as they are.
        Returns conflict notes (to be added to 'Cần kiểm tra').
        """
        notes = self._write_vendor_and_date(row, rec, overwrite_blank_only=True)
        self._write_content(row, rec, qpn_png, improvement_jpg, fill_temporary)
        if status_text:
            self._set_cell(row, "status", status_text)
        if note_text or notes:
            self._set_cell(row, "note", "; ".join(x for x in [note_text, *notes] if x))
        return notes

    def mark_rows_red(self, rows: List[int]) -> None:
        """Highlight duplicate Management Number rows (values untouched)."""
        if not rows:
            return
        red = PatternFill("solid", fgColor="FFC7CE")
        last_col = max(self.columns.values()) if self.columns else self.ws.max_column
        self._ensure_backup()
        for r in rows:
            for c in range(1, last_col + 1):
                cell = self.ws.cell(row=r, column=c)
                try:
                    cell.fill = red
                except AttributeError:      # MergedCell
                    pass
        self._dirty = True

    def fill_missing_fields(self, row: int, rec, missing: List[str], qpn_png: Optional[Path] = None,
                            improvement_jpg=None) -> List[str]:
        """Partial update of an EXISTING row: write ONLY the fields listed in ``missing``.

        Populated cells are preserved even when the report holds another value (Vendor / Ngày phát sinh keep
        their conflict notes as before).  Returns the conflict notes.
        """
        notes = self._write_vendor_and_date(row, rec, overwrite_blank_only=True)
        heights = [self.ws.row_dimensions[row].height or 15.0]
        for f in ("model", "item", "defect_content", "root_cause", "improvement"):
            if f in missing:
                val = getattr(rec, f, "") or ""
                if val:
                    self._set_cell(row, f, val)
                    heights.append(self._text_height_pt(row, f, val))
        self._write_images_with_row_height(row, heights, qpn_png, improvement_jpg,
                                           want_qpn="qpn" in missing, want_imp="improvement_image" in missing)
        return notes

    def _refresh_images(self) -> None:
        """Work around openpyxl closing in-memory image buffers after a save.

        Images read from an existing workbook keep a BytesIO ``ref`` that openpyxl
        closes during ``save``; the next save would fail with
        'I/O operation on closed file'.  We cache the bytes once and hand a fresh
        buffer to openpyxl before every save.
        """
        for ws in self.wb.worksheets:
            for img in getattr(ws, "_images", []):
                ref = img.ref
                if isinstance(ref, (str, Path)):
                    continue
                cached = getattr(img, "_re_bytes", None)
                if cached is None:
                    try:
                        if hasattr(ref, "getvalue"):
                            cached = ref.getvalue()
                        else:
                            ref.seek(0)
                            cached = ref.read()
                    except Exception:
                        continue
                    img._re_bytes = cached  # type: ignore[attr-defined]
                img.ref = BytesIO(cached)

    def save(self) -> Path:
        if self.probe:
            raise TemplateError("Workbook mở ở chế độ quét (probe) – không ghi")
        if not self._dirty and not self._placed:
            return self.output                           # skip-only batch: the master is not rewritten
        tmp = self.output.with_name(self.output.stem + ".saving.xlsx")
        problems = self.image_bounds_report()
        if problems:                                   # never persist an overflowing picture (#39)
            raise TemplateError("Ảnh vượt ra ngoài ô đích: " + "; ".join(problems))
        self._refresh_images()
        try:
            self.wb.save(tmp)                           # openpyxl closes the ZipFile it writes
        except OSError as e:
            self._discard_tmp(tmp)
            if is_sharing_violation(e):
                raise ExcelLockedError(tmp, e) from e
            raise
        try:
            os.replace(tmp, self.output)                # atomic on the same volume; the master is never half-written
        except OSError as e:
            self._discard_tmp(tmp)                      # the original master stays valid, backup stays valid
            if is_sharing_violation(e):
                LOG.warning("EXCEL_COMMIT_LOCKED target=%s error=%r", self.output, e)
                raise ExcelLockedError(self.output, e) from e
            raise
        self._dirty = False
        return self.output

    @staticmethod
    def _discard_tmp(tmp: Path) -> None:
        try:
            if tmp.exists():
                tmp.unlink()
        except OSError:
            pass

    def close(self) -> None:
        try:
            self.wb.close()
        except Exception:
            pass


def validate_template(template: Path) -> Tuple[bool, str]:
    """Quick diagnostic: can we find the destination sheet (by structure) and its key columns?"""
    try:
        wb = load_workbook(template, read_only=False)
    except Exception as e:  # noqa: BLE001
        return False, f"Không mở được file: {e}"
    try:
        candidates, notes = [], []
        for ws in wb.worksheets:
            key = norm_key(ws.title)
            if key in EXCLUDED_SHEET_KEYS or any(k in key for k in MAPPING_SHEET_KEYS):
                continue
            found = set()
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
                for c in row:
                    if isinstance(c.value, str) and len(c.value.strip()) <= 60:
                        k = norm_key(c.value)
                        for f in FIELD_ALIASES:
                            if _alias_score(k, f):
                                found.add(f)
            missing = [f for f in REQUIRED_FIELDS if f not in found]
            if missing:
                notes.append(f"'{ws.title}': thiếu " + ", ".join(FIELD_LABELS[f] for f in missing))
            else:
                candidates.append((ws.title, len(found)))
        if len(candidates) == 1:
            return True, f"OK – sheet '{candidates[0][0]}' ({candidates[0][1]} cột nhận diện)"
        if candidates:
            return False, "Cần kiểm tra: nhiều sheet có cấu trúc bảng kiểm chứng: " + ", ".join(f"'{n}'" for n, _ in candidates)
        return False, "Không tìm thấy sheet chứa bảng kiểm chứng. " + "; ".join(notes)
    finally:
        wb.close()
