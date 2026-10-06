"""PROMPT-004 – final production extraction + fixed-master rules (§20 regression checklist).

Everything runs on the structural reproduction of the real production layout (tests/test_content_region.py)
plus purpose-built decks; no real report files, no Ollama.
"""
import datetime as dt
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import PatternFill
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

import app.batch_processor as bp
from app.batch_processor import BatchOptions, BatchProcessor
from app.classifier import heuristic_classify
from app.content_region import ROLE_CAPTION, ROLE_CONTENT, ROLE_SIDEBAR, ROLE_TITLE, classify_blocks
from app.excel_writer import BACKUP_DIR_NAME, ExcelWriter, TemplateError
from app.extractor import (defect_content_from_filename, derive_occurrence_date, extract_record,
                           management_number_from_filename)
from app.improvement_pictures import AMBIGUOUS_REASON, select_after_pictures
from app.pptx_parser import parse_pptx

from tests.test_content_region import (CAUSE_BLOCK_1, CAUSE_BLOCK_2, COL, IMP_ITEM_1, IMP_ITEM_2, INSPECTION_TEXT,
                                       LONG_TERM_TEXT, SHEET, TEMP_TEXT, _furniture, _images_at, _pic, _prefill,
                                       _shape, _tb, make_real_layout_report, real_deck)  # noqa: F401

MGMT = "260920045-VOC"
DECK_NAME = "(CTMS)_20601_260920045-VOC_ Đối sách LỖI MẺ XƯỚC, BONG SƠN 24.9.2026 SEV.pptx"


def _image_contains_rgb(image, rgb, tolerance=3):
    with Image.open(image.ref) as pil:
        pixels = pil.convert("RGB").tobytes()
    return any(all(abs(pixels[i + channel] - rgb[channel]) <= tolerance for channel in range(3))
               for i in range(0, len(pixels), 3))


def _run(files, template, out, **kw):
    events = []
    opts = BatchOptions(files=[Path(f) for f in files], template=template, output_file=out, use_ollama=False, **kw)
    proc = BatchProcessor(opts, on_file=lambda i, s, d: events.append((i, s, d)))
    return proc.run(), events, proc


def _slide_text(report, cls, kind):
    rec = extract_record(report, cls, {"rear": "Rear"}, ["A185"])
    return rec.root_cause if kind == "cause" else rec.improvement


# ================================================================== §1-§3 filename rules
@pytest.mark.parametrize("name,mgmt,date_txt", [
    ("(CTMS)_11108_260918080-VOC_ Đối sách LỖI BONG ATN, MỤN 22.9.2026 SEV.pptx", "260918080-VOC", "18/09/2026"),
    ("260922134 Đối sách lỗi Mẻ xước 23.9.2026.pptx", "260922134", "22/09/2026"),
    ("(CTMS) 260923045_ LỖI BONG ATN 22.9.2026 (DRT).pptx", "260923045", "23/09/2026"),
])
def test_management_number_both_forms_and_date_from_key_only(name, mgmt, date_txt):
    assert management_number_from_filename(name) == mgmt
    d, why = derive_occurrence_date(mgmt)
    assert why == "" and f"{d:%d/%m/%Y}" == date_txt          # 22.9.2026 in the name never wins


def test_defect_content_from_filename_strips_trailing_metadata():
    assert defect_content_from_filename("... LỖI BONG ATN, MỤN 22.9.2026 SEV.pptx") == "BONG ATN, MỤN"
    assert defect_content_from_filename("X_ LỖI MẺ XƯỚC, BONG SƠN 24.9.2026 SEV (DRT).pptx") == "MẺ XƯỚC, BONG SƠN"
    assert defect_content_from_filename("X_ LỖI BONG ATN (A185 Rear) 22.9.2026.pptx") == "BONG ATN"
    assert defect_content_from_filename("khong_co_marker 22.9.2026.pptx") == ""


# ================================================================== §4-§7 content region
def test_root_cause_title_and_sidebar_excluded_blocks_retained_in_order(real_deck):
    r, cls, rec = real_deck
    roles = {b.text.split("\n")[0]: b.role for b in classify_blocks(r.slide(2))}
    assert roles["1. NGUYÊN NHÂN"] == ROLE_TITLE and roles["Nguyên nhân"] == ROLE_SIDEBAR
    assert roles["Nguyên nhân trong kiểm tra:"] == ROLE_CONTENT
    assert rec.root_cause == CAUSE_BLOCK_1 + "\n\n" + CAUSE_BLOCK_2            # both blocks, source order, verbatim
    assert "1. NGUYÊN NHÂN" not in rec.root_cause and "CTMS" not in rec.root_cause
    assert rec.root_cause.count("Nguyên nhân trong kiểm tra:") == 2            # repeated valid subheading kept


def test_improvement_title_sidebar_and_graphical_captions_excluded(real_deck):
    r, cls, rec = real_deck
    roles = classify_blocks(r.slide(4))
    by = {b.text.split("\n")[0]: b for b in roles}
    assert by["3. CẢI TIẾN TRONG SẢN XUẤT"].role == ROLE_TITLE
    assert by["Cải tiến trong kiểm tra"].role == ROLE_SIDEBAR                   # circular sidebar label
    caps = [b for b in roles if b.role == ROLE_CAPTION]
    assert {b.text for b in caps} == {"Trước cải tiến", "Sau cải tiến"} and len(caps) == 4
    imp = rec.improvement
    for furniture in ("3. CẢI TIẾN TRONG SẢN XUẤT", "Trước cải tiến", "Sau cải tiến", "Cải tiến trong kiểm tra\n"):
        assert furniture not in imp


def test_business_truoc_sau_lines_and_all_items_retained(real_deck):
    r, cls, rec = real_deck
    imp = rec.improvement
    assert "+ Trước: Jig nén bằng nhôm" in imp and "+ Sau: Bọc silicon 2mm" in imp
    assert "+ Trước: Tool miết đầu nhọn" in imp and "+ Sau: Tool miết đầu tròn R3" in imp
    assert imp.index(IMP_ITEM_1) < imp.index(IMP_ITEM_2)                       # multiple items, in order
    assert IMP_ITEM_1 in imp and IMP_ITEM_2 in imp                             # verbatim, not summarised


def test_long_term_retained_temporary_excluded(real_deck):
    r, cls, rec = real_deck
    assert LONG_TERM_TEXT in rec.improvement and "5. ĐỐI SÁCH LÂU DÀI" not in rec.improvement
    assert "Sorting 100%" not in rec.improvement and TEMP_TEXT not in rec.improvement
    assert "2. XỬ LÝ TẠM THỜI" not in rec.improvement


# ================================================================== §8-§10 After-only pictures
def _deck(path: Path, slides):
    """slides: list of callables(slide, W) building one slide each (slide 1 is a cover)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W = prs.slide_width
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    _tb(s, "BÁO CÁO ĐỐI SÁCH LỖI MẺ XƯỚC REAR A185", 0.5, 1, 12, 1, 28, True)
    _tb(s, f"Model: A185\nItem: Rear\nManagement No: {MGMT}", 0.5, 2.5, 6, 2, 16)
    for fn in slides:
        fn(prs.slides.add_slide(blank), W)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def _cause(s, W):
    _tb(s, "1. NGUYÊN NHÂN", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Nguyên nhân", 0.15, 2.5, 1.3, 1.3)
    _tb(s, CAUSE_BLOCK_1, 1.7, 1.1, 11, 2.2, 13)
    _furniture(s, W)


def _prod(n_before, n_after, text=IMP_ITEM_1, title="3. CẢI TIẾN TRONG SẢN XUẤT", extra=None):
    def build(s, W):
        _tb(s, title, 0.5, 0.3, 8, 0.7, 24, True)
        _shape(s, MSO_SHAPE.OVAL, "Cải tiến trong sản xuất", 0.15, 3.0, 1.3, 1.3)
        _tb(s, text, 1.7, 1.0, 5.6, 2.8, 12)
        _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Trước cải tiến", 7.5, 1.0, 1.5, 0.35)
        for i in range(n_before):
            s.shapes.add_picture(_pic("#ffe0b2", "b"), Inches(7.5 + 1.4 * i), Inches(1.4), Inches(1.3), Inches(1.2))
        s.shapes.add_picture(_pic("#999999", "arrow", (400, 60)), Inches(7.5), Inches(2.62), Inches(2.0), Inches(0.3))
        _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 2.95, 1.5, 0.35)
        for i in range(n_after):
            s.shapes.add_picture(_pic("#c8e6c9", "a"), Inches(7.5 + 1.4 * i), Inches(3.35), Inches(1.3), Inches(1.2))
        if extra:
            extra(s, W)
        _furniture(s, W)
    return build


def _control(s, W):
    _tb(s, "4. CẢI TIẾN KIỂM SOÁT", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Cải tiến kiểm soát", 0.15, 3.0, 1.3, 1.3)
    _tb(s, INSPECTION_TEXT, 1.7, 1.0, 5.6, 2.5, 12)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 1.0, 1.5, 0.35)
    s.shapes.add_picture(_pic("#b3e5fc", "control"), Inches(7.5), Inches(1.4), Inches(4), Inches(2.5))
    _furniture(s, W)


def _select(path):
    r = parse_pptx(path)
    cls = heuristic_classify(r)
    sel = select_after_pictures(r, cls.improvement_image_slides or cls.improvement_slides)
    return r, cls, sel


def test_one_before_one_after_gives_only_after(tmp_path):
    r, cls, sel = _select(_deck(tmp_path / DECK_NAME, [_cause, _prod(1, 1)]))
    assert len(sel.after) == 1 and [x.kind for x in sel.rejected if x.kind == "before"] == ["before"]
    assert sel.after[0].block.top > sel.rejected[0].block.top                    # the lower (After) picture


def test_two_before_three_after_gives_exactly_three_after(tmp_path):
    r, cls, sel = _select(_deck(tmp_path / DECK_NAME, [_cause, _prod(2, 3)]))
    assert len(sel.after) == 3 and sum(1 for x in sel.rejected if x.kind == "before") == 2
    assert [a.block.left for a in sel.after] == sorted(a.block.left for a in sel.after)     # source order
    assert not sel.reasons


def test_multiple_items_all_after_pictures_in_order(real_deck):
    r, cls, rec = real_deck
    assert len(rec.after_pictures) == 5 and rec.after_picture_slides == [4]
    orders = [p.block.order for p in rec.after_pictures]
    assert orders == sorted(orders)                                              # item 1 (3) then item 2 (2)
    sel = select_after_pictures(r, [4])
    assert sum(1 for x in sel.rejected if x.kind == "before") == 3               # 2 + 1 Before, none inserted


def test_multiple_slides_keep_report_order(tmp_path):
    deck = _deck(tmp_path / DECK_NAME, [_cause, _prod(1, 2), _prod(1, 1, text=IMP_ITEM_2)])
    r, cls, sel = _select(deck)
    assert [a.slide for a in sel.after] == [3, 3, 4] and sel.slides_with_after == [3, 4]


def test_control_inspection_pictures_arrows_logos_excluded(tmp_path):
    deck = _deck(tmp_path / DECK_NAME, [_cause, _prod(1, 2), _control])
    r, cls, sel = _select(deck)
    assert [a.slide for a in sel.after] == [3, 3]                                 # control slide contributes nothing
    reasons = [x.reason for x in sel.rejected]
    assert any("inspection/control" in why for why in reasons)
    assert any("arrow" in why or "decor" in why or "ratio" in why for why in reasons)   # arrow strip
    assert all(x.slide != 4 for x in sel.after)
    logos = [x for x in sel.rejected if x.block.top < r.slide_height * 0.12 and x.block.left > r.slide_width * 0.78]
    assert logos and all(x.kind == "excluded" for x in logos)


def test_sidebar_inspection_label_does_not_exclude_production_slide(real_deck):
    r, cls, rec = real_deck                                                      # slide 4 sidebar says "Cải tiến trong kiểm tra"
    assert rec.after_picture_slides == [4] and len(rec.after_pictures) == 5


def test_ambiguous_picture_omitted_with_review_reason(real_deck):
    r, cls, rec = real_deck
    sel = select_after_pictures(r, [7])
    assert sel.after == [] and AMBIGUOUS_REASON.format(n=7) in sel.reasons
    assert AMBIGUOUS_REASON.format(n=7) == "Cần kiểm tra: Không xác định chắc chắn ảnh Sau cải tiến tại slide 7"
    assert any("tại slide 7" in x for x in rec.review_reasons)


def test_before_pictures_never_reach_excel(template, tmp_path):
    deck = _deck(tmp_path / DECK_NAME, [_cause, _prod(2, 3)])
    _prefill(template, [{"mgmt": MGMT}])
    out = tmp_path / "out" / "k.xlsx"
    s, ev, proc = _run([deck], template, out)
    ws = load_workbook(out)[SHEET]
    imgs = _images_at(ws, 4, COL["image"])
    assert len(imgs) == 1  # the three After pictures of this defect share one visual-region crop
    assert _image_contains_rgb(imgs[0], (0xC8, 0xE6, 0xC9))
    assert not _image_contains_rgb(imgs[0], (0xFF, 0xE0, 0xB2))


# ================================================================== §11-§17 fixed master
def _manual_png(tmp_path, name="manual.png", color="red"):
    png = tmp_path / name
    Image.new("RGB", (40, 30), color).save(png)
    return png


def _add_image(template, anchor, png):
    wb = load_workbook(template)
    img = XLImage(str(png))
    img.anchor = anchor
    wb[SHEET].add_image(img)
    wb.save(template)


def test_complete_row_skipped_before_parse_and_qwen(template, tmp_path, monkeypatch):
    deck = make_real_layout_report(tmp_path / DECK_NAME)
    _prefill(template, [{"mgmt": MGMT, "vendor": "Doaltech", "model": "A185", "item": "Rear", "defect": "x",
                         "cause": "y", "improvement": "z", "date": dt.date(2026, 9, 20)}])
    _add_image(template, "H4", _manual_png(tmp_path))
    _add_image(template, "K4", _manual_png(tmp_path, "m2.png", "blue"))
    calls = {"parse": 0, "classify": 0}
    monkeypatch.setattr(bp, "parse_pptx", lambda p: calls.__setitem__("parse", calls["parse"] + 1))
    monkeypatch.setattr(bp, "classify", lambda *a, **k: calls.__setitem__("classify", calls["classify"] + 1))
    out = tmp_path / "out" / "k.xlsx"
    s, ev, proc = _run([deck], template, out)
    assert s.skipped == 1 and calls == {"parse": 0, "classify": 0}
    assert proc.results[0].error.startswith("Bỏ qua — đã cập nhật")
    assert s.backup_file == "" and not (out.parent / BACKUP_DIR_NAME).exists()   # skip-only batch: no backup
    ws = load_workbook(out)[SHEET]
    assert len(_images_at(ws, 4, COL["qpn"])) == 1 and len(_images_at(ws, 4, COL["image"])) == 1


def test_partial_row_fills_only_missing_and_preserves_populated(template, tmp_path):
    deck = make_real_layout_report(tmp_path / DECK_NAME)
    _prefill(template, [{"mgmt": MGMT, "vendor": "Teawon", "model": "A185", "item": "Rear", "defect": "nhập tay",
                         "improvement": "đối sách nhập tay", "date": dt.date(2026, 9, 20)}])
    _add_image(template, "H4", _manual_png(tmp_path))                            # QPN present
    out = tmp_path / "out" / "k.xlsx"
    s, ev, proc = _run([deck], template, out)
    fr = proc.results[0]
    assert set(fr.filled_fields) == {"root_cause", "improvement_image"}
    ws = load_workbook(out)[SHEET]
    assert ws.cell(row=4, column=COL["cause"]).value == CAUSE_BLOCK_1 + "\n\n" + CAUSE_BLOCK_2
    assert ws.cell(row=4, column=COL["improvement"]).value == "đối sách nhập tay"      # populated -> preserved
    assert ws.cell(row=4, column=COL["vendor"]).value == "Teawon"                      # not overwritten by Doaltech
    assert ws.cell(row=4, column=COL["defect"]).value == "nhập tay"
    assert len(_images_at(ws, 4, COL["qpn"])) == 1                                     # existing QPN untouched
    with Image.open(_images_at(ws, 4, COL["qpn"])[0].ref) as pil:
        assert pil.convert("RGB").getpixel((2, 2)) == (255, 0, 0)
    assert len(_images_at(ws, 4, COL["image"])) == 2  # one complete rendered crop per production defect
    assert all(ws.cell(row=4, column=k).value == "OK" for k in range(12, 20))


def test_manually_cleared_field_is_refilled(template, tmp_path):
    deck = make_real_layout_report(tmp_path / DECK_NAME)
    _prefill(template, [{"mgmt": MGMT}])
    out = tmp_path / "out" / "k.xlsx"
    s1, _, _ = _run([deck], template, out)
    wb = load_workbook(out)
    ws = wb[SHEET]
    cause = ws.cell(row=4, column=COL["cause"]).value
    assert cause
    ws.cell(row=4, column=COL["cause"]).value = None                              # user clears it by hand
    wb.save(out)
    s2, _, proc2 = _run([deck], template, out)
    assert proc2.results[0].filled_fields == ["root_cause"]
    ws2 = load_workbook(out)[SHEET]
    assert ws2.cell(row=4, column=COL["cause"]).value == cause
    assert len(_images_at(ws2, 4, COL["image"])) == 2  # existing defect crops are not re-inserted


def test_existing_improvement_images_preserved_when_only_text_missing(template, tmp_path):
    deck = make_real_layout_report(tmp_path / DECK_NAME)
    _prefill(template, [{"mgmt": MGMT, "defect": "d", "improvement": "i"}])
    _add_image(template, "H4", _manual_png(tmp_path))
    _add_image(template, "K4", _manual_png(tmp_path, "m2.png", "blue"))         # manual improvement picture
    out = tmp_path / "out" / "k.xlsx"
    s, ev, proc = _run([deck], template, out)
    assert "improvement_image" not in proc.results[0].filled_fields and "root_cause" in proc.results[0].filled_fields
    ws = load_workbook(out)[SHEET]
    imgs = _images_at(ws, 4, COL["image"])
    assert len(imgs) == 1
    with Image.open(imgs[0].ref) as pil:
        assert pil.convert("RGB").getpixel((2, 2)) == (0, 0, 255)


def test_duplicate_rows_topmost_updated_extras_red_values_unchanged(template, tmp_path):
    deck = make_real_layout_report(tmp_path / DECK_NAME)
    _prefill(template, [{"mgmt": "260901001-VOC", "defect": "other"},
                        {"mgmt": MGMT},
                        {"mgmt": "260902002-VOC"},
                        {"mgmt": MGMT, "defect": "giữ nguyên", "cause": "giữ nguyên"},
                        {"mgmt": MGMT}])
    wb = load_workbook(template)
    wb[SHEET].cell(row=7, column=12).value = "=1+1"                              # formula in a duplicate row
    wb.save(template)
    out = tmp_path / "out" / "k.xlsx"
    s, ev, proc = _run([deck], template, out)
    fr = proc.results[0]
    assert fr.excel_row == 5 and fr.duplicate_rows == [7, 8] and fr.status in ("completed", "needs_review")
    assert not any("tô đỏ" in r or "xuất hiện" in r for r in fr.review_reasons)  # diagnostics only
    diag = bp.format_file_diagnostics(fr)
    assert "Dòng sử dụng       : 5" in diag and "Management Number bị trùng tại dòng: 7, 8" in diag
    assert "Đã đánh dấu đỏ các dòng trùng." in diag
    ws = load_workbook(out)[SHEET]
    assert ws.cell(row=5, column=COL["cause"]).value and len(_images_at(ws, 5, COL["image"])) == 2
    for r in (7, 8):
        assert ws.cell(row=r, column=COL["mgmt"]).fill.fgColor.rgb.endswith("FFC7CE")
        assert len(_images_at(ws, r, COL["image"])) == 0
    assert ws.cell(row=7, column=COL["defect"]).value == "giữ nguyên" and ws.cell(row=7, column=COL["cause"]).value == "giữ nguyên"
    assert ws.cell(row=8, column=COL["cause"]).value is None
    assert ws.cell(row=7, column=12).value == "=1+1"                                # formula preserved
    assert not str(ws.cell(row=4, column=COL["mgmt"]).fill.fgColor.rgb).endswith("FFC7CE")
    assert not str(ws.cell(row=6, column=COL["mgmt"]).fill.fgColor.rgb).endswith("FFC7CE")
    assert all(ws.cell(row=5, column=k).value == "OK" for k in range(12, 20))


def test_backup_before_first_modification_only(template, tmp_path):
    deck = make_real_layout_report(tmp_path / DECK_NAME)
    _prefill(template, [{"mgmt": MGMT}, {"mgmt": "260902002-VOC"}])
    _add_image(template, "H4", _manual_png(tmp_path))                            # deck has no QPN slide
    out = tmp_path / "out" / "k.xlsx"
    s, ev, proc = _run([deck], template, out)
    bdir = out.parent / BACKUP_DIR_NAME
    backups = sorted(bdir.glob("k_backup_*.xlsx"))
    assert len(backups) == 1 and s.backup_file == str(backups[0])
    wsb = load_workbook(backups[0])[SHEET]
    assert wsb.cell(row=4, column=COL["cause"]).value is None                     # pre-modification state
    assert load_workbook(out)[SHEET].cell(row=4, column=COL["cause"]).value
    # second run: everything complete -> skip, no new backup, master not rewritten
    mtime = out.stat().st_mtime_ns
    s2, _, _ = _run([deck], template, out)
    assert s2.skipped == 1 and s2.backup_file == "" and sorted(bdir.glob("k_backup_*.xlsx")) == backups
    assert out.stat().st_mtime_ns == mtime


def test_backup_taken_for_red_marking_only(template, tmp_path):
    deck = make_real_layout_report(tmp_path / DECK_NAME)
    _prefill(template, [{"mgmt": MGMT, "vendor": "Doaltech", "model": "A185", "item": "Rear", "defect": "x",
                         "cause": "y", "improvement": "z", "date": dt.date(2026, 9, 20)}, {"mgmt": MGMT}])
    _add_image(template, "H4", _manual_png(tmp_path))
    _add_image(template, "K4", _manual_png(tmp_path, "m2.png", "blue"))
    out = tmp_path / "out" / "k.xlsx"
    s, ev, proc = _run([deck], template, out)
    assert s.skipped == 1 and proc.results[0].duplicate_rows == [5]
    assert s.backup_file and Path(s.backup_file).exists()                          # red marking is a modification


def test_backup_failure_blocks_modification(template, tmp_path, monkeypatch):
    _prefill(template, [{"mgmt": MGMT}])
    out = tmp_path / "out" / "k.xlsx"
    w = ExcelWriter(template, out)
    monkeypatch.setattr(bp.shutil if hasattr(bp, "shutil") else __import__("shutil"), "copyfile",
                        lambda *a, **k: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(TemplateError, match="sao lưu"):
        w._set_cell(4, "root_cause", "x")
    assert w.ws.cell(row=4, column=COL["cause"]).value is None and w.backup_path is None and not w._dirty
    w.close()
    assert load_workbook(out)[SHEET].cell(row=4, column=COL["cause"]).value is None


def test_probe_writer_never_backs_up_or_writes(template, tmp_path):
    _prefill(template, [{"mgmt": MGMT}])
    w = ExcelWriter(template, tmp_path / "out" / "k.xlsx", probe=True)
    with pytest.raises(TemplateError):
        w._set_cell(4, "root_cause", "x")
    assert w.backup_path is None and not (tmp_path / "out").exists()
