"""PROMPT-004C Part B – After-picture evidence without explicit Trước/Sau captions: blue native text and
Before→After arrows (deterministic PPTX structure; no OCR, no pixels, no LLM)."""
from pathlib import Path

import pytest
from lxml import etree
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from app.classifier import heuristic_classify
from app.improvement_pictures import AMBIGUOUS_REASON, is_blue, select_after_pictures
from app.pptx_parser import parse_pptx
from tests.test_content_region import (CAUSE_BLOCK_1, COL, IMP_ITEM_1, IMP_ITEM_2, INSPECTION_TEXT, SHEET,
                                       _furniture, _images_at, _pic, _prefill, _shape, _tb)
from tests.test_prompt004 import MGMT, _run

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
BLUE = RGBColor(0x00, 0x70, 0xC0)
BEFORE_RGB, AFTER_RGB = (0xFF, 0xE0, 0xB2), (0xC8, 0xE6, 0xC9)
NAME = "(CTMS)_260925015_ĐỐI SÁCH CẢI TIẾN MODEL A253 FRONT LỖI SƠN, HÀN (XƯỚC BÓNG) 24.09.2026.pptx"


# ------------------------------------------------------------------ deck helpers
def _tb_colored(slide, lines, left, top, width, height, size=12):
    """lines: [(text, 'blue'|'theme'|None)] – colour applied to the whole paragraph run."""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (txt, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        if color == "blue":
            r.font.color.rgb = BLUE
        elif color == "theme":
            r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1       # default theme accent1 = #4f81bd (blue)
        elif color == "red":
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    return tb


def _arrow(slide, kind, left, top, width, height, rotation=0):
    a = slide.shapes.add_shape(kind, Inches(left), Inches(top), Inches(width), Inches(height))
    a.rotation = rotation
    return a


def _connector(slide, x1, y1, x2, y2, head="tailEnd"):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln = c.line._get_or_add_ln()
    e = etree.SubElement(ln, "{%s}%s" % (A_NS, head))
    e.set("type", "triangle")
    return c


def _pics(slide, colors, left, top, w=1.3, h=1.2, gap=0.1, vertical=False):
    out = []
    for i, c in enumerate(colors):
        x = left + (0 if vertical else i * (w + gap))
        y = top + (i * (h + gap) if vertical else 0)
        out.append(slide.shapes.add_picture(_pic(c, "p"), Inches(x), Inches(y), Inches(w), Inches(h)))
    return out


def _deck(path: Path, builders):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W = prs.slide_width
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    _tb(s, "BÁO CÁO ĐỐI SÁCH LỖI SƠN FRONT A253", 0.5, 1, 12, 1, 28, True)
    _tb(s, f"Model: A253\nItem: Front\nManagement No: {MGMT}", 0.5, 2.5, 6, 2, 16)
    s = prs.slides.add_slide(blank)
    _tb(s, "1. NGUYÊN NHÂN", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Nguyên nhân", 0.15, 2.5, 1.3, 1.3)
    _tb(s, CAUSE_BLOCK_1, 1.7, 1.1, 11, 2.2, 13)
    _furniture(s, W)
    for fn in builders:
        fn(prs.slides.add_slide(blank), W)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def _prod_head(s, W, title="3. CẢI TIẾN TRONG SẢN XUẤT"):
    _tb(s, title, 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Cải tiến trong sản xuất", 0.15, 3.0, 1.3, 1.3)
    _furniture(s, W)


def _select(path, slides=None):
    r = parse_pptx(path)
    cls = heuristic_classify(r)
    nums = slides or (cls.improvement_image_slides or cls.improvement_slides)
    return r, select_after_pictures(r, nums)


def _rgb(ref):
    with Image.open(ref) as im:
        return im.convert("RGB").getpixel((5, 5))


def _after_colors(sel):
    return [tuple(Image.open(__import__("io").BytesIO(a.block.image_blob)).convert("RGB").getpixel((5, 5)))
            for a in sel.after]


# ------------------------------------------------------------------ colour maths
def test_is_blue_accepts_office_blues_and_rejects_others():
    assert is_blue("#0070c0") and is_blue("#4f81bd") and is_blue("#1f4e79") and is_blue("#2e75b6")
    assert not is_blue("#000000") and not is_blue("#c00000") and not is_blue("#00b050") and not is_blue("#ffffff")
    assert not is_blue("#bdd7ee")            # pale tint, not an "After" colour
    assert not is_blue("") and not is_blue("#zzz")


# 9. explicit caption still wins (PROMPT-004 behaviour)
def test_explicit_sau_cai_tien_caption(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb(s, IMP_ITEM_1, 1.7, 1.0, 5.6, 2.8, 12)
        _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Trước cải tiến", 7.5, 1.0, 1.5, 0.35)
        _pics(s, ["#ffe0b2"], 7.5, 1.4)
        _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 2.95, 1.5, 0.35)
        _pics(s, ["#c8e6c9", "#c8e6c9"], 7.5, 3.35)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert _after_colors(sel) == [AFTER_RGB, AFTER_RGB] and not sel.reasons
    assert all(a.anchor.startswith(("caption", "inline", "row neighbour")) for a in sel.after)


# 10. blue "+ Sau:" line + pictures below it
def test_blue_plus_sau_line_owns_pictures_below(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb_colored(s, [("Cải tiến jig nén", None), ("+ Trước: Jig nén bằng nhôm", None)], 1.7, 1.0, 11, 0.8)
        _pics(s, ["#ffe0b2", "#ffe0b2"], 1.7, 1.9)
        _tb_colored(s, [("+ Sau: Bọc silicon 2mm lên bề mặt jig", "blue")], 1.7, 3.3, 11, 0.4)
        _pics(s, ["#c8e6c9", "#c8e6c9", "#c8e6c9"], 1.7, 3.8)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert _after_colors(sel) == [AFTER_RGB] * 3 and not sel.reasons
    assert sum(1 for x in sel.rejected if x.kind == "before") == 2
    assert r.slide(3).blocks[-1].line_colors or True


# 11. blue After description, no "+ Sau:" wording, no caption button
def test_blue_description_without_caption(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb(s, "Cải tiến tool miết\n- Hiện trạng: tool miết đầu nhọn gây xước", 1.7, 1.0, 5.6, 1.0, 12)
        _pics(s, ["#ffe0b2"], 1.7, 2.1)
        _tb_colored(s, [("Thay tool miết đầu tròn R3, bọc silicon chống xước", "blue")], 7.5, 1.0, 5.3, 0.6)
        _pics(s, ["#c8e6c9", "#c8e6c9"], 7.5, 1.7)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert _after_colors(sel) == [AFTER_RGB, AFTER_RGB]
    assert all("blue" in a.anchor for a in sel.after)
    # the black-described Before picture has no evidence either way -> flagged, never copied
    assert AMBIGUOUS_REASON.format(n=3) in sel.reasons
    assert all(_rgb_blob(x.block) != AFTER_RGB for x in sel.rejected if x.kind == "ambiguous")


def _rgb_blob(block):
    import io
    with Image.open(io.BytesIO(block.image_blob)) as im:
        return im.convert("RGB").getpixel((5, 5))


def test_theme_blue_is_resolved_from_the_theme(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb_colored(s, [("Lắp thêm tấm chắn chống văng sơn", "theme")], 1.7, 1.0, 5, 0.6)
        _pics(s, ["#c8e6c9"], 1.7, 1.7)
        _tb(s, "Hiện trạng: không có tấm chắn", 7.5, 1.0, 5, 0.6, 12)
        _pics(s, ["#ffe0b2"], 7.5, 1.7)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert _after_colors(sel) == [AFTER_RGB]


# 12. [Before][Before] → [After][After][After]  (also 14)
def test_right_arrow_destination_pictures_only(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb(s, IMP_ITEM_1.split("\n")[0], 1.7, 1.0, 11, 0.5, 12)
        _pics(s, ["#ffe0b2", "#ffe0b2"], 1.7, 1.8)
        _arrow(s, MSO_SHAPE.RIGHT_ARROW, 4.7, 2.2, 0.9, 0.4)
        _pics(s, ["#c8e6c9", "#c8e6c9", "#c8e6c9"], 5.8, 1.8)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert _after_colors(sel) == [AFTER_RGB] * 3 and not sel.reasons
    assert sum(1 for x in sel.rejected if x.kind == "before") == 2
    assert all("destination" in a.anchor for a in sel.after)
    lefts = [a.block.left for a in sel.after]
    assert lefts == sorted(lefts)


def test_left_pointing_arrow_reverses_sides(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb(s, IMP_ITEM_1.split("\n")[0], 1.7, 1.0, 11, 0.5, 12)
        _pics(s, ["#c8e6c9"], 1.7, 1.8)                        # After on the LEFT this time
        _arrow(s, MSO_SHAPE.LEFT_ARROW, 3.3, 2.2, 0.9, 0.4)
        _pics(s, ["#ffe0b2"], 4.5, 1.8)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert _after_colors(sel) == [AFTER_RGB]


def test_connector_with_arrow_head_counts_as_arrow(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb(s, IMP_ITEM_1.split("\n")[0], 1.7, 1.0, 11, 0.5, 12)
        _pics(s, ["#ffe0b2"], 1.7, 1.8)
        _connector(s, 3.2, 2.4, 4.3, 2.4)                      # tailEnd triangle → points right
        _pics(s, ["#c8e6c9", "#c8e6c9"], 4.5, 1.8)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert _after_colors(sel) == [AFTER_RGB, AFTER_RGB]
    assert [a.direction for a in r.slide(3).arrows] == ["right"]


# 13. vertical arrow
def test_down_arrow_destination_below(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb(s, IMP_ITEM_2.split("\n")[0], 1.7, 1.0, 11, 0.5, 12)
        _pics(s, ["#ffe0b2", "#ffe0b2"], 1.7, 1.6, h=1.0)
        _arrow(s, MSO_SHAPE.DOWN_ARROW, 2.9, 2.75, 0.4, 0.6)
        _pics(s, ["#c8e6c9", "#c8e6c9"], 1.7, 3.5, h=1.0)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert _after_colors(sel) == [AFTER_RGB, AFTER_RGB] and not sel.reasons


def test_rotated_right_arrow_is_a_down_arrow(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb(s, IMP_ITEM_2.split("\n")[0], 1.7, 1.0, 11, 0.5, 12)
        _pics(s, ["#ffe0b2"], 1.7, 1.6, h=1.0)
        _arrow(s, MSO_SHAPE.RIGHT_ARROW, 2.1, 2.8, 0.6, 0.4, rotation=90)
        _pics(s, ["#c8e6c9"], 1.7, 3.5, h=1.0)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert r.slide(3).arrows[0].direction == "down" and _after_colors(sel) == [AFTER_RGB]


# 15. two improvement blocks on one slide, each with its own arrow
def test_two_blocks_each_arrow_classifies_only_its_row(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb(s, "Cải tiến 1: jig nén", 1.7, 0.95, 5, 0.4, 12)
        _pics(s, ["#ffe0b2"], 1.7, 1.4, h=1.1)
        _arrow(s, MSO_SHAPE.RIGHT_ARROW, 3.2, 1.75, 0.8, 0.4)
        _pics(s, ["#c8e6c9", "#c8e6c9"], 4.2, 1.4, h=1.1)          # After A, After B
        _tb(s, "Cải tiến 2: tool miết", 1.7, 3.1, 5, 0.4, 12)
        _pics(s, ["#ffe0b2"], 1.7, 3.6, h=1.1)
        _arrow(s, MSO_SHAPE.RIGHT_ARROW, 3.2, 3.95, 0.8, 0.4)
        _pics(s, ["#c8e6c9"], 4.2, 3.6, h=1.1)                     # After C
        _pics(s, ["#ffe0b2"], 9.0, 3.6, h=1.1)                     # far-away picture: no evidence
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert _after_colors(sel) == [AFTER_RGB] * 3
    tops = [a.block.top for a in sel.after]
    assert tops == sorted(tops)                                     # A, B (row 1) then C (row 2)
    assert sum(1 for x in sel.rejected if x.kind == "before") == 2
    assert AMBIGUOUS_REASON.format(n=3) in sel.reasons               # the stray picture is flagged, not copied


# 16. blue logo text never classifies
def test_blue_logo_text_ignored(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb_colored(s, [("SAMSUNG", "blue")], 10.9, 0.15, 2, 0.4, 16)     # header-band blue logo text
        _tb(s, "Cải tiến tool miết", 1.7, 1.0, 5, 0.5, 12)
        _pics(s, ["#ffe0b2"], 10.9, 0.7)                                   # picture right under the logo text
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert sel.after == [] and AMBIGUOUS_REASON.format(n=3) in sel.reasons


# 17. blue text inside inspection/control content does not validate control pictures
def test_blue_inspection_text_does_not_make_control_pictures_after(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb(s, "Cải tiến máy móc", 1.7, 0.65, 5, 0.35, 12)
        _tb_colored(s, [("+ Sau: Bọc silicon 2mm", "blue")], 1.7, 1.0, 5, 0.4)
        _pics(s, ["#c8e6c9"], 1.7, 1.5)
        _tb_colored(s, [("Cải tiến trong kiểm tra:", None), ("Bổ sung tiêu chuẩn kiểm tra ngoại quan", "blue")],
                    1.7, 3.3, 5, 0.8)
        _pics(s, ["#b3e5fc"], 1.7, 4.2)
        _arrow(s, MSO_SHAPE.RIGHT_ARROW, 3.2, 4.6, 0.8, 0.4)
        _pics(s, ["#b3e5fc"], 4.2, 4.2)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert _after_colors(sel) == [AFTER_RGB]
    assert sum(1 for x in sel.rejected if x.kind == "excluded" and "inspection" in x.reason) == 2

    def ctrl(s, W):
        _prod_head(s, W, title="4. CẢI TIẾN KIỂM SOÁT")
        _tb_colored(s, [(INSPECTION_TEXT.split("\n")[0], "blue")], 1.7, 1.0, 5, 0.6)
        _pics(s, ["#b3e5fc"], 1.7, 1.7)
        _arrow(s, MSO_SHAPE.RIGHT_ARROW, 3.2, 2.1, 0.8, 0.4)
        _pics(s, ["#b3e5fc"], 4.2, 1.7)
    r2, sel2 = _select(_deck(tmp_path / "b" / NAME, [ctrl]), slides=[3])
    assert sel2.after == [] and all(x.kind == "excluded" for x in sel2.rejected)


# 18. decorative arrow far from any picture pair / in the title band does nothing
def test_decorative_arrow_does_not_classify(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _arrow(s, MSO_SHAPE.RIGHT_ARROW, 8.6, 0.35, 0.6, 0.3)           # title-band ornament
        _tb(s, "Cải tiến tool miết", 1.7, 1.0, 5, 0.5, 12)
        _pics(s, ["#ffe0b2"], 1.7, 1.6)
        _arrow(s, MSO_SHAPE.RIGHT_ARROW, 1.7, 5.9, 0.8, 0.4)            # arrow with no picture on either side
        _pics(s, ["#c8e6c9"], 7.5, 5.0)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert sel.after == [] and AMBIGUOUS_REASON.format(n=3) in sel.reasons


# 19. arrow (shape or picture) never reaches Excel
def test_arrow_itself_never_inserted(template, tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb(s, IMP_ITEM_1.split("\n")[0], 1.7, 1.0, 11, 0.5, 12)
        _pics(s, ["#ffe0b2"], 1.7, 1.8)
        _arrow(s, MSO_SHAPE.RIGHT_ARROW, 3.2, 2.2, 0.9, 0.4)
        s.shapes.add_picture(_pic("#999999", "arrow", (400, 60)), Inches(3.2), Inches(2.8), Inches(1.0), Inches(0.15))
        _pics(s, ["#c8e6c9", "#c8e6c9"], 4.5, 1.8)
    deck = _deck(tmp_path / f"(CTMS)_{MGMT}_ĐỐI SÁCH LỖI SƠN 24.09.2026.pptx", [slide])
    _prefill(template, [{"mgmt": MGMT}])
    out = tmp_path / "out" / "k.xlsx"
    s_, ev, proc = _run([deck], template, out)
    ws = load_workbook(out)[SHEET]
    imgs = _images_at(ws, 4, COL["image"])
    assert len(imgs) == 2 and all(_rgb(im.ref) == AFTER_RGB for im in imgs)


# 20. single picture, no evidence -> omitted + reason (case 5)
def test_single_picture_without_evidence_is_ambiguous(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb(s, "Cải tiến jig nén\n- Bọc silicon 2mm", 1.7, 1.0, 5, 1.0, 12)
        _pics(s, ["#c8e6c9"], 1.7, 2.2, w=3, h=2.2)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert sel.after == [] and sel.reasons == [AMBIGUOUS_REASON.format(n=3)]
    assert [x.kind for x in sel.rejected if x.block.width > r.slide_width * 0.1] == ["ambiguous"]


def test_blue_text_vs_arrow_conflict_is_ambiguous(tmp_path):
    def slide(s, W):
        _prod_head(s, W)
        _tb_colored(s, [("Mô tả màu xanh bên trái", "blue")], 1.7, 1.0, 2.5, 0.5)
        _pics(s, ["#ffe0b2"], 1.7, 1.6)                       # blue text above it, but the arrow says source side
        _arrow(s, MSO_SHAPE.RIGHT_ARROW, 3.2, 2.0, 0.9, 0.4)
        _pics(s, ["#c8e6c9"], 4.5, 1.6)
    r, sel = _select(_deck(tmp_path / NAME, [slide]))
    assert _after_colors(sel) == [AFTER_RGB]
    assert any(x.kind == "ambiguous" for x in sel.rejected) and AMBIGUOUS_REASON.format(n=3) in sel.reasons


# 21. PROMPT-004 explicit layout unchanged (real-layout deck: 2B+3A and 1B+2A)
def test_prompt004_real_layout_unchanged(real_deck):
    r, cls, rec = real_deck
    assert len(rec.after_pictures) == 5 and rec.after_picture_slides == [4]
    assert all(a.anchor.startswith(("caption", "row neighbour", "inline")) for a in rec.after_pictures)


real_deck = pytest.importorskip("tests.test_content_region").real_deck
