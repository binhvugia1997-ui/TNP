"""Headless integration checks for the local React + pywebview application facade.

These exercise the production GuiController and service contracts with temporary report/Excel fixtures.  The
synthetic processor below covers lifecycle/stop semantics only; it is not evidence of PPTX extraction or Windows
WebView/native-dialog acceptance.
"""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import threading
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.application_service import ApplicationService, ServiceError
from app.batch_processor import BatchSummary
from app.bridge import BridgeService
from app.config import AppConfig
from app.content_learning import ContentCandidate
from app.desktop import FRONTEND_DIST, FRONTEND_INDEX, FRONTEND_URL
from app.excel_writer import ExcelLockedError
from app.gui_controller import GuiController
from app.image_learning import ImageCandidate
from app.prescan import (ACTION_FAST_SKIP, ACTION_INVALID_MGMT, ACTION_MASTER_COMPLETE, ACTION_OUTSIDE_PERIOD,
                         ACTION_PROCESS, ACTION_PROCESS_NEW_ROW, ACTION_SOURCE_DUPLICATE)


def _config(sample_tree, tmp_path):
    output = tmp_path / "output" / "Kiem_chung_09_2026.xlsx"
    return {
        "paths": {"reportFolder": str(sample_tree["reports"]), "template": str(sample_tree["template"]),
                  "output": str(output)},
        "period": {"mode": "all", "month": "9", "year": "2026", "from": "", "to": ""},
        "forceReprocess": False,
        "ollama": {"host": "127.0.0.1", "port": "11434", "model": "qwen3:4b"},
        "update": {"path": "", "autoCheck": False},
    }


def _make_service(sample_tree, tmp_path, processor_factory=None):
    dto = _config(sample_tree, tmp_path)
    settings_path = tmp_path / "config" / "settings.json"
    cfg = AppConfig(last_report_folder=dto["paths"]["reportFolder"],
                    last_template=dto["paths"]["template"], last_output_file=dto["paths"]["output"],
                    period_mode="all")
    kwargs = {"processor_factory": processor_factory} if processor_factory else {}
    controller = GuiController(cfg, config_path=settings_path,
                               learning_dir_override=tmp_path / "learning_data", **kwargs)
    service = ApplicationService(controller, config_path=settings_path,
                                 manual_fields_path=tmp_path / "config" / "manual_fields.json")
    return service, dto


def _scan(service, config):
    result = service.scan_reports(config)
    assert result["scan"]["scanned"] is True
    assert result["reports"]
    return result


def _report_for(result, name):
    return next(report for report in result["reports"] if report["fileName"] == name)


def test_default_service_persistent_state_stays_inside_disposable_pytest_root(tmp_path):
    from app import runtime_paths

    test_root = Path(os.environ[runtime_paths.TEST_RUNTIME_ROOT_ENV]).resolve()
    service = ApplicationService()
    assert runtime_paths.portable_root() == test_root
    assert service._manual_fields_path == test_root / "config" / "manual_fields.json"
    assert runtime_paths.logs_dir(create=False) == test_root / "logs"
    assert runtime_paths.output_dir(create=False) == test_root / "Output"
    assert runtime_paths.learning_dir(create=False) == test_root / "learning_data"

    report_folder = tmp_path / "reports"
    output_file = tmp_path / "output" / "result.xlsx"
    service.controller.report_folder = str(report_folder)
    service.controller.template = str(tmp_path / "template.xlsx")
    service.controller.output = str(output_file)
    service.controller.save_settings()
    assert runtime_paths.config_file() == test_root / "config" / "config.json"
    assert AppConfig.load().last_report_folder == str(report_folder)
    assert AppConfig.load().last_output_file == str(output_file)

    report_key = str((report_folder / "260901001-VOC_report.pptx").resolve())
    service._manual_fields = {report_key: {"vendor": "Disposable Vendor"}}
    service._pending_manual = {report_key: {"vendor": "Disposable Vendor"}}
    service._save_manual_fields()
    persisted = json.loads(service._manual_fields_path.read_text(encoding="utf-8"))
    assert persisted["pending"] == [report_key]
    assert service._manual_fields_path.is_relative_to(test_root)
    reloaded = ApplicationService()
    assert reloaded._pending_manual == {report_key: {"vendor": "Disposable Vendor"}}


def test_desktop_bundle_url_is_relative_and_rooted_only_at_dist():
    app_root = Path(__file__).resolve().parents[1] / "app"
    assert FRONTEND_URL == "../frontend/dist/index.html"
    assert not FRONTEND_URL.startswith(("/", "file:", "http:"))
    assert (app_root / FRONTEND_URL).resolve() == FRONTEND_INDEX
    assert FRONTEND_INDEX.parent == FRONTEND_DIST


def test_frontend_favicon_is_a_local_relative_vite_asset():
    frontend = Path(__file__).resolve().parents[1] / "frontend"
    html = (frontend / "index.html").read_text(encoding="utf-8")
    match = re.search(r'<link rel="icon" type="image/svg\+xml" href="([^"]+)"', html)
    assert match is not None
    href = match.group(1)
    assert not href.startswith(("/", "//", "http:", "https:"))
    icon = (frontend / href).resolve()
    assert icon.is_file() and icon.parent.is_relative_to((frontend / "src" / "assets").resolve())
    assert "<svg" in icon.read_text(encoding="utf-8")
    vite = (frontend / "vite.config.ts").read_text(encoding="utf-8")
    assert "base: './'" in vite and "publicDir: false" in vite


def test_bridge_required_contract_scan_details_and_status_keys(sample_tree, tmp_path):
    service, config = _make_service(sample_tree, tmp_path)
    bridge = BridgeService(service)

    for method in ("get_app_version", "choose_pptx_file", "choose_report_folder", "get_current_config"):
        assert callable(getattr(bridge, method))
    assert bridge.get_app_version()["ok"] is True
    assert bridge.get_current_config()["data"]["app"]["runtime"] == "pywebview-development"

    state = _scan(service, config)
    assert len(state["reports"]) == 4
    assert all(row["path"] == "" for row in state["reports"])  # source paths are intentionally not shown per row
    source_path = str(sample_tree["files"][0].resolve())
    service.controller.log_lines.append(f"Synthetic source diagnostic: {source_path}")
    redacted_state = service.dashboard_state()
    assert source_path not in json.dumps(redacted_state, ensure_ascii=False)
    assert sample_tree["files"][0].name in redacted_state["logs"][-1]["text"]
    first = state["reports"][0]
    details = service.report_details(first["id"])
    assert details["id"] == first["id"]
    assert isinstance(details["causes"], list) and isinstance(details["afterImages"], list)
    assert details["slides"] == 0  # not processed yet; scanning never opens the PPTX

    types_source = (Path(__file__).resolve().parents[1] / "frontend" / "src" / "types.ts").read_text(encoding="utf-8")
    status_type = types_source.split("export type StatusKey =", 1)[1].split("export const STATUS_LABEL", 1)[0]
    actual_keys = set(re.findall(r"'([^']+)'", status_type))
    label_block = types_source.split("export const STATUS_LABEL:", 1)[1].split("export const BUCKET_OF", 1)[0]
    label_keys = set(re.findall(r"^\s+([a-z_]+):", label_block, re.MULTILINE))
    backend_values = {
        ApplicationService._status_key(ACTION_PROCESS, False, "waiting"),
        ApplicationService._status_key(ACTION_PROCESS_NEW_ROW, False, "waiting"),
        ApplicationService._status_key(ACTION_MASTER_COMPLETE, False, "waiting"),
        ApplicationService._status_key(ACTION_OUTSIDE_PERIOD, False, "waiting"),
        ApplicationService._status_key(ACTION_SOURCE_DUPLICATE, False, "waiting"),
        ApplicationService._status_key(ACTION_FAST_SKIP, False, "waiting"),
        ApplicationService._status_key(ACTION_INVALID_MGMT, False, "waiting"),
        ApplicationService._status_key(ACTION_PROCESS, False, "reading"),
        ApplicationService._status_key(ACTION_PROCESS, True, "waiting"),
        *(ApplicationService._status_key(ACTION_PROCESS, False, final)
          for final in ("completed", "completed_new", "needs_review", "not_written", "error", "skipped")),
    }
    assert actual_keys <= label_keys
    assert backend_values <= actual_keys
    assert {row["status"] for row in state["reports"]} <= actual_keys

    missing = bridge.get_report_details("999")
    assert missing["ok"] is False and missing["error"]["code"] == "REPORT_NOT_FOUND"


class _GatedProcessor:
    """Tiny asynchronous worker fixture: can stop only between files, as the real BatchProcessor promises."""

    instances = []

    def __init__(self, options, on_prescan=None, on_file=None, on_batch=None, on_done=None, on_log=None):
        self.options = options
        self.on_prescan = on_prescan
        self.on_file = on_file
        self.on_batch = on_batch
        self.on_done = on_done
        self.on_log = on_log
        self.results = []
        self.summary = None
        self._thread = None
        self._stop = threading.Event()
        self.first_file_started = threading.Event()
        self.finish_current_file = threading.Event()
        self.processed = []
        self.__class__.instances.append(self)

    def start(self):
        self._thread = threading.Thread(target=self.run, name="fake-service-worker", daemon=True)
        self._thread.start()

    def run(self):
        total = len(self.options.files)
        if self.on_log:
            self.on_log("FAKE_WORKER_STARTED")
        for index, path in enumerate(self.options.files):
            if self.on_file:
                self.on_file(index, "reading", "")
                self.on_file(index, "extracting", "")
            if index == 0:
                self.first_file_started.set()
                self.finish_current_file.wait(timeout=5)
            self.processed.append(path)
            if self.on_file:
                self.on_file(index, "completed", "")
            if self.on_batch:
                self.on_batch(len(self.processed), total)
            if self._stop.is_set():
                break
        self.summary = BatchSummary(total=total, completed=len(self.processed), stopped=self._stop.is_set())
        if self.on_done:
            self.on_done(self.summary)

    def request_stop(self):
        self._stop.set()

    def is_running(self):
        return bool(self._thread and self._thread.is_alive())


def test_processing_worker_progress_logs_and_stop_after_current_file(sample_tree, tmp_path):
    _GatedProcessor.instances.clear()
    service, config = _make_service(sample_tree, tmp_path, processor_factory=_GatedProcessor)
    _scan(service, config)

    started = service.start_processing(config, use_ollama=False)
    worker = _GatedProcessor.instances[-1]
    assert started["job"]["status"] in ("processing", "stopping")
    assert worker.first_file_started.wait(timeout=3)

    stopping = service.stop_after_current()
    assert stopping["job"]["status"] == "stopping"
    assert any("dừng sau khi" in entry["text"].lower() for entry in stopping["logs"])
    worker.finish_current_file.set()
    worker._thread.join(timeout=5)
    assert not worker._thread.is_alive()

    final = service.dashboard_state()
    assert len(worker.processed) == 1  # no unsafe thread kill and no second file after stop was requested
    assert final["job"]["status"] == "done"
    assert final["job"]["stopped"] is True
    assert final["job"]["doneCount"] == 1 and final["job"]["percent"] == pytest.approx(25.0)
    assert final["job"]["error"] == ""
    assert any(entry["text"] == "FAKE_WORKER_STARTED" for entry in final["logs"])
    assert service.controller.request_stop() is False  # final state is idle; stop is no longer accepted


def test_real_application_service_runs_production_batch_and_applies_manual_fields(sample_tree, tmp_path):
    service, config = _make_service(sample_tree, tmp_path)
    token = service.register_pptx_selection(str(sample_tree["files"][0]))["token"]
    scanned = service.scan_reports(config, token)
    report = _report_for(scanned, sample_tree["files"][0].name)

    manual = service.save_manual_fields(report["id"], {"vendor": "Manual Vendor LLC",
                                                        "occurrence_date": "05/10/2026"})
    assert manual["status"] == "pending"  # saved before an output row exists; applied on real processing

    service.start_processing(config, use_ollama=False)
    processor = service.controller.processor
    processor._thread.join(timeout=60)
    assert not processor._thread.is_alive()
    final = service.dashboard_state()
    assert final["job"]["status"] == "done" and final["job"]["doneCount"] == 1
    assert final["job"]["percent"] == pytest.approx(100.0)
    assert service.controller.summary is not None and service.controller.summary.failed == 0
    assert len(processor.results) == 1
    result = processor.results[0]
    assert result.slide_count == 8 and result.status in {"completed", "needs_review"}
    assert result.vendor == "Manual Vendor LLC" and result.occurrence_date == "05/10/2026"

    output = Path(config["paths"]["output"])
    assert output.is_file()
    history_path = output.parent / "logs" / "history.json"
    assert history_path.is_file() and history_path.is_relative_to(tmp_path)
    workbook = load_workbook(output, data_only=True)
    try:
        sheet = workbook["Kiểm chứng"]
        matching_rows = [row for row in range(4, sheet.max_row + 1)
                         if sheet.cell(row=row, column=2).value == result.management_number]
        assert len(matching_rows) == 1
        row = matching_rows[0]
        assert sheet.cell(row=row, column=3).value == "Manual Vendor LLC"
        assert sheet.cell(row=row, column=4).value.strftime("%d/%m/%Y") == "05/10/2026"
    finally:
        workbook.close()

    details = service.report_details(report["id"])
    assert details["slides"] == 8 and details["causes"] and details["countermeasures"]
    final_report = _report_for(final, report["fileName"])
    assert final_report["path"] == "" and final_report["vendor"] == "Manual Vendor LLC"


def test_configuration_ollama_update_state_and_disabled_install(sample_tree, tmp_path, monkeypatch):
    service, config = _make_service(sample_tree, tmp_path)
    bridge = BridgeService(service)
    calls = []
    monkeypatch.setattr(service.controller, "check_ollama_async", lambda **kw: calls.append(("ollama", kw)))
    monkeypatch.setattr(service.controller, "refresh_models_async", lambda: calls.append(("models", {})))
    monkeypatch.setattr(service.controller, "check_update_async", lambda **kw: calls.append(("update", kw)))

    config["ollama"] = {"host": "http://127.0.0.1:11434", "port": "11434", "model": "qwen3:4b"}
    config["update"] = {"path": str(tmp_path / "lan-updates"), "autoCheck": True}
    saved = bridge.save_configuration(config)
    assert saved["ok"] is True
    current = saved["data"]
    assert current["ollama"]["host"] == "127.0.0.1" and current["ollama"]["port"] == "11434"
    assert current["ollama"]["model"] == "qwen3:4b" and current["update"]["autoCheck"] is True

    checked = bridge.check_ollama_connection(config)
    refreshed = bridge.refresh_ollama_models(config)
    update = bridge.check_update(config)
    assert checked["ok"] and refreshed["ok"] and update["ok"]
    assert [name for name, _ in calls] == ["ollama", "models", "update"]
    assert calls[0][1] == {"local_first": True}
    assert calls[-1][1] == {"startup": False}

    install = bridge.install_update()
    assert install["ok"] is True and install["data"]["status"] == "unavailable"
    assert service.current_config()["app"]["version"] == "1.3.5"


def test_bridge_native_selection_tokens_dto_validation_and_error_sanitizing(sample_tree, tmp_path, monkeypatch):
    service, config = _make_service(sample_tree, tmp_path)
    bridge = BridgeService(service)

    class DialogWindow:
        def __init__(self, choices):
            self.choices = list(choices)
            self.calls = []

        def create_file_dialog(self, kind, **kwargs):
            self.calls.append((kind, kwargs))
            return self.choices.pop(0)

    selected_file = sample_tree["files"][0]
    window = DialogWindow([[str(selected_file)], [str(sample_tree["reports"])]])
    bridge.bind_window(window)
    monkeypatch.setattr(bridge, "_webview_constant", lambda name: name)

    selected = bridge.choose_pptx_file()
    assert selected["ok"] and selected["data"]["name"] == selected_file.name
    assert "token" in selected["data"] and "path" not in selected["data"]
    folder = bridge.choose_report_folder()
    assert folder["ok"] and Path(folder["data"]["path"]) == sample_tree["reports"].resolve()
    assert window.calls == [
        ("OPEN_DIALOG", {"file_types": ("PowerPoint presentation (*.pptx)",)}),
        ("FOLDER_DIALOG", {}),
    ]

    scanned = bridge.scan_reports(config, selected["data"]["token"])
    assert scanned["ok"] and len(scanned["data"]["reports"]) == 1
    assert scanned["data"]["reports"][0]["path"] == ""
    service.clear_pptx_selection()
    expired = bridge.scan_reports(config, selected["data"]["token"])
    assert expired["ok"] is False and expired["error"]["code"] == "INVALID_SELECTION"

    invalid_dto = bridge.save_configuration(["not", "an", "object"])
    assert invalid_dto["ok"] is False and invalid_dto["error"]["code"] == "INVALID_DTO"
    invalid_boolean = bridge.save_configuration({**config, "forceReprocess": "false"})
    assert invalid_boolean["ok"] is False and invalid_boolean["error"]["code"] == "INVALID_DTO"
    invalid_flag = bridge.start_processing(config, use_ollama="false")
    assert invalid_flag["ok"] is False and invalid_flag["error"]["code"] == "INVALID_DTO"
    invalid_port = bridge.use_discovered_server("127.0.0.1", "11434")
    assert invalid_port["ok"] is False and invalid_port["error"]["code"] == "INVALID_DTO"

    def explode():
        raise RuntimeError(r"C:\private\reports\secret.pptx")
    monkeypatch.setattr(service, "app_version", explode)
    internal = bridge.get_app_version()
    assert internal["ok"] is False and internal["error"]["code"] == "INTERNAL_ERROR"
    assert "secret.pptx" not in internal["error"]["message"]

    def source_error():
        raise ServiceError(f"Cannot read {selected_file.resolve()}", "READ_FAILED")
    monkeypatch.setattr(service, "current_config", source_error)
    safe_error = bridge.get_current_config()
    assert safe_error["ok"] is False and str(selected_file.resolve()) not in safe_error["error"]["message"]
    assert selected_file.name in safe_error["error"]["message"]
    assert not hasattr(bridge, "open_path")  # fixed backend-owned actions only; no arbitrary path endpoint


def test_xlsx_and_learning_export_dialog_filters_paths_and_save_validation(sample_tree, tmp_path, monkeypatch):
    service, _ = _make_service(sample_tree, tmp_path)
    bridge = BridgeService(service)
    folder = tmp_path / "native selections"
    nested = folder / "nested"
    nested.mkdir(parents=True)
    template = folder / "Verification.xlsx"
    template.write_bytes(b"xlsx")
    output = folder / "Kiem_chung.xlsx"
    export = folder / "learning_data_export.zip"
    direct_export = folder / "learning_data_export_direct.zip"
    raw_template = nested / ".." / template.name
    raw_output = nested / ".." / output.name
    raw_export = nested / ".." / export.name
    raw_direct_export = nested / ".." / direct_export.name

    class DialogWindow:
        def __init__(self, choices):
            self.choices = list(choices)
            self.calls = []

        def create_file_dialog(self, kind, **kwargs):
            self.calls.append((kind, kwargs))
            return self.choices.pop(0)

    window = DialogWindow([[str(raw_template)], [str(raw_output)], [str(raw_export)], [str(raw_direct_export)]])
    bridge.bind_window(window)
    monkeypatch.setattr(bridge, "_webview_constant", lambda name: name)

    template_result = bridge.choose_template_file()
    output_result = bridge.choose_output_file()
    export_result = bridge.choose_learning_export()
    exported_paths = []
    monkeypatch.setattr(service, "export_learning_data",
                        lambda path: exported_paths.append(path) or {"saved": path})
    direct_export_result = bridge.export_learning_data()

    assert template_result == {"ok": True, "data": {"cancelled": False, "path": str(template.resolve()), "kind": "template"}}
    assert output_result == {"ok": True, "data": {"cancelled": False, "path": str(output.resolve())}}
    assert export_result == {"ok": True, "data": {"cancelled": False, "path": str(export.resolve())}}
    assert direct_export_result == {"ok": True, "data": {"saved": str(direct_export.resolve())}}
    assert exported_paths == [str(direct_export.resolve())]
    assert window.calls == [
        ("OPEN_DIALOG", {"file_types": ("Excel workbook (*.xlsx)",)}),
        ("SAVE_DIALOG", {"save_filename": "Kiem_chung_Ket_qua.xlsx", "file_types": ("Excel workbook (*.xlsx)",)}),
        ("SAVE_DIALOG", {"save_filename": "learning_data_export.zip", "file_types": ("ZIP archive (*.zip)",)}),
        ("SAVE_DIALOG", {"save_filename": "learning_data_export.zip", "file_types": ("ZIP archive (*.zip)",)}),
    ]

    invalid_template = folder / "wrong.xlsm"
    invalid_template.write_bytes(b"not-xlsx")
    invalid_output = folder / "wrong.xlsm"
    directory_with_xlsx_suffix = folder / "not-a-file.xlsx"
    directory_with_xlsx_suffix.mkdir()
    for method, chosen in (("choose_template_file", invalid_template),
                           ("choose_template_file", folder / "missing.xlsx"),
                           ("choose_output_file", invalid_output),
                           ("choose_output_file", directory_with_xlsx_suffix)):
        window.choices = [[str(chosen)]]
        result = getattr(bridge, method)()
        assert result["ok"] is False and result["error"]["code"] == "INVALID_SELECTION"


@pytest.mark.parametrize("method,expected_call", [
    ("choose_pptx_file", ("OPEN_DIALOG", {"file_types": ("PowerPoint presentation (*.pptx)",)})),
    ("choose_template_file", ("OPEN_DIALOG", {"file_types": ("Excel workbook (*.xlsx)",)})),
    ("choose_output_file", ("SAVE_DIALOG", {"save_filename": "Kiem_chung_Ket_qua.xlsx",
                                               "file_types": ("Excel workbook (*.xlsx)",)})),
    ("choose_learning_export", ("SAVE_DIALOG", {"save_filename": "learning_data_export.zip",
                                                  "file_types": ("ZIP archive (*.zip)",)})),
    ("export_learning_data", ("SAVE_DIALOG", {"save_filename": "learning_data_export.zip",
                                                "file_types": ("ZIP archive (*.zip)",)})),
    ("choose_report_folder", ("FOLDER_DIALOG", {})),
    ("choose_update_folder", ("FOLDER_DIALOG", {})),
])
def test_native_picker_cancellation_is_normal_and_filters_are_single_label(sample_tree, tmp_path, monkeypatch,
                                                                            method, expected_call):
    service, _ = _make_service(sample_tree, tmp_path)
    bridge = BridgeService(service)

    class CancelWindow:
        def __init__(self):
            self.calls = []

        def create_file_dialog(self, kind, **kwargs):
            self.calls.append((kind, kwargs))
            return None

    window = CancelWindow()
    bridge.bind_window(window)
    monkeypatch.setattr(bridge, "_webview_constant", lambda name: name)
    result = getattr(bridge, method)()
    assert result == {"ok": True, "data": {"cancelled": True}}
    assert window.calls == [expected_call]


def test_manual_vendor_date_validation_lock_retry_and_persistence(sample_tree, tmp_path, monkeypatch):
    service, config = _make_service(sample_tree, tmp_path)
    output = Path(config["paths"]["output"])
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(sample_tree["template"], output)
    state = _scan(service, config)
    report = _report_for(state, sample_tree["files"][0].name)
    locked = {"value": True}
    written = []

    class FakeExcelWriter:
        columns = {"vendor": 3, "occurrence_date": 4}

        def __init__(self, template, destination):
            self.destination = Path(destination)

        def find_rows_by_management_number(self, management_number):
            assert management_number
            return [4]

        def apply_manual_fields(self, row, fields):
            written.append((row, dict(fields)))

        def save(self):
            if locked["value"]:
                raise ExcelLockedError(self.destination, OSError("sharing violation"))

        def close(self):
            pass

    monkeypatch.setattr("app.application_service.ExcelWriter", FakeExcelWriter)
    result = service.save_manual_fields(report["id"], {"vendor": "Vendor A\nVendor B",
                                                        "occurrence_date": "05/10/2026"})
    assert result["status"] == "locked"
    assert "Excel đang bị khóa" in result["message"]
    row = _report_for(result["state"], report["fileName"])
    assert row["vendor"] == "Vendor A\nVendor B" and row["occurrenceDate"] == "05/10/2026"
    assert row["manualPending"] is True
    payload = json.loads((tmp_path / "config" / "manual_fields.json").read_text(encoding="utf-8"))
    saved_override = next(values for values in payload["overrides"].values() if values)
    assert saved_override["occurrence_date"] == "2026-10-05"

    with pytest.raises(ServiceError) as err:
        service.save_manual_fields(report["id"], {"occurrence_date": "31/02/2026"})
    assert err.value.code == "INVALID_DATE"
    with pytest.raises(ServiceError) as err:
        service.save_manual_fields(report["id"], {"path": r"C:\private\anything"})
    assert err.value.code == "INVALID_MANUAL_FIELDS"

    locked["value"] = False
    retried = service.retry_manual_fields(report["id"])
    assert retried["status"] == "applied"
    assert written == [(4, {"vendor": "Vendor A\nVendor B", "occurrence_date": "2026-10-05"})] * 2
    updated = _report_for(retried["state"], report["fileName"])
    assert updated["manualPending"] is False and updated["occurrenceDate"] == "05/10/2026"
    persisted = json.loads((tmp_path / "config" / "manual_fields.json").read_text(encoding="utf-8"))
    assert persisted["pending"] == []


def _learning_candidates(sample_tree):
    path = Path(sample_tree["files"][0])
    image = ImageCandidate(
        candidate_id="260918080-VOC|S5|#21|1", management_number="260918080-VOC",
        source_name=path.name, source_file=str(path), slide=5, picture_id=21, order=1,
        bounds=(100, 100, 300, 200), slide_size=(1000, 800), block_bounds=(0, 0, 1000, 800),
        features={}, evidence=["test candidate"], confidence=0.8,
    )
    content = ContentCandidate(
        candidate_id="260918080-VOC|S5|SH22|1", management_number="260918080-VOC",
        source_name=path.name, source_file=str(path), slide=5, shape_id=22, order=1, shape_kind="text",
        bounds=(100, 100, 400, 200), slide_size=(1000, 800), text="Cải tiến trong sản xuất",
        features={}, evidence=["test content candidate"], confidence=0.7,
    )
    return image, content


def _png_bytes():
    image = Image.new("RGB", (48, 32), "#20a080")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def test_learning_candidates_counts_label_review_saved_reapply_retry_and_model_update(
        sample_tree, tmp_path, monkeypatch):
    service, config = _make_service(sample_tree, tmp_path)
    image, content = _learning_candidates(sample_tree)
    service.controller.processor = SimpleNamespace(results=[SimpleNamespace(image_candidates=[image],
                                                                            content_candidates=[content])])
    monkeypatch.setattr(service.controller, "candidate_blob", lambda candidate: _png_bytes(), raising=False)
    output = Path(config["paths"]["output"])
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(sample_tree["template"], output)
    service.controller.template = str(sample_tree["template"])
    service.controller.output = str(output)

    from app.image_review import ReapplyResult
    import app.image_review as image_review

    image_results = [
        ReapplyResult(prepared_rows=[4], errors=["Excel locked"], locked_path=str(output)),
        ReapplyResult(prepared_rows=[4], updated_rows=[4], pictures=1),
    ]
    content_results = [
        ReapplyResult(prepared_rows=[4], errors=["Excel locked"], locked_path=str(output)),
        ReapplyResult(prepared_rows=[4], updated_rows=[4]),
    ]
    reapply_calls = {"image": 0, "content": 0}

    def reapply_images(*args, **kwargs):
        reapply_calls["image"] += 1
        return image_results.pop(0)

    def reapply_content(*args, **kwargs):
        reapply_calls["content"] += 1
        return content_results.pop(0)

    monkeypatch.setattr(image_review, "reapply_labels", reapply_images)
    monkeypatch.setattr(image_review, "reapply_content_labels", reapply_content)

    initial = service.learning_state()
    assert initial["counts"] == {"image": {"total": 1, "labeled": 0},
                                 "content": {"total": 1, "labeled": 0}}
    assert initial["images"][0]["src"].startswith("data:image/jpeg;base64,")
    assert initial["contents"][0]["text"] == content.text

    labeled_image = service.set_learning_label("image", image.candidate_id, "AFTER", "verified image")
    labeled_content = service.set_learning_label("content", content.candidate_id, "IMPROVEMENT_CONTENT",
                                                 "verified text")
    assert labeled_image["counts"]["image"] == {"total": 1, "labeled": 1}
    assert labeled_content["counts"]["content"] == {"total": 1, "labeled": 1}
    assert labeled_image["images"][0]["labelPending"] is True
    assert labeled_content["contents"][0]["labelPending"] is True

    saved_image = service.save_learning_labels("image", {image.candidate_id: "image note saved"})
    saved_content = service.save_learning_labels("content", {content.candidate_id: "content note saved"})
    assert saved_image["state"]["excelLastResult"]["kind"] == "locked"
    assert saved_content["state"]["excelLastResult"]["kind"] == "locked"
    assert saved_image["state"]["excelPending"] == 1
    assert saved_content["state"]["excelPending"] == 2
    assert saved_image["state"]["images"][0]["labelPending"] is False
    labels = service.controller.learning.store.latest()
    content_labels = service.controller.learning.content.store.latest()
    assert labels[image.candidate_id]["label"] == "AFTER"
    assert labels[image.candidate_id]["note"] == "image note saved"
    assert content_labels[content.candidate_id]["label"] == "IMPROVEMENT_CONTENT"
    assert content_labels[content.candidate_id]["note"] == "content note saved"

    retried = service.retry_learning_excel()
    assert retried["state"]["excelLastResult"]["kind"] == "ok"
    assert retried["state"]["excelPending"] == 0
    assert reapply_calls == {"image": 2, "content": 2}
    assert len(service.controller.learning.store.latest()) == 1  # retry changes Excel only, never duplicates labels
    assert len(service.controller.learning.content.store.latest()) == 1

    exported = tmp_path / "exports" / "learning_data_export.zip"
    export_result = service.export_learning_data(str(exported))
    assert export_result["imageLabelCount"] == 1 and export_result["contentLabelCount"] == 1
    with zipfile.ZipFile(exported) as archive:
        names = set(archive.namelist())
        assert {"manifest.json", "image_labels.jsonl", "content_labels.jsonl"} <= names
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest["format"] == "tnp-learning-export-v1"
        assert json.loads(archive.read("image_labels.jsonl"))["label"] == "AFTER"
        assert json.loads(archive.read("content_labels.jsonl"))["label"] == "IMPROVEMENT_CONTENT"
        assert str(sample_tree["files"][0]) not in archive.read("image_labels.jsonl").decode("utf-8")

    entered, release = threading.Event(), threading.Event()

    def slow_train_models():
        entered.set()
        assert release.wait(timeout=5)
        return True, "Test model update completed."

    monkeypatch.setattr(service.controller, "train_models", slow_train_models, raising=False)
    training = service.train_models()
    assert entered.wait(timeout=3)
    assert training["training"] is True and service.learning_state()["training"] is True
    release.set()
    service._training_thread.join(timeout=5)
    final_learning = service.learning_state()
    assert final_learning["training"] is False
    assert final_learning["trainingMessage"] == "Test model update completed."


def test_fixed_output_log_and_learning_actions_do_not_accept_paths_from_ui(sample_tree, tmp_path, monkeypatch):
    service, config = _make_service(sample_tree, tmp_path)
    output = Path(config["paths"]["output"])
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(b"xlsx placeholder")
    logs = output.parent / "logs"
    logs.mkdir()
    (logs / "app.log").write_text("log", encoding="utf-8")
    service.controller.output = str(output)

    opened = []
    monkeypatch.setattr(service, "_open_local_path",
                        lambda path, code, message: opened.append((Path(path), code)) or {"message": "opened"})
    assert service.open_output_file()["message"] == "opened"
    assert service.open_output_folder()["message"] == "opened"
    assert service.open_log_file()["message"] == "opened"
    assert service.open_log_folder()["message"] == "opened"
    learning_dir = service.controller.learning_folder()
    learning_dir.mkdir(parents=True, exist_ok=True)
    assert service.open_learning_folder()["message"] == "opened"
    assert [code for _, code in opened] == ["OUTPUT_NOT_FOUND", "OUTPUT_NOT_FOUND", "LOG_NOT_FOUND",
                                           "LOG_NOT_FOUND", "LEARNING_NOT_FOUND"]
    assert opened[0][0] == output and opened[1][0] == output.parent


def test_status_maps_all_backend_final_statuses():
    for key in ("completed", "completed_new", "needs_review", "not_written", "error", "skipped",
                "outside_period", "source_duplicate", "fast_skip"):
        value = ApplicationService._status_key(ACTION_PROCESS, False, key)
        assert value in {"completed", "needs_review", "error", "skipped", "outside_period",
                         "source_duplicate", "fast_skip"}
    assert ApplicationService._status_key(ACTION_PROCESS, False, "reading") == "processing"
    assert ApplicationService._status_key(ACTION_PROCESS, True, "waiting") == "excluded"


def test_no_scan_and_invalid_file_selection_return_stable_bridge_errors(sample_tree, tmp_path):
    service, config = _make_service(sample_tree, tmp_path)
    bridge = BridgeService(service)
    no_scan = bridge.get_report_details("0")
    assert no_scan["ok"] is False and no_scan["error"]["code"] == "REPORT_NOT_FOUND"
    with pytest.raises(ServiceError) as err:
        service.register_pptx_selection(str(sample_tree["template"]))
    assert err.value.code == "INVALID_SELECTION"
    with pytest.raises(ServiceError) as err:
        service.export_learning_data(str(tmp_path / "learning.json"))
    assert err.value.code == "INVALID_EXPORT_PATH"
