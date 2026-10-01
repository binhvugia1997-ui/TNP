"""Regression tests for requirement #16 – improvement / QPN images must fit the destination width."""
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

from app.excel_writer import (EMU_PER_PX, IMAGE_MARGIN_PX, MAX_IMAGE_HEIGHT_PT, MAX_ROW_HEIGHT_PT, PX_PER_PT,
                              ExcelWriter, col_width_to_px, plan_image_grid)
from app.extractor import ExtractedRecord
from app.qpn_renderer import trim_white_margins

MGMT = "260918080-VOC"
IMG_COL, QPN_COL = 11, 8


def _png(path: Path, w: int, h: int, color=(200, 30, 30)) -> Path:
    Image.new("RGB", (w, h), color).save(path)
    return path


def _record(**kw):
    rec = ExtractedRecord(management_number=MGMT, model="A185", item="Rear", defect_content="Xước",
                          root_cause="NGUYÊN NHÂN\n- a", improvement="CẢI TIẾN\n- b", qpn_slide=2)
    for k, v in kw.items():
        setattr(rec, k, v)
    return rec


def _prepare(template: Path, row=4, img_width_chars=None):
    wb = load_workbook(template)
    ws = wb["Kiểm chứng"]
    ws.cell(row=row, column=2, value=MGMT)
    if img_width_chars is not None:
        ws.column_dimensions[get_column_letter(IMG_COL)].width = img_width_chars
    wb.save(template)


def _geom(img):
    """(x_px, y_px, w_px, h_px, col0, row0) of an image anchored with offsets."""
    a = img.anchor          # displayed size lives in the anchor ext (img.width is the pixel size after reload)
    return (a._from.colOff // EMU_PER_PX, a._from.rowOff // EMU_PER_PX, a.ext.cx // EMU_PER_PX,
            a.ext.cy // EMU_PER_PX, a._from.col, a._from.row)


def _images_in(ws, row, col):
    return [i for i in ws._images if i.anchor._from.row == row - 1 and i.anchor._from.col == col - 1]


# ---------------------------------------------------------------- pure layout engine
def test_col_width_to_px_follows_excel_formula():
    assert col_width_to_px(8.43) == round(8.43 * 7 + 5)
    assert col_width_to_px(40) > col_width_to_px(20)


def test_narrow_source_image_is_widened_to_slot_width():
    plan = plan_image_grid([(120, 90)], area_px=300)
    (x, y, w, h), = plan["items"]
    assert plan["columns"] == 1
    assert w == 300 - 2 * IMAGE_MARGIN_PX                      # fills usable width, not kept tiny
    assert abs(w / h - 120 / 90) < 0.03                        # aspect preserved
    assert x == IMAGE_MARGIN_PX and y == IMAGE_MARGIN_PX


def test_tall_source_image_is_capped_at_max_height_without_distortion():
    plan = plan_image_grid([(300, 2400)], area_px=400)
    (x, y, w, h), = plan["items"]
    assert h <= MAX_IMAGE_HEIGHT_PT * PX_PER_PT + 1
    assert abs(w / h - 300 / 2400) < 0.01                      # aspect preserved
    assert x >= IMAGE_MARGIN_PX and x + w <= 400 - IMAGE_MARGIN_PX   # centred inside the slot
    assert plan["height_pt"] <= MAX_ROW_HEIGHT_PT


def test_wide_source_image_is_shrunk_to_slot_width():
    plan = plan_image_grid([(4000, 1000)], area_px=350)
    (x, y, w, h), = plan["items"]
    assert w == 350 - 2 * IMAGE_MARGIN_PX and abs(w / h - 4.0) < 0.05


def test_five_images_independent_no_overlap_within_area_and_row_cap():
    sizes = [(640, 480), (200, 300), (1200, 400), (800, 800), (500, 350)]
    plan = plan_image_grid(sizes, area_px=360)
    items = plan["items"]
    assert len(items) == 5
    for (w0, h0), (x, y, w, h) in zip(sizes, items):
        assert abs((w / h) - (w0 / h0)) < 0.05                 # never distorted
        assert x >= IMAGE_MARGIN_PX and x + w <= 360 - IMAGE_MARGIN_PX
        assert h >= 40                                         # readable, no thumbnails
    for i in range(5):
        for j in range(i + 1, 5):
            xi, yi, wi, hi = items[i]
            xj, yj, wj, hj = items[j]
            assert xi + wi <= xj or xj + wj <= xi or yi + hi <= yj or yj + hj <= yi, "images overlap"
    assert plan["height_pt"] <= MAX_ROW_HEIGHT_PT
    # source order kept: reading order top->bottom, left->right matches input order
    order = sorted(range(5), key=lambda k: (items[k][1], items[k][0]))
    assert order == [0, 1, 2, 3, 4]


def test_two_small_images_prefer_single_column_full_width():
    plan = plan_image_grid([(400, 200), (400, 200)], area_px=300)
    assert plan["columns"] == 1
    assert all(w == 300 - 2 * IMAGE_MARGIN_PX for (_, _, w, _) in plan["items"])
    assert plan["items"][1][1] > plan["items"][0][1] + plan["items"][0][3]      # stacked below with gap


def test_layout_is_deterministic():
    sizes = [(640, 480), (200, 300), (1200, 400)]
    assert plan_image_grid(sizes, 320) == plan_image_grid(sizes, 320)


# ---------------------------------------------------------------- workbook integration
def test_images_fit_destination_column_and_adapt_to_column_width(template, tmp_path):
    pics = [_png(tmp_path / f"p{i}.png", 240, 90) for i in range(2)]      # narrow but short -> always fits
    out_a = tmp_path / "a.xlsx"
    _prepare(template, img_width_chars=30)
    w = ExcelWriter(template, out_a)
    w.update_record(4, _record(), improvement_jpg=pics)
    w.save()
    wsa = load_workbook(out_a)["Kiểm chứng"]
    imgs_a = _images_in(wsa, 4, IMG_COL)
    assert len(imgs_a) == 2
    area_a = w.image_area_px(4, "improvement_image")
    assert area_a == col_width_to_px(30)
    for img in imgs_a:
        x, y, iw, ih, _, _ = _geom(img)
        assert iw == area_a - 2 * IMAGE_MARGIN_PX and x + iw <= area_a          # no spill into next column
        assert abs(iw / ih - 240 / 90) < 0.03
    assert wsa.row_dimensions[4].height >= imgs_a[-1].anchor._from.rowOff / EMU_PER_PX / PX_PER_PT

    # user widens the column -> displayed width follows (no hard-coded pixels)
    out_b = tmp_path / "b.xlsx"
    _prepare(template, img_width_chars=60)
    w2 = ExcelWriter(template, out_b)
    w2.update_record(4, _record(), improvement_jpg=pics)
    w2.save()
    wsb = load_workbook(out_b)["Kiểm chứng"]
    imgs_b = _images_in(wsb, 4, IMG_COL)
    assert _geom(imgs_b[0])[2] > _geom(imgs_a[0])[2]
    assert _geom(imgs_b[0])[2] == col_width_to_px(60) - 2 * IMAGE_MARGIN_PX


def test_row_height_raised_but_never_reduced_and_other_rows_untouched(template, tmp_path):
    _prepare(template)
    wb = load_workbook(template)
    ws = wb["Kiểm chứng"]
    ws.row_dimensions[4].height = 380
    ws.row_dimensions[5].height = 33
    wb.save(template)
    out = tmp_path / "o.xlsx"
    w = ExcelWriter(template, out)
    w.update_record(4, _record(), improvement_jpg=[_png(tmp_path / "s.png", 100, 50)])
    w.save()
    ws2 = load_workbook(out)["Kiểm chứng"]
    assert ws2.row_dimensions[4].height == 380
    assert ws2.row_dimensions[5].height == 33


def test_rerun_keeps_images_identical(template, tmp_path):
    from app.excel_writer import ExcelWriter as W
    _prepare(template)
    out = tmp_path / "o.xlsx"
    pics = [_png(tmp_path / f"p{i}.png", 150 + 40 * i, 120) for i in range(3)]
    w = W(template, out)
    w.update_record(4, _record(), improvement_jpg=pics)
    w.save()
    ws1 = load_workbook(out)["Kiểm chứng"]
    g1 = sorted(_geom(i) for i in _images_in(ws1, 4, IMG_COL))
    h1 = ws1.row_dimensions[4].height

    # second run on the already-complete row: fill_missing_fields must not add/move/resize anything
    w2 = W(template, out)
    assert "improvement_image" not in w2.missing_managed_fields(4)
    w2.fill_missing_fields(4, _record(), w2.missing_managed_fields(4), improvement_jpg=pics)
    w2.save()
    ws2 = load_workbook(out)["Kiểm chứng"]
    g2 = sorted(_geom(i) for i in _images_in(ws2, 4, IMG_COL))
    assert g2 == g1 and len(g2) == 3
    assert ws2.row_dimensions[4].height == h1

    # --force style full rewrite: still exactly 3 images with the same geometry (no progressive resize)
    w3 = W(template, out)
    w3.update_record(4, _record(), improvement_jpg=pics)
    w3.save()
    ws3 = load_workbook(out)["Kiểm chứng"]
    assert sorted(_geom(i) for i in _images_in(ws3, 4, IMG_COL)) == g1
    assert ws3.row_dimensions[4].height == h1


def test_qpn_fits_its_own_column_and_whitespace_is_trimmed(template, tmp_path):
    # a "slide" with content in the middle and wide white margins
    slide = Image.new("RGB", (1920, 1080), "white")
    for x in range(400, 1500):
        for y in range(200, 900, 1):
            slide.putpixel((x, y), (0, 0, 120))
    trimmed = trim_white_margins(slide)
    assert trimmed.width < 1920 and trimmed.height < 1080
    assert trimmed.width >= 1100 and trimmed.height >= 700                  # content never cropped
    qpn = tmp_path / "qpn.png"
    trimmed.save(qpn)

    _prepare(template)
    out = tmp_path / "o.xlsx"
    w = ExcelWriter(template, out)
    w.update_record(4, _record(), qpn_png=qpn, improvement_jpg=[_png(tmp_path / "i.png", 300, 300)])
    w.save()
    ws = load_workbook(out)["Kiểm chứng"]
    (q,) = _images_in(ws, 4, QPN_COL)
    (imp,) = _images_in(ws, 4, IMG_COL)
    qa = w.image_area_px(4, "qpn")
    _, _, qw, qh, qc, _ = _geom(q)
    _, _, iw, _, ic, _ = _geom(imp)
    assert qw == qa - 2 * IMAGE_MARGIN_PX and abs(qw / qh - trimmed.width / trimmed.height) < 0.03
    assert qc != ic                                                          # separate geometry
    assert iw == w.image_area_px(4, "improvement_image") - 2 * IMAGE_MARGIN_PX


# ---------------------------------------------------------------- explicit regression list (#16)
def test_horizontal_position_margin_and_centering():
    # full-width picture: starts exactly at the left margin and ends at the right margin
    (x, y, w, h), = plan_image_grid([(400, 100)], area_px=300)["items"]
    assert x == IMAGE_MARGIN_PX and x + w == 300 - IMAGE_MARGIN_PX
    # height-capped picture is narrower than the slot -> centred (left/right gap differ by <= 1 px)
    (x, y, w, h), = plan_image_grid([(100, 2000)], area_px=400)["items"]
    assert abs(x - (400 - x - w)) <= 1 and x > IMAGE_MARGIN_PX


def test_portrait_image_scales_to_width_and_increases_required_height():
    land = plan_image_grid([(400, 200)], area_px=300)
    port = plan_image_grid([(300, 400)], area_px=300)          # portrait, below the 300 pt cap
    assert port["items"][0][2] == land["items"][0][2] == 300 - 2 * IMAGE_MARGIN_PX   # both width-fitted
    assert port["items"][0][3] > land["items"][0][3]
    assert port["height_pt"] > land["height_pt"]
    assert abs(port["items"][0][2] / port["items"][0][3] - 300 / 400) < 0.02


def test_required_row_height_comes_from_final_resized_heights():
    sizes = [(400, 200), (400, 300)]
    plan = plan_image_grid(sizes, area_px=300)
    items = plan["items"]
    expected_px = IMAGE_MARGIN_PX + items[0][3] + 6 + items[1][3] + IMAGE_MARGIN_PX
    assert abs(plan["height_pt"] * PX_PER_PT - expected_px) <= 2
    assert plan["height_pt"] <= MAX_ROW_HEIGHT_PT


def test_row_height_in_workbook_equals_layout_height(template, tmp_path):
    _prepare(template)
    out = tmp_path / "o.xlsx"
    pics = [_png(tmp_path / "a.png", 400, 200), _png(tmp_path / "b.png", 400, 300)]
    w = ExcelWriter(template, out)
    plan = plan_image_grid([(400, 200), (400, 300)], w.image_area_px(4, "improvement_image"))
    w.update_record(4, _record(), improvement_jpg=pics)
    w.save()
    ws = load_workbook(out)["Kiểm chứng"]
    # rec text is short, so the image layout drives the row height (+2 pt breathing space)
    assert abs(ws.row_dimensions[4].height - (plan["height_pt"] + 2)) < 0.01


def test_no_hard_coded_screenshot_width_in_writer():
    src = Path("app/excel_writer.py").read_text(encoding="utf-8")
    assert "col_width_to_px(" in src
    # the legacy fixed pixel widths must not be used anywhere for picture sizing
    for token in ("= 1116", "1116,", "max(40, ", "_embed_image"):
        assert token not in src, token
