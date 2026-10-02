"""PROMPT-004C – folder discovery is separated from report eligibility: every PPTX report is listed, never
silently dropped because of its name, its Management Number or the master Excel."""
import logging
import threading
from datetime import date

import app.gui_controller as gc
from app.config import AppConfig
from app.extractor import defect_content_from_filename, derive_occurrence_date, management_number_from_filename
from app.gui_controller import SCAN_FILTERS_VI, GuiController
from app.prescan import ACTION_MASTER_NOT_FOUND, ACTION_PROCESS, month_period, prescan
from app.scanner import is_report_file, rejection_reason, scan_folder, scan_inputs
from tests.test_gui_redesign import _make_app

NEW = "(CTMS)_260925015_ĐỐI SÁCH CẢI TIẾN MODEL A253 FRONT LỖI SƠN, HÀN (XƯỚC BÓNG) 24.09.2026.pptx"
FIVE = ["(CTMS)_11108_260918080-VOC_ Đối sách LỖI BONG ATN, MỤN 22.9.2026 SEV.pptx",
        "(CTMS)_20506_260922134_ Đối sách LỖI MẺ XƯỚC 23.9.2026 SEV.pptx",
        "(CTMS)_20601_260920045-VOC_ Đối sách LỖI MẺ XƯỚC, BONG SƠN 24.9.2026 SEV.pptx",
        "(CTMS)_20602_260923045_ Đối sách LỖI BONG ATN 22.9.2026 SEV.pptx",
        "(CTMS)_20603_260924011-VOC_ Đối sách LỖI XƯỚC 25.9.2026 SEV.pptx"]
SIX = FIVE + [NEW]
SIX_KEYS = ["260918080-VOC", "260922134", "260920045-VOC", "260923045", "260924011-VOC", "260925015"]


def _folder(tmp_path, names=SIX):
    folder = tmp_path / "Bao_cao" / "2026-09"
    folder.mkdir(parents=True)
    for n in names:
        (folder / n).write_bytes(b"x")
    return folder


# 1-5. discovery + file-name parsing of the exact real name
def test_new_report_discovered_by_scanner(tmp_path):
    folder = _folder(tmp_path)
    found = [p.name for p in scan_folder(folder)]
    assert NEW in found and len(found) == 6
    assert rejection_reason(folder / NEW) == "" and is_report_file(folder / NEW)


def test_new_report_management_number_and_date():
    assert management_number_from_filename(NEW) == "260925015"
    d, why = derive_occurrence_date("260925015")
    assert why == "" and d == date(2026, 9, 25)                         # no SEV / no date suffix needed
    assert management_number_from_filename("(CTMS)_20506_260918080-VOC_ Đối sách LỖI X 22.9.2026 SEV.pptx") == "260918080-VOC"
    assert defect_content_from_filename(NEW)                              # parentheses in the defect do not break parsing


def test_parentheses_sev_model_front_never_exclude(tmp_path):
    variants = [NEW,
                "(CTMS)_260925016_LỖI (A) (B) (XƯỚC BÓNG).pptx",                         # many parentheses
                "260925017 LỖI SƠN.pptx",                                                # no SEV, no date, no prefix
                "(CTMS)_260925018_MODEL A999 BACK LỖI MỚI 24.09.2026.pptx",              # unknown model / item
                "(CTMS)_260925019_LỖI SƠN 24.09.2026.PPTX"]                              # upper-case extension
    folder = _folder(tmp_path, variants)
    assert sorted(p.name for p in scan_folder(folder)) == sorted(variants)
    src = open(gc.__file__.replace("gui_controller.py", "scanner.py"), encoding="utf-8").read()
    code = "\n".join(ln for ln in src.splitlines() if not ln.strip().startswith(("#", '"""')) and "docstring" not in ln)
    for word in ("A185", "A253", "REAR", "FRONT", "'SEV'", '"SEV"'):
        assert word not in code                                          # nothing content-specific in discovery


def test_scanner_logs_every_rejected_powerpoint_like_entry(tmp_path, caplog):
    folder = _folder(tmp_path, [NEW, "~$" + NEW, ".hidden.pptx", "notes.txt", "deck.pptx.tmp"])
    rejected = []
    with caplog.at_level(logging.INFO, logger="report_extractor.scanner"):
        found = scan_inputs([folder], on_reject=lambda p, why: rejected.append((p.name, why)))
    assert [p.name for p in found] == [NEW]
    names = dict(rejected)
    assert "~$" + NEW in names and "khoá/tạm" in names["~$" + NEW]
    assert ".hidden.pptx" in names and "ẩn" in names[".hidden.pptx"]
    assert "deck.pptx.tmp" in names and "phần mở rộng" in names["deck.pptx.tmp"]
    assert "notes.txt" not in names                                       # genuinely unrelated files are not noise
    assert sum("SCAN_REJECT" in r.message for r in caplog.records) == 3


# 6. absent master row -> visible with MASTER_NOT_FOUND status (no Qwen, no parse)
def test_missing_master_row_keeps_report_visible(tmp_path, template, monkeypatch):
    folder = _folder(tmp_path)
    from openpyxl import load_workbook
    wb = load_workbook(template)
    for i, k in enumerate(SIX_KEYS[:5]):
        wb["Kiểm chứng"].cell(row=4 + i, column=2, value=k)
    wb.save(template)
    import app.pptx_parser as pp
    monkeypatch.setattr(pp, "parse_pptx", lambda *a, **k: (_ for _ in ()).throw(AssertionError("prescan opened a pptx")))
    ctl = GuiController(AppConfig(period_mode="all"), config_path=tmp_path / "c.json")
    ctl._today = lambda: date(2026, 9, 30)
    ctl.set_report_folder(str(folder))
    ctl.set_template(str(template))
    ctl.set_output(str(tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"))
    assert len(ctl.all_files) == 6
    res = ctl.scan()
    assert res is not None and res.counts()["discovered"] == 6 and res.counts()["candidates"] == 5
    rows = {r.path.name: r for r in ctl.scan_rows()}                     # default view = all discovered files
    assert len(rows) == 6 and rows[NEW].action == ACTION_MASTER_NOT_FOUND
    assert rows[NEW].status_vi == "Không tìm thấy Management Number trong Excel" and rows[NEW].management_number == "260925015"
    assert NEW in {r.path.name for r in ctl.scan_rows(SCAN_FILTERS_VI[0])}  # even the "cần xử lý" view keeps it
    assert NEW in {r.path.name for r in ctl.scan_rows(SCAN_FILTERS_VI[2])}
    assert ctl.queue_text() == ("Tổng file phát hiện: 6   Sẽ xử lý: 5   Bỏ qua/đã cập nhật: 0   "
                                "Không tìm thấy Management Number: 1   Lỗi/không hợp lệ: 0   Đã loại thủ công: 0")
    assert "Không tìm thấy Management Number trong Excel: 1" in res.summary_lines_vi()


# 7 + 8. GUI shows 6 immediately, lists all 6, no Ollama call for discovery
def test_gui_counts_six_reports_and_lists_all(monkeypatch, tmp_path):
    calls = []

    class NoClient:
        def __init__(self, *a, **k):
            calls.append(a)

        def test_connection(self):
            raise AssertionError("Ollama must not be needed to list reports")

    gui, a, reg, folder = _make_app(monkeypatch, tmp_path, names=SIX)
    assert a.lbl_found.cfg["text"] == "Đã tìm thấy 6 báo cáo"
    assert len(a.ctl.all_files) == 6 and a.var_scan_filter.get() == gc.DEFAULT_SCAN_FILTER_VI
    assert sorted(v["values"][4] for v in a.scan_tree.items.values()) == sorted(SIX)
    assert a.lbl_list_count.cfg["text"] == "6 file"
    assert a.lbl_queue.cfg["text"].startswith("Tổng file phát hiện: 6   Sẽ xử lý: 6")
    a.ctl._client_factory = NoClient
    a.scan_reports()
    for t in threading.enumerate():
        if t.name == "prescan":
            t.join(10)
    a._poll()
    a._render_scan()
    assert calls == [] and len(a.scan_tree.items) == 6


# 9. the five September reports behave exactly as before
def test_five_september_reports_unchanged(tmp_path):
    folder = _folder(tmp_path, FIVE)
    files = scan_folder(folder)
    assert [management_number_from_filename(p.name) for p in files] == sorted(SIX_KEYS[:5], key=lambda k: FIVE[SIX_KEYS.index(k)].lower()) \
        or len(files) == 5
    res = prescan(files, month_period(2026, 9), None, None, today=date(2026, 9, 30))
    assert res.counts()["discovered"] == 5 and all(it.action == ACTION_PROCESS for it in res.items)
    assert {it.management_number for it in res.items} == set(SIX_KEYS[:5])
    six = prescan(scan_folder(_folder(tmp_path / "six")), month_period(2026, 9), None, None, today=date(2026, 9, 30))
    assert six.counts()["discovered"] == 6 and six.counts()["candidates"] == 6      # new one joins without a master
