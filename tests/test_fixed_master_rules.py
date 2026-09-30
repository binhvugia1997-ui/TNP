"""FINAL fixed-master rules (match mode):

1. Management Number not in the master -> nothing written, 'Không tìm thấy Management Number', rerun possible.
2. Row already complete -> 'Bỏ qua — đã cập nhật' without parsing / Qwen / re-extraction / image rewrite.
3. Failed / interrupted / not written / Cần kiểm tra -> eligible again.
4. Row partially complete -> auto-fill ONLY the blank fields, never touch populated ones.
"""
import datetime as dt
from pathlib import Path

from openpyxl import load_workbook

from app import batch_processor as bp
from app.batch_processor import BatchOptions, BatchProcessor, filter_review_reasons
from app.excel_writer import ExcelWriter

SHEET = "Kiểm chứng"
COL = {"stt": 1, "mgmt": 2, "vendor": 3, "date": 4, "model": 5, "item": 6, "defect": 7, "qpn": 8, "cause": 9,
       "improvement": 10, "image": 11}


def _prefill(template: Path, rows):
    wb = load_workbook(template)
    ws = wb[SHEET]
    for i, values in enumerate(rows):
        r = 4 + i
        ws.cell(row=r, column=1, value=i + 1)
        for key, val in values.items():
            ws.cell(row=r, column=COL[key], value=val)
        for k in range(12, 20):
            ws.cell(row=r, column=k, value="OK")
    wb.save(template)


def _run(files, template, out, **kw):
    events = []
    opts = BatchOptions(files=[Path(f) for f in files], template=template, output_file=out, use_ollama=False, **kw)
    proc = BatchProcessor(opts, on_file=lambda i, s, d: events.append((i, s, d)))
    return proc.run(), events, proc


def _images_at(ws, row, col):
    return [im for im in ws._images if im.anchor._from.row + 1 == row and im.anchor._from.col + 1 == col]


# ------------------------------------------------------------------ rule 1
def test_rule1_not_found_no_parse_no_write(sample_tree, template, tmp_path, monkeypatch):
    _prefill(template, [{"mgmt": "999999999-VOC"}])
    calls = []
    monkeypatch.setattr(bp, "parse_pptx", lambda p: calls.append(p) or (_ for _ in ()).throw(AssertionError("parsed")))
    out = tmp_path / "out" / "k.xlsx"
    summary, events, proc = _run([sample_tree["files"][0]], template, out)
    assert summary.not_written == 1 and summary.failed == 0 and calls == []          # decided before parsing
    assert events[-1][1] == "not_written"
    assert "Không tìm thấy Management Number 260918080-VOC" in proc.results[0].error
    assert proc.results[0].management_number == "260918080-VOC"
    ws = load_workbook(out)[SHEET]
    assert ws.cell(row=5, column=2).value is None                                     # no row created
    # user adds the key -> rerun works (real parsing this time)
    monkeypatch.undo()
    _prefill(template, [{"mgmt": "999999999-VOC"}, {"mgmt": "260918080-VOC"}])
    summary2, _, proc2 = _run([sample_tree["files"][0]], template, tmp_path / "out2" / "k.xlsx")
    assert summary2.completed == 1 and proc2.results[0].excel_row == 5


# ------------------------------------------------------------------ rule 2 + 3
def test_rule2_complete_row_is_skipped_without_qwen_or_extraction(sample_tree, template, tmp_path, monkeypatch):
    _prefill(template, [{"mgmt": "260918080-VOC"}])
    out = tmp_path / "out" / "k.xlsx"
    s1, _, _ = _run([sample_tree["files"][0]], template, out)
    assert s1.completed == 1
    ws = load_workbook(out)[SHEET]
    before = [ws.cell(row=4, column=c).value for c in range(1, 20)]
    n_img = len(ws._images)

    parsed, classified = [], []
    monkeypatch.setattr(bp, "parse_pptx", lambda p: parsed.append(p))
    monkeypatch.setattr(bp, "classify", lambda *a, **k: classified.append(a))
    s2, events, proc = _run([sample_tree["files"][0]], template, out)          # same output workbook
    assert s2.skipped == 1 and s2.completed == 0 and parsed == [] and classified == []
    assert events[-1][1] == "skipped" and events[-1][2].startswith("Bỏ qua — đã cập nhật")
    assert proc.results[0].excel_row == 4
    ws2 = load_workbook(out)[SHEET]
    assert [ws2.cell(row=4, column=c).value for c in range(1, 20)] == before
    assert len(ws2._images) == n_img                                            # QPN / images not rewritten
    assert bp.STAGE_LABELS_VI["skipped"] == "Bỏ qua — đã cập nhật"


def test_rule2_skip_is_row_based_not_fingerprint_based(sample_tree, template, tmp_path):
    """A different copy of the same report (other file name/bytes) is also skipped when the row is complete;
    conversely a previously processed file is re-processed when its row was cleared (rule 3)."""
    _prefill(template, [{"mgmt": "260918080-VOC"}])
    out = tmp_path / "out" / "k.xlsx"
    _run([sample_tree["files"][0]], template, out)
    copy = tmp_path / "260918080-VOC_copy.pptx"
    copy.write_bytes(Path(sample_tree["files"][0]).read_bytes() + b"\0")
    s, _, _ = _run([copy], template, out)
    assert s.skipped == 1
    # clear the cause cell -> eligible again, only that field gets filled
    wb = load_workbook(out)
    wb[SHEET].cell(row=4, column=COL["cause"]).value = None
    wb.save(out)
    s3, _, proc3 = _run([sample_tree["files"][0]], template, out)
    assert s3.completed == 1 and proc3.results[0].filled_fields == ["root_cause"]


def test_rule3_needs_review_row_stays_eligible(sample_tree, template, tmp_path):
    _prefill(template, [{"mgmt": "260918083-VOC"}])
    out = tmp_path / "out" / "k.xlsx"
    f = [p for p in sample_tree["files"] if "260918083" in p.name]
    s1, _, _ = _run(f, template, out)
    assert s1.needs_review == 1                                  # no QPN in this sample -> Cần kiểm tra
    s2, _, _ = _run(f, template, out)
    assert s2.needs_review == 1 and s2.skipped == 0              # still processed on the next run


# ------------------------------------------------------------------ rule 4
def test_rule4_partial_row_fills_only_missing_fields(sample_tree, template, tmp_path):
    _prefill(template, [{"mgmt": "260918080-VOC", "vendor": "Teawon", "date": dt.datetime(2026, 9, 18),
                         "qpn": "(ảnh QPN thủ công)", "improvement": "Đối sách đã nhập tay", "model": "A185",
                         "item": "Rear"}])
    # put a manual picture in the QPN cell so the image field counts as complete
    wb = load_workbook(template)
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image
    png = tmp_path / "manual.png"
    Image.new("RGB", (40, 30), "red").save(png)
    img = XLImage(str(png))
    img.anchor = "H4"
    wb[SHEET].add_image(img)
    wb.save(template)

    out = tmp_path / "out" / "k.xlsx"
    summary, events, proc = _run([sample_tree["files"][0]], template, out)
    fr = proc.results[0]
    assert summary.completed == 0 and summary.needs_review == 1          # vendor conflict Teawon vs Doaltech -> review
    assert sorted(fr.filled_fields) == ["defect_content", "improvement_image", "root_cause"]
    ws = load_workbook(out)[SHEET]
    r = 4
    assert ws.cell(row=r, column=COL["mgmt"]).value == "260918080-VOC"
    assert ws.cell(row=r, column=COL["vendor"]).value == "Teawon"                 # kept, not overwritten
    assert ws.cell(row=r, column=COL["date"]).value == dt.datetime(2026, 9, 18)
    assert ws.cell(row=r, column=COL["model"]).value == "A185"
    assert ws.cell(row=r, column=COL["qpn"]).value == "(ảnh QPN thủ công)"       # QPN preserved
    assert len(_images_at(ws, r, COL["qpn"])) == 1                                # manual picture, no new one
    assert ws.cell(row=r, column=COL["improvement"]).value == "Đối sách đã nhập tay"
    assert ws.cell(row=r, column=COL["defect"]).value                              # filled
    assert ws.cell(row=r, column=COL["cause"]).value                               # filled
    assert len(_images_at(ws, r, COL["image"])) == 1                              # improvement image inserted
    assert all(ws.cell(row=r, column=k).value == "OK" for k in range(12, 20))     # WEEK untouched
    assert any("Vendor trong Excel: Teawon" in x for x in fr.review_reasons)
    # rerun: now complete -> skipped
    s2, _, _ = _run([sample_tree["files"][0]], template, out)
    assert s2.skipped == 1


def test_rule4_populated_text_not_overwritten_and_no_review_for_it(sample_tree, template, tmp_path):
    """Report has no QPN, but the row already has a QPN picture -> no 'Không tìm thấy QPN' review."""
    f = [p for p in sample_tree["files"] if "260918083" in p.name]
    _prefill(template, [{"mgmt": "260918083-VOC", "defect": "Lỗi nhập tay"}])
    wb = load_workbook(template)
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image
    png = tmp_path / "manual.png"
    Image.new("RGB", (40, 30), "blue").save(png)
    img = XLImage(str(png))
    img.anchor = "H4"
    wb[SHEET].add_image(img)
    wb.save(template)
    out = tmp_path / "out" / "k.xlsx"
    summary, _, proc = _run(f, template, out)
    fr = proc.results[0]
    assert summary.completed == 1, fr.review_reasons
    assert not any("QPN" in r or "Nội dung lỗi" in r for r in fr.review_reasons)
    ws = load_workbook(out)[SHEET]
    assert ws.cell(row=4, column=COL["defect"]).value == "Lỗi nhập tay"
    assert "defect_content" not in fr.filled_fields and "qpn" not in fr.filled_fields


def test_force_reprocess_still_rewrites_managed_fields(sample_tree, template, tmp_path):
    _prefill(template, [{"mgmt": "260918080-VOC", "model": "WRONG"}])
    out = tmp_path / "out" / "k.xlsx"
    s, _, _ = _run([sample_tree["files"][0]], template, out, force_reprocess=True)
    assert s.completed == 1
    ws = load_workbook(out)[SHEET]
    assert ws.cell(row=4, column=COL["model"]).value == "A185"
    assert ws.cell(row=4, column=COL["mgmt"]).value == "260918080-VOC"
    assert all(ws.cell(row=4, column=k).value == "OK" for k in range(12, 20))


def test_missing_managed_fields_and_reason_filter(template, tmp_path):
    _prefill(template, [{"mgmt": "260918080-VOC", "vendor": "Doaltech", "cause": "x"}])
    w = ExcelWriter(template, tmp_path / "o.xlsx")
    missing = w.missing_managed_fields(4)
    assert "vendor" not in missing and "root_cause" not in missing
    assert {"model", "item", "defect_content", "improvement", "qpn", "improvement_image", "occurrence_date"} == set(missing)
    w.close()
    reasons = ["Không tìm thấy QPN trong báo cáo", "Không tìm thấy Nguyên nhân trong báo cáo",
               "Không có hình ảnh cải tiến", "Không xác định được Vendor từ nội dung báo cáo", "Phân loại mơ hồ"]
    assert filter_review_reasons(reasons, ["root_cause"]) == ["Không tìm thấy Nguyên nhân trong báo cáo", "Phân loại mơ hồ"]
