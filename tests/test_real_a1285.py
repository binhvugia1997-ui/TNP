"""Regression tests built from the structure of the real A1285 report / Kiem_chung template
(as reported from the Windows diagnostics)."""
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from app.classifier import (first_heading_kind, heuristic_classify, is_cover_slide,
                            merge_llm_with_heuristic, section_kind_of_heading)
from app.excel_writer import ExcelWriter, week_index
from app.extractor import ExtractedRecord, extract_record
from app.main import write_utf8_report
from app.pptx_parser import parse_pptx


# --------------------------------------------------------------------------- template
def make_real_like_template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kiểm chứng"
    ws["A1"] = "CÔNG TY TNHH ..."
    ws.merge_cells("A2:S2")
    ws["A2"] = "BẢNG KIỂM CHỨNG ĐỐI SÁCH CẢI TIẾN NHÀ CUNG CẤP"
    hdr = ["STT", "Management\nnumber", "Vendor", "Ngày phát sinh", "Model", "Item",
           "Nội dung lỗi", "QPN", "Nguyên nhân", "Nội dung đối sách cải tiến", "Hình ảnh cải tiến"]
    for i, h in enumerate(hdr, start=1):
        ws.cell(row=4, column=i, value=h)
        ws.merge_cells(start_row=4, start_column=i, end_row=5, end_column=i)
    ws.merge_cells("L4:S4")
    ws["L4"] = "Theo dõi cải tiến"
    for k in range(1, 9):
        ws.cell(row=5, column=11 + k, value=f"WEEK + {k}")
    wb.save(path)
    return path


def test_week_index_variants():
    for txt in ("WEEK + 1", "WEEK +1", "week+1", "W1", "+1", "Tuần 1", "T+1", "1W", "1"):
        assert week_index(__import__("app.pptx_parser", fromlist=["norm_key"]).norm_key(txt)) == 1, txt
    assert week_index("week 10") is None and week_index("model") is None


def test_real_template_theo_doi_cai_tien_weeks(tmp_path):
    tpl = make_real_like_template(tmp_path / "Kiem_chung.xlsx")
    w = ExcelWriter(tpl, tmp_path / "out.xlsx")
    assert w.header_row == 4 and w.data_start == 6
    assert w.columns["management_number"] == 2            # multi-line header – must not regress
    assert w.columns["vendor"] == 3 and w.columns["occurrence_date"] == 4
    expected = dict(zip([f"week_{k}" for k in range(1, 9)], range(12, 20)))   # L..S
    for f, col in expected.items():
        assert w.columns.get(f) == col, (f, w.columns.get(f))
    # the merged parent is not mistaken for a data column
    assert 12 not in [c for f, c in w.columns.items() if not f.startswith("week_")]
    row = w.append_record(ExtractedRecord(management_number="X-1", model="A185", item="Rear",
                                          root_cause="nn", improvement="ds"))
    w.save()
    ws = load_workbook(tmp_path / "out.xlsx")["Kiểm chứng"]
    assert row == 6
    assert ws.cell(row=6, column=3).value is None and ws.cell(row=6, column=4).value is None
    for col in range(12, 20):                                # detected but still BLANK
        assert ws.cell(row=6, column=col).value is None
    assert ws["L4"].value == "Theo dõi cải tiến" and ws["L5"].value == "WEEK + 1"


# --------------------------------------------------------------------------- report
def _png(color):
    im = Image.new("RGB", (300, 200), color)
    b = BytesIO()
    im.save(b, "PNG")
    b.seek(0)
    return b


def _tb(slide, text, top, size=14, bold=False, left=0.5, width=8.0, height=1.0):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
    return tb


def make_a1285_like(path: Path, defect_text_lines: bool = False) -> Path:
    """8 slides: cover, 1. HIỆN TRẠNG (picture only), NGUYÊN NHÂN, PHƯƠNG PHÁP XỬ LÝ TẠM THỜI,
    improvements 5-7 (pictures), 8 = HIỆU QUẢ + ĐỐI SÁCH LÂU DÀI (picture)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)                                           # 1 cover
    _tb(s, "BÁO CÁO CẢI TIẾN LỖI XƯỚC REAR", 1.0, 30, True)
    _tb(s, "Model: SM-A185      Item: Rear\nVendor: ABC VINA\nNgày: 18/09/2026", 3.0, 16)
    s.shapes.add_picture(_png("#ffffff"), Inches(9), Inches(5.5), Inches(3), Inches(1.5))   # logo
    s = prs.slides.add_slide(blank)                                           # 2 hiện trạng
    _tb(s, "1. HIỆN TRẠNG", 0.3, 24, True)
    if defect_text_lines:
        _tb(s, "Xước: 15ea\nMẻ: 13ea", 1.2, 14)
    s.shapes.add_picture(_png("#ffd6d6"), Inches(0.5), Inches(1.2), Inches(12), Inches(5.5))  # QPN screenshot
    s = prs.slides.add_slide(blank)                                           # 3 nguyên nhân
    _tb(s, "2. NGUYÊN NHÂN", 0.3, 24, True)
    _tb(s, "Nguyên nhân trong sản xuất\n- Khay chứa không có lót", 1.2, 14)
    s = prs.slides.add_slide(blank)                                           # 4 tạm thời
    _tb(s, "3. PHƯƠNG PHÁP XỬ LÝ TẠM THỜI", 0.3, 24, True)
    _tb(s, "- Sorting 100% hàng tồn: 2.350 pcs", 1.2, 14)
    s.shapes.add_picture(_png("#fff2cc"), Inches(0.5), Inches(3), Inches(6), Inches(3))
    for n, title in ((5, "4. CẢI TIẾN TRONG SẢN XUẤT"), (6, "5. CẢI TIẾN TRONG KIỂM TRA"), (7, "6. CẢI TIẾN JIG")):
        s = prs.slides.add_slide(blank)                                       # 5-7 cải tiến
        _tb(s, title, 0.3, 24, True)
        _tb(s, f"- Trước: cũ ({n})\n- Sau: mới ({n})", 1.2, 14)
        s.shapes.add_picture(_png("#c8e6c9"), Inches(7), Inches(1.2), Inches(5.5), Inches(3))
    s = prs.slides.add_slide(blank)                                           # 8 hiệu quả + lâu dài
    _tb(s, "7. HIỆU QUẢ SAU CẢI TIẾN", 0.3, 24, True)
    _tb(s, "Theo dõi 4 tuần: chưa có dữ liệu", 1.0, 14)
    s.shapes.add_picture(_png("#b3e5fc"), Inches(7), Inches(1.0), Inches(5.5), Inches(2.5))
    _tb(s, "ĐỐI SÁCH LÂU DÀI", 4.0, 20, True)
    _tb(s, "- Cập nhật SOP Rev.03, triển khai ngang dòng A175", 4.7, 14)
    prs.save(str(path))
    return path


def test_hien_trang_is_defect_heading():
    assert section_kind_of_heading("1. HIỆN TRẠNG") == "defect"
    assert section_kind_of_heading("PHƯƠNG PHÁP XỬ LÝ TẠM THỜI") == "temporary"
    assert section_kind_of_heading("7. HIỆU QUẢ SAU CẢI TIẾN") == "verify"


def test_a1285_structure_classification(tmp_path):
    r = parse_pptx(make_a1285_like(tmp_path / "A1285.pptx"))
    assert is_cover_slide(r, r.slides[0])
    assert first_heading_kind(r.slides[1]) == "defect"
    c = heuristic_classify(r)
    assert c.defect_slide == 2
    assert c.cause_slides == [3]
    assert c.temporary_slides == [4]
    assert 1 not in c.improvement_slides and 2 not in c.improvement_slides
    assert [5, 6, 7] == [n for n in c.improvement_slides if n < 8]
    assert 8 in c.improvement_slides                        # ĐỐI SÁCH LÂU DÀI text lives there
    # images: only slides whose own section is improvement
    assert c.improvement_image_slides == [5, 6, 7]
    # QPN: no "Quality Problem Notice" text anywhere -> HIỆN TRẠNG page with picture (deterministic,
    # source 'defect_structure'); resolved structurally -> no 'Cần kiểm tra' for the QPN itself
    assert c.qpn_slide == 2 and c.qpn_source == "defect_structure"
    assert not any("QPN" in a or "Quality Problem Notice" in a for a in c.ambiguities)
    assert any(n.startswith("QPN = trang hiện trạng") for n in c.notes)


def test_a1285_record_contents(tmp_path):
    r = parse_pptx(make_a1285_like(tmp_path / "A1285.pptx"))
    rec = extract_record(r, heuristic_classify(r))
    assert rec.model == "A185" and rec.item == "Rear"
    assert rec.qpn_slide == 2
    assert "Khay chứa không có lót" in rec.root_cause
    for n in (5, 7):
        assert f"- Sau: mới ({n})" in rec.improvement
    assert "- Sau: mới (6)" not in rec.improvement            # PROMPT-001: "5. CẢI TIẾN TRONG KIỂM TRA" -> zero text
    assert "ĐỐI SÁCH LÂU DÀI" in rec.improvement and "Cập nhật SOP Rev.03" in rec.improvement   # production action body
    assert "Sorting 100%" not in rec.improvement and "chưa có dữ liệu" not in rec.improvement
    assert "BÁO CÁO CẢI TIẾN" not in rec.improvement        # cover text not copied
    # defect text exists only inside the picture -> blank + explicit review reason (no fabrication)
    assert rec.defect_content == ""
    assert "Nội dung lỗi chỉ có trong hình ảnh QPN, không có text để sao chép – cần bổ sung thủ công" in rec.review_reasons
    assert not any("Không tìm thấy QPN" in x or "chỉ do AI" in x for x in rec.review_reasons)
    assert rec.vendor == "" and rec.occurrence_date is None and all(v == "" for v in rec.weeks.values())


def test_a1285_defect_text_when_present(tmp_path):
    r = parse_pptx(make_a1285_like(tmp_path / "A1285b.pptx", defect_text_lines=True))
    rec = extract_record(r, heuristic_classify(r))
    assert rec.defect_content == "Xước: 15ea\nMẻ: 13ea"
    assert not any("Nội dung lỗi" in x for x in rec.review_reasons)


def test_llm_cannot_widen_image_slides_to_cover_or_hien_trang(tmp_path):
    r = parse_pptx(make_a1285_like(tmp_path / "A1285.pptx"))
    heur = heuristic_classify(r)
    raw = {"qpn_slide": None, "defect_slide": 2, "cause_slides": [3], "temporary_slides": [4],
           "improvement_slides": [1, 2, 5, 6, 7, 8], "improvement_image_slides": [1, 2, 5, 6, 7, 8]}
    c = merge_llm_with_heuristic(raw, heur, r)
    assert c.improvement_image_slides == [5, 6, 7]
    assert 2 not in c.improvement_slides
    assert c.qpn_slide == 2 and c.qpn_source == "defect_structure"
    assert not any("chỉ do AI" in a for a in c.ambiguities)


def test_utf8_report_file(tmp_path):
    p = write_utf8_report("4. CẢI TIẾN – Nguyên nhân – Kiểm chứng", tmp_path / "d" / "x.txt")
    raw = p.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    assert p.read_text(encoding="utf-8-sig") == "4. CẢI TIẾN – Nguyên nhân – Kiểm chứng"
