"""PROMPT-006B – learning / correction for improvement CONTENT regions.  Deterministic extractor = baseline;
learning only decides whether a block is included (verbatim text, source order)."""
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pptx.enum.shapes import MSO_SHAPE

import app
from app import content_learning as cl
from app import image_learning as il
from app import updater as up
from app.batch_processor import BatchOptions, BatchProcessor
from app.classifier import heuristic_classify
from app.content_learning import (CONTENT_FEATURE_NAMES, CONTENT_FEATURE_SCHEMA, CONTENT_LABELS, MSG_MODEL_UNAVAILABLE,
                                  MSG_NOT_ENOUGH, ContentLearning, build_content_candidates, decide_content,
                                  explain_content, rebuild_improvement_text)
from app.extractor import collect_sections, extract_record
from app.image_learning import ImageLearning, TrainingError, save_model, train_model
from app.image_review import reapply_content_labels
from app.pptx_parser import parse_pptx
from tests.test_after_evidence import NAME, _deck, _pics, _prod_head, _tb, _tb_colored
from tests.test_content_region import (COL, IMP_ITEM_1, IMP_ITEM_2, INSPECTION_TEXT, SHEET, _images_at, _prefill,
                                       _shape, make_real_layout_report)
from tests.test_prompt004 import MGMT

UNRELATED = "Ghi chú nội bộ: lịch họp QA 25/09 phòng 3 (không liên quan cải tiến)"
BLUE_SAU = "+ Sau: Bọc silicon 2mm lên bề mặt jig nén"


# ------------------------------------------------------------------ decks
def slide_items(s, W):                 # two improvement items + unrelated note + inline Trước/Sau
    _prod_head(s, W)
    _tb(s, IMP_ITEM_1, 1.7, 1.0, 5.6, 2.8, 12)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 2.95, 1.5, 0.35)
    _pics(s, ["#c8e6c9"], 7.5, 3.35)
    _tb(s, IMP_ITEM_2, 1.7, 4.0, 5.6, 2.4, 12)
    _tb(s, UNRELATED, 7.5, 5.5, 5.3, 0.5, 11)


def slide_blue(s, W):                  # blue "+ Sau:" line as its own block
    _prod_head(s, W)
    _tb(s, "Cải tiến jig nén\n+ Trước: Jig nén bằng nhôm", 1.7, 1.0, 11, 0.8, 12)
    _tb_colored(s, [(BLUE_SAU, "blue")], 1.7, 2.0, 11, 0.4)
    _pics(s, ["#c8e6c9", "#c8e6c9"], 1.7, 2.6)


def slide_temp(s, W):                  # temporary action slide – excluded by business rules
    _tb(s, "2. XỬ LÝ TẠM THỜI", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Xử lý tạm thời", 0.15, 2.5, 1.3, 1.3)
    _tb(s, "- Sàng lọc 100% tồn kho, cách ly 120 pcs lỗi", 1.7, 1.1, 11, 1.5, 13)


def slide_inspection(s, W):
    _prod_head(s, W, title="4. CẢI TIẾN TRONG KIỂM TRA")
    _tb(s, INSPECTION_TEXT, 1.7, 1.0, 5.6, 2.5, 12)


def slide_continuation(s, W):          # heading-less second production slide
    _prod_head(s, W, title="3. CẢI TIẾN TRONG SẢN XUẤT (tiếp)")
    _tb(s, "- Bổ sung tấm chắn chống văng sơn tại buồng sơn số 2", 1.7, 1.0, 11, 0.6, 12)


def _build(tmp_path, builders, name=NAME):
    r = parse_pptx(_deck(tmp_path / name, builders))
    cls = heuristic_classify(r)
    sections = collect_sections(r, cls)
    slides = sorted(set(cls.improvement_slides) | {s.slide for s in sections if s.kind in ("improvement", "standard")}
                    | set(cls.temporary_slides))
    cands = decide_content(build_content_candidates(r, sections, slides, MGMT, str(r.path)))
    return r, cls, sections, cands


def _by_text(cands, start):
    return next(c for c in cands if c.text.startswith(start))


def _synthetic(n_inc=12, n_exc=12, schema=CONTENT_FEATURE_SCHEMA):
    recs = []
    for i in range(n_inc + n_exc):
        inc = i < n_inc
        f = {k: 0.0 for k in CONTENT_FEATURE_NAMES}
        f.update({"relative_x": 0.15 if inc else 0.6 + 0.01 * i, "relative_y": 0.2 + 0.01 * i,
                  "relative_width": 0.4 if inc else 0.3, "text_length": 0.4 if inc else 0.1,
                  "paragraph_count": 4.0 if inc else 1.0, "under_improvement_title": 1.0,
                  "contains_sau": 1.0 if inc else 0.0, "in_improvement_item": 1.0 if inc else 0.0})
        recs.append({"schema_version": schema, "candidate_id": f"X|S3|SH{i}|{i}",
                     "label": "IMPROVEMENT_CONTENT" if inc else "EXCLUDE_CONTENT", "features": f})
    return recs


# ================================================================== §21 features / hard exclusions / baseline
def test_version_build_007():
    assert app.__version__ == "1.2.1" and app.BUILD_NUMBER == 12 and app.BUILD_ID == "012"
    init = (Path(__file__).resolve().parent.parent / "backup" / "ReportExtractor_v1.0.4_Build004_STABLE" / "app"
            / "__init__.py").read_text(encoding="utf-8")
    assert '__version__ = "1.0.4"' in init and "BUILD_NUMBER = 4" in init


def test_candidates_schema_identity_and_hard_exclusions(tmp_path):
    r, cls, sections, cands = _build(tmp_path, [slide_items])
    assert cands and all(set(c.features) == set(CONTENT_FEATURE_NAMES) for c in cands)
    assert CONTENT_FEATURE_SCHEMA == 1 and set(CONTENT_FEATURE_NAMES) != set(il.FEATURE_NAMES)
    for c in cands:
        assert c.candidate_id == f"{MGMT}|S{c.slide}|SH{c.shape_id}|{c.order}" and str(tmp_path) not in c.candidate_id
        assert 0.0 <= c.confidence <= 1.0 and c.slide_size == (r.slide_width, r.slide_height)
    title = _by_text(cands, "3. CẢI TIẾN")
    sidebar = _by_text(cands, "Cải tiến trong sản xuất")
    caption = _by_text(cands, "Sau cải tiến")
    footer = _by_text(cands, "CTMS")
    deco = _by_text(cands, "▶")
    for c in (title, sidebar, caption, footer, deco):
        assert c.hard_excluded and c.decision == "exclude" and c.confidence == 0.0 and c.detail_kind == "DECORATION"
    assert title.features["role_title"] and sidebar.features["role_sidebar"] and caption.features["role_caption"]
    assert footer.features["role_furniture"]
    # user label can never resurrect a hard exclusion
    out = decide_content([title], None, {title.candidate_id: "IMPROVEMENT_CONTENT"})
    assert out[0].decision == "exclude" and out[0].decision_source == "rules"


def test_items_truoc_sau_and_unrelated_text(tmp_path):
    r, cls, sections, cands = _build(tmp_path, [slide_items])
    i1, i2, un = _by_text(cands, "Cải tiến lỗi Mẻ xước"), _by_text(cands, "Cải tiến lỗi bong sơn"), _by_text(cands, "Ghi chú")
    for c in (i1, i2):
        assert c.decision == "include" and c.confidence >= 0.9 and c.baseline_included
        assert c.features["contains_truoc"] and c.features["contains_sau"] and c.features["under_improvement_title"]
        assert c.detail_kind in ("AFTER_TEXT", "IMPROVEMENT_BODY")
    assert i1.features["reading_order"] < i2.features["reading_order"] and i1.features["near_picture"]
    # the unrelated note sits in the production section by geometry -> deterministic extractor keeps it (baseline)
    assert un.baseline_included and un.decision == "include"
    text = explain_content(i1)
    assert text.splitlines()[0] == "Slide 3" and "Kết quả: Nội dung cải tiến" in text and "Độ tin cậy: Cao" in text
    assert "- nằm trong khu vực CẢI TIẾN TRONG SẢN XUẤT" in text and "- thứ tự đọc phù hợp" in text


def test_blue_sau_line_block(tmp_path):
    r, cls, sections, cands = _build(tmp_path, [slide_blue])
    blue = _by_text(cands, "+ Sau:")
    assert blue.features["has_blue_text"] and blue.features["contains_sau"] and blue.decision == "include"
    assert "- có chữ xanh (mô tả Sau)" in explain_content(blue) and "- có dòng + Sau:" in explain_content(blue)


def test_temporary_and_inspection_regions(tmp_path):
    r, cls, sections, cands = _build(tmp_path, [slide_items, slide_temp, slide_inspection])
    tmp = _by_text(cands, "- Sàng lọc")
    assert tmp.hard_excluded and tmp.deterministic_kind == "excluded_section" and tmp.detail_kind == "TEMPORARY_CONTENT"
    assert decide_content([tmp], None, {tmp.candidate_id: "IMPROVEMENT_CONTENT"})[0].decision == "exclude"
    insp = _by_text(cands, "Cải tiến trong kiểm tra:")
    assert not insp.hard_excluded and insp.decision == "exclude" and insp.confidence <= cl.THRESHOLD_EXCLUDE
    assert insp.detail_kind == "CONTROL_CONTENT" and insp.features["under_inspection_title"]
    assert "cải tiến kiểm tra / kiểm soát" in explain_content(insp)


def test_multi_slide_order_and_baseline_unchanged(tmp_path):
    r = parse_pptx(_deck(tmp_path / NAME, [slide_items, slide_continuation]))
    cls = heuristic_classify(r)
    base = extract_record(r, cls)
    lrn = ImageLearning(tmp_path / "lrn")
    rec = extract_record(r, cls, learning=lrn)
    assert rec.improvement == base.improvement and rec.content_candidates and base.content_candidates == []
    assert rec.improvement.index("Cải tiến lỗi Mẻ xước") < rec.improvement.index("Cải tiến lỗi bong sơn") \
        < rec.improvement.index("- Bổ sung tấm chắn")
    assert rebuild_improvement_text(r, collect_sections(r, cls), rec.content_candidates) is None
    # the real-layout deck (cause / temporary / production / inspection / long-term) is also byte-identical
    deck = make_real_layout_report(tmp_path / "(CTMS)_1_260920045-VOC_ Đối sách LỖI MẺ XƯỚC 24.9.2026.pptx")
    r2 = parse_pptx(deck)
    c2 = heuristic_classify(r2)
    assert extract_record(r2, c2, learning=lrn).improvement == extract_record(r2, c2).improvement


# ================================================================== user corrections (block-level)
def test_user_exclude_and_include_corrections_rebuild_verbatim_in_order(tmp_path):
    r = parse_pptx(_deck(tmp_path / NAME, [slide_items, slide_inspection, slide_continuation]))
    cls = heuristic_classify(r)
    lrn = ImageLearning(tmp_path / "lrn")
    rec = extract_record(r, cls, learning=lrn)
    assert UNRELATED in rec.improvement and INSPECTION_TEXT not in rec.improvement
    un = _by_text(rec.content_candidates, "Ghi chú")
    insp = _by_text(rec.content_candidates, "Cải tiến trong kiểm tra:")
    lrn.content.store.label(un, "EXCLUDE_CONTENT")                 # A + B + C -> A + B
    lrn.content.store.label(insp, "IMPROVEMENT_CONTENT")           # omitted D -> included
    rec2 = extract_record(r, cls, learning=lrn)
    t = rec2.improvement
    assert UNRELATED not in t and INSPECTION_TEXT in t
    assert IMP_ITEM_1 in t and IMP_ITEM_2 in t                      # verbatim, untouched
    assert t.index(IMP_ITEM_1) < t.index(IMP_ITEM_2) < t.index("Cải tiến trong kiểm tra:") < t.index("- Bổ sung tấm chắn")
    c_un = next(c for c in rec2.content_candidates if c.candidate_id == un.candidate_id)
    assert c_un.decision == "exclude" and c_un.decision_source == "user" and lrn.content.model is None
    assert any(e.startswith("người dùng đã xác nhận: Không lấy") for e in c_un.evidence)
    # no model involved; relabel back restores the baseline text exactly
    lrn.content.store.label(un, "IMPROVEMENT_CONTENT")
    lrn.content.store.label(insp, "EXCLUDE_CONTENT")
    assert extract_record(r, cls, learning=lrn).improvement == extract_record(r, cls).improvement


def test_content_label_store_relabel_duplicate_persistence(tmp_path):
    r, cls, sections, cands = _build(tmp_path, [slide_items])
    un = _by_text(cands, "Ghi chú")
    store = ContentLearning(tmp_path / "lrn").store
    assert store.path.name == "content_labels.jsonl" and store.labels == CONTENT_LABELS
    rec = store.label(un, "EXCLUDE_CONTENT")
    assert rec["schema_version"] == 1 and rec["previous_label"] is None and rec["text"].startswith("Ghi chú")
    assert rec["shape_id"] == un.shape_id and "picture_id" not in rec and set(rec["features"]) == set(CONTENT_FEATURE_NAMES)
    assert store.label(un, "EXCLUDE_CONTENT") is None             # duplicate protection
    rec2 = store.label(un, "IMPROVEMENT_CONTENT")
    assert rec2["previous_label"] == "EXCLUDE_CONTENT" and rec2["timestamp"]
    with pytest.raises(ValueError):
        store.label(un, "AFTER")                                   # image label never accepted here
    again = ContentLearning(tmp_path / "lrn")                      # restart
    assert again.overrides() == {un.candidate_id: "IMPROVEMENT_CONTENT"}
    assert again.counts() == {"total": 1, "include": 1, "exclude": 0}
    assert again.status_text().startswith("Nội dung cải tiến: 1   Không lấy: 0")
    # image dataset untouched / separate file
    assert not (tmp_path / "lrn" / "image_labels.jsonl").exists()
    lines = (tmp_path / "lrn" / "content_labels.jsonl").read_text(encoding="utf-8").splitlines()
    assert len(lines) == 2 and all(json.loads(l)["candidate_id"] == un.candidate_id for l in lines)


# ================================================================== learning
def test_content_training_minimums_and_single_class(tmp_path):
    c = ContentLearning(tmp_path / "lrn")
    ok, msg = c.train()
    assert not ok and msg.startswith(MSG_NOT_ENOUGH) and "0/" not in msg
    with pytest.raises(TrainingError, match="Chưa đủ dữ liệu học nội dung"):
        train_model(_synthetic(12, 4), min_examples=cl.MIN_EXAMPLES, min_per_class=cl.MIN_PER_CLASS,
                    msg=MSG_NOT_ENOUGH, schema=CONTENT_FEATURE_SCHEMA, labels=CONTENT_LABELS,
                    feature_names=CONTENT_FEATURE_NAMES, positive="IMPROVEMENT_CONTENT")
    with pytest.raises(TrainingError):
        train_model(_synthetic(25, 0), min_examples=cl.MIN_EXAMPLES, min_per_class=cl.MIN_PER_CLASS,
                    schema=CONTENT_FEATURE_SCHEMA, labels=CONTENT_LABELS, feature_names=CONTENT_FEATURE_NAMES,
                    positive="IMPROVEMENT_CONTENT")
    assert cl.MIN_EXAMPLES == 20 and cl.MIN_PER_CLASS == 5 and cl.MIN_EXAMPLES > il.MIN_EXAMPLES


def test_content_training_both_classes_persist_reload(tmp_path):
    c = ContentLearning(tmp_path / "lrn")
    for rec in _synthetic(12, 12):
        cand = cl.ContentCandidate(candidate_id=rec["candidate_id"], management_number="X", source_name="x.pptx",
                                   source_file="", slide=3, shape_id=1, order=1, shape_kind="paragraph",
                                   bounds=(0, 0, 1, 1), slide_size=(10, 10), text="t", features=rec["features"])
        c.store.label(cand, rec["label"])
    ok, msg = c.train()
    assert ok and msg.startswith("Đã cập nhật mô hình nội dung từ 24 mẫu")
    p = tmp_path / "lrn" / "model" / "content_model.json"
    meta = json.loads(p.read_text(encoding="utf-8"))
    assert meta["schema_version"] == 1 and meta["feature_names"] == list(CONTENT_FEATURE_NAMES) and meta["n_examples"] == 24
    assert not (tmp_path / "lrn" / "model" / "image_model.json").exists()
    again = ContentLearning(tmp_path / "lrn")
    assert again.model is not None and again.model.probability(_synthetic()[0]["features"]) > 0.75
    assert again.model.probability(_synthetic()[-1]["features"]) < 0.35
    assert "Mô hình: đã huấn luyện (24 mẫu)" in again.status_text()


@pytest.mark.parametrize("damage", ["corrupt", "schema", "features"])
def test_content_model_corrupt_or_incompatible_fallback(tmp_path, damage):
    model = train_model(_synthetic(12, 12), schema=CONTENT_FEATURE_SCHEMA, labels=CONTENT_LABELS,
                        feature_names=CONTENT_FEATURE_NAMES, positive="IMPROVEMENT_CONTENT")
    p = save_model(model, tmp_path / "lrn", cl.CONTENT_MODEL_FILE)
    if damage == "corrupt":
        p.write_text("{oops", encoding="utf-8")
    else:
        d = json.loads(p.read_text(encoding="utf-8"))
        if damage == "schema":
            d["schema_version"] = 2
        else:
            d["feature_names"] = list(il.FEATURE_NAMES)             # an IMAGE model in the content slot
        p.write_text(json.dumps(d), encoding="utf-8")
    c = ContentLearning(tmp_path / "lrn")
    assert c.model is None and c.model_status in ("corrupt", "incompatible") and MSG_MODEL_UNAVAILABLE in c.status_text()
    r = parse_pptx(_deck(tmp_path / NAME, [slide_items]))
    cls = heuristic_classify(r)
    lrn = ImageLearning(tmp_path / "lrn")
    assert extract_record(r, cls, learning=lrn).improvement == extract_record(r, cls).improvement


def _forced_model(bias):
    n = len(CONTENT_FEATURE_NAMES)
    return il.ImageModel(weights=[0.0] * n, bias=bias, means=[0.0] * n, scales=[1.0] * n,
                         feature_names=list(CONTENT_FEATURE_NAMES), schema_version=CONTENT_FEATURE_SCHEMA,
                         trained_at="t", n_examples=30, n_after=15, n_non_after=15)


def test_model_only_acts_on_uncertain_band_and_never_on_hard_exclusions(tmp_path):
    r, cls, sections, cands = _build(tmp_path, [slide_items, slide_temp, slide_inspection])
    # push one block into the uncertain band artificially (the synthetic decks are all strongly classified)
    un = _by_text(cands, "Ghi chú")
    un.confidence, un.baseline_included = 0.5, False
    out = decide_content(cands, _forced_model(-10.0))
    assert all(c.decision == "exclude" and c.learned_probability is None for c in out if c.hard_excluded)
    assert all(c.learned_probability is None for c in out if c.confidence >= cl.THRESHOLD_INCLUDE)
    c_un = next(c for c in out if c.candidate_id == un.candidate_id)
    assert c_un.decision == "exclude" and c_un.decision_source == "model" and c_un.learned_probability < 0.01
    out = decide_content(cands, _forced_model(10.0))
    c_un = next(c for c in out if c.candidate_id == un.candidate_id)
    assert c_un.decision == "include" and c_un.decision_source == "model"
    assert all(c.decision == "exclude" for c in out if c.hard_excluded)             # +10 never flips a hard exclusion
    assert _by_text(out, "- Sàng lọc").decision == "exclude"
    insp = _by_text(out, "Cải tiến trong kiểm tra:")
    assert insp.decision == "exclude" and insp.decision_source == "rules"            # strong rule evidence wins
    out = decide_content(cands, _forced_model(0.0))
    c_un = next(c for c in out if c.candidate_id == un.candidate_id)
    assert c_un.decision == "review" and cl.effective_included(c_un) is False       # falls back to the extractor
    # exact user correction beats the model
    out = decide_content(cands, _forced_model(10.0), {un.candidate_id: "EXCLUDE_CONTENT"})
    c_un = next(c for c in out if c.candidate_id == un.candidate_id)
    assert c_un.decision == "exclude" and c_un.decision_source == "user"


def test_schema_mismatch_records_ignored():
    xs, ys = il.training_examples(_synthetic(6, 6) + _synthetic(6, 6, schema=2), CONTENT_FEATURE_SCHEMA,
                                  CONTENT_LABELS, CONTENT_FEATURE_NAMES, "IMPROVEMENT_CONTENT")
    assert len(xs) == 12 and sum(ys) == 6


# ================================================================== Excel re-apply
def test_reapply_changes_only_the_improvement_cell(template, tmp_path):
    deck = _deck(tmp_path / f"(CTMS)_{MGMT}_ĐỐI SÁCH LỖI SƠN 24.09.2026.pptx", [slide_items, slide_inspection])
    _prefill(template, [{"mgmt": MGMT, "vendor": "Mtech", "qpn": "QPN nhập tay"}])
    wb = load_workbook(template)
    ws = wb[SHEET]
    ws.cell(row=4, column=12, value="=1+1")                         # a formula in the WEEK area
    wb.save(template)
    out = tmp_path / "out" / "k.xlsx"
    lrn_dir = tmp_path / "lrn"
    proc = BatchProcessor(BatchOptions(files=[deck], template=template, output_file=out, use_ollama=False,
                                       learning_dir=lrn_dir))
    proc.run()
    fr = proc.results[0]
    assert fr.content_candidates and all(isinstance(x, str) for x in fr.to_dict()["content_candidates"])
    wb = load_workbook(out)
    wb[SHEET].cell(row=4, column=30, value="Ghi chú tay")           # manual column typed after the batch
    wb.save(out)
    ws = load_workbook(out)[SHEET]
    before = {c.coordinate: c.value for c in ws[4] if c.column != COL["improvement"]}
    n_imgs = len(_images_at(ws, 4, COL["image"]))
    assert UNRELATED in ws.cell(row=4, column=COL["improvement"]).value
    for p in (out.parent / "backup").glob("*.xlsx"):
        p.unlink()
    lrn = ImageLearning(lrn_dir)
    un = _by_text(fr.content_candidates, "Ghi chú")
    insp = _by_text(fr.content_candidates, "Cải tiến trong kiểm tra:")
    lrn.content.store.label(un, "EXCLUDE_CONTENT")
    lrn.content.store.label(insp, "IMPROVEMENT_CONTENT")
    res = reapply_content_labels(template, out, [un, insp], lrn)
    assert res.ok and res.updated_rows == [4], res.errors
    ws2 = load_workbook(out)[SHEET]
    val = ws2.cell(row=4, column=COL["improvement"]).value
    assert UNRELATED not in val and INSPECTION_TEXT in val and IMP_ITEM_1 in val
    assert {c.coordinate: c.value for c in ws2[4] if c.column != COL["improvement"]} == before
    assert ws2.cell(row=4, column=12).value == "=1+1" and ws2.cell(row=4, column=30).value == "Ghi chú tay"
    assert ws2.cell(row=4, column=COL["vendor"]).value == "Mtech" and ws2.cell(row=4, column=COL["qpn"]).value == "QPN nhập tay"
    assert len(_images_at(ws2, 4, COL["image"])) == n_imgs
    assert len(list((out.parent / "backup").glob("*_backup_*.xlsx"))) == 1   # backup before re-apply
    assert un.decision == "exclude" and un.decision_source == "user"
    # next batch run (force) produces the corrected text directly
    proc2 = BatchProcessor(BatchOptions(files=[deck], template=template, output_file=out, use_ollama=False,
                                        learning_dir=lrn_dir, force_reprocess=True))
    proc2.run()
    assert UNRELATED not in load_workbook(out)[SHEET].cell(row=4, column=COL["improvement"]).value


def test_reapply_missing_row_touches_nothing(template, tmp_path):
    deck = _deck(tmp_path / f"(CTMS)_{MGMT}_ĐỐI SÁCH LỖI SƠN 24.09.2026.pptx", [slide_items])
    _prefill(template, [{"mgmt": "111111111"}])
    out = tmp_path / "out" / "k.xlsx"
    out.parent.mkdir(parents=True)
    out.write_bytes(template.read_bytes())
    r, cls, sections, cands = _build(tmp_path, [slide_items])
    lrn = ImageLearning(tmp_path / "lrn")
    un = _by_text(cands, "Ghi chú")
    un.source_file = str(deck)
    lrn.content.store.label(un, "EXCLUDE_CONTENT")
    stamp = out.stat().st_mtime_ns
    res = reapply_content_labels(template, out, [un], lrn)
    assert not res.ok and "không tìm thấy dòng Management Number" in res.errors[0]
    assert out.stat().st_mtime_ns == stamp and not list((out.parent / "backup").glob("*.xlsx"))


# ================================================================== combined training / GUI / updater
def test_train_all_reports_both_datasets_separately(tmp_path):
    lrn = ImageLearning(tmp_path / "lrn")
    ok, msg = lrn.train_all()
    lines = msg.splitlines()
    assert not ok and lines[0].startswith("Ảnh cải tiến: chưa đủ dữ liệu — 0/10") \
        and lines[1].startswith("Nội dung cải tiến: chưa đủ dữ liệu — 0/20")
    for rec in _synthetic(12, 12):
        cand = cl.ContentCandidate(candidate_id=rec["candidate_id"], management_number="X", source_name="x.pptx",
                                   source_file="", slide=3, shape_id=1, order=1, shape_kind="paragraph",
                                   bounds=(0, 0, 1, 1), slide_size=(10, 10), text="t", features=rec["features"])
        lrn.content.store.label(cand, rec["label"])
    ok, msg = lrn.train_all()
    lines = msg.splitlines()
    assert ok and lines[0].startswith("Ảnh cải tiến: chưa đủ dữ liệu") and lines[1] == "Nội dung cải tiến: đã cập nhật mô hình — 24 mẫu."
    assert lrn.model is None and lrn.content.model is not None


def test_gui_content_review_window(monkeypatch, tmp_path):
    from tests.test_gui_redesign import _make_app
    gui, a, reg, folder = _make_app(monkeypatch, tmp_path)
    a.ctl.learning_dir = tmp_path / "lrn"
    a.ctl._learning = None
    assert a.lbl_learning_content.cfg["text"].startswith("Nội dung:  Nội dung cải tiến: 0   Không lấy: 0")
    a.open_content_review()
    assert reg["toplevels"] == []
    r, cls, sections, cands = _build(tmp_path, [slide_items])

    class FR:
        content_candidates = cands
        image_candidates = []
    class P:
        results = [FR()]
    a.ctl.processor = P()
    a.ctl.state = "idle"
    a.open_content_review()
    assert len(reg["toplevels"]) == 1
    labels = [w.k.get("text") for w in reg["widgets"] if isinstance(w, reg["classes"]["Button"])]
    for t in ("Nội dung cải tiến", "Không lấy", "Mục trước", "Mục tiếp", "Lưu xác nhận"):
        assert t in labels
    assert "Management Number" in a.creview_head.cfg["text"]
    info = "".join(a.creview_info.lines)
    assert "Kết quả:" in info and "Tiêu đề slide:" in info and "Vị trí trên slide" in info
    assert a.creview_text.lines and a.btn_creview_save.cfg["state"] == "disabled"
    # hard-excluded title/sidebar/footer are not offered
    assert all(not c.hard_excluded for c in a.creview_cands)
    a._creview_set_label("EXCLUDE_CONTENT")
    assert a.btn_creview_save.cfg["state"] == "normal" and "(chưa lưu)" in "".join(a.creview_info.lines)
    a._creview_move(1)
    assert a.lbl_creview_pos.cfg["text"].startswith("Mục 2/")
    a._creview_save()
    assert a.ctl.learning.content.counts()["exclude"] == 1
    assert a.lbl_learning_content.cfg["text"].startswith("Nội dung:  Nội dung cải tiến: 0   Không lấy: 1")
    with pytest.raises(ValueError):
        a.ctl.set_pending_content_label(cands[0], "AFTER")


def test_updater_preserves_content_learning_files(tmp_path, monkeypatch):
    from tests.test_updater import make_portable, make_update_folder
    portable = make_portable(tmp_path / "app")
    ld = portable / "learning_data"
    (ld / "model").mkdir(parents=True)
    (ld / "content_labels.jsonl").write_text("{}\n", encoding="utf-8")
    (ld / "image_labels.jsonl").write_text("{}\n", encoding="utf-8")
    (ld / "model" / "content_model.json").write_text("{}", encoding="utf-8")
    folder = make_update_folder(tmp_path / "Update")
    check = up.check_for_update(str(folder), current_build=4)
    staged = up.stage_update(check, portable)
    assert up.apply_update(Path(staged.manifest_path), alive=lambda pid: False, sleep=lambda s: None,
                           spawn=lambda cmd, **kw: None) == 0
    assert (ld / "content_labels.jsonl").exists() and (ld / "image_labels.jsonl").exists()
    assert (ld / "model" / "content_model.json").exists()
