"""Phase 4 – production workflow: Management Number from file name -> existing row,
Ngày phát sinh derived from the Management Number, Vendor from the improvement text.
"""
import datetime as dt
from pathlib import Path

from openpyxl import load_workbook

from app.batch_processor import BatchOptions, BatchProcessor
from app.classifier import heuristic_classify
from app.excel_writer import ExcelWriter
from app.extractor import (derive_occurrence_date, extract_record, find_vendor_candidates,
                           management_number_from_filename)
from app.pptx_parser import parse_pptx

REAL_NAME = "(CTMS)_11107_251119092-VOC_ Đối sách lỗi xước Rear A185 17.11.2025.pptx"


# ---------------------------------------------------------------- management number / date
def test_management_number_from_real_filename():
    assert management_number_from_filename(REAL_NAME) == "251119092-VOC"
    assert management_number_from_filename("260918080-VOC_report2.pptx") == "260918080-VOC"
    assert not management_number_from_filename("report3.pptx")
    assert not management_number_from_filename("A185 17.11.2025.pptx")               # a date is not a key


def test_management_number_never_taken_from_slide_text(report_factory):
    # slide text says 260918080-VOC; the file name has a different key -> file name wins
    p = report_factory("(CTMS)_1_251119092-VOC_x.pptx", mgmt="260918080-VOC")
    rec = extract_record(parse_pptx(p), heuristic_classify(parse_pptx(p)))
    assert rec.management_number == "251119092-VOC"
    # no key in the file name -> blank + review, even though the slide contains one
    p2 = report_factory("noname.pptx", mgmt="260918080-VOC")
    rec2 = extract_record(parse_pptx(p2), heuristic_classify(parse_pptx(p2)))
    assert rec2.management_number == ""
    assert "Không xác định được Management Number từ tên file" in rec2.review_reasons


def test_occurrence_date_from_management_number():
    assert derive_occurrence_date("251119092-VOC") == (dt.date(2025, 11, 19), "")
    assert derive_occurrence_date("260918080-VOC") == (dt.date(2026, 9, 18), "")
    d, why = derive_occurrence_date("251319092-VOC")          # month 13
    assert d is None and why
    d, why = derive_occurrence_date("250230001-VOC")          # 30 Feb
    assert d is None and why
    assert derive_occurrence_date("")[0] is None


def test_occurrence_date_is_not_filename_date(report_factory):
    p = report_factory(REAL_NAME, mgmt="251119092-VOC")
    rec = extract_record(parse_pptx(p), heuristic_classify(parse_pptx(p)))
    assert rec.occurrence_date == dt.date(2025, 11, 19)
    assert rec.occurrence_date_text == "19/11/2025"          # not 17.11.2025 from the name


# ---------------------------------------------------------------- vendor
def test_vendor_phrase_candidates_are_raw_source_spelling():
    # raw 'công đoạn assy X' names (used only to report unknown candidates)
    assert find_vendor_candidates("Tại công đoạn assy Taewon, thao tác lấy hàng ...") == ["Taewon"]
    assert find_vendor_candidates("Công đoạn Assy Sung Kwang Vina thực hiện") == ["Sung Kwang Vina"]
    assert find_vendor_candidates("Vendor: CÔNG TY TNHH ABC VINA\nNhà cung cấp Taewon") == []
    assert find_vendor_candidates("Tại công đoạn lắp ráp Rear, thao tác ...") == []


def test_vendor_from_improvement_not_from_cover(a185_report):
    rep = parse_pptx(a185_report)
    rec = extract_record(rep, heuristic_classify(rep))
    assert rec.vendors == ["Doaltech"]                 # "Công đoạn Assy Daoltech" -> canonical spelling
    assert rec.vendor == "Doaltech"
    assert "ABC VINA" not in rec.vendor                # cover 'Vendor:' line is never used
    assert not any("Vendor" in r for r in rec.review_reasons)


def test_vendor_conflict_and_missing(report_factory):
    from types import SimpleNamespace
    from app import extractor as ex
    empty = SimpleNamespace(all_text=lambda: "")
    rec = ex.ExtractedRecord(management_number="251119092-VOC",
                             improvement="Tại công đoạn assy Taewon ... Tại công đoạn assy Daoltech")
    ex.extract_vendor(rec, empty)
    assert rec.vendors == ["Teawon", "Doaltech"]       # several vendors are normal, no conflict
    assert rec.vendor == "Teawon\nDoaltech" and rec.review_reasons == []
    rec2 = ex.ExtractedRecord(management_number="251119092-VOC", improvement="Không có cụm từ nào")
    ex.extract_vendor(rec2, empty)
    assert rec2.vendor == "" and rec2.vendors == []
    assert "Không xác định được Vendor từ nội dung báo cáo" in rec2.review_reasons


# ---------------------------------------------------------------- excel: locate existing row
def _prefill(template: Path, rows):
    """Put pre-existing Management Numbers (and optional vendor/date) in the template's data rows."""
    wb = load_workbook(template)
    ws = wb["Kiểm chứng"]
    for i, (mgmt, vendor, date) in enumerate(rows):
        r = 4 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=mgmt)
        if vendor is not None:
            ws.cell(row=r, column=3, value=vendor)
        if date is not None:
            ws.cell(row=r, column=4, value=date)
        for k in range(12, 20):
            ws.cell(row=r, column=k, value="OK")          # WEEK cells must stay untouched
    wb.save(template)
    return template


def test_find_rows_exact_match_only(template, tmp_path):
    _prefill(template, [("251119092-VOC", None, None), (" 260918080-voc ", None, None),
                        ("260918080-VOC", None, None)])
    w = ExcelWriter(template, tmp_path / "o.xlsx")
    assert w.find_rows_by_management_number("251119092-VOC") == [4]
    assert w.find_rows_by_management_number("260918080-VOC") == [5, 6]        # trim + case-insensitive
    assert w.find_rows_by_management_number("25111909-VOC") == []            # no fuzzy / prefix matching
    assert w.find_rows_by_management_number("") == []


def _run(files, template, out, **kw):
    events = []
    opts = BatchOptions(files=[Path(f) for f in files], template=template, output_file=out,
                        use_ollama=False, **kw)
    proc = BatchProcessor(opts, on_file=lambda i, s, d: events.append((i, s, d)))
    return proc.run(), events, proc


def test_match_mode_updates_existing_row_and_preserves_key(sample_tree, template, tmp_path):
    _prefill(template, [("999999999-VOC", None, None), ("260918080-VOC", None, None)])
    out = tmp_path / "out" / "k.xlsx"
    summary, events, proc = _run([sample_tree["files"][0]], template, out)
    assert summary.completed == 1 and summary.not_written == 0 and summary.needs_review == 0
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=5, column=2).value == "260918080-VOC"                  # key untouched
    assert ws.cell(row=5, column=5).value == "A185" and ws.cell(row=5, column=6).value == "Rear"
    assert ws.cell(row=5, column=3).value == "Doaltech"
    d = ws.cell(row=5, column=4).value
    assert isinstance(d, dt.datetime) and d.date() == dt.date(2026, 9, 18)
    assert ws.cell(row=5, column=4).number_format == "DD/MM/YYYY"
    assert "checklist đầu ca" in ws.cell(row=5, column=10).value   # long-term content (title excluded)
    assert all(ws.cell(row=5, column=k).value == "OK" for k in range(12, 20))  # WEEK untouched
    # row 4 (other key) and no new row appended
    assert ws.cell(row=4, column=5).value is None and ws.cell(row=6, column=2).value is None
    assert proc.results[0].excel_row == 5


def test_match_mode_not_found_duplicate_and_missing_key(sample_tree, template, tmp_path, report_factory):
    _prefill(template, [("260918081-VOC", None, None), ("260918081-VOC", None, None)])
    nokey = report_factory("bao_cao.pptx")
    files = [sample_tree["files"][0], sample_tree["files"][1], nokey]      # absent (-> new row) / duplicate / no key
    out = tmp_path / "out" / "k.xlsx"
    summary, events, proc = _run(files, template, out)
    assert summary.not_written == 1 and summary.completed == 1 and summary.needs_review == 1 and summary.new_rows == 1
    msgs = [r.error for r in proc.results]
    assert proc.results[0].new_row and proc.results[0].excel_row == 6                 # appended after the table
    assert "Không xác định được Management Number từ tên file" in msgs[2]
    # duplicate key: topmost row is the canonical destination, the other row stays untouched but is marked red
    dup = proc.results[1]
    assert dup.status == "needs_review" and dup.excel_row == 4
    assert any("xuất hiện 2 dòng (4, 5)" in r and "tô đỏ" in r for r in dup.review_reasons)
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=4, column=5).value == "A185" and ws.cell(row=5, column=5).value is None
    assert ws.cell(row=5, column=2).fill.fgColor.rgb.endswith("FFC7CE") and not str(ws.cell(row=4, column=2).fill.fgColor.rgb).endswith("FFC7CE")
    assert ws.cell(row=6, column=2).value == "260918080-VOC" and ws.cell(row=6, column=5).value == "A185"
    for r in (7, 8):                                                       # nothing written anywhere else
        assert ws.cell(row=r, column=2).value is None and ws.cell(row=r, column=10).value is None


def test_vendor_and_date_conflicts_keep_existing_value(sample_tree, template, tmp_path):
    _prefill(template, [("260918080-VOC", "Teawon", dt.datetime(2026, 9, 1))])
    out = tmp_path / "out" / "k.xlsx"
    summary, events, proc = _run([sample_tree["files"][0]], template, out)
    assert summary.needs_review == 1
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=4, column=3).value == "Teawon"                        # kept, not overwritten
    assert ws.cell(row=4, column=4).value == dt.datetime(2026, 9, 1)
    reasons = " ".join(proc.results[0].review_reasons)
    assert "Vendor trong Excel: Teawon" in reasons and "Vendor trong báo cáo: Doaltech" in reasons
    assert "01/09/2026" in reasons and "18/09/2026" in reasons


def test_vendor_and_date_same_value_no_review(sample_tree, template, tmp_path):
    _prefill(template, [("260918080-VOC", "daoltech ", "18/09/2026")])
    out = tmp_path / "out" / "k.xlsx"
    summary, events, proc = _run([sample_tree["files"][0]], template, out)
    assert summary.completed == 1
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=4, column=3).value == "daoltech "                     # existing cell untouched
    assert ws.cell(row=4, column=4).value == "18/09/2026"


def test_data_sheet_is_never_used(sample_tree, template, tmp_path):
    """A 'Data' sheet with Defect details / vendor / dates must not feed any field."""
    wb = load_workbook(template)
    ds = wb.create_sheet("Data")
    ds.append(["Management number", "Defect details", "Vendor", "Date"])
    ds.append(["260918080-VOC", "DATA SHEET DEFECT TEXT", "DataVendor", "01/01/2020"])
    wb.save(template)
    _prefill(template, [("260918080-VOC", None, None)])
    out = tmp_path / "out" / "k.xlsx"
    summary, _, proc = _run([sample_tree["files"][0]], template, out)
    assert summary.completed == 1
    ws = load_workbook(out)["Kiểm chứng"]
    for c in range(1, 20):
        v = ws.cell(row=4, column=c).value
        assert v is None or ("DATA SHEET" not in str(v) and "DataVendor" not in str(v))
    assert ws.cell(row=4, column=3).value == "Doaltech"
    assert ws.cell(row=4, column=4).value.date() == dt.date(2026, 9, 18)
    # the Data sheet itself is left as is
    assert load_workbook(out)["Data"]["B2"].value == "DATA SHEET DEFECT TEXT"


def test_append_mode_still_available(sample_tree, tmp_path):
    out = tmp_path / "out" / "k.xlsx"
    summary, _, _ = _run([sample_tree["files"][0]], sample_tree["template"], out, row_mode="append")
    assert summary.completed == 1
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=4, column=2).value == "260918080-VOC" and ws.cell(row=4, column=3).value == "Doaltech"




# ---------------------------------------------------------------- both Management Number formats
def test_management_number_both_formats_realistic_filenames():
    assert management_number_from_filename("(CTMS)_11107_260601038-VOC_ Đối sách lỗi xước Rear A185 05.06.2026.pptx") == "260601038-VOC"
    assert management_number_from_filename("(CTMS)_11108_260601017_ Đối sách lỗi mẻ Main A175 05.06.2026.pptx") == "260601017"
    assert management_number_from_filename("260601038-VOC.pptx") == "260601038-VOC"
    assert management_number_from_filename("260601017.pptx") == "260601017"
    assert management_number_from_filename(r"D:\real_data\reports\260601017 - A185.pptx") == "260601017"
    assert management_number_from_filename("Đối sách 260601017-VOC (rev2).pptx") == "260601017-VOC"   # suffix kept


def test_management_number_negative_tokens():
    # CTMS prefix, file-name dates (dotted / compact / ISO), model numbers, timestamps: never a key
    for name in ["(CTMS)_11107_ Đối sách A185 17.11.2025.pptx",
                 "17112025_A185_Rear.pptx",
                 "20260601_report.pptx",
                 "A185_2026-06-01.pptx",
                 "SM-A185_20251117123456.pptx",          # 14-digit timestamp
                 "12345678-VOC.pptx",                     # only 8 digits
                 "abc260601017.pptx",                     # glued to letters -> not delimited
                 "991399001.pptx",                        # 9 digits but month 13 -> not YYMMDD
                 "report3.pptx"]:
        assert management_number_from_filename(name) == "", name


def test_occurrence_date_both_formats():
    assert derive_occurrence_date("260601017") == (dt.date(2026, 6, 1), "")
    assert derive_occurrence_date("260601038-VOC") == (dt.date(2026, 6, 1), "")
    assert derive_occurrence_date("260632017")[0] is None


def test_bare_key_matches_excel_row_without_stripping_suffix(sample_tree, template, tmp_path, report_factory):
    # Excel stores the bare key as a NUMBER (typical when typed by hand) and the suffixed one as text
    _prefill(template, [(260601017, None, None), ("260601038-VOC", None, None), ("260601017-VOC", None, None)])
    w = ExcelWriter(template, tmp_path / "o.xlsx")
    assert w.find_rows_by_management_number("260601017") == [4]            # not row 6 (suffix differs)
    assert w.find_rows_by_management_number("260601038-VOC") == [5]
    assert w.find_rows_by_management_number("260601038") == []             # bare key never matches suffixed row
    # end-to-end: bare-number file name -> its own row, date derived 01/06/2026
    p = report_factory("(CTMS)_11108_260601017_ Đối sách A185 05.06.2026.pptx", mgmt="260601017")
    summary, _, proc = _run([p], template, tmp_path / "out" / "k.xlsx")
    assert summary.completed == 1 and proc.results[0].excel_row == 4
    ws = load_workbook(tmp_path / "out" / "k.xlsx")["Kiểm chứng"]
    assert ws.cell(row=4, column=2).value == 260601017                     # key cell untouched
    assert ws.cell(row=4, column=4).value.date() == dt.date(2026, 6, 1)
    assert ws.cell(row=6, column=5).value is None


# ---------------------------------------------------------------- destination sheet found by structure
import pytest  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from app.excel_writer import TemplateError, validate_template  # noqa: E402

DEST_HEADERS = ["STT", "Management number", "Tên vendor", "Ngày phát sinh", "Model", "Item",
                "Nội dung lỗi", "QPN", "Nguyên nhân", "Nội dung đối sách cải tiến", "Hình ảnh cải tiến"]


def _dest_sheet(wb, title, mgmt="260918080-VOC", weeks=True):
    ws = wb.create_sheet(title)
    ws["A1"] = "BẢNG KIỂM CHỨNG"
    for i, h in enumerate(DEST_HEADERS, start=1):
        ws.cell(row=2, column=i, value=h)
        ws.merge_cells(start_row=2, start_column=i, end_row=3, end_column=i)
    if weeks:
        ws.merge_cells(start_row=2, start_column=12, end_row=2, end_column=19)
        ws.cell(row=2, column=12, value="Theo dõi cải tiến")
        for k in range(1, 9):
            ws.cell(row=3, column=11 + k, value=f"WEEK +{k}")
    ws.cell(row=4, column=1, value=1)
    ws.cell(row=4, column=2, value=mgmt)
    for k in range(12, 20):
        ws.cell(row=4, column=k, value="OK")
    return ws


def _data_sheet(wb):
    ds = wb.create_sheet("Data")
    ds.append(["Management number", "Defect details", "Vendor", "Model", "Item", "Date"])
    ds.append(["260918080-VOC", "DATA SHEET DEFECT TEXT", "DataVendor", "A185", "Rear", "01/01/2020"])
    return ds


def _workbook(path, dest_titles, with_data=True, first_title=None):
    wb = Workbook()
    wb.active.title = first_title or "Trang tính1"          # an unrelated first sheet
    wb.active["A1"] = "ghi chú"
    if with_data:
        _data_sheet(wb)
    for t in dest_titles:
        _dest_sheet(wb, t)
    m = wb.create_sheet("Phân loại")
    m.append(["Model", "Item", "Từ khóa"])
    m.append(["A185", "Rear", "rear; nắp lưng"])
    wb.save(path)
    return path


@pytest.mark.parametrize("title", ["Sheet1", "Form", "Mau", "Bảng theo dõi đối sách lỗi"])
def test_destination_sheet_found_by_structure_regardless_of_name(sample_tree, tmp_path, title):
    tpl = _workbook(tmp_path / "t.xlsx", [title])
    assert "Kiểm chứng" not in load_workbook(tpl).sheetnames
    ok, msg = validate_template(tpl)
    assert ok and title in msg
    w = ExcelWriter(tpl, tmp_path / "probe.xlsx")
    assert w.ws.title == title
    assert {"management_number", "qpn", "root_cause", "improvement", "week_1", "week_8"} <= set(w.columns)
    out = tmp_path / "out" / "k.xlsx"
    summary, _, proc = _run([sample_tree["files"][0]], tpl, out)
    assert summary.completed == 1 and proc.results[0].excel_row == 4
    ws = load_workbook(out)[title]
    assert ws.cell(row=4, column=2).value == "260918080-VOC"
    assert ws.cell(row=4, column=5).value == "A185" and ws.cell(row=4, column=6).value == "Rear"
    assert ws.cell(row=4, column=3).value == "Doaltech"
    assert ws.cell(row=4, column=4).value.date() == dt.date(2026, 9, 18)
    assert "checklist đầu ca" in ws.cell(row=4, column=10).value
    assert all(ws.cell(row=4, column=k).value == "OK" for k in range(12, 20))
    assert ws.cell(row=5, column=2).value is None                     # no appended row


def test_data_sheet_never_chosen_even_when_it_is_the_only_match(sample_tree, tmp_path):
    # Data + one valid destination -> destination wins, Data untouched, Data values never copied
    tpl = _workbook(tmp_path / "t.xlsx", ["Form"])
    out = tmp_path / "out" / "k.xlsx"
    summary, _, _ = _run([sample_tree["files"][0]], tpl, out)
    assert summary.completed == 1
    wb = load_workbook(out)
    assert wb["Data"]["B2"].value == "DATA SHEET DEFECT TEXT" and wb["Data"]["C2"].value == "DataVendor"
    ws = wb["Form"]
    for c in range(1, 20):
        v = ws.cell(row=4, column=c).value
        assert v is None or ("DATA SHEET" not in str(v) and "DataVendor" not in str(v))
    # a 'Data' sheet that even carries the full destination header set is still excluded
    wb2 = Workbook()
    wb2.active.title = "Trang tính1"
    _dest_sheet(wb2, "Data")
    p2 = tmp_path / "only_data.xlsx"
    wb2.save(p2)
    with pytest.raises(TemplateError) as ei:
        ExcelWriter(p2, tmp_path / "o2.xlsx")
    assert "Không tìm thấy sheet chứa bảng kiểm chứng" in str(ei.value)
    assert not validate_template(p2)[0]


def test_multiple_destination_sheets_are_not_guessed(sample_tree, tmp_path):
    tpl = _workbook(tmp_path / "t.xlsx", ["Kiểm chứng", "Form"])
    with pytest.raises(TemplateError) as ei:
        ExcelWriter(tpl, tmp_path / "o.xlsx")
    msg = str(ei.value)
    assert msg.startswith("Cần kiểm tra") and "'Kiểm chứng'" in msg and "'Form'" in msg
    ok, vmsg = validate_template(tpl)
    assert not ok and "Cần kiểm tra" in vmsg and "'Form'" in vmsg
    # batch: nothing written, every file reported with the reason
    out = tmp_path / "out" / "k.xlsx"
    summary, events, _ = _run([sample_tree["files"][0]], tpl, out)
    assert summary.failed == 1 and summary.completed == 0
    assert "Cần kiểm tra" in summary.errors[0]["error"]
    assert all(s == "error" for i, s, d in events)


def test_no_destination_sheet_gives_clear_diagnostic(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Kiểm chứng"                                  # right name, wrong structure
    for i, h in enumerate(["STT", "Model", "Item", "Nội dung lỗi"], start=1):
        ws.cell(row=1, column=i, value=h)
    _data_sheet(wb)
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    with pytest.raises(TemplateError) as ei:
        ExcelWriter(p, tmp_path / "o.xlsx")
    msg = str(ei.value)
    assert "Không tìm thấy sheet chứa bảng kiểm chứng" in msg
    assert "'Kiểm chứng': thiếu cột" in msg
    for label in ("Management number", "QPN", "Nguyên nhân", "Nội dung đối sách cải tiến"):
        assert label in msg
    ok, vmsg = validate_template(p)
    assert not ok and "QPN" in vmsg


# ---------------------------------------------------------------- controlled multi-value Vendor
from types import SimpleNamespace  # noqa: E402

from app import extractor as ex  # noqa: E402
from app.extractor import CANONICAL_VENDORS, canonical_vendors, canonicalize_vendor_value  # noqa: E402


def _vendor_rec(improvement, other=""):
    rec = ex.ExtractedRecord(management_number="260601038-VOC", improvement=improvement)
    ex.extract_vendor(rec, SimpleNamespace(all_text=lambda: other))
    return rec


def test_it_alone_is_not_a_vendor():
    assert canonical_vendors("Bộ phận IT kiểm tra; công đoạn IT; ITEM; IT.") == []
    rec = _vendor_rec("Tại công đoạn IT thao tác lại")
    assert rec.vendors == [] and rec.vendor == ""
    assert "IT" not in rec.vendor_candidates
    assert "Không xác định được Vendor từ nội dung báo cáo" in rec.review_reasons


def test_it_vendors_keep_their_prefix():
    # the real 260601038-VOC case: 'Assy IT' must never be reduced to 'IT'
    assert canonical_vendors("Tại công đoạn Assy IT, thao tác lấy hàng") == ["Assy IT"]
    assert canonical_vendors("Công đoạn CNC IT") == ["CNC IT"]
    assert canonical_vendors("Công đoạn Sơn IT") == ["Sơn IT"]
    assert canonical_vendors("Công đoạn Son IT") == ["Sơn IT"]          # accent-tolerant, canonical output
    assert canonical_vendors("Công đoạn Nhựa IT") == ["Nhựa IT"]
    assert canonical_vendors("công đoạn  assy   it") == ["Assy IT"]     # case / whitespace tolerant
    rec = _vendor_rec("Tại công đoạn Assy IT thao tác")
    assert rec.vendors == ["Assy IT"] and rec.review_reasons == []      # no 'unknown candidate IT' noise


def test_mtech_variants_stay_distinct():
    assert canonical_vendors("Vendor Mtech") == ["Mtech"]
    assert canonical_vendors("Vendor Mtech VN") == ["Mtech VN"]
    assert canonical_vendors("Vendor Mtech 4") == ["Mtech 4"]
    assert canonical_vendors("Mtech VN và Mtech 4 và Mtech") == ["Mtech VN", "Mtech 4", "Mtech"]


def test_all_canonical_vendors_recognised():
    text = "Tesung; Doaltech; APV; An Lập; Yongsong; HK; Teawon; công đoạn assy Taewon"
    assert canonical_vendors(text) == ["Tesung", "Doaltech", "APV", "An Lập", "Yongsong", "HK", "Teawon"]
    for v in CANONICAL_VENDORS:
        assert canonical_vendors(f"Tại công đoạn {v} thao tác") == [v]
        assert canonical_vendors(f"tại công đoạn {v.lower()} thao tác") == [v]
        assert canonical_vendors(f"TẠI CÔNG ĐOẠN {v.upper()} THAO TÁC") == [v]


def test_multiple_vendors_dedupe_and_source_order():
    text = ("Công đoạn Sơn IT: kiểm tra bề mặt\nTại công đoạn Assy IT: thay khay\n"
            "Vendor Mtech bổ sung film\nSơn IT kiểm tra lại\nAssy IT xác nhận")
    rec = _vendor_rec(text)
    assert rec.vendors == ["Sơn IT", "Assy IT", "Mtech"]
    assert rec.vendor == "Sơn IT\nAssy IT\nMtech"                        # one per line
    assert not any("Vendor" in r for r in rec.review_reasons)           # multiple vendors are normal


def test_unknown_vendor_candidate_not_invented():
    rec = _vendor_rec("Tại công đoạn assy Sung Kwang Vina thao tác; công đoạn Assy IT kiểm tra")
    assert rec.vendors == ["Assy IT"]                                   # unknown name never written
    assert any("Sung Kwang Vina" in r and "không có trong danh sách chuẩn" in r for r in rec.review_reasons)
    rec2 = _vendor_rec("Tại công đoạn assy Sung Kwang Vina thao tác")
    assert rec2.vendors == [] and "Sung Kwang Vina" in rec2.vendor_candidates
    assert "Không xác định được Vendor từ nội dung báo cáo" in rec2.review_reasons
    assert any("Sung Kwang Vina" in r for r in rec2.review_reasons)


def test_excel_vendor_cell_normalisation():
    assert canonicalize_vendor_value("Mtech / Tesung") == (["Mtech", "Tesung"], [])
    assert canonicalize_vendor_value("Mtech\nAssy IT; tesung, son it") == (["Mtech", "Assy IT", "Tesung", "Sơn IT"], [])
    assert canonicalize_vendor_value("Công ty ABC") == ([], ["Công ty ABC"])
    assert canonicalize_vendor_value("") == ([], [])


def test_multi_vendor_written_one_per_line_with_wrap(sample_tree, template, tmp_path, report_factory):
    _prefill(template, [("260918080-VOC", None, None)])
    out = tmp_path / "out" / "k.xlsx"
    summary, _, _ = _run([sample_tree["files"][0]], template, out)      # sample has 'Công đoạn Assy Daoltech'
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=4, column=3).value == "Doaltech" and ws.cell(row=4, column=3).alignment.wrap_text
    # a multi-vendor record through the writer directly
    w = ExcelWriter(template, tmp_path / "w.xlsx")
    rec = ex.ExtractedRecord(management_number="260918080-VOC", vendors=["Sơn IT", "Assy IT", "Mtech"],
                             vendor="Sơn IT\nAssy IT\nMtech", model="A185", item="Rear")
    notes = w.update_record(4, rec)
    w.save()
    c = load_workbook(tmp_path / "w.xlsx")["Kiểm chứng"].cell(row=4, column=3)
    assert c.value == "Sơn IT\nAssy IT\nMtech" and c.alignment.wrap_text and notes == []


def test_equivalent_multi_vendor_sets_do_not_conflict(template, tmp_path):
    _prefill(template, [("260601038-VOC", "tesung; mtech", None), ("260601017", "Mtech\nTesung", None)])
    w = ExcelWriter(template, tmp_path / "w.xlsx")
    rec = ex.ExtractedRecord(management_number="260601038-VOC", vendors=["Mtech", "Tesung"], vendor="Mtech\nTesung")
    assert w.update_record(4, rec) == []
    assert w.update_record(5, rec) == []
    w.save()
    ws = load_workbook(tmp_path / "w.xlsx")["Kiểm chứng"]
    assert ws.cell(row=4, column=3).value == "tesung; mtech"            # existing cell preserved verbatim
    assert ws.cell(row=5, column=3).value == "Mtech\nTesung"


def test_different_multi_vendor_sets_preserve_excel_and_review(template, tmp_path):
    _prefill(template, [("260601038-VOC", "Mtech / Tesung", None)])
    w = ExcelWriter(template, tmp_path / "w.xlsx")
    rec = ex.ExtractedRecord(management_number="260601038-VOC", vendors=["Mtech", "Assy IT", "Tesung"],
                             vendor="Mtech\nAssy IT\nTesung")
    notes = w.update_record(4, rec)
    w.save()
    assert len(notes) == 1
    assert "Vendor trong Excel: Mtech / Tesung" in notes[0]
    assert "Vendor trong báo cáo: Mtech / Assy IT / Tesung" in notes[0]
    assert load_workbook(tmp_path / "w.xlsx")["Kiểm chứng"].cell(row=4, column=3).value == "Mtech / Tesung"


def test_missing_qpn_and_defect_stay_blank_with_review(report_factory, template, tmp_path):
    p = report_factory("(CTMS)_1_260601017_ Đối sách A185.pptx", mgmt="260601017", with_qpn=False)
    rep = parse_pptx(p)
    rec = extract_record(rep, heuristic_classify(rep))
    assert rec.qpn_slide in (None, 0) or "Không tìm thấy QPN trong báo cáo" in rec.review_reasons
    assert "Không tìm thấy QPN trong báo cáo" in rec.review_reasons
    assert rec.defect_content == "" and "Không tìm thấy Nội dung lỗi trong báo cáo" in rec.review_reasons
    assert "Nội dung lỗi" in rec.blank_fields
    # end-to-end: row written with blanks, status Cần kiểm tra, nothing guessed
    _prefill(template, [("260601017", None, None)])
    out = tmp_path / "out" / "k.xlsx"
    summary, _, proc = _run([p], template, out)
    assert summary.needs_review == 1 and summary.completed == 0
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=4, column=7).value in (None, "")                  # Nội dung lỗi blank
    assert ws.cell(row=4, column=5).value == "A185"                      # what exists is still written
    reasons = proc.results[0].review_reasons
    assert "Không tìm thấy QPN trong báo cáo" in reasons and "Không tìm thấy Nội dung lỗi trong báo cáo" in reasons


# ---------------------------------------------------------------- historical workbook aliases
def test_historical_aliases_map_to_canonical():
    assert canonical_vendors("Lắp ráp IT") == ["Assy IT"]
    assert canonical_vendors("lắp  ráp  it") == ["Assy IT"]               # case / whitespace tolerant
    assert canonical_vendors("Sơn Intops") == ["Sơn IT"]
    assert canonical_vendors("CNC Intops") == ["CNC IT"]
    assert canonical_vendors("NCC JT") == ["JT Tech Vina"]
    assert canonical_vendors("JT Tech Vina") == ["JT Tech Vina"]
    assert "JT Tech Vina" in CANONICAL_VENDORS
    # alias never yields a second / partial vendor
    assert canonical_vendors("Tại công đoạn lắp ráp IT thao tác") == ["Assy IT"]
    assert canonical_vendors("Tại công đoạn lắp ráp Rear") == []


def test_jt_vendor_detected_in_report_text():
    rec = _vendor_rec("Tại công đoạn Assy JT Tech Vina kiểm tra lại; NCC JT bổ sung khay")
    assert rec.vendors == ["JT Tech Vina"] and rec.review_reasons == []


def test_excel_alias_cells_normalise_for_comparison():
    assert canonicalize_vendor_value("Mtech\nTesung\nLắp ráp IT") == (["Mtech", "Tesung", "Assy IT"], [])
    assert canonicalize_vendor_value("Sơn Intops\nCNC Intops") == (["Sơn IT", "CNC IT"], [])
    assert canonicalize_vendor_value("NCC JT") == (["JT Tech Vina"], [])
    assert canonicalize_vendor_value("Công ty Mtech Việt Nam") == ([], ["Công ty Mtech Việt Nam"])   # a sentence is not a vendor list


def _upd(template, tmp_path, excel_vendor, ppt_vendors, name="w.xlsx"):
    _prefill(template, [("260601038-VOC", excel_vendor, None)])
    w = ExcelWriter(template, tmp_path / name)
    rec = ex.ExtractedRecord(management_number="260601038-VOC", vendors=list(ppt_vendors),
                             vendor="\n".join(ppt_vendors))
    notes = w.update_record(4, rec)
    w.save()
    return notes, load_workbook(tmp_path / name)["Kiểm chứng"].cell(row=4, column=3).value


def test_alias_sets_equivalent_no_conflict_cell_preserved(template, tmp_path):
    notes, cell = _upd(template, tmp_path, "Mtech\nTesung\nLắp ráp IT", ["Mtech", "Tesung", "Assy IT"])
    assert notes == [] and cell == "Mtech\nTesung\nLắp ráp IT"           # historical spelling kept verbatim
    notes, cell = _upd(template, tmp_path, "NCC JT", ["JT Tech Vina"], "w2.xlsx")
    assert notes == [] and cell == "NCC JT"
    notes, cell = _upd(template, tmp_path, "Sơn Intops\nCNC Intops", ["Sơn IT", "CNC IT"], "w3.xlsx")
    assert notes == [] and cell == "Sơn Intops\nCNC Intops"


def test_order_difference_is_not_a_conflict(template, tmp_path):
    notes, cell = _upd(template, tmp_path, "Assy IT; Tesung, mtech", ["Mtech", "Tesung", "Assy IT"])
    assert notes == [] and cell == "Assy IT; Tesung, mtech"


def test_genuine_difference_still_conflicts_despite_aliases(template, tmp_path):
    notes, cell = _upd(template, tmp_path, "Sơn Intops\nCNC Intops", ["HK", "Sơn IT"])
    assert cell == "Sơn Intops\nCNC Intops"                                # Excel kept
    assert len(notes) == 1
    assert "Vendor trong Excel: Sơn IT / CNC IT" in notes[0]
    assert "Vendor trong báo cáo: HK / Sơn IT" in notes[0]


def test_blank_excel_vendor_gets_canonical_names(template, tmp_path):
    notes, cell = _upd(template, tmp_path, None, ["Mtech", "Tesung", "Assy IT"])
    assert notes == [] and cell == "Mtech\nTesung\nAssy IT"


# ---------------------------------------------------------------- Nội dung lỗi from the file name
from app.extractor import defect_content_from_filename  # noqa: E402


@pytest.mark.parametrize("name,expected", [
    ("(CTMS)_260923045_ĐỐI SÁCH CẢI TIẾN MODEL A185 REAR LỖI BONG ATN, MỤN 22.9.2026 SEV.pptx", "BONG ATN, MỤN"),
    ("(CTMS)_260924036_ĐỐI SÁCH CẢI TIẾN MODEL A185 REAR LỖI SƠN, MẺ, LỆCH TAPE 24.9.2026.pptx", "SƠN, MẺ, LỆCH TAPE"),
    ("(CTMS)_20506_260918080-VOC_ĐỐI SÁCH CẢI TIẾN MODEL A185 REAR LỖI XƯỚC, LỆCH ANT, MẺ 16.9.2026.pptx", "XƯỚC, LỆCH ANT, MẺ"),
    ("(CTMS)_260925011_ĐỐI SÁCH CẢI TIẾN MODEL A185 REAR LỖI BẨN KEO 25.9.2026 (Dreamtech).pptx", "BẨN KEO"),
    ("(CTMS)_260926002-VOC_ĐỐI SÁCH CẢI TIẾN MODEL A185 REAR LỖI NỨT (Dreamtech) 26.9.2026 SEV.pptx", "NỨT"),
    # older June pattern: "lỗi <defect> <Item> <Model> <date>"
    ("(CTMS)_11203_260601038-VOC_ Đối sách lỗi xước Rear A185 05.06.2026.pptx", "xước"),
    ("(CTMS)_11205_260601017_ Đối sách lỗi xước, mẻ Rear A185 05.06.2026.pptx", "xước, mẻ"),
    ("(CTMS)_11107_251119092-VOC_ Đối sách lỗi xước Rear A185 (Dreamtech) 17.11.2025.pptx", "xước"),
    # full Windows path is fine
    (r"D:\real_data\reports\(CTMS)_260923045_ĐỐI SÁCH CẢI TIẾN MODEL A185 REAR LỖI BONG ATN, MỤN 22.9.2026 SEV.pptx", "BONG ATN, MỤN"),
])
def test_defect_content_from_filename(name, expected):
    assert defect_content_from_filename(name) == expected


@pytest.mark.parametrize("name", [
    "(CTMS)_20506_260918080-VOC_A185_Rear.pptx",          # no marker
    "(CTMS)_1_260601017_ Đối sách A185.pptx",
    "LỖI 22.9.2026.pptx",                                  # marker but only a date after it
    "LỖI SEV.pptx",
    "report3.pptx",
    "Bao cao chat luong A185 Rear 05.06.2026.pptx",
])
def test_defect_content_from_filename_unreliable_is_blank(name):
    assert defect_content_from_filename(name) == ""


def test_defect_filename_never_adds_quantities():
    assert "ea" not in defect_content_from_filename("(CTMS)_260923045_ĐỐI SÁCH A185 REAR LỖI BONG ATN, MỤN 22.9.2026.pptx")


def test_defect_filename_primary_then_slide_text_fallback(report_factory):
    # file name has the marker -> file name wins over slide text (slide has "Xước: 15ea" table)
    p = report_factory("(CTMS)_20506_260918080-VOC_ĐỐI SÁCH CẢI TIẾN MODEL A185 REAR LỖI XƯỚC, LỆCH ANT, MẺ 16.9.2026.pptx")
    rep = parse_pptx(p)
    rec = extract_record(rep, heuristic_classify(rep))
    assert rec.defect_content == "XƯỚC, LỆCH ANT, MẺ"
    assert rec.management_number == "260918080-VOC" and rec.occurrence_date_text == "18/09/2026"   # not 16.9.2026
    # no marker -> existing slide-text extraction
    p2 = report_factory("(CTMS)_20506_260918080-VOC_A185_Rear.pptx")
    rep2 = parse_pptx(p2)
    rec2 = extract_record(rep2, heuristic_classify(rep2))
    assert "Xước: 15ea" in rec2.defect_content
    # neither -> blank + review
    p3 = report_factory("(CTMS)_1_260601017_ Bao cao A185.pptx", with_qpn=False)
    rep3 = parse_pptx(p3)
    rec3 = extract_record(rep3, heuristic_classify(rep3))
    assert rec3.defect_content == "" and "Không tìm thấy Nội dung lỗi trong báo cáo" in rec3.review_reasons
