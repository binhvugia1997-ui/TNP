"""PROMPT-025: multi-improvement segmentation + accurate After-evidence regions.

One slide may carry SEVERAL logical improvement items (Report -> Slide[] -> ImprovementItem[] ->
Before/After regions). Segmentation runs before evidence cropping on multiple signals (headings,
Trước/Sau markers, arrows, geometry, vertical gaps – no OCR, no screenshot/red-rectangle detection).
The real case: one slide with item #1 "Lỗi xước rear ..." (Before left, blue arrow middle, After =
THREE pictures right, "Sau cải進" caption beneath) and item #2 "Cải tiến lỗi lệch ATN ...".
Expected: exactly 2 ImprovementItems, 2 item-scoped AfterVisualRegions, 2 crops; item #1's crop must
NOT contain item #2's heading/body; 3 pictures in one After block = ONE region.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.batch_processor import BatchOptions, BatchProcessor
from app.classifier import heuristic_classify
from app.extractor import extract_record
from app.image_extractor import export_after_pictures
from app.image_learning import build_candidates
from app.improvement_items import (item_index_for_owner, production_item_regions,
                                   segment_report_items)
from app.improvement_pictures import (SEMANTIC_PRODUCTION, select_after_pictures, slide_items)
from app.improvement_visual import build_improvement_visual_regions
from app.pptx_parser import parse_pptx
from app.qpn_renderer import SlideRenderer
from tests.test_after_evidence import _deck, _pics, _prod_head, _tb
from tests.test_content_region import COL, SHEET
from tests.test_prompt004 import MGMT
from tests.test_prompt020_report_scoping import _image_has_rgb, _row_images
from tools.make_samples import make_template

NAME = "(CTMS)_20601_260920045-VOC_ Đối sách LỖI XƯỚC REAR, LỆCH ATN 24.9.2026.pptx"

ITEM1_HEAD = "Lỗi xước rear (Áp dụng cải tiến 17/9 – Công đoạn Assy Daoltech):"
ITEM1_BODY = ("- Cải tiến jig nén: thay vật liệu nhôm bằng thép SKD11\n"
              "- Bọc silicon 2mm toàn bộ mặt tiếp xúc của jig, bo tròn cạnh R1.0")
ITEM2_HEAD = "Cải tiến lỗi lệch ATN (Áp dụng cải tiến 20/9 – Công đoạn Gia công):"
ITEM2_BODY = ("- Cải tiến gá kẹp: cố định vị trí gá bằng chốt định vị\n"
              "- Thêm chốt định vị, gá cố định vị trí khi ép nhựa")
ITEM2_DEFECT_HEAD = "Lỗi lệch ATN sau ép nhựa"          # defect-style heading containing "sau"

BEFORE_RGB = (255, 224, 178)                            # #ffe0b2
AFTER1_RGB = (200, 230, 201)                            # #c8e6c9  (item #1 After)
AFTER2_RGB = (179, 229, 252)                            # #b3e5fc  (item #2 After)
ARROW_BLUE = RGBColor(0x00, 0x70, 0xC0)

EMU = 914400


# ------------------------------------------------------------------ fixture builders
def _caption(slide, text, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    shape.text_frame.text = text
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0x1F, 0x5F, 0xA0)
    shape.line.width = Pt(1.25)
    return shape


def _blue_arrow(slide, left, top, width, height):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = ARROW_BLUE
    arrow.line.fill.background()
    return arrow


def _item_text(slide, heading, body, top, head_h=0.4, body_h=0.6, bold_head=True, mid_block=False):
    """Item text: heading line + descriptive body (the Trước:/Sau: markers are the caption buttons
    of the visual row, per the reported structure).

    ``mid_block``: the colon-terminated heading is a NON-bold line inside a combined frame (real decks
    repeat the item-heading style inside body frames) – the hardest segmentation variant."""
    if mid_block:
        lines = [body.split("\n")[0], heading, *body.split("\n")[1:]]
        _tb(slide, "\n".join(lines), 1.7, top, 5.6, 1.05, 12, False)
        return
    _tb(slide, heading, 1.7, top, 5.6, head_h, 12, bold_head)
    _tb(slide, body, 1.7, top + head_h + 0.04, 5.6, body_h, 12, False)


def _visual_row(slide, top, n_before=2, n_after=3, after_color="#c8e6c9", below_caption=True):
    """Authored visual row: Before group left, blue transition arrow middle, After group right,
    optional 'Sau cải進' caption directly beneath the After group.  No containing rectangle."""
    _caption(slide, "Trước cải進", 1.7, top, 1.5, 0.32)
    _pics(slide, ["#ffe0b2"] * n_before, 1.7, top + 0.37, w=1.3, h=0.85, gap=0.1)
    _blue_arrow(slide, 4.75, top + 0.55, 1.4, 0.4)
    _caption(slide, "Sau cải進", 6.5, top, 1.5, 0.32)
    _pics(slide, [after_color] * n_after, 6.5, top + 0.37, w=1.3, h=0.85, gap=0.1)
    if below_caption:
        _caption(slide, "Sau cải進", 7.35, top + 1.32, 1.95, 0.34)


def _two_item_slide(item2_head=ITEM2_HEAD, mid_block_item2=False, n_after_item2=2):
    def build(s, W):
        _prod_head(s, W)
        _item_text(s, ITEM1_HEAD, ITEM1_BODY, 1.02)
        _visual_row(s, 2.38, n_before=2, n_after=3, after_color="#c8e6c9", below_caption=True)
        _item_text(s, item2_head, ITEM2_BODY, 4.45, mid_block=mid_block_item2)
        _visual_row(s, 5.55, n_before=1, n_after=n_after_item2, after_color="#b3e5fc",
                    below_caption=False)
    return build


def _three_item_slide():
    def build(s, W):
        _prod_head(s, W)
        _item_text(s, ITEM1_HEAD, ITEM1_BODY, 0.98, head_h=0.36, body_h=0.56)
        _visual_row(s, 1.98, n_before=1, n_after=2, after_color="#c8e6c9", below_caption=False)
        _item_text(s, ITEM2_HEAD, ITEM2_BODY, 3.28, head_h=0.36, body_h=0.56)
        _visual_row(s, 4.28, n_before=1, n_after=2, after_color="#b3e5fc", below_caption=False)
        _item_text(s, "Cải tiến lỗi mẻ xước gá ép (Áp dụng cải tiến 25/9 – Công đoạn Lắp ráp):",
                   "- Bổ sung chốt dẫn hướng\n+ Sau: Gá có dẫn hướng, không mẻ xước", 5.58,
                   head_h=0.36, body_h=0.56)
        _visual_row(s, 6.18, n_before=1, n_after=1, after_color="#ffe082", below_caption=False)
    return build


def _single_item_slide():
    def build(s, W):
        _prod_head(s, W)
        _item_text(s, ITEM1_HEAD, ITEM1_BODY, 1.02)
        _visual_row(s, 2.38, n_before=2, n_after=3, after_color="#c8e6c9", below_caption=True)
    return build


def _temporary_slide():
    def build(s, W):
        _tb(s, "2. XỬ LÝ TẠM THỜI", 0.5, 0.3, 8, 0.7, 24, True)
        _tb(s, "- Sorting 100% lô hàng tồn tại kho: 1.250 pcs\n- Tăng cường kiểm tra ngoại quan tại OQC",
           1.7, 1.1, 11, 1.0, 13)
        _pics(s, ["#fff2cc"], 1.7, 2.4, w=3, h=1.8)
    return build


# ------------------------------------------------------------------ helpers
def _report(path: Path):
    return parse_pptx(str(path))


def _selection(report):
    cls = heuristic_classify(report)
    slides = cls.improvement_image_slides or cls.improvement_slides
    return cls, slides, select_after_pictures(report, slides)


def _record(path: Path):
    report = _report(path)
    cls = heuristic_classify(report)
    return report, cls, extract_record(report, cls, {"rear": "Rear"}, ["A185"])


def _regions(report, sel):
    return build_improvement_visual_regions(report, sel.after, management_number=MGMT)


def _export(report, sel, out_dir, renderer=None, **kw):
    return export_after_pictures(report, sel.after, Path(out_dir), renderer=renderer,
                                 management_number=MGMT, **kw)


def _builtin_renderer():
    return SlideRenderer(prefer=("builtin",), width_px=1200)


def _crop_has_rgb(path, rgb, tolerance=4, minimum=15):
    with Image.open(path) as im:
        pixels = im.convert("RGB").tobytes()
    matches = sum(1 for i in range(0, len(pixels), 3)
                  if all(abs(pixels[i + c] - rgb[c]) <= tolerance for c in range(3)))
    return matches >= minimum


def _run_batch(files, template, output, **kw):
    processor = BatchProcessor(BatchOptions(
        files=[Path(f) for f in files], template=template, output_file=output,
        use_ollama=False, use_fast_cache=False, use_image_learning=False, **kw))
    summary = processor.run()
    assert summary.failed == 0
    return processor


# ================================================================== §30 base fixture: exactly 2 items / 2 regions / 2 crops
def test_two_items_segmented_two_regions_two_crops(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    report = _report(deck)
    s = report.slide(3)

    items = [it for it in slide_items(s) if it.region_kind == "item" and it.semantic_role == SEMANTIC_PRODUCTION]
    assert len(items) == 2
    assert items[0].heading == ITEM1_HEAD and items[1].heading == ITEM2_HEAD
    # item #1's span ends at item #2's heading (authored whitespace stays inside the gap)
    assert items[0].bounds[1] + items[0].bounds[3] <= items[1].bounds[1]

    cls, slides, sel = _selection(report)
    assert slides == [3]
    after_by_owner = {}
    for ref in sel.after:
        after_by_owner.setdefault(ref.owner_id, []).append(ref)
    assert len(after_by_owner) == 2
    owners = [it.owner_id for it in items]
    assert set(after_by_owner) == set(owners)
    assert len(after_by_owner[owners[0]]) == 3          # THREE After pictures of item #1
    assert len(after_by_owner[owners[1]]) == 2
    assert all(ref.excel_output_eligible and ref.confident_owner for ref in sel.after)
    befores = [ref for ref in sel.rejected if ref.kind == "before"]
    assert len(befores) == 3                            # 2 (item #1) + 1 (item #2)
    assert {ref.owner_id for ref in befores} == set(owners)

    regions = _regions(report, sel)
    assert len(regions) == 2
    assert [r.item_index for r in regions] == [0, 1]
    assert [r.improvement_item_id for r in regions] == owners
    assert all(r.after_block_index == 0 for r in regions)
    assert len(regions[0].pictures) == 3                # 3 pictures = ONE region (never 3)
    assert len(regions[1].pictures) == 2

    grouped, problems = _export(report, sel, tmp_path / "crops", renderer=_builtin_renderer())
    assert problems == []
    assert len(grouped.groups) == 2 and len(grouped) == 2
    assert all(len(group) == 1 for group in grouped.groups)
    assert len(set(grouped.owners)) == 2                # one Excel image row per item


# ================================================================== §34 next-item text never enters item #1's crop
def test_item1_crop_excludes_item2_heading_and_body(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    report = _report(deck)
    _, _, sel = _selection(report)
    regions = _regions(report, sel)
    r1, r2 = regions

    s = report.slide(3)
    item2_text_boxes = [b for b in s.blocks
                        if b.is_text and (ITEM2_HEAD in b.text or "gá kẹp" in b.text or "chốt định vị" in b.text)]
    assert item2_text_boxes
    # geometric guarantee: item #1's crop ends above every item #2 text box
    for box in item2_text_boxes:
        assert r1.bbox[3] <= box.top, (r1.bbox, box.top)
    # and no item #2 text object was absorbed into the region
    included_texts = [" ".join(b.text.split()) for b in r1.included_objects]
    assert not any(ITEM2_HEAD in t or "gá klem" in t or "chốt định vị" in t for t in included_texts)
    included_texts2 = [" ".join(b.text.split()) for b in r2.included_objects]
    assert not any(ITEM1_HEAD in t or "SKD11" in t for t in included_texts2)

    # pixel guarantee: crop 1 shows item #1's After colour only, crop 2 item #2's
    grouped, problems = _export(report, sel, tmp_path / "crops2", renderer=_builtin_renderer())
    assert problems == [] and len(grouped.groups) == 2
    assert _crop_has_rgb(grouped.groups[0][0], AFTER1_RGB)
    assert not _crop_has_rgb(grouped.groups[0][0], AFTER2_RGB)
    assert not _crop_has_rgb(grouped.groups[0][0], BEFORE_RGB)
    assert _crop_has_rgb(grouped.groups[1][0], AFTER2_RGB)
    assert not _crop_has_rgb(grouped.groups[1][0], AFTER1_RGB)
    assert not _crop_has_rgb(grouped.groups[1][0], BEFORE_RGB)


# ================================================================== §15 whitespace padding policy
def test_whitespace_padding_stays_between_items(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    report = _report(deck)
    _, _, sel = _selection(report)
    regions = _regions(report, sel)
    r1 = regions[0]
    s = report.slide(3)
    item2_head_box = next(b for b in s.blocks if b.is_text and ITEM2_HEAD in b.text)
    # authored whitespace: the padded crop may use the gap but must not touch item #2
    assert r1.bbox[3] <= item2_head_box.top
    # the crop still contains the full After group incl. the 'Sau cải進' caption beneath
    after_pics = [ref.block for ref in sel.after if ref.owner_id == r1.improvement_item_id]
    assert r1.bbox[1] <= min(p.top for p in after_pics)
    assert r1.bbox[3] >= max(p.bottom for p in after_pics)
    below_caption = next(b for b in s.blocks if b.is_text and " ".join(b.text.split()) == "Sau cải進"
                         and b.top > 3.5 * EMU)
    assert r1.bbox[3] >= below_caption.bottom


# ================================================================== §14 caption / annotation association
def test_caption_and_annotation_association_item_scoped(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    report = _report(deck)
    _, _, sel = _selection(report)
    regions = _regions(report, sel)
    r1, r2 = regions

    def texts(region):
        return [" ".join(b.text.split()) for b in region.included_objects if b.text.strip()]

    def caption_ids(region, caption_text):
        # a caption may appear twice (vector style record + its text block) – count distinct shapes
        return {obj.shape_id for obj in region.included_objects
                if " ".join(obj.text.split()) == caption_text}

    # item #1's region carries its own 'Sau cải進' captions (above + beneath the After group)
    assert len(caption_ids(r1, "Sau cải進")) == 2
    assert "Trước cải進" not in texts(r1)                      # Before caption excluded
    # item #2's captions are NOT absorbed into item #1's region
    assert len(caption_ids(r2, "Sau cải進")) == 1
    assert "Trước cải進" not in texts(r2)
    # the transition arrow is a signal but stays out of the After crop
    arrow_right = (4.75 + 1.4) * EMU                            # arrow spans x 4.75–6.15in
    assert r1.bbox[0] > arrow_right
    excluded_text = " ".join(r1.excluded_objects)
    assert "Right Arrow" in excluded_text and "transition" in excluded_text


# ================================================================== §16 Before exclusion
def test_before_pictures_never_enter_after_crops(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    report = _report(deck)
    _, _, sel = _selection(report)
    before_refs = [ref for ref in sel.rejected if ref.kind == "before"]
    assert len(before_refs) == 3
    assert all(ref.temporal_role == "BEFORE" and not ref.excel_output_eligible for ref in before_refs)
    regions = _regions(report, sel)
    for region in regions:
        assert all(ref.kind == "after" for ref in region.pictures)
        # every Before caption of the slide is explicitly excluded from every region
        assert any("reason=before-caption" in d for d in region.excluded_objects)
    grouped, problems = _export(report, sel, tmp_path / "crops3", renderer=_builtin_renderer())
    assert problems == []
    for group in grouped.groups:
        for path in group:
            assert not _crop_has_rgb(path, BEFORE_RGB)


# ================================================================== §20 report isolation
def test_report_isolation_two_structurally_identical_reports(tmp_path):
    deck_a = _deck(tmp_path / "a" / NAME, [_two_item_slide()])
    deck_b = _deck(tmp_path / "b" / NAME, [_two_item_slide()])
    report_a, _, sel_a = _report(deck_a), *_selection(_report(deck_a))[1:]
    report_b = _report(deck_b)
    _, _, sel_b = _selection(report_b)
    # owner ids stay report-local: identical headings, identical owner ids, different report scopes
    owners_a = {ref.owner_id for ref in sel_a.after}
    owners_b = {ref.owner_id for ref in sel_b.after}
    assert owners_a == owners_b and len(owners_a) == 2
    assert {ref.report_scope_id for ref in sel_a.after} != {ref.report_scope_id for ref in sel_b.after}
    # regions never mix reports: every region keeps its own scope
    regions_a = _regions(report_a, sel_a)
    regions_b = _regions(report_b, sel_b)
    assert len(regions_a) == 2 and len(regions_b) == 2
    assert {r.report_scope_id for r in regions_a} == {sel_a.after[0].report_scope_id}
    assert {r.report_scope_id for r in regions_b} == {sel_b.after[0].report_scope_id}


# ================================================================== §23 learning identity per item
def test_learning_identity_distinguishes_report_slide_item(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    report = _report(deck)
    _, slides, sel = _selection(report)
    candidates = build_candidates(report, slides, sel, MGMT, str(deck))
    owners = {it.owner_id for it in production_item_regions(slide_items(report.slide(3)))}
    assert len(owners) == 2
    by_owner = {}
    for cand in candidates:
        if cand.logical_item_owner in owners:
            by_owner.setdefault(cand.logical_item_owner, []).append(cand)
    assert set(by_owner) == owners
    all_ids = [cand.candidate_id for cands in by_owner.values() for cand in cands]
    assert len(all_ids) == len(set(all_ids))                 # stable, collision-free identity
    for owner, cands in by_owner.items():
        indexes = {cand.item_index for cand in cands}
        assert indexes == {item_index_for_owner(slide_items(report.slide(3)), owner)}
        assert all(cand.management_number == MGMT for cand in cands)
        assert all(cand.source_name == NAME for cand in cands)
        headings = {cand.owner_heading for cand in cands}
        assert headings == ({ITEM1_HEAD} if ITEM1_HEAD in owner or "xuoc" in owner else {ITEM2_HEAD})
    # candidate_id stays backward-compatible (report|slide|shape|order) – migration safe
    sample = next(iter(by_owner.values()))[0]
    assert sample.candidate_id.startswith(f"{MGMT}|S3|#")


# ================================================================== §21 render once, crop per item
def test_slide_rendered_once_for_two_item_regions(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    report = _report(deck)
    _, _, sel = _selection(report)
    assert len(_regions(report, sel)) == 2

    class CountingRenderer(SlideRenderer):
        def __init__(self):
            super().__init__(prefer=("builtin",), width_px=1200)
            self.calls = []

        def render(self, report, slide_numbers, out_dir):
            self.calls.append(tuple(slide_numbers))
            return super().render(report, slide_numbers, out_dir)

    renderer = CountingRenderer()
    grouped, problems = _export(report, sel, tmp_path / "crops4", renderer=renderer)
    assert problems == [] and len(grouped.groups) == 2
    assert renderer.calls == [(3,)]                          # ONE render shared by both item crops


# ================================================================== §43 cooperative cancellation hook
def test_cancellation_hook_between_item_regions(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    report = _report(deck)
    _, _, sel = _selection(report)
    calls = {"n": 0}

    def should_cancel():
        calls["n"] += 1
        return calls["n"] > 1                                 # stop before the 2nd item region

    grouped, problems = _export(report, sel, tmp_path / "crops5", renderer=_builtin_renderer(),
                                should_cancel=should_cancel)
    assert len(grouped.groups) == 1                          # first item finished, nothing half-done
    assert any("REGION_CANCELLED" in p and "đã dừng" in p for p in problems)
    assert all(path.exists() for group in grouped.groups for path in group)


# ================================================================== §26 Excel mapping: one row, both items in order, 2 image groups
def test_excel_mapping_one_row_two_image_groups(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    template = make_template(tmp_path / "Verification.xlsx")
    out = tmp_path / "out" / "result.xlsx"
    processor = _run_batch([deck], template, out)
    assert processor.summary.completed + processor.summary.needs_review == 1   # row written (review flags OK)

    saved = load_workbook(out)[SHEET]
    cell = saved.cell(row=4, column=COL["improvement"])
    text = cell.value or ""
    # no item disappears, no cross-item merge: both items' texts, in authored order
    assert ITEM1_HEAD in text and ITEM2_HEAD in text
    assert "thay vật liệu nhôm bằng thép SKD11" in text
    assert "cố định vị trí gá bằng chốt định vị" in text
    assert text.index(ITEM1_HEAD) < text.index("SKD11") < text.index(ITEM2_HEAD) < text.index("chốt định vị")
    # captions / title never leak into the improvement cell
    for junk in ("Trước cải進", "Sau cải進", "3. CẢI TIẾN TRONG SẢN XUẤT", "CTMS"):
        assert junk not in text
    # 2 item-scoped image groups in the improvement_image cell (one horizontal row per item)
    images = _row_images(saved, 4)
    assert len(images) == 2
    assert _image_has_rgb(images[0], AFTER1_RGB) and not _image_has_rgb(images[0], AFTER2_RGB)
    assert _image_has_rgb(images[1], AFTER2_RGB) and not _image_has_rgb(images[1], AFTER1_RGB)
    assert not any(_image_has_rgb(image, BEFORE_RGB) for image in images)
    # the FileResult log carries the item model too
    result = processor.results[0]
    assert [item["index"] for item in result.improvement_items] == [0, 1]
    assert {item["slide"] for item in result.improvement_items} == {3}


# ================================================================== §25 force reprocess replaces stale crops
def test_force_reprocess_replaces_stale_item_crops(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    template = make_template(tmp_path / "Verification.xlsx")
    out = tmp_path / "out" / "result.xlsx"
    _run_batch([deck], template, out)
    saved = load_workbook(out)[SHEET]
    assert len(_row_images(saved, 4)) == 2
    assert _image_has_rgb(_row_images(saved, 4)[0], AFTER1_RGB)

    # rebuild the same report with different After evidence (stale crop must be replaced)
    def rebuilt(s, W):
        _prod_head(s, W)
        _item_text(s, ITEM1_HEAD, ITEM1_BODY, 1.02)
        _visual_row(s, 2.38, n_before=2, n_after=3, after_color="#d1c4e9", below_caption=True)
        _item_text(s, ITEM2_HEAD, ITEM2_BODY, 4.45)
        _visual_row(s, 5.55, n_before=1, n_after=2, after_color="#ffe082", below_caption=False)

    _deck(deck, [rebuilt])
    processor = _run_batch([deck], template, out, force_reprocess=True)
    assert processor.summary.completed + processor.summary.needs_review == 1
    saved = load_workbook(out)[SHEET]
    images = _row_images(saved, 4)
    assert len(images) == 2                                     # still exactly 2 item groups
    assert _image_has_rgb(images[0], (209, 196, 233)) and not _image_has_rgb(images[0], AFTER1_RGB)
    assert _image_has_rgb(images[1], (255, 224, 130)) and not _image_has_rgb(images[1], AFTER2_RGB)


# ================================================================== §25 file-atomic stop with multi-item slide
def test_stop_after_current_file_keeps_multi_item_row_complete(tmp_path):
    deck_a = _deck(tmp_path / "a" / NAME, [_two_item_slide()])
    other = "(CTMS)_20601_260920046-VOC_ Đối sách LỖI KHÁC 24.9.2026.pptx"
    deck_b = _deck(tmp_path / "b" / other, [_single_item_slide()])
    template = make_template(tmp_path / "Verification.xlsx")
    out = tmp_path / "out" / "result.xlsx"
    events = []
    opts = BatchOptions(files=[deck_a, deck_b], template=template, output_file=out,
                        use_ollama=False, use_fast_cache=False, use_image_learning=False)
    proc = BatchProcessor(opts, on_file=lambda i, s, d: events.append((i, s, d)))
    proc.on_file = lambda i, s, d: (events.append((i, s, d)), proc.request_stop()) if (i == 0 and s == "reading") else events.append((i, s, d))
    summary = proc.run()
    assert summary.stopped is True
    assert summary.completed + summary.needs_review == 1        # current file finished, batch stopped
    assert out.exists()
    saved = load_workbook(out)[SHEET]
    assert len(_row_images(saved, 4)) == 2                      # file-atomic: both item crops present


# ================================================================== §27 segmentation diagnostics / item model
def test_item_model_diagnostics_and_item_scoped_text(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    report, cls, rec = _record(deck)
    from app.report_identity import report_scope_key
    report_scope = report_scope_key(str(deck))
    assert len(rec.improvement_items) == 2
    first, second = rec.improvement_items
    assert (first["slide"], first["index"]) == (3, 0)
    assert (second["slide"], second["index"]) == (3, 1)
    assert first["itemId"] != second["itemId"]
    assert first["heading"] == ITEM1_HEAD and second["heading"] == ITEM2_HEAD
    assert first["scopedItemId"] == f"{report_scope}|S3|{first['itemId']}"
    assert second["scopedItemId"] == f"{report_scope}|S3|{second['itemId']}"
    assert len(first["afterPictures"]) == 3 and len(second["afterPictures"]) == 2
    assert len(first["beforePictures"]) == 2 and len(second["beforePictures"]) == 1
    assert any("Sau cải進" in cap for cap in first["captions"])
    # item-scoped text: each item keeps its own lines, never the next item's
    assert ITEM1_HEAD in first["text"] and "SKD11" in first["text"]
    assert ITEM2_HEAD not in first["text"] and "chốt định vị" not in first["text"]
    assert ITEM2_HEAD in second["text"] and "chốt định vị" in second["text"]
    assert ITEM1_HEAD not in second["text"] and "SKD11" not in second["text"]
    # Excel contract preserved: the improvement cell is the deterministic section join
    assert rec.improvement.index(ITEM1_HEAD) < rec.improvement.index(ITEM2_HEAD)
    assert "SKD11" in rec.improvement and "chốt định vị" in rec.improvement


def test_segment_report_items_logs_and_orders(tmp_path):
    deck = _deck(tmp_path / NAME, [_two_item_slide()])
    report = _report(deck)
    items_by_slide = segment_report_items(report, [3], log=True)
    assert list(items_by_slide) == [3]
    items = items_by_slide[3]
    assert [it.index for it in items] == [0, 1]
    assert items[0].contains_y(items[0].bounds[1] + 10)
    assert not items[0].contains_y(items[1].bounds[1] + 10)
    # whitespace lines between the spans belong to no item (but the section join keeps them)
    assert items[0].bounds[1] + items[0].bounds[3] <= items[1].bounds[1]


# ================================================================== §31 required fixtures: 3-item and 1-item slides
def test_three_item_slide_three_regions(tmp_path):
    deck = _deck(tmp_path / NAME, [_three_item_slide()])
    report = _report(deck)
    _, _, sel = _selection(report)
    items = production_item_regions(slide_items(report.slide(3)))
    assert len(items) == 3
    regions = _regions(report, sel)
    assert len(regions) == 3
    assert [r.item_index for r in regions] == [0, 1, 2]
    assert [len(r.pictures) for r in regions] == [2, 2, 1]
    grouped, problems = _export(report, sel, tmp_path / "crops6", renderer=_builtin_renderer())
    assert problems == [] and len(grouped.groups) == 3
    # each crop keeps its own After colour (no cross-item merge)
    assert _crop_has_rgb(grouped.groups[0][0], AFTER1_RGB)
    assert _crop_has_rgb(grouped.groups[1][0], AFTER2_RGB)
    assert _crop_has_rgb(grouped.groups[2][0], (255, 224, 130))
    assert not _crop_has_rgb(grouped.groups[0][0], AFTER2_RGB)
    assert not _crop_has_rgb(grouped.groups[1][0], AFTER1_RGB)


def test_single_item_slide_one_region_and_text_equals_section(tmp_path):
    deck = _deck(tmp_path / NAME, [_single_item_slide()])
    report, cls, rec = _record(deck)
    items = production_item_regions(slide_items(report.slide(3)))
    assert len(items) == 1
    assert len(rec.improvement_items) == 1
    item = rec.improvement_items[0]
    assert item["index"] == 0 and len(item["afterPictures"]) == 3
    # single-item slide: item text == section text (same lines, same order – no loss, no duplication)
    def _lines(text):
        return [ln for ln in (text or "").splitlines() if ln.strip()]
    assert _lines(item["text"]) == _lines(rec.improvement)
    _, _, sel = _selection(report)
    regions = _regions(report, sel)
    assert len(regions) == 1 and len(regions[0].pictures) == 3
    grouped, problems = _export(report, sel, tmp_path / "crops7", renderer=_builtin_renderer())
    assert problems == [] and len(grouped.groups) == 1


# ================================================================== §32 heading-style variants (real failure classes)
def test_defect_style_item2_heading_with_sau_inside(tmp_path):
    """Item #2's heading opens with a defect name and contains 'sau' – the exact reported Windows case."""
    deck = _deck(tmp_path / NAME, [_two_item_slide(item2_head=ITEM2_DEFECT_HEAD)])
    report = _report(deck)
    items = production_item_regions(slide_items(report.slide(3)))
    assert [it.heading for it in items] == [ITEM1_HEAD, ITEM2_DEFECT_HEAD]
    _, _, sel = _selection(report)
    regions = _regions(report, sel)
    assert len(regions) == 2 and [r.item_index for r in regions] == [0, 1]
    # item #1's crop still excludes item #2's heading/body
    s = report.slide(3)
    item2_boxes = [b for b in s.blocks if b.is_text and (ITEM2_DEFECT_HEAD in b.text or "gá klem" in b.text)]
    assert item2_boxes and all(regions[0].bbox[3] <= b.top for b in item2_boxes)
    grouped, problems = _export(report, sel, tmp_path / "crops8", renderer=_builtin_renderer())
    assert problems == [] and len(grouped.groups) == 2
    assert _crop_has_rgb(grouped.groups[0][0], AFTER1_RGB)
    assert not _crop_has_rgb(grouped.groups[0][0], AFTER2_RGB)


def test_mid_block_non_bold_item2_heading(tmp_path):
    """Item #2's colon-terminated heading is a non-bold line inside a combined body frame."""
    deck = _deck(tmp_path / NAME, [_two_item_slide(mid_block_item2=True)])
    report = _report(deck)
    items = production_item_regions(slide_items(report.slide(3)))
    assert [it.heading for it in items] == [ITEM1_HEAD, ITEM2_HEAD]
    _, _, sel = _selection(report)
    regions = _regions(report, sel)
    assert len(regions) == 2 and [r.item_index for r in regions] == [0, 1]
    assert len(regions[0].pictures) == 3 and len(regions[1].pictures) == 2
    grouped, problems = _export(report, sel, tmp_path / "crops9", renderer=_builtin_renderer())
    assert problems == [] and len(grouped.groups) == 2
    assert _crop_has_rgb(grouped.groups[0][0], AFTER1_RGB)
    assert not _crop_has_rgb(grouped.groups[0][0], AFTER2_RGB)


# ================================================================== §17 temporary-action exclusion in multi-item reports
def test_temporary_action_excluded_in_multi_item_report(tmp_path):
    deck = _deck(tmp_path / NAME, [_temporary_slide(), _two_item_slide()])
    report, cls, rec = _record(deck)
    assert "Sorting 100%" not in rec.improvement
    assert "Tăng cường kiểm tra ngoại quan" not in rec.improvement
    assert "Sorting 100%" in rec.temporary_excluded
    # the temporary slide contributes no improvement items and no After evidence
    assert {item["slide"] for item in rec.improvement_items} == {4}
    _, slides, sel = _selection(report)
    assert slides == [4]
    assert len(rec.improvement_items) == 2
