"""Reviewed scan list before Start: manual exclusion / restore / rescan / stale invalidation (controller layer)."""
import json
import shutil
import types
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

import app.batch_processor as bp
from app.config import AppConfig
from app.gui_controller import (NO_SCAN_VI, RUNNING_VI, SCAN_FILTERS_VI, STALE_LIST_VI, USER_EXCLUDED, GuiController)
from app.prescan import (ACTION_OUTSIDE_PERIOD, ACTION_PROCESS, ACTION_PROCESS_NEW_ROW, ACTION_SOURCE_DUPLICATE,
                         CACHE_FILE_NAME)

SHEET = "Kiểm chứng"
RUN_TODAY = date(2026, 9, 20)


def _fake(folder: Path, name: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    p = folder / name
    p.write_bytes(b"not a pptx")
    return p


@pytest.fixture
def world(sample_tree, template, tmp_path, monkeypatch):
    """Folder with 2 real reports (keys partially in master), 1 new-key real report, skips of every kind."""
    reports = tmp_path / "Bao_cao"
    sep = reports / "2026-09"
    sep.mkdir(parents=True)
    shutil.copy(sample_tree["files"][0], sep / "260918080-VOC_a.pptx")          # existing partial row
    shutil.copy(sample_tree["files"][1], sep / "260918081-VOC_b.pptx")          # existing partial row
    shutil.copy(sample_tree["files"][2], sep / "260918082-VOC_c.pptx")          # NEW key -> PROCESS_NEW_ROW
    _fake(reports / "2026-10", "261002001-VOC_oct.pptx")                        # outside period
    _fake(sep / "dup", "260918080-VOC_copy.pptx")                               # source duplicate (older)
    import os
    os.utime(sep / "dup" / "260918080-VOC_copy.pptx", (1_600_000_000, 1_600_000_000))
    _fake(sep, "khong_ma.pptx")                                                 # invalid key
    wb = load_workbook(template)
    ws = wb[SHEET]
    ws.cell(row=4, column=2, value="260918080-VOC")
    ws.cell(row=5, column=2, value="260918081-VOC")
    wb.save(template)
    out = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"

    class _D(date):
        @classmethod
        def today(cls):
            return RUN_TODAY
    monkeypatch.setattr(bp, "date", _D)
    ctl = GuiController(AppConfig(), config_path=tmp_path / "c.json")
    ctl._today = lambda: RUN_TODAY
    ctl.set_report_folder(str(reports))
    ctl.set_template(str(template))
    ctl.set_output(str(out))
    return types.SimpleNamespace(ctl=ctl, reports=reports, template=template, out=out)


def _by_name(ctl, filter_name=SCAN_FILTERS_VI[1]):
    return {r.path.name: r for r in ctl.scan_rows(filter_name)}


def _spies(monkeypatch):
    calls = {"parse": [], "classify": 0}
    real = bp.parse_pptx
    real_cls = bp.classify

    def parse(p):
        calls["parse"].append(Path(p).name)
        return real(p)

    def classify(report, client, model):
        calls["classify"] += 1
        assert client is None
        return real_cls(report, client, model)
    monkeypatch.setattr(bp, "parse_pptx", parse)
    monkeypatch.setattr(bp, "classify", classify)
    return calls


def _run(ctl):
    assert ctl.start(use_ollama=False)
    ctl.processor._thread.join(120)
    ctl.pump()


# ---------------------------------------------------------------- 1, 2, 31: list content
def test_scan_list_shows_candidates_and_skips_without_opening_pptx(world, monkeypatch):
    calls = _spies(monkeypatch)
    ctl = world.ctl
    res = ctl.scan()
    assert calls["parse"] == [] and not world.out.exists()                       # nothing opened / written
    rows = _by_name(ctl)
    assert len(rows) == 6 and len(ctl.all_files) == 6
    assert rows["260918080-VOC_a.pptx"].action == ACTION_PROCESS and rows["260918080-VOC_a.pptx"].will_process
    assert rows["260918082-VOC_c.pptx"].action == ACTION_PROCESS_NEW_ROW
    assert rows["260918082-VOC_c.pptx"].status_vi == "Sẽ thêm mới vào Excel"
    assert rows["261002001-VOC_oct.pptx"].status_vi == "Ngoài thời gian xử lý"
    assert rows["khong_ma.pptx"].status_vi.startswith("Không xác định được Management Number")
    dup = [r for r in rows.values() if r.action == ACTION_SOURCE_DUPLICATE]
    assert len(dup) == 1 and dup[0].path.parent.name == "dup"
    assert rows["260918080-VOC_a.pptx"].occurrence_date == "18/09/2026"
    assert str(rows["260918080-VOC_a.pptx"].path) == str(world.reports / "2026-09" / "260918080-VOC_a.pptx")   # 31
    assert rows["260918080-VOC_a.pptx"].as_values(1)[-1].endswith("260918080-VOC_a.pptx")
    # filters only change the display
    assert {r.path.name for r in ctl.scan_rows(SCAN_FILTERS_VI[0])} == {"260918080-VOC_a.pptx", "260918081-VOC_b.pptx",
                                                                        "260918082-VOC_c.pptx"}
    assert len(ctl.scan_rows(SCAN_FILTERS_VI[2])) == 3
    assert len(ctl.final_queue()) == 3 and res.counts()["candidates"] == 3
    assert ctl.queue_text() == "Cần xử lý sau khi quét: 3   Đã loại thủ công: 0   Sẽ xử lý: 3"


# ---------------------------------------------------------------- 3-13, 29: exclusion
def test_exclude_one_and_many_never_touches_files_or_excel(world, monkeypatch):
    calls = _spies(monkeypatch)
    ctl = world.ctl
    ctl.scan()
    rows = _by_name(ctl)
    new_key = rows["260918082-VOC_c.pptx"]
    srcs = {p: (p.stat().st_size, p.stat().st_mtime_ns, p.read_bytes()) for p in ctl.all_files}
    assert ctl.exclude([new_key.index]) == ""                                     # 3
    assert ctl.exclude([rows["261002001-VOC_oct.pptx"].index])                    # not a candidate -> problem text
    assert _by_name(ctl)["260918082-VOC_c.pptx"].state == USER_EXCLUDED
    assert _by_name(ctl)["260918082-VOC_c.pptx"].status_vi == "Đã loại thủ công"
    assert ctl.queue_text() == "Cần xử lý sau khi quét: 3   Đã loại thủ công: 1   Sẽ xử lý: 2"
    assert ctl.exclude([rows["260918080-VOC_a.pptx"].index, rows["260918081-VOC_b.pptx"].index]) == ""   # 4
    assert ctl.queue_counts()["will_process"] == 0 and ctl.final_queue() == []
    assert ctl.restore([rows["260918080-VOC_a.pptx"].index]) == ""
    _run(ctl)
    assert calls["parse"] == ["260918080-VOC_a.pptx"] and calls["classify"] == 1  # 9, 10: excluded never opened
    assert len(ctl.files) == 1 and ctl.progress.total == 1
    assert {p: (p.stat().st_size, p.stat().st_mtime_ns, p.read_bytes()) for p in ctl.all_files} == srcs   # 7, 8
    assert all(p.exists() for p in ctl.all_files)
    ws = load_workbook(world.out)[SHEET]
    assert ws.cell(row=4, column=5).value == "A185"                                # processed one
    assert ws.cell(row=5, column=5).value is None                                  # 11: excluded existing row untouched
    assert ws.cell(row=6, column=2).value is None                                  # 12, 29: no new row for excluded key
    cache = json.loads((world.out.parent / "logs" / CACHE_FILE_NAME).read_text(encoding="utf-8"))
    assert set(cache["entries"]) == {"260918080-VOC"}                             # 13


def test_delete_key_and_context_menu_share_the_controller_action(world, monkeypatch):
    """The view's Delete binding and the context-menu command both call ``exclude_selected`` -> ``ctl.exclude``."""
    ctl = world.ctl
    ctl.scan()
    idx = _by_name(ctl)["260918081-VOC_b.pptx"].index
    src = Path("app/gui.py").read_text(encoding="utf-8")
    assert 'bind("<Delete>", lambda _e: self.exclude_selected())' in src           # 5
    assert 'add_command(label="Xóa khỏi danh sách xử lý", command=self.exclude_selected)' in src   # 6
    assert 'text="Xóa khỏi danh sách", command=self.exclude_selected' in src
    assert src.count("self.ctl.exclude(") == 1                                      # one controller call site
    assert ctl.exclude([idx]) == "" and _by_name(ctl)["260918081-VOC_b.pptx"].excluded


# ---------------------------------------------------------------- 14, 15, 30: queue / progress
def test_queue_denominator_and_eta_use_final_list(world, tmp_path, monkeypatch):
    ctl = world.ctl
    # 20 candidates (fake new keys + 3 real) – exclude 3 -> 17
    for i in range(17):
        _fake(world.reports / "2026-09" / "more", f"2609{i + 1:02d}9{i:02d}-VOC_x.pptx")
    ctl.scan()
    assert ctl.queue_counts()["candidates"] == 20
    cands = [r for r in ctl.scan_rows(SCAN_FILTERS_VI[0])][:3]
    assert ctl.exclude([r.index for r in cands]) == ""
    assert ctl.queue_counts()["will_process"] == 17
    t = [0.0]
    ctl._clock = lambda: t[0]
    assert ctl.start(use_ollama=False, in_thread=False)
    ctl.pump()
    assert ctl.progress.total == 17 and len(ctl.files) == 17                        # 14
    assert ctl.progress.text == "Đã xử lý: 17 / 17 — 100%"
    assert len(ctl.report_durations) <= 10 and ctl.summary.total == 17             # 15: ETA samples from queue only
    # 30: retained new key -> row appended only when processed
    ws = load_workbook(world.out)[SHEET]
    keys = [ws.cell(row=r, column=2).value for r in range(4, ws.max_row + 1) if ws.cell(row=r, column=2).value]
    new_total = sum(1 for r in ctl.scan_rows() if r.action == ACTION_PROCESS_NEW_ROW)
    new_excluded = sum(1 for r in cands if r.action == ACTION_PROCESS_NEW_ROW)
    assert len(keys) == 2 + new_total - new_excluded                                  # existing 2 + new rows processed
    assert len(set(keys)) == len(keys)                                               # no duplicate rows


# ---------------------------------------------------------------- 16-19: restore
def test_restore_only_reverses_manual_exclusion(world):
    ctl = world.ctl
    ctl.scan()
    rows = _by_name(ctl)
    i = rows["260918082-VOC_c.pptx"].index
    ctl.exclude([i])
    assert ctl.restore([i]) == "" and _by_name(ctl)["260918082-VOC_c.pptx"].will_process        # 16
    oct_i = rows["261002001-VOC_oct.pptx"].index
    dup_i = next(r.index for r in rows.values() if r.action == ACTION_SOURCE_DUPLICATE)
    inv_i = rows["khong_ma.pptx"].index
    for j, act in ((oct_i, ACTION_OUTSIDE_PERIOD), (dup_i, ACTION_SOURCE_DUPLICATE), (inv_i, "INVALID_MANAGEMENT_NUMBER")):
        assert ctl.restore([j]) == "Chỉ khôi phục được file đã loại thủ công."                  # 17, 18, 19
        assert next(r for r in ctl.scan_rows() if r.index == j).action == act
    assert len(ctl.final_queue()) == 3


# ---------------------------------------------------------------- 20-26: rescan / stale
def test_rescan_resets_exclusions_and_inputs_invalidate_list(world, tmp_path):
    ctl = world.ctl
    ctl.scan()
    i = _by_name(ctl)["260918080-VOC_a.pptx"].index
    ctl.exclude([i])
    assert ctl.queue_counts()["excluded"] == 1
    ctl.scan()                                                                        # 20
    assert ctl.queue_counts()["excluded"] == 0 and not ctl.scan_stale
    # 21 folder
    other = tmp_path / "Other"
    other.mkdir()
    ctl.set_report_folder(str(other))
    assert ctl.scan_stale and STALE_LIST_VI in ctl.validate() and ctl.scan_message == STALE_LIST_VI
    assert not ctl.start(use_ollama=False)                                            # 26
    ctl.set_report_folder(str(world.reports))
    ctl.scan()
    assert not ctl.scan_stale
    # 22 Excel
    ctl.set_output(str(tmp_path / "Output" / "Kiem_chung_T09_2026.xlsx"))
    assert ctl.scan_stale
    ctl.scan()
    # 23 month
    ctl.set_period("month", 9, 2026)
    assert ctl.scan_stale
    ctl.scan()
    # 24 range
    ctl.set_period("range", start="01/09/2026", end="30/09/2026")
    assert ctl.scan_stale
    ctl.scan()
    # unchanged value -> not stale
    ctl.set_period("range", start="01/09/2026", end="30/09/2026")
    assert not ctl.scan_stale
    # 25 force
    ctl.set_force(True)
    assert ctl.scan_stale and not ctl.start(use_ollama=False)
    ctl.scan()
    assert not ctl.scan_stale and ctl.validate() == []


# ---------------------------------------------------------------- 27, 28: no mutation while running
def test_removal_disabled_while_processing(world, monkeypatch):
    ctl = world.ctl
    ctl.scan()
    rows = _by_name(ctl)
    gate = __import__("threading").Event()
    real = bp.parse_pptx

    def slow(p):
        gate.wait(10)
        return real(p)
    monkeypatch.setattr(bp, "parse_pptx", slow)
    assert ctl.start(use_ollama=False)
    try:
        assert ctl.exclude([rows["260918081-VOC_b.pptx"].index]) == RUNNING_VI        # 27
        assert ctl.exclude([rows["260918080-VOC_a.pptx"].index]) == RUNNING_VI        # 28: active report
        assert ctl.restore([rows["260918080-VOC_a.pptx"].index]) == RUNNING_VI
        assert ctl.queue_counts()["excluded"] == 0
    finally:
        gate.set()
        ctl.processor._thread.join(120)
        ctl.pump()
    assert ctl.state == "idle" and len(ctl.files) == 3
    assert ctl.exclude([]) != RUNNING_VI


def test_exclude_without_scan_is_a_clear_message(tmp_path):
    ctl = GuiController(AppConfig(period_mode="all"), config_path=tmp_path / "c.json")
    assert ctl.exclude([0]) == NO_SCAN_VI and ctl.restore([0]) == NO_SCAN_VI
    assert ctl.scan_rows() == [] and ctl.final_queue() == []
