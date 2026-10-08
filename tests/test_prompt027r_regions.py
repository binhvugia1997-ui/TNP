"""PROMPT-027R: After-region evidence on a real multi-item slide — membership, item boundaries and diagnostics.

The reported failure has three structural faces, each reproduced here on a synthetic slide that follows the real layout
(Mục #1 with several After pictures and the heading of Mục #2 directly below it):

* Bug A — the production region of an item can hold only PART of the authored After block: a caption can be
  attributed across an item boundary, a block can be split by a global distance threshold, a picture can be lost to
  the wrong temporal class.
* Bug B — the previous item's crop reaches the trailing text of the NEXT item heading: the final crop must be clamped
  by the structural heading boundary against the FINAL crop's horizontal extent, not against the pictures' extent.
* Member/evidence truth — the Learning overlay, the Excel crop and the DTO must describe ONE production region.

Everything is headless and deterministic.  The built-in renderer is forced, so no test depends on PowerPoint.
"""
from __future__ import annotations

import json
import logging
from pathlib import Path

from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree

from app.improvement_pictures import SEMANTIC_PRODUCTION, slide_items
from app.improvement_visual import (CLUSTER_COL_GAP_FRACTION, CLUSTER_ROW_GAP_FRACTION, PICTURE_GAP_RATIO,
                                    build_improvement_visual_regions, cluster_after_pictures, crop_box_px)
from app.pptx_parser import parse_pptx
from app.qpn_renderer import SlideRenderer
from tests.test_after_evidence import _deck, _pics, _prod_head, _tb
from tests.test_prompt004 import MGMT
from tests.test_prompt025_multi_item import (ITEM1_BODY, ITEM1_HEAD, ITEM2_BODY, ITEM2_HEAD, NAME, _crop_has_rgb,
                                             _export, _regions, _report, _selection)

EMU = 914400
ORANGE = (0xFF, 0xE0, 0xB2)                       # Before pictures
GREEN = ((0x2E, 0x7D, 0x32), (0x43, 0xA0, 0x47), (0x66, 0xBB, 0x6A))   # item #1 After pictures A, B, C
BLUE = ((0x15, 0x65, 0xC0), (0x42, 0xA5, 0xF5))   # item #2 After pictures
ITEM1_AFTER_LEFTS = (6.5, 7.9, 9.3)               # 0.1in authored gaps, picture width 1.3in
ITEM2_BODY_INLINE = "- Cải tiến gá kẹp: cố định vị trí\n+ Sau: Thêm chốt định vị"


def _hex(rgb) -> str:
    return "#%02x%02x%02x" % tuple(rgb)


def _caption(slide, text, left, top, width=1.5, height=0.32):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    shape.text_frame.text = text
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(0x1F, 0x5F, 0xA0)
    shape.line.width = Pt(1.25)
    return shape


def _arrow(slide, left, top, width, height, rgb=(0x00, 0x70, 0xC0)):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(width), Inches(height))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(*rgb)
    arrow.line.fill.background()
    return arrow


def _real_layout(item1_body=ITEM1_BODY, after1=ITEM1_AFTER_LEFTS, after2=(6.5, 7.9), head2_left=1.7,
                 head2_width=5.6, head2_top=4.12, extra=None):
    """Mục #1 (heading, body, Before pair → arrow → three After pictures, 'Sau cải tiến' above and beneath) followed
    immediately by Mục #2 (heading, body, Before picture, two After pictures).  ``extra(slide, W)`` adds objects."""
    def build(s, W):
        _prod_head(s, W)
        _tb(s, ITEM1_HEAD, 1.7, 1.02, 5.6, 0.4, 12, True)
        _tb(s, item1_body, 1.7, 1.46, 5.6, 0.6, 12)
        _caption(s, "Trước cải tiến", 1.7, 2.38)
        _pics(s, [_hex(ORANGE)] * 2, 1.7, 2.75, w=1.3, h=0.85, gap=0.1)
        _arrow(s, 4.75, 2.93, 1.4, 0.4)
        _caption(s, "Sau cải tiến", 6.5, 2.38)
        for left, rgb in zip(after1, GREEN):
            _pics(s, [_hex(rgb)], left, 2.75, w=1.3, h=0.85)
        _caption(s, "Sau cải tiến", 7.35, 3.70, 1.95, 0.34)
        _tb(s, ITEM2_HEAD, head2_left, head2_top, head2_width, 0.4, 12, True)
        _tb(s, ITEM2_BODY, 1.7, head2_top + 0.44, 5.6, 0.6, 12)
        _caption(s, "Trước cải tiến", 1.7, 5.05)
        _pics(s, [_hex(ORANGE)], 1.7, 5.45, w=1.3, h=0.85)
        _caption(s, "Sau cải tiến", 6.5, 5.05)
        for left, rgb in zip(after2, BLUE):
            _pics(s, [_hex(rgb)], left, 5.45, w=1.3, h=0.85)
        if extra is not None:
            extra(s, W)
    return build


def _build_report(tmp_path: Path, build, folder: str = "in"):
    deck = _deck(tmp_path / folder / NAME, [build])
    report = _report(deck)
    _, _, selection = _selection(report)
    return report, selection, deck


def _by_owner(regions):
    out = {}
    for region in regions:
        out.setdefault(region.improvement_item_id, []).append(region)
    return out


def _item_owner(report, heading_fragment: str) -> str:
    items = [it for it in slide_items(report.slide(3)) if it.region_kind == "item"
             and it.semantic_role == SEMANTIC_PRODUCTION]
    return next(it.owner_id for it in items if heading_fragment in it.heading)


def _colours_in(path, palette):
    return {rgb: _crop_has_rgb(path, rgb) for rgb in palette}


def _text_box_rect(report, text_fragment: str):
    block = next(b for b in report.slide(3).blocks if b.is_text and text_fragment in b.text)
    return block.left, block.top, block.right, block.bottom


# ================================================================== CASE 1 — same item, transitively connected block
def test_case1_three_connected_after_pictures_form_one_region(tmp_path):
    report, selection, _ = _build_report(tmp_path, _real_layout())
    owner1 = _item_owner(report, "Lỗi xước rear")
    regions = _by_owner(_regions(report, selection))[owner1]
    assert len(regions) == 1, "A-B-C at authored 0.1in gaps are ONE After block, not three"
    assert sorted(ref.block.shape_id for ref in regions[0].pictures) == sorted(
        ref.block.shape_id for ref in selection.after if ref.owner_id == owner1)
    assert len(regions[0].pictures) == 3


def test_case1b_wide_authored_gaps_still_form_one_block(tmp_path):
    """Bug A root cause: a global 8%-of-slide gap split one authored After block into three regions.  Gaps up to
    one picture width, with nothing in between, are the same block (picture-relative, not a larger global threshold)."""
    report, selection, _ = _build_report(tmp_path, _real_layout(after1=(6.5, 8.9, 11.3)), "wide")
    owner1 = _item_owner(report, "Lỗi xước rear")
    regions = _by_owner(_regions(report, selection))[owner1]
    assert len(regions) == 1 and len(regions[0].pictures) == 3
    # the global absolute tolerance was NOT raised to achieve this (PROMPT-027R §7: no blind threshold increase)
    assert CLUSTER_ROW_GAP_FRACTION == 0.08 and CLUSTER_COL_GAP_FRACTION == 0.08 and PICTURE_GAP_RATIO == 1.0


def test_cluster_function_is_transitive_and_respects_structural_gaps(tmp_path):
    report, selection, _ = _build_report(tmp_path, _real_layout(), "unit")
    owner1 = _item_owner(report, "Lỗi xước rear")
    members = [ref for ref in selection.after if ref.owner_id == owner1]
    clusters = cluster_after_pictures(members, [], report.slide_width, report.slide_height)
    assert len(clusters) == 1 and len(clusters[0]) == 3
    # a structural object between A and C separates them unless the chain connects them through B
    first, middle, last = sorted(members, key=lambda r: r.block.left)
    blocker = (int(middle.block.left), int(middle.block.top), int(middle.block.right), int(middle.block.bottom))
    split = cluster_after_pictures([first, last], [], report.slide_width, report.slide_height,
                                   obstacles=[blocker])
    assert len(split) == 2, "a picture-sized structural object in the gap must keep the two blocks apart"


# ================================================================== CASE 2 — genuinely separate After blocks
def _two_blocks_one_item(tmp_path, blocker_between: bool):
    def layout(s, W):
        _prod_head(s, W)
        _tb(s, ITEM1_HEAD, 1.7, 1.02, 5.6, 0.4, 12, True)
        _tb(s, ITEM1_BODY, 1.7, 1.46, 5.6, 0.6, 12)
        _caption(s, "Trước cải tiến", 1.7, 2.38)
        _pics(s, [_hex(ORANGE)], 1.7, 2.75, w=1.3, h=0.85)
        _caption(s, "Sau cải tiến", 6.5, 2.38)
        _pics(s, [_hex(GREEN[0])], 6.5, 2.75, w=1.3, h=0.85)
        _pics(s, [_hex(GREEN[1])], 7.9, 2.75, w=1.3, h=0.85)
        # second After block: its own caption ABOVE it, placed at a picture-scale distance on the right
        if blocker_between:
            _pics(s, [_hex(ORANGE)], 9.3, 2.75, w=1.3, h=0.85)               # a Before picture separates the blocks
            _caption(s, "Sau cải tiến", 10.6, 2.38)
            _pics(s, [_hex(GREEN[2])], 10.6, 2.75, w=1.3, h=0.85)
        else:
            _caption(s, "Sau cải tiến", 9.5, 2.38)
            _pics(s, [_hex(GREEN[2])], 9.5, 2.75, w=1.3, h=0.85)
    return layout


def test_case2_two_separate_after_blocks_of_one_item_stay_two_regions(tmp_path):
    report, selection, _ = _build_report(tmp_path, _two_blocks_one_item(tmp_path, blocker_between=False), "two")
    owner1 = _item_owner(report, "Lỗi xước rear")
    regions = _by_owner(_regions(report, selection))[owner1]
    assert len(regions) == 2, "two After captions above two blocks are two After blocks of ONE item"
    assert [r.after_block_index for r in regions] == [0, 1]
    assert regions[0].block_id != regions[1].block_id
    assert {len(r.pictures) for r in regions} == {2, 1}


def test_case2b_a_before_picture_between_blocks_keeps_them_apart(tmp_path):
    report, selection, _ = _build_report(tmp_path, _two_blocks_one_item(tmp_path, blocker_between=True), "two_b")
    owner1 = _item_owner(report, "Lỗi xước rear")
    regions = _by_owner(_regions(report, selection))[owner1]
    assert len(regions) == 2
    for region in regions:
        assert all(ref.temporal_role == "AFTER" for ref in region.pictures)


# ================================================================== CASE 3 — Mục #1 stops before Mục #2 heading
def test_case3_item1_region_stops_before_item2_heading(tmp_path, caplog):
    with caplog.at_level(logging.INFO, logger="report_extractor.images"):
        report, selection, _ = _build_report(tmp_path, _real_layout(), "case3")
        regions = _regions(report, selection)
    owner1 = _item_owner(report, "Lỗi xước rear")
    item1 = _by_owner(regions)[owner1][0]
    heading2_top = _text_box_rect(report, ITEM2_HEAD)[1]
    # the invariant of the reported slide: region(item #1).bottom < heading(item #2).top
    assert item1.bbox[3] < heading2_top, (item1.bbox[3], heading2_top)
    # every Mục #1 After picture is inside the region
    for ref in item1.pictures:
        assert item1.bbox[1] <= ref.block.top and ref.block.bottom <= item1.bbox[3]
    # the boundary that limits the item is the next ITEM heading, and it is recorded
    assert item1.boundary_source == "next_item_heading"
    assert item1.boundary_top == heading2_top
    spans = [r.getMessage() for r in caplog.records if r.getMessage().startswith("ITEM_SPAN ")]
    assert any("item_index=0" in line and "next_item_id=" in line and "boundary_source=next_item_heading" in line
               and "next_heading_y=" in line for line in spans), spans
    evidence = [r.getMessage() for r in caplog.records if r.getMessage().startswith("REGION_EVIDENCE ")]
    assert any("boundary_source=next_item_heading" in line for line in evidence)


def test_case3b_narrow_next_heading_straddled_by_an_annotation_is_never_leaked(tmp_path, caplog):
    """The reported leak: the crop's horizontal extent reaches the trailing text of the next heading, while the
    pictures themselves do not overlap that heading horizontally, so the picture-based clamp never fired."""
    def marker(s, W):
        oval = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.0), Inches(3.5), Inches(0.8), Inches(0.8))
        oval.fill.background()
        oval.line.color.rgb = RGBColor(0xDA, 0x23, 0x2D)
        oval.line.width = Pt(2)

    layout = _real_layout(item1_body=ITEM1_BODY, head2_width=4.6, extra=marker)
    with caplog.at_level(logging.INFO, logger="report_extractor.images"):
        report, selection, _ = _build_report(tmp_path, layout, "m5")
        regions = _regions(report, selection)
    owner1 = _item_owner(report, "Lỗi xước rear")
    item1 = _by_owner(regions)[owner1][0]
    heading2_top = _text_box_rect(report, ITEM2_HEAD)[1]
    assert item1.bbox[3] < heading2_top, "the crop must end above the next heading even when x-disjoint from pictures"
    assert any(r.getMessage().startswith("REGION_CLAMP ") and "boundary_source=next_item_heading" in r.getMessage()
               for r in caplog.records)


# ================================================================== CASE 4 — invisible No Fill / No Line container
def test_case4_invisible_container_is_neither_drawn_nor_a_crop_boundary(tmp_path):
    def invisible(s, W):
        frame = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.2), Inches(0.9), Inches(7.6), Inches(3.9))
        sp_pr = frame._element.spPr
        for tag in ("a:solidFill", "a:noFill", "a:ln"):
            for element in sp_pr.findall(qn(tag)):
                sp_pr.remove(element)
        etree.SubElement(sp_pr, qn("a:noFill"))
        line = etree.SubElement(sp_pr, qn("a:ln"))
        etree.SubElement(line, qn("a:noFill"))
        # a text-bearing invisible container far outside the After footprint
        frame_text = s.shapes.add_textbox(Inches(0.3), Inches(6.7), Inches(3.0), Inches(0.35))
        frame_text.text_frame.text = "ghi chú ẩn"

    plain_report, plain_sel, _ = _build_report(tmp_path, _real_layout(), "plain")
    plain = _by_owner(_regions(plain_report, plain_sel))[_item_owner(plain_report, "Lỗi xước rear")][0]
    report, selection, _ = _build_report(tmp_path, _real_layout(extra=invisible), "invisible")
    region = _by_owner(_regions(report, selection))[_item_owner(report, "Lỗi xước rear")][0]
    assert region.bbox == plain.bbox, "an invisible container must not be used as the crop boundary"
    for obj in region.included_objects:
        assert (obj.fill_visible or obj.line_visible or obj.text.strip()), "invisible shapes are never painted"
        assert "ghi chú ẩn" not in obj.text


# ================================================================== CASE 5 — legitimate annotations are included
def test_case5_legitimate_annotations_of_the_block_are_included(tmp_path):
    def annotations(s, W):
        # an overlay mark on the first After picture (overlay rule) and a measurement label under it (containment)
        mark = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.45), Inches(2.7), Inches(0.5), Inches(0.5))
        mark.fill.background()
        mark.line.color.rgb = RGBColor(0xDA, 0x23, 0x2D)
        mark.line.width = Pt(3)
        label = s.shapes.add_textbox(Inches(6.5), Inches(3.62), Inches(1.2), Inches(0.28))
        label.text_frame.text = "120 mm"
        label.text_frame.paragraphs[0].runs[0].font.size = Pt(9)

    report, selection, _ = _build_report(tmp_path, _real_layout(extra=annotations), "annot")
    region = _by_owner(_regions(report, selection))[_item_owner(report, "Lỗi xước rear")][0]
    texts = [" ".join(obj.text.split()) for obj in region.included_objects]
    assert "120 mm" in texts, "a measurement label that lies on the After block belongs to it"
    assert any(obj.prst == "ellipse" for obj in region.included_objects), "the overlay mark belongs to the block"
    assert any(" ".join(obj.text.split()) == "Sau cải tiến" for obj in region.included_objects)


# ================================================================== CASE 6 — neighbouring / next-item text excluded
def test_case6_neighbouring_item_body_text_never_enters_the_crop(tmp_path):
    """Item #2's body paragraph overlapped the After caption band of item #2 in the reported layout: its full box was
    absorbed by intersection and dragged item #2's Before picture into the crop.  Text must lie mostly in the footprint."""
    layout = _real_layout(head2_top=4.12)
    report, selection, _ = _build_report(tmp_path, layout, "text")
    regions = _by_owner(_regions(report, selection))
    owner2 = _item_owner(report, "lệch ATN")
    item2 = regions[owner2][0]
    body2 = _text_box_rect(report, "gá kẹp")
    assert all("gá kẹp" not in " ".join(obj.text.split()) for obj in item2.included_objects)
    assert not any(body2 == (obj.left, obj.top, obj.right, obj.bottom) for obj in item2.included_objects)
    before2 = next(b for b in report.slide(3).pictures if abs(b.left - 1.7 * EMU) < 2000 and b.top > 5 * EMU)
    assert item2.bbox[0] >= before2.right, "the Before picture of Mục #2 stays out of its After crop"
    assert any("outside-visual-footprint" in detail for detail in item2.excluded_objects)


def test_case6b_inline_plus_sau_line_cannot_claim_a_before_picture(tmp_path):
    """Temporal regression: an open-ended '+ Sau:' line above a Before picture must not turn it into After when the
    arrow says the picture is on the Before side (the inline line and the structural arrow disagree → ambiguous)."""
    layout = _real_layout(item1_body=ITEM1_BODY + "\n+ Sau: Bọc silicon 2mm toàn bộ mặt tiếp xúc")
    report, selection, _ = _build_report(tmp_path, layout, "inline")
    after_colours = set()
    for ref in selection.after:
        pixels = ref.block.image_blob
        with Image.open(__import__("io").BytesIO(pixels)) as im:
            after_colours.add(tuple(im.convert("RGB").getpixel((5, 5))))
    assert ORANGE not in after_colours, "a Before picture was classified as After"
    assert after_colours <= set(GREEN) | set(BLUE)
    befores = [ref for ref in selection.rejected if ref.kind == "before"]
    assert len(befores) == 3 and all(ref.temporal_role == "BEFORE" for ref in befores)


# ================================================================== CASE 7 — candidate vs region bounds (DTO)
def test_case7_candidate_and_region_bounds_are_reported_separately(sample_tree, tmp_path):
    from tests.test_prompt024r_learning_workspace import _service_with_candidates
    from tests.test_prompt027_rendering import _candidate_from_region

    deck = _deck(tmp_path / "in" / NAME, [_real_layout()])
    report = _report(deck)
    _, _, selection = _selection(report)
    regions = _regions(report, selection)
    region = regions[0]
    candidate = _candidate_from_region(deck, report, region, index=1)
    service, _ = _service_with_candidates(sample_tree, tmp_path, [candidate])
    dto = service.learning_state()["images"][0]
    # candidate overlay = the picture; evidence overlay = the production region (never derived from each other)
    assert dto["targetBbox"]["width"] < dto["evidenceRegionBbox"]["width"]
    assert dto["evidenceRegionBbox"] == {"x": region.bbox[0], "y": region.bbox[1],
                                         "width": region.width, "height": region.height}
    # PROMPT-027R §24: the production membership is exposed for acceptance, as display/debug data
    assert dto["evidenceRegionPictureIds"] == [int(ref.block.shape_id) for ref in region.pictures]
    assert dto["evidenceRegionPictureCount"] == len(region.pictures) == 3
    assert dto["evidenceRegionId"] == region.block_id
    assert dto["evidenceRegionBlockIndex"] == region.after_block_index
    assert dto["evidenceRegionBoundarySource"] == region.boundary_source
    assert str(Path(deck).resolve()) not in json.dumps(dto, ensure_ascii=False)


# ================================================================== CASE 8 — the production crop holds every member
def test_case8_production_crop_contains_every_member_and_no_neighbour(tmp_path):
    report, selection, _ = _build_report(tmp_path, _real_layout(), "crop")
    grouped, problems = _export(report, selection, tmp_path / "crops", renderer=SlideRenderer(
        prefer=("builtin",), width_px=1200))
    assert problems == []
    owner1 = _item_owner(report, "Lỗi xước rear")
    owner2 = _item_owner(report, "lệch ATN")
    order = {owner: index for index, owner in enumerate([owner1, owner2])}
    crop1 = grouped.groups[order[owner1]][0]
    crop2 = grouped.groups[order[owner2]][0]
    for rgb in GREEN:
        assert _crop_has_rgb(crop1, rgb), "every member picture of Mục #1 must be in its crop"
    for rgb in BLUE + (ORANGE,):
        assert not _crop_has_rgb(crop1, rgb), "no neighbouring item picture (or Before picture) in Mục #1's crop"
    for rgb in BLUE:
        assert _crop_has_rgb(crop2, rgb)
    for rgb in GREEN + (ORANGE,):
        assert not _crop_has_rgb(crop2, rgb)


# ================================================================== CASE 9 — report isolation
def test_case9_identical_reports_never_share_regions_or_pictures(tmp_path):
    report_a, sel_a, deck_a = _build_report(tmp_path / "a", _real_layout(), "one")
    report_b, sel_b, deck_b = _build_report(tmp_path / "b", _real_layout(), "two")
    regions_a = _regions(report_a, sel_a)
    regions_b = _regions(report_b, sel_b)
    assert {r.report_scope_id for r in regions_a} == {sel_a.after[0].report_scope_id}
    assert {r.report_scope_id for r in regions_b} == {sel_b.after[0].report_scope_id}
    assert {r.report_scope_id for r in regions_a}.isdisjoint({r.report_scope_id for r in regions_b})
    assert {r.block_id for r in regions_a}.isdisjoint({r.block_id for r in regions_b})
    for region in regions_a:
        assert all(ref.report_scope_id == region.report_scope_id for ref in region.pictures)


# ================================================================== §8 group ancestry
def test_grouped_after_block_keeps_absolute_geometry_and_records_ancestry(tmp_path):
    def grouped_layout(s, W):
        _prod_head(s, W)
        _tb(s, ITEM1_HEAD, 1.7, 1.02, 5.6, 0.4, 12, True)
        _tb(s, ITEM1_BODY, 1.7, 1.46, 5.6, 0.6, 12)
        _caption(s, "Trước cải tiến", 1.7, 2.38)
        _pics(s, [_hex(ORANGE)], 1.7, 2.75, w=1.3, h=0.85)
        _caption(s, "Sau cải tiến", 6.5, 2.38)
        group = s.shapes.add_group_shape()
        for left, rgb in zip(ITEM1_AFTER_LEFTS, GREEN):
            group.shapes.add_picture(_pic_stream(rgb), Inches(left), Inches(2.75), Inches(1.3), Inches(0.85))
        _tb(s, ITEM2_HEAD, 1.7, 4.12, 5.6, 0.4, 12, True)

    report = parse_pptx(str(_deck(tmp_path / "grp" / NAME, [grouped_layout])))
    pictures = [b for b in report.slide(3).blocks if b.kind == "picture" and b.group_path]
    assert len(pictures) == 3, "group children are still flattened into the picture set"
    assert len({b.group_path for b in pictures}) == 1, "the three children share one authored group"
    assert all(b.parent_group_id == pictures[0].group_path[-1] and b.group_root_id == pictures[0].group_path[0]
               for b in pictures)
    # PROMPT-011: absolute geometry is unchanged by the ancestry bookkeeping
    assert sorted(round(b.left / EMU, 2) for b in pictures) == [6.5, 7.9, 9.3]
    _, _, selection = _selection(report)
    regions = _regions(report, selection)
    assert regions and len(regions[0].pictures) == 3


def _pic_stream(rgb):
    from io import BytesIO
    bio = BytesIO()
    Image.new("RGB", (480, 360), _hex(rgb)).save(bio, "PNG")
    bio.seek(0)
    return bio


# ================================================================== §12/§16 next-item callout is not an item
def test_callout_lying_on_a_picture_is_not_an_improvement_item(tmp_path):
    def callout(s, W):
        label = s.shapes.add_textbox(Inches(9.45), Inches(3.00), Inches(0.82), Inches(0.30))
        label.text_frame.text = "CHECK"
        label.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        label.text_frame.paragraphs[0].runs[0].font.bold = True

    report, _, _ = _build_report(tmp_path, _real_layout(extra=callout), "callout")
    headings = [it.heading for it in slide_items(report.slide(3)) if it.region_kind == "item"]
    assert "CHECK" not in headings


# ================================================================== §3/§5 diagnostics
def test_region_diagnostics_expose_members_clusters_and_boundaries(tmp_path, caplog):
    with caplog.at_level(logging.INFO):
        report, selection, _ = _build_report(tmp_path, _real_layout(), "diag")
        _export(report, selection, tmp_path / "diag-crops", renderer=SlideRenderer(prefer=("builtin",), width_px=1200))
    text = "\n".join(r.getMessage() for r in caplog.records)
    for token in ("PICTURE_DECISION ", "REGION_MEMBER ", "REGION_CLUSTERS ", "ITEM_SPAN ", "REGION_EVIDENCE ",
                  "REGION_CROP "):
        assert token in text, token
    member = next(r.getMessage() for r in caplog.records if r.getMessage().startswith("REGION_MEMBER "))
    for field in ("scope=", "slide=", "item_index=", "item=", "shape=", "source_order=", "bbox=", "temporal=",
                  "semantic=", "owner=", "confident=", "eligible=", "group="):
        assert field in member, field
    decision = next(r.getMessage() for r in caplog.records if r.getMessage().startswith("PICTURE_DECISION "))
    assert "excel_output_eligible=" in decision and "logical_item_owner=" in decision
    # no absolute path leaks into the diagnostics
    assert str(Path(report.path).resolve()) not in text


def test_obstacle_boundary_never_removes_a_member(tmp_path):
    """Members win: when a boundary would cut through a member picture the boundary is reported, the member stays."""
    from app.improvement_visual import _Boundary, _clamp_bottom
    members_bottom = 4_000_000
    bbox = (1_000_000, 2_000_000, 5_000_000, 5_000_000)
    bottom, applied, conflicts = _clamp_bottom(bbox, members_bottom,
                                               [_Boundary(1_000_000, 3_500_000, 6_000_000, 4_000_000, "next_item_heading")],
                                               guard_y=10_000)
    assert bottom == bbox[3] and applied is None and len(conflicts) == 1
    bottom, applied, conflicts = _clamp_bottom(bbox, members_bottom,
                                               [_Boundary(1_000_000, 4_500_000, 6_000_000, 4_600_000, "next_item_heading")],
                                               guard_y=10_000)
    assert applied is not None and bottom == 4_490_000 and not conflicts


def test_crop_conversion_still_uses_actual_rendered_dimensions():
    assert crop_box_px((0, 0, 1000, 1000), 2000, 2000, 400, 300) == (0, 0, 200, 150)


# ================================================================== §25 one production truth for overlay and Excel
def test_case10_reviewer_decision_moves_the_production_region_and_the_dto_with_it(tmp_path, monkeypatch):
    """A reviewer's decision is authoritative for Excel (PROMPT-006): excluding one picture removes it from the
    production region, and the Learning DTO's evidence region is built from the SAME refs — never a second truth."""
    from app.application_service import ApplicationService
    from app.image_learning import select_with_learning

    from app.image_learning import build_candidates

    report, selection, deck = _build_report(tmp_path, _real_layout(), "learn")
    owner1 = _item_owner(report, "Lỗi xước rear")
    deterministic = _by_owner(_regions(report, selection))[owner1][0]
    picture_c = max(deterministic.pictures, key=lambda ref: ref.block.left)
    # the candidate id exactly as the production pipeline names it (labels are keyed by it)
    candidate = next(c for c in build_candidates(report, [3], selection, MGMT, str(deck))
                     if c.picture_id == picture_c.block.shape_id)

    class Reviewed:
        """The learning state of a reviewer who ignores this one picture (no model yet)."""
        enabled = True
        model = None

        def overrides(self):
            return {candidate.candidate_id: "IGNORE"}

        def apply(self, cands):                       # the same decision step as ImageLearning.apply
            from app.image_learning import decide
            return decide(cands, self.model, self.overrides())

    refs, _cands, _reasons = select_with_learning(report, [3], selection, Reviewed(), MGMT, str(deck))
    production = _by_owner(build_regions(report, refs))[owner1][0]
    assert picture_c.block.shape_id not in [r.block.shape_id for r in production.pictures]
    assert len(production.pictures) == len(deterministic.pictures) - 1

    service = ApplicationService.__new__(ApplicationService)
    service._slide_preview_cache = {}
    service._preview_cache = {}
    dto_regions = service._slide_evidence_regions(report, 3, candidate, Reviewed())
    item1 = next(r for r in dto_regions if r["itemId"] == owner1)
    assert item1["bbox"] == {"x": production.bbox[0], "y": production.bbox[1],
                             "width": production.width, "height": production.height}
    assert sorted(item1["pictureShapeIds"]) == sorted(r.block.shape_id for r in production.pictures)
    # the learning identity separates the two decisions, so no cached region can outlive the label
    assert service._learning_identity(Reviewed()) != service._learning_identity(None)


def build_regions(report, refs):
    return build_improvement_visual_regions(report, refs, management_number=MGMT)
