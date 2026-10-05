"""PROMPT-006 – improvement-image learning & correction mode (§24 backup tool, §25 features, §26 labeling,
§27 learning, §28 portable/update).  The deterministic PROMPT-004C detector stays the baseline; nothing here
talks to Ollama."""
import json
import shutil
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pptx.enum.shapes import MSO_SHAPE

import app
from app import image_learning as il
from app import updater as up
from app.batch_processor import BatchOptions, BatchProcessor
from app.image_learning import (FEATURE_NAMES, IMAGE_FEATURE_SCHEMA, MSG_MODEL_UNAVAILABLE, MSG_NOT_ENOUGH,
                                ImageCandidate, ImageLearning, LabelStore, TrainingError, build_candidates, decide,
                                explain, load_model, save_model, train_model)
from app.image_review import reapply_labels
from app.improvement_pictures import AMBIGUOUS_REASON
from tests.test_after_evidence import (AFTER_RGB, BEFORE_RGB, NAME, _arrow, _deck, _pics, _prod_head, _rgb,
                                       _select, _tb, _tb_colored)
from tests.test_content_region import COL, IMP_ITEM_1, INSPECTION_TEXT, SHEET, _images_at, _prefill, _shape
from tests.test_prompt004 import MGMT
from tests.test_updater import make_portable, make_update_folder

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
import make_stable_backup as msb  # noqa: E402

BACKUP_DIR = Path(__file__).resolve().parent.parent / "backup" / "ReportExtractor_v1.0.4_Build004_STABLE"


# ------------------------------------------------------------------ slide builders
def slide_caption(s, W):                       # explicit Trước / Sau captions
    _prod_head(s, W)
    _tb(s, IMP_ITEM_1, 1.7, 1.0, 5.6, 2.8, 12)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Trước cải tiến", 7.5, 1.0, 1.5, 0.35)
    _pics(s, ["#ffe0b2"], 7.5, 1.4)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 2.95, 1.5, 0.35)
    _pics(s, ["#c8e6c9", "#c8e6c9"], 7.5, 3.35)


def slide_blue_ambiguous(s, W):                # blue After text + one picture without any evidence
    _prod_head(s, W)
    _tb(s, "Cải tiến tool miết\n- Hiện trạng: tool miết đầu nhọn gây xước", 1.7, 1.0, 5.6, 1.0, 12)
    _pics(s, ["#ffe0b2"], 1.7, 2.1)
    _tb_colored(s, [("Thay tool miết đầu tròn R3, bọc silicon chống xước", "blue")], 7.5, 1.0, 5.3, 0.6)
    _pics(s, ["#c8e6c9", "#c8e6c9"], 7.5, 1.7)


def slide_arrow(s, W):                         # [Before][Before] -> [After][After][After]
    _prod_head(s, W)
    _tb(s, IMP_ITEM_1.split("\n")[0], 1.7, 1.0, 11, 0.5, 12)
    _pics(s, ["#ffe0b2", "#ffe0b2"], 1.7, 1.8)
    _arrow(s, MSO_SHAPE.RIGHT_ARROW, 4.7, 2.2, 0.9, 0.4)
    _pics(s, ["#c8e6c9", "#c8e6c9", "#c8e6c9"], 5.8, 1.8)


def slide_inspection(s, W):                    # inspection/control improvement slide – never After
    _prod_head(s, W, title="4. CẢI TIẾN KIỂM TRA / KIỂM SOÁT")
    _tb(s, INSPECTION_TEXT, 1.7, 1.0, 5.6, 2.0, 12)
    _pics(s, ["#c8e6c9", "#c8e6c9"], 7.5, 1.4)


def slide_lonely(s, W):                        # single picture, no evidence at all -> ambiguous
    _prod_head(s, W)
    _tb(s, "Cải tiến jig định vị\n- Tăng số điểm định vị từ 2 lên 4", 1.7, 1.0, 5.6, 1.0, 12)
    _pics(s, ["#c8e6c9"], 7.5, 1.4)


def _cands(tmp_path, builders, mn=MGMT, name=NAME):
    r, sel = _select(_deck(tmp_path / name, builders))
    slides = list(range(3, 3 + len(builders)))
    c = build_candidates(r, slides, sel, mn, str(r.path))
    return r, sel, decide(c)


def _content(cands):
    return [c for c in cands if not c.hard_excluded]


def _synthetic_records(n_after=6, n_non=6, schema=IMAGE_FEATURE_SCHEMA):
    """Separable synthetic examples: After pictures sit right of centre with blue text, Non-After left."""
    recs = []
    for i in range(n_after + n_non):
        is_after = i < n_after
        f = {k: 0.0 for k in FEATURE_NAMES}
        f.update({"relative_x": 0.6 + 0.02 * i if is_after else 0.1 + 0.02 * i, "relative_y": 0.3,
                  "relative_width": 0.1, "relative_height": 0.15, "relative_area": 0.015,
                  "blue_after_text": 1.0 if is_after else 0.0, "picture_count": 3.0,
                  "picture_order": 3.0 if is_after else 1.0, "is_last_in_block": 1.0 if is_after else 0.0,
                  "production_block": 1.0, "caption_distance": 1.0})
        recs.append({"schema_version": schema, "candidate_id": f"X|S3|#{i}|{i}", "label": "AFTER" if is_after else
                     ("BEFORE" if i % 2 else "IGNORE"), "features": f})
    return recs


# ================================================================== §24 stable backup tool
def test_24_stable_backup_exists_and_is_clean():
    assert BACKUP_DIR.is_dir() and msb.verify_backup(BACKUP_DIR) == []
    init = (BACKUP_DIR / "app" / "__init__.py").read_text(encoding="utf-8")
    assert '__version__ = "1.0.4"' in init and "BUILD_NUMBER = 4" in init       # backup frozen at 1.0.4 / 004
    assert app.__version__ == "1.2.1" and app.BUILD_NUMBER == 12             # dev version moved on
    assert not any(p.name in ("sample_data", ".venv", "Output", "logs", "config", "learning_data")
                   for p in BACKUP_DIR.rglob("*") if p.is_dir())


def test_24_backup_tool_refuses_to_overwrite_and_excludes_data(tmp_path):
    root = tmp_path / "proj"
    for rel in msb.REQUIRED + ("tests/test_x.py", "tools/t.py"):
        (root / rel).parent.mkdir(parents=True, exist_ok=True)
        (root / rel).write_text("x", encoding="utf-8")
    tracked = list(msb.REQUIRED) + ["tests/test_x.py", "tools/t.py", "sample_data/a.pptx", "Output/r.xlsx",
                                    "logs/app.log", "config/config.json", ".venv/lib/x.py", "learning_data/image_labels.jsonl",
                                    "backup/old/app/gui.py", "app/__pycache__/x.pyc", "release/x.zip"]
    for rel in tracked:
        (root / rel).parent.mkdir(parents=True, exist_ok=True)
        (root / rel).write_bytes(b"d")
    target = msb.make_backup(root, "SNAP", files=tracked)
    copied = sorted(p.relative_to(target).as_posix() for p in target.rglob("*") if p.is_file())
    assert copied == sorted(list(msb.REQUIRED) + ["tests/test_x.py", "tools/t.py"])
    assert msb.verify_backup(target) == []
    with pytest.raises(msb.BackupError, match="không ghi đè"):
        msb.make_backup(root, "SNAP", files=tracked)
    with pytest.raises(msb.BackupError, match="Thiếu file nguồn"):
        msb.make_backup(root, "SNAP2", files=["tests/test_x.py"])
    assert not (root / "backup" / "SNAP2").exists()
    (target / "sample_data").mkdir()
    (target / "x.xlsx").write_bytes(b"d")
    assert len(msb.verify_backup(target)) == 2


# ================================================================== §25 candidate features (synthetic PPTX)
def test_25_candidates_have_normalised_features_and_stable_ids(tmp_path):
    r, sel, cands = _cands(tmp_path, [slide_caption, slide_blue_ambiguous])
    assert cands and all(set(c.features) == set(FEATURE_NAMES) for c in cands)
    for c in cands:
        assert c.candidate_id == f"{MGMT}|S{c.slide}|#{c.picture_id}|{c.order}"
        assert str(tmp_path) not in c.candidate_id and "/" not in c.candidate_id and "\\" not in c.candidate_id
        assert 0 <= c.features["relative_x"] <= 1 and 0 <= c.features["relative_width"] <= 1
        assert 0.0 <= c.confidence <= 1.0 and c.slide_size == (r.slide_width, r.slide_height)
        assert c.nearest_heading.startswith("3. CẢI TIẾN")
    # explicit captions
    s3 = [c for c in _content(cands) if c.slide == 3]
    befores = [c for c in s3 if c.features["explicit_before"]]
    afters = [c for c in s3 if c.features["explicit_after"]]
    assert len(befores) == 1 and len(afters) == 2
    assert befores[0].nearest_caption == "Trước cải tiến" and befores[0].confidence <= 0.1
    assert all(a.confidence >= 0.9 and a.decision == "include" for a in afters)
    assert befores[0].decision == "exclude" and befores[0].deterministic_kind == "before"
    # blue text + ambiguous picture
    s4 = [c for c in _content(cands) if c.slide == 4]
    blue = [c for c in s4 if c.features["blue_after_text"]]
    amb = [c for c in s4 if c.deterministic_kind == "ambiguous"]
    assert len(blue) == 2 and all(c.decision == "include" and c.nearest_text_color for c in blue)
    assert len(amb) == 1 and amb[0].decision == "review" and amb[0].confidence == 0.5
    assert "Hiện trạng" in amb[0].nearby_text
    # the logo in the header band is a hard exclusion, never a review item
    logos = [c for c in cands if c.hard_excluded]
    assert logos and all(c.decision == "exclude" and c.confidence == 0.0 and c.features["logo_decorative"] for c in logos)


def test_25_arrow_and_inspection_evidence(tmp_path):
    r, sel, cands = _cands(tmp_path, [slide_arrow, slide_inspection])
    s3 = [c for c in _content(cands) if c.slide == 3]
    dest = [c for c in s3 if c.features["arrow_destination"]]
    src = [c for c in s3 if c.features["arrow_source"]]
    assert len(dest) == 3 and len(src) == 2 and all(c.features["arrow_horizontal"] for c in dest + src)
    assert all(c.decision == "include" and c.confidence >= il.THRESHOLD_INCLUDE for c in dest)
    assert all(c.decision == "exclude" and c.confidence <= il.THRESHOLD_EXCLUDE for c in src)
    assert [c.features["picture_order"] for c in sorted(s3, key=lambda c: c.order)] == [1, 2, 3, 4, 5]
    assert sum(c.features["is_last_in_block"] for c in s3) == 1 and sum(c.features["is_first_in_block"] for c in s3) == 1
    s4 = [c for c in _content(cands) if c.slide == 4]
    assert s4 and all(c.features["inspection_control"] and not c.features["production_block"] for c in s4)
    assert all(c.decision == "exclude" and c.confidence <= 0.1 for c in s4)


def test_25_explanation_lines_are_structured_vietnamese(tmp_path):
    r, sel, cands = _cands(tmp_path, [slide_caption])
    a = next(c for c in cands if c.decision == "include")
    text = explain(a)
    lines = text.split("\n")
    assert lines[0].startswith("Ảnh slide 3") and lines[1] == "Kết quả: Sau cải tiến"
    assert lines[2].startswith("Độ tin cậy: Cao") and "Bằng chứng:" in lines
    assert any(l.startswith("- có chú thích 'Sau cải tiến'") for l in lines)
    amb = _cands(tmp_path / "b", [slide_lonely])[2]
    rv = [c for c in amb if c.decision == "review"]
    assert rv and "Kết quả: Cần kiểm tra" in explain(rv[0]) and "Độ tin cậy: Thấp" in explain(rv[0])


def test_25_candidates_without_learning_never_change_the_baseline(tmp_path):
    """Deterministic baseline (PROMPT-004C) == learning layer with no labels and no model."""
    for builders in ([slide_caption], [slide_blue_ambiguous], [slide_arrow], [slide_lonely], [slide_inspection]):
        r, sel = _select(_deck(tmp_path / NAME, builders))
        refs, cands, extra = il.select_with_learning(r, [3], sel, ImageLearning(tmp_path / "lrn"), MGMT, str(r.path))
        assert [(x.slide, x.block.shape_id) for x in refs] == [(x.slide, x.block.shape_id) for x in sel.after]
        assert extra == []


# ================================================================== §26 labeling / persistence / override
def test_26_label_store_append_supersede_and_duplicate_protection(tmp_path):
    r, sel, cands = _cands(tmp_path, [slide_blue_ambiguous])
    amb = next(c for c in cands if c.decision == "review")
    store = LabelStore(tmp_path / "learning_data")
    assert store.current_label(amb.candidate_id) == il.UNLABELED
    rec = store.label(amb, "AFTER")
    assert rec and rec["schema_version"] == 1 and rec["previous_label"] is None and rec["timestamp"]
    assert store.label(amb, "AFTER") is None                       # identical relabel -> no duplicate record
    rec2 = store.label(amb, "BEFORE", note="nhầm")
    assert rec2["previous_label"] == "AFTER" and rec2["note"] == "nhầm"
    lines = (tmp_path / "learning_data" / "image_labels.jsonl").read_text(encoding="utf-8").splitlines()
    assert len(lines) == 2 and all(json.loads(l)["candidate_id"] == amb.candidate_id for l in lines)
    assert store.current_label(amb.candidate_id) == "BEFORE"
    assert store.counts() == {"total": 1, "positive": 0, "negative": 1, "after": 0, "non_after": 1, "before": 1,
                              "control": 0, "ignore": 0}
    with pytest.raises(ValueError):
        store.label(amb, "MAYBE")
    stored = json.loads(lines[0])
    assert set(stored["features"]) == set(FEATURE_NAMES) and stored["management_number"] == MGMT
    assert stored["source_name"] == NAME and "/" not in stored["source_name"]
    # a damaged line never hides the others
    with open(store.path, "a", encoding="utf-8") as fh:
        fh.write("{broken json\n")
    assert store.counts()["total"] == 1


def test_26_user_label_is_authoritative_in_current_result_without_training(tmp_path):
    lrn = ImageLearning(tmp_path / "lrn")
    r, sel = _select(_deck(tmp_path / NAME, [slide_blue_ambiguous]))
    refs, cands, _ = il.select_with_learning(r, [3], sel, lrn, MGMT, str(r.path))
    assert len(refs) == 2 and AMBIGUOUS_REASON.format(n=3) in sel.reasons
    amb = next(c for c in cands if c.decision == "review")
    lrn.store.label(amb, "AFTER")
    refs2, cands2, _ = il.select_with_learning(r, [3], sel, lrn, MGMT, str(r.path))
    assert len(refs2) == 3 and lrn.model is None                   # no model involved
    c2 = next(c for c in cands2 if c.candidate_id == amb.candidate_id)
    assert c2.decision == "include" and c2.decision_source == "user" and c2.user_label == "AFTER"
    assert any(e.startswith("người dùng đã xác nhận") for e in c2.evidence)
    # overriding a deterministic AFTER to IGNORE removes it
    a0 = next(c for c in cands2 if c.decision_source == "rules" and c.decision == "include")
    lrn.store.label(a0, "IGNORE")
    refs3, _, _ = il.select_with_learning(r, [3], sel, lrn, MGMT, str(r.path))
    assert len(refs3) == 2 and all(x.block.shape_id != a0.picture_id for x in refs3)
    # a user label never resurrects a hard exclusion (logo / arrow / icon) – false inclusion is the worse error
    logo = next(c for c in cands2 if c.hard_excluded)
    out = decide([logo], None, {logo.candidate_id: "AFTER"})
    assert out[0].decision == "exclude" and out[0].decision_source == "rules"


def test_26_labels_persist_across_restart_and_relabel_audit(tmp_path):
    r, sel, cands = _cands(tmp_path, [slide_lonely])
    amb = next(c for c in cands if c.decision == "review")
    ImageLearning(tmp_path / "lrn").store.label(amb, "CONTROL")
    again = ImageLearning(tmp_path / "lrn")                        # "restart"
    assert again.overrides() == {amb.candidate_id: "CONTROL"}
    again.store.label(amb, "AFTER")
    recs = again.store.records()
    assert [x["label"] for x in recs] == ["CONTROL", "AFTER"] and recs[1]["previous_label"] == "CONTROL"
    assert ImageLearning(tmp_path / "lrn").overrides() == {amb.candidate_id: "AFTER"}
    assert again.status_text().startswith("Mẫu đã xác nhận: 1   Sau cải tiến: 1   Không phải Sau: 0")


def test_26_batch_pipeline_uses_labels_and_reapply_updates_only_image_cell(template, tmp_path):
    deck = _deck(tmp_path / f"(CTMS)_{MGMT}_ĐỐI SÁCH LỖI SƠN 24.09.2026.pptx", [slide_blue_ambiguous])
    _prefill(template, [{"mgmt": MGMT, "vendor": "Mtech", "model": "A253", "qpn": "QPN nhập tay"}])
    out = tmp_path / "out" / "k.xlsx"
    lrn_dir = tmp_path / "lrn"
    opts = BatchOptions(files=[deck], template=template, output_file=out, use_ollama=False, learning_dir=lrn_dir)
    proc = BatchProcessor(opts)
    proc.run()
    fr = proc.results[0]
    assert fr.image_candidates and any(c.decision == "review" for c in fr.image_candidates)
    ws = load_workbook(out)[SHEET]
    assert len(_images_at(ws, 4, COL["image"])) == 2
    assert "Không xác định chắc chắn" in " ".join(fr.review_reasons)
    # the batch_result.json entry stays compact (ids only)
    assert all(isinstance(x, str) for x in fr.to_dict()["image_candidates"])
    # user confirms the ambiguous picture as After -> re-apply rewrites ONLY the picture cell
    lrn = ImageLearning(lrn_dir)
    amb = next(c for c in fr.image_candidates if c.decision == "review")
    lrn.store.label(amb, "AFTER")
    before_cells = {c.coordinate: c.value for c in ws[4] if c.column != COL["image"]}
    res = reapply_labels(template, out, [amb], lrn)
    assert res.ok and res.updated_rows == [4] and res.pictures == 3, res.errors
    ws2 = load_workbook(out)[SHEET]
    imgs = _images_at(ws2, 4, COL["image"])
    assert len(imgs) == 3 and sorted(_rgb(im.ref) for im in imgs) == sorted([AFTER_RGB, AFTER_RGB, BEFORE_RGB])
    assert {c.coordinate: c.value for c in ws2[4] if c.column != COL["image"]} == before_cells
    assert ws2.cell(row=4, column=COL["vendor"]).value == "Mtech" and ws2.cell(row=4, column=COL["qpn"]).value == "QPN nhập tay"
    assert list((out.parent / "backup").glob("*_backup_*.xlsx"))      # backup before modification
    assert amb.decision == "include" and amb.decision_source == "user"
    # a second batch run (new processor) now includes the confirmed picture directly
    proc2 = BatchProcessor(BatchOptions(files=[deck], template=template, output_file=out, use_ollama=False,
                                        learning_dir=lrn_dir, force_reprocess=True))
    proc2.run()
    fr2 = proc2.results[0]
    assert len(fr2.after_pictures) == 3 and "Không xác định chắc chắn" not in " ".join(fr2.review_reasons)


def test_26_reapply_reports_missing_row_and_touches_nothing(template, tmp_path):
    deck = _deck(tmp_path / f"(CTMS)_{MGMT}_ĐỐI SÁCH LỖI SƠN 24.09.2026.pptx", [slide_lonely])
    _prefill(template, [{"mgmt": "999999999"}])
    out = tmp_path / "out" / "k.xlsx"
    shutil.copy(template, out.parent.mkdir(parents=True, exist_ok=True) or out)
    r, sel = _select(deck)
    lrn = ImageLearning(tmp_path / "lrn")
    _, cands, _ = il.select_with_learning(r, [3], sel, lrn, MGMT, str(deck))
    amb = next(c for c in cands if c.decision == "review")
    lrn.store.label(amb, "AFTER")
    stamp = out.stat().st_mtime_ns
    res = reapply_labels(template, out, [amb], lrn)
    assert not res.ok and "không tìm thấy dòng Management Number" in res.errors[0]
    assert out.stat().st_mtime_ns == stamp and not list((out.parent / "backup").glob("*.xlsx"))


def test_26_thumbnail_is_gui_only_and_optional(tmp_path):
    r, sel, cands = _cands(tmp_path, [slide_caption])
    c = next(c for c in cands if c.decision == "include")
    blob = next(p.image_blob for p in r.slide(3).pictures if p.shape_id == c.picture_id)
    path = il.save_thumbnail(c, blob, tmp_path / "lrn")
    assert path and Path(path).parent.name == "thumbnails" and Path(path).stat().st_size > 0
    assert il.save_thumbnail(c, None, tmp_path / "lrn") == ""
    assert il.save_thumbnail(c, b"not an image", tmp_path / "lrn") == ""
    assert "thumbnail" not in FEATURE_NAMES                        # classifier never sees pixels


# ================================================================== §27 learning
def test_27_insufficient_data_and_single_class_refuse_training(tmp_path):
    lrn = ImageLearning(tmp_path / "lrn")
    ok, msg = lrn.train()
    assert not ok and msg.startswith(MSG_NOT_ENOUGH) and lrn.model is None
    assert MSG_NOT_ENOUGH in lrn.status_text()
    with pytest.raises(TrainingError, match="Chưa đủ dữ liệu học"):
        train_model(_synthetic_records(n_after=12, n_non=0))       # one class only
    with pytest.raises(TrainingError):
        train_model(_synthetic_records(n_after=5, n_non=2))        # below MIN_PER_CLASS
    assert not il.model_path(tmp_path / "lrn").exists()


def test_27_training_persist_reload_and_metadata(tmp_path):
    recs = _synthetic_records(8, 8)
    model = train_model(recs)
    assert model.schema_version == IMAGE_FEATURE_SCHEMA and model.feature_names == list(FEATURE_NAMES)
    assert model.n_examples == 16 and model.n_after == 8 and model.n_non_after == 8 and model.trained_at
    assert model.train_accuracy >= 0.9
    p_after = model.probability(recs[0]["features"])
    p_non = model.probability(recs[-1]["features"])
    assert p_after > 0.75 > 0.35 > p_non
    path = save_model(model, tmp_path / "lrn")
    meta = json.loads(path.read_text(encoding="utf-8"))
    assert meta["schema_version"] == 1 and meta["n_examples"] == 16 and meta["feature_names"] == list(FEATURE_NAMES)
    loaded, status = load_model(tmp_path / "lrn")
    assert status == "ok" and loaded.probability(recs[0]["features"]) == pytest.approx(p_after)
    # same data -> same model (deterministic)
    assert train_model(recs).weights == model.weights


def test_27_trained_via_store_and_status_text(tmp_path):
    lrn = ImageLearning(tmp_path / "lrn")
    for rec in _synthetic_records(6, 6):
        c = ImageCandidate(candidate_id=rec["candidate_id"], management_number="X", source_name="x.pptx",
                           source_file="", slide=3, picture_id=1, order=1, bounds=(0, 0, 1, 1), slide_size=(10, 10),
                           block_bounds=(0, 0, 1, 1), features=rec["features"])
        lrn.store.label(c, rec["label"])
    ok, msg = lrn.train()
    assert ok and msg.startswith("Đã cập nhật mô hình ảnh từ 12 mẫu") and lrn.model_status == "ok"
    assert "Mô hình: đã huấn luyện" in lrn.status_text() and "Mẫu đã xác nhận: 12" in lrn.status_text()
    assert ImageLearning(tmp_path / "lrn").model is not None       # reload on restart


@pytest.mark.parametrize("damage", ["corrupt", "incompatible", "missing_fields"])
def test_27_corrupt_or_incompatible_model_falls_back_to_rules(tmp_path, damage):
    model = train_model(_synthetic_records(6, 6))
    path = save_model(model, tmp_path / "lrn")
    if damage == "corrupt":
        path.write_text("{not json", encoding="utf-8")
    elif damage == "incompatible":
        d = json.loads(path.read_text(encoding="utf-8"))
        d["schema_version"] = 99
        path.write_text(json.dumps(d), encoding="utf-8")
    else:
        d = json.loads(path.read_text(encoding="utf-8"))
        d["feature_names"] = d["feature_names"][:3]
        path.write_text(json.dumps(d), encoding="utf-8")
    lrn = ImageLearning(tmp_path / "lrn")
    assert lrn.model is None and lrn.model_status in ("corrupt", "incompatible")
    assert MSG_MODEL_UNAVAILABLE in lrn.status_text()
    r, sel = _select(_deck(tmp_path / NAME, [slide_blue_ambiguous]))
    refs, cands, _ = il.select_with_learning(r, [3], sel, lrn, MGMT, str(r.path))   # processing continues
    assert len(refs) == len(sel.after) and all(c.learned_probability is None for c in cands)


def test_27_model_only_resolves_uncertain_candidates_never_hard_exclusions(tmp_path):
    lrn = ImageLearning(tmp_path / "lrn")
    # a model that says "everything is After"
    model = il.ImageModel(weights=[0.0] * len(FEATURE_NAMES), bias=10.0, means=[0.0] * len(FEATURE_NAMES),
                          scales=[1.0] * len(FEATURE_NAMES), feature_names=list(FEATURE_NAMES),
                          schema_version=IMAGE_FEATURE_SCHEMA, trained_at="t", n_examples=20, n_after=10, n_non_after=10)
    save_model(model, tmp_path / "lrn")
    lrn.reload_model()
    assert lrn.model is not None
    r, sel = _select(_deck(tmp_path / NAME, [slide_blue_ambiguous, slide_caption, slide_inspection]))
    refs, cands, extra = il.select_with_learning(r, [3, 4, 5], sel, lrn, MGMT, str(r.path))
    by_src = {}
    for c in cands:
        by_src.setdefault((c.decision, c.decision_source), []).append(c)
    assert all(c.decision == "exclude" and c.learned_probability is None for c in cands if c.hard_excluded)
    assert all(c.decision == "exclude" for c in cands if c.deterministic_kind == "before")        # 0.05 stays out
    assert all(c.decision == "exclude" for c in cands if c.features["inspection_control"])         # never After
    resolved = [c for c in cands if c.decision_source == "model"]
    assert len(resolved) == 1 and resolved[0].slide == 3 and resolved[0].decision == "include"
    assert resolved[0].learned_probability > il.THRESHOLD_INCLUDE
    assert any(e.startswith("mô hình học") for e in resolved[0].evidence)
    assert len([x for x in refs if x.slide == 3]) == 3 and len([x for x in refs if x.slide == 5]) == 0
    # the opposite model excludes the ambiguous picture with an explicit review note
    model.bias = -10.0
    save_model(model, tmp_path / "lrn")
    lrn.reload_model()
    refs2, cands2, extra2 = il.select_with_learning(r, [3], sel, lrn, MGMT, str(r.path))
    assert len(refs2) == 2 and extra2 and "bị loại theo mô hình học" in extra2[0]
    # an undecided model (p in the middle) keeps the picture in review (omitted, flagged)
    model.bias = 0.0
    save_model(model, tmp_path / "lrn")
    lrn.reload_model()
    refs3, cands3, extra3 = il.select_with_learning(r, [3], sel, lrn, MGMT, str(r.path))
    assert len(refs3) == 2 and any(c.decision == "review" and c.learned_probability == 0.5 for c in cands3)


def test_27_user_override_beats_model(tmp_path):
    lrn = ImageLearning(tmp_path / "lrn")
    model = il.ImageModel(weights=[0.0] * len(FEATURE_NAMES), bias=10.0, means=[0.0] * len(FEATURE_NAMES),
                          scales=[1.0] * len(FEATURE_NAMES), feature_names=list(FEATURE_NAMES),
                          schema_version=IMAGE_FEATURE_SCHEMA, trained_at="t", n_examples=20, n_after=10, n_non_after=10)
    save_model(model, tmp_path / "lrn")
    lrn.reload_model()
    r, sel = _select(_deck(tmp_path / NAME, [slide_blue_ambiguous]))
    _, cands, _ = il.select_with_learning(r, [3], sel, lrn, MGMT, str(r.path))
    amb = next(c for c in cands if c.decision_source == "model")
    lrn.store.label(amb, "IGNORE")
    refs, cands2, _ = il.select_with_learning(r, [3], sel, lrn, MGMT, str(r.path))
    c = next(c for c in cands2 if c.candidate_id == amb.candidate_id)
    assert c.decision == "exclude" and c.decision_source == "user" and len(refs) == 2


def test_27_schema_mismatch_records_are_ignored_for_training():
    recs = _synthetic_records(6, 6) + _synthetic_records(6, 6, schema=2)
    xs, ys = il.training_examples(recs)
    assert len(xs) == 12
    with pytest.raises(TrainingError):
        train_model(_synthetic_records(6, 6, schema=2))


def test_27_thresholds_and_confidence_bands():
    assert il.THRESHOLD_INCLUDE > 0.5 > il.THRESHOLD_EXCLUDE
    assert il.confidence_band(0.95) == "high" and il.confidence_band(0.05) == "high"
    assert il.confidence_band(0.5) == "low" and il.confidence_band(0.65) == "medium"


# ================================================================== §28 portable / update
def test_28_learning_dir_is_beside_portable_root_and_preserved_by_updater(tmp_path, monkeypatch):
    import app.runtime_paths as rp
    monkeypatch.undo()                                             # use the real resolver for this check
    monkeypatch.setattr(rp, "portable_root", lambda: tmp_path / "Portable")
    d = rp.learning_dir()
    assert d == tmp_path / "Portable" / "learning_data" and d.is_dir()
    assert il.learning_dir().parent == tmp_path / "Portable"
    assert "learning_data" in up.PRESERVED_ENTRIES and up.is_preserved("learning_data")
    assert up.is_preserved("Learning_Data")


def test_28_apply_update_keeps_learning_data_and_rollback_restores_it(tmp_path, monkeypatch):
    portable = make_portable(tmp_path / "app")
    ld = portable / "learning_data"
    (ld / "model").mkdir(parents=True)
    (ld / "thumbnails").mkdir()
    (ld / "image_labels.jsonl").write_text('{"candidate_id": "a|S3|#1|1", "label": "AFTER"}\n', encoding="utf-8")
    (ld / "model" / "image_model.json").write_text('{"schema_version": 1}', encoding="utf-8")
    folder = make_update_folder(tmp_path / "Update")
    check = up.check_for_update(str(folder), current_build=4)
    staged = up.stage_update(check, portable)
    rc = up.apply_update(Path(staged.manifest_path), alive=lambda pid: False, sleep=lambda s: None,
                         spawn=lambda cmd, **kw: None)
    assert rc == 0 and (portable / up.APP_EXE).read_bytes() == b"MZ-new-1.0.5"
    assert (ld / "image_labels.jsonl").read_text(encoding="utf-8").startswith('{"candidate_id"')
    assert (ld / "model" / "image_model.json").exists() and (ld / "thumbnails").is_dir()
    # rollback after a failed apply keeps the user data as well
    (ld / "image_labels.jsonl").write_text("second\n", encoding="utf-8")
    staged2 = up.stage_update(check, portable)
    orig = up._copy
    calls = {"n": 0}

    def flaky_copy(src, dst):
        calls["n"] += 1
        if calls["n"] == 3:
            raise OSError(5, "disk I/O error")
        orig(src, dst)
    monkeypatch.setattr(up, "_copy", flaky_copy)
    rc2 = up.apply_update(Path(staged2.manifest_path), alive=lambda pid: False, sleep=lambda s: None,
                          spawn=lambda cmd, **kw: None)
    assert rc2 == 1 and (ld / "image_labels.jsonl").read_text(encoding="utf-8") == "second\n"
    assert (ld / "model" / "image_model.json").exists()


def test_28_release_artifact_never_contains_learning_data(tmp_path):
    import importlib
    bp = importlib.import_module("build_portable")
    assert "learning_data" in bp.FORBIDDEN_IN_ARTIFACT
    folder = tmp_path / "ReportExtractor"
    (folder / "_internal").mkdir(parents=True)
    (folder / up.APP_EXE).write_bytes(b"MZ")
    (folder / "learning_data").mkdir()
    (folder / "learning_data" / "image_labels.jsonl").write_text("{}", encoding="utf-8")
    problems = bp.validate_artifact(folder)
    assert any("learning_data" in p for p in problems)
    gi = (Path(__file__).resolve().parent.parent / ".gitignore").read_text(encoding="utf-8")
    assert "learning_data/" in gi.splitlines()


def test_28_gui_controller_learning_api(tmp_path, monkeypatch):
    from app.config import AppConfig
    from app.gui_controller import GuiController
    c = GuiController(AppConfig(), config_path=tmp_path / "cfg" / "config.json", learning_dir_override=tmp_path / "root")
    assert c.learning_status_text().startswith("Mẫu đã xác nhận: 0   Sau cải tiến: 0   Không phải Sau: 0")
    assert c.learning_folder() == tmp_path / "root"                   # override = the learning_data folder itself
    assert c.review_candidates() == [] and "Chưa có ảnh" in c.review_summary_text()
    ok, msg = c.train_image_model()
    assert not ok and msg.startswith(MSG_NOT_ENOUGH)
    ok, msg = c.save_confirmations([])
    assert not ok and msg == "Chưa chọn nhãn nào."
    assert c.build_options(False).learning_dir == tmp_path / "root"
    # labels chosen in the review window are written on "Lưu xác nhận"
    r, sel, cands = _cands(tmp_path, [slide_lonely])
    amb = next(x for x in cands if x.decision == "review")
    c.set_pending_label(amb, "AFTER")
    with pytest.raises(ValueError):
        c.set_pending_label(amb, "NOPE")
    ok, msg = c.save_confirmations(cands)
    assert ok and msg.startswith("Đã lưu 1 nhãn") and c.pending_labels == {}
    assert c.learning_counts()["after"] == 1
    ok, msg = c.export_learning_data(str(tmp_path / "exp" / "labels.json"))
    assert ok and json.loads((tmp_path / "exp" / "labels.json").read_text(encoding="utf-8"))[0]["label"] == "AFTER"


# ================================================================== GUI card + review window (fake tk)
def test_28_gui_learning_card_and_review_window(monkeypatch, tmp_path):
    from tests.test_gui_redesign import _make_app
    gui, a, reg, folder = _make_app(monkeypatch, tmp_path)
    a.ctl.learning_dir = tmp_path / "lrn"
    a.ctl._learning = None
    texts = [w.k.get("text") for w in reg["widgets"] if isinstance(w, reg["classes"]["Button"])]
    for t in ("Kiểm tra ảnh cải tiến", "Kiểm tra nội dung cải tiến", "Cập nhật mô hình học", "Mở thư mục dữ liệu học",
              "Xuất dữ liệu học"):
        assert t in texts
    assert a.lbl_learning.cfg["text"].startswith("Ảnh:  Mẫu đã xác nhận: 0")
    assert "Xóa toàn bộ" not in texts
    # no batch yet -> info box, no window
    a.open_image_review()
    assert reg["toplevels"] == []
    # fake a finished batch with candidates
    r, sel, cands = _cands(tmp_path, [slide_blue_ambiguous])

    class FR:
        image_candidates = cands
    class P:
        results = [FR()]
    a.ctl.processor = P()
    a.ctl.state = "idle"
    a.open_image_review()
    assert len(reg["toplevels"]) == 1 and reg["toplevels"][0].cfg.get("title") is None
    labels = [w.k.get("text") for w in reg["widgets"] if isinstance(w, reg["classes"]["Button"]) and w.master is not None]
    for t in ("Sau cải tiến", "Trước cải tiến", "Kiểm tra / Kiểm soát", "Không lấy", "Ảnh trước", "Ảnh tiếp", "Lưu xác nhận"):
        assert t in labels
    body = "".join(a.review_text.lines)
    assert "Kết quả: Cần kiểm tra" in body and "Management Number" in a.review_head.cfg["text"]
    assert a.btn_review_save.cfg["state"] == "disabled"
    a._review_set_label("AFTER")
    assert a.btn_review_save.cfg["state"] == "normal" and "(chưa lưu)" in "".join(a.review_text.lines)
    a._review_move(1)
    assert a.lbl_review_pos.cfg["text"].startswith("Ảnh 2/")
    a._review_move(-1)
    a._review_save()
    assert a.ctl.learning_counts()["after"] == 1 and a.lbl_learning.cfg["text"].startswith("Ảnh:  Mẫu đã xác nhận: 1")
    a.train_models()                                               # insufficient data -> two messages, no crash
    assert any(l.startswith("Ảnh cải tiến: chưa đủ dữ liệu — 1/10") for l in a.ctl.log_lines)
    assert any(l.startswith("Nội dung cải tiến: chưa đủ dữ liệu — 0/20") for l in a.ctl.log_lines)
