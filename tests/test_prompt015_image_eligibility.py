"""PROMPT-015: strict final-image eligibility, learning-independent gates, and owner-grouped XLSX rows."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE

from app.batch_processor import BatchOptions, BatchProcessor
from app.excel_writer import (EMU_PER_PX, IMAGE_MARGIN_PX, ExcelWriter, assert_image_inside_area)
from app.extractor import ExtractedRecord
from app.image_extractor import GroupedImagePaths, export_after_pictures
from app.image_learning import ImageLearning, build_candidates, select_with_learning
from app.improvement_pictures import (SEMANTIC_INSPECTION, SEMANTIC_PRODUCTION, SEMANTIC_TEMPORARY,
                                      SEMANTIC_VERIFICATION, group_refs, select_after_pictures)
from app.pptx_parser import parse_pptx
from tests.test_after_evidence import _deck, _pics, _prod_head, _select, _shape, _tb
from tests.test_content_region import COL, SHEET
from tests.test_prompt004 import MGMT

IMG_COL = COL["image"]
NAME = f"(CTMS)_{MGMT}_countermeasure_report.pptx"


def _mixed_production_slide(slide, width):
    _prod_head(slide, width)

    # Production item A: one Before image and two After images owned by the same item.
    _tb(slide, "Cải tiến jig định vị", 1.7, 0.9, 5.6, 0.4, 12)
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Trước cải tiến", 7.5, 0.9, 1.5, 0.35)
    _pics(slide, ["#ffe0b2"], 7.5, 1.3, w=1.7, h=0.6)
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 9.4, 0.9, 1.5, 0.35)
    _pics(slide, ["#c8e6c9", "#a5d6a7"], 9.4, 1.3, w=1.7, h=0.6)

    # Production item B is a distinct owner. Its After picture is eligible; the far-right picture
    # has no temporal evidence and must remain out even though a production owner is nearby.
    _tb(slide, "Cải tiến quy trình cấp liệu", 1.7, 2.35, 5.6, 0.4, 12)
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 2.35, 1.5, 0.35)
    _pics(slide, ["#81c784"], 7.5, 2.75, w=1.7, h=0.6)
    _pics(slide, ["#ffcdd2"], 11.5, 2.65, w=1.7, h=0.6)

    # These separate sections explicitly carry After captions, but their semantic roles are not
    # production improvement and therefore cannot enter the final workbook image cell.
    _tb(slide, "Cải tiến trong kiểm tra", 1.7, 3.8, 5.6, 0.4, 12)
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 3.8, 1.5, 0.35)
    _pics(slide, ["#b3e5fc"], 7.5, 4.2, w=1.7, h=0.6)

    _tb(slide, "Xử lý tạm thời", 1.7, 5.1, 5.6, 0.4, 12)
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 5.1, 1.5, 0.35)
    _pics(slide, ["#fff9c4"], 7.5, 5.5, w=1.7, h=0.6)

    _tb(slide, "Duy trì và áp dụng cải tiến", 1.7, 6.3, 5.6, 0.4, 12)
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 6.3, 1.5, 0.35)
    _pics(slide, ["#d1c4e9"], 7.5, 6.7, w=1.7, h=0.6)

    # Tiny decorative image; it must not become a candidate for Excel output.
    from tests.test_after_evidence import _pic
    slide.shapes.add_picture(_pic("#777777", "decorative"), 0, 0, width=int(0.08 * 914400), height=int(0.08 * 914400))


def _simple_production_after_slide(slide, width):
    _prod_head(slide, width)
    _tb(slide, "Cải tiến jig định vị", 1.7, 0.9, 5.6, 0.4, 12)
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 0.9, 1.5, 0.35)
    _pics(slide, ["#81c784"], 7.5, 1.3, w=1.7, h=0.6)


def _ownerless_after_slide(slide, width):
    _prod_head(slide, width)
    # A caption proves temporal After, but there is no confidently owned production item.
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 5.0, 1.0, 1.5, 0.35)
    _pics(slide, ["#ce93d8"], 5.0, 1.4, w=1.6, h=0.7)


def _color(block):
    with Image.open(BytesIO(block.image_blob)) as image:
        return image.convert("RGB").getpixel((5, 5))


def _image_record():
    return ExtractedRecord(management_number=MGMT, model="A185", item="Rear", defect_content="Defect",
                           root_cause="Cause", improvement="Improvement")


def _workbook_images(ws, row, col):
    return [image for image in ws._images
            if image.anchor._from.row == row - 1 and image.anchor._from.col == col - 1]


def _image_geometry(image):
    anchor = image.anchor
    return (anchor._from.colOff // EMU_PER_PX, anchor._from.rowOff // EMU_PER_PX,
            anchor.ext.cx // EMU_PER_PX, anchor.ext.cy // EMU_PER_PX)


def test_mixed_slide_requires_after_production_semantics_and_confident_owner(tmp_path):
    deck = _deck(tmp_path / NAME, [_mixed_production_slide, _ownerless_after_slide])
    report = parse_pptx(deck)
    selected = select_after_pictures(report, [3, 4])

    assert len(selected.after) == 3
    assert all(ref.excel_output_eligible and ref.temporal_role == "AFTER"
               and ref.semantic_role == SEMANTIC_PRODUCTION and ref.confident_owner and ref.owner_id
               for ref in selected.after)
    assert selected.slides_with_after == [3]

    # Three eligible photos belong to two logical production items, not one broad slide section.
    grouped = group_refs(selected.after)
    assert len(grouped) == 2 and sorted(map(len, grouped.values())) == [1, 2]
    ordered = [ref for refs in grouped.values() for ref in refs]
    assert [ref.source_order for ref in ordered] == sorted(ref.source_order for ref in ordered)

    rejected = selected.rejected
    by_color = {_color(ref.block): ref for ref in rejected if ref.block.image_blob}
    assert by_color[(255, 224, 178)].temporal_role == "BEFORE"                 # Before
    assert by_color[(179, 229, 252)].semantic_role == SEMANTIC_INSPECTION       # control / inspection
    assert by_color[(255, 249, 196)].semantic_role == SEMANTIC_TEMPORARY         # temporary action
    assert by_color[(209, 196, 233)].semantic_role == SEMANTIC_VERIFICATION     # follow-up / verification
    assert by_color[(255, 205, 210)].temporal_role == "UNKNOWN"                 # production, no After evidence
    assert by_color[(255, 205, 210)].confident_owner and not by_color[(255, 205, 210)].excel_output_eligible
    assert all(not ref.excel_output_eligible for ref in rejected)
    assert any(ref.exclusion_reason for ref in rejected)

    # Learning labels may describe relevance, but must not manufacture temporal or semantic eligibility.
    learning = ImageLearning(tmp_path / "learning")
    ambiguous = by_color[(255, 205, 210)]
    inspection = by_color[(179, 229, 252)]
    initial_candidates = build_candidates(report, [3, 4], selected, MGMT, str(deck))
    amb_candidate = next(c for c in initial_candidates if c.picture_id == ambiguous.block.shape_id and c.slide == 3)
    insp_candidate = next(c for c in initial_candidates if c.picture_id == inspection.block.shape_id and c.slide == 3)
    learning.store.label(amb_candidate, "AFTER")
    learning.store.label(insp_candidate, "AFTER")
    learned_refs, candidates, extra = select_with_learning(report, [3, 4], selected, learning, MGMT, str(deck))
    assert learning.store.counts()["after"] == 2                    # labels remain saved and trainable
    assert len(learned_refs) == 3 and all(ref.excel_output_eligible for ref in learned_refs)
    assert all((ref.slide, ref.block.shape_id) not in {
        (ambiguous.slide, ambiguous.block.shape_id), (inspection.slide, inspection.block.shape_id)}
               for ref in learned_refs)
    amb_candidate = next(c for c in candidates if c.picture_id == ambiguous.block.shape_id and c.slide == 3)
    assert amb_candidate.decision == "include" and amb_candidate.decision_source == "user"
    assert amb_candidate.temporal_role == "UNKNOWN" and not amb_candidate.excel_output_eligible
    assert any("không đủ điều kiện Excel" in reason for reason in extra)


def test_no_eligible_after_pictures_means_no_image_cell_output(template, tmp_path):
    deck = _deck(tmp_path / NAME, [_ownerless_after_slide])
    report = parse_pptx(deck)
    selected = select_after_pictures(report, [3])
    assert selected.after == [] and not selected.slides_with_after
    assert any(ref.temporal_role == "AFTER" and not ref.confident_owner for ref in selected.rejected)

    paths, problems = export_after_pictures(report, selected.after, tmp_path / "assets")
    assert not paths and not problems
    output = tmp_path / "no-images.xlsx"
    writer = ExcelWriter(template, output)
    writer.update_record(4, _image_record(), improvement_jpg=paths)
    writer.save()
    saved = load_workbook(output)[SHEET]
    assert _workbook_images(saved, 4, IMG_COL) == []


def test_force_reprocess_clears_old_improvement_images_when_none_remain(template, tmp_path):
    """A full force rewrite replaces the image cell with today's eligibility result, including no images."""
    deck = _deck(tmp_path / NAME, [_simple_production_after_slide])
    wb = load_workbook(template)
    wb[SHEET].cell(row=4, column=COL["mgmt"], value=MGMT)
    wb.save(template)
    output = tmp_path / "force-reprocess.xlsx"
    learning_dir = tmp_path / "learning"

    first = BatchProcessor(BatchOptions(files=[deck], template=template, output_file=output, use_ollama=False,
                                        learning_dir=learning_dir, use_fast_cache=False))
    first_summary = first.run()
    assert first_summary.failed == 0
    assert len(_workbook_images(load_workbook(output)[SHEET], 4, IMG_COL)) == 1

    # Replace the source with a temporally labelled picture that has no confident item owner.
    _deck(deck, [_ownerless_after_slide])
    second = BatchProcessor(BatchOptions(files=[deck], template=template, output_file=output, use_ollama=False,
                                         force_reprocess=True, learning_dir=learning_dir, use_fast_cache=False))
    second_summary = second.run()
    assert second_summary.failed == 0 and second.results[0].after_pictures == []
    saved = load_workbook(output)[SHEET]
    assert _workbook_images(saved, 4, IMG_COL) == []


def test_grouped_owner_rows_preserve_order_aspect_and_non_overlap(template, tmp_path):
    source_sizes = [(400, 200), (200, 400), (1600, 900)]
    paths = []
    for index, size in enumerate(source_sizes):
        path = tmp_path / f"eligible-{index}.png"
        Image.new("RGB", size, (60 + index * 40, 150, 210)).save(path)
        paths.append(path)
    grouped_paths = GroupedImagePaths([paths[:2], paths[2:]], owners=["item-a", "item-b"])
    output = tmp_path / "grouped.xlsx"
    writer = ExcelWriter(template, output)
    writer.update_record(4, _image_record(), improvement_jpg=grouped_paths)
    writer.save()

    ws = load_workbook(output)[SHEET]
    images = _workbook_images(ws, 4, IMG_COL)
    assert len(images) == 3
    geometry = [_image_geometry(image) for image in images]             # insertion/source order
    (x0, y0, w0, h0), (x1, y1, w1, h1), (x2, y2, w2, h2) = geometry
    assert x0 < x1 and abs((y0 + h0 / 2) - (y1 + h1 / 2)) <= 1        # one horizontal row for item A
    assert y2 > max(y0, y1) and x2 == IMAGE_MARGIN_PX                  # item B starts a new row
    for (src_w, src_h), (_, _, width, height) in zip(source_sizes, geometry):
        assert abs((width / height) / (src_w / src_h) - 1.0) < 0.06     # proportional scaling
    area_w = writer.image_area_px(4, "improvement_image")
    area_h = writer._area_height_px(4, None)
    assert_image_inside_area(geometry, area_w, area_h)
    for i, (x_i, y_i, w_i, h_i) in enumerate(geometry):
        for x_j, y_j, w_j, h_j in geometry[i + 1:]:
            assert x_i + w_i <= x_j or x_j + w_j <= x_i or y_i + h_i <= y_j or y_j + h_j <= y_i
