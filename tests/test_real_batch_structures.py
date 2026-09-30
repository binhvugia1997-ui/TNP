"""Regression fixtures modelled on the five real production reports of the first qwen3:4b batch
(260601030-VOC, 260601037-VOC, 260601038-VOC, 260601039-VOC, 260601017).

Common real layout: slide 1 cover; slide 2 = QPN form as a flat picture (no 'Quality Problem
Notice' text visible to python-pptx) + NGUYÊN NHÂN text; slide 3 = XỬ LÝ TẠM THỜI; slides 4..6 =
CẢI TIẾN with pictures (slide 5 is a heading-less continuation); slide 7 = ĐỐI SÁCH LÂU DÀI (text only).
"""
import datetime as dt
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches

from make_samples import _pic, _textbox  # noqa: E402  (tools/ on sys.path via conftest)

from app.batch_processor import BatchOptions, BatchProcessor
from app.classifier import classify, heuristic_classify, merge_llm_with_heuristic
from app.extractor import extract_record
from app.pptx_parser import parse_pptx

W = Inches(13.333)


def build_real_like(path: Path, *, improvement_vendor_text: str, long_term: bool = True,
                    with_qpn_text_in_alt: bool = False) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, Inches(7.5)
    blank = prs.slide_layouts[6]
    # 1 cover – no section heading; several dates that must NEVER become Ngày phát sinh
    s = prs.slides.add_slide(blank)
    _textbox(s, "BÁO CÁO ĐỐI SÁCH LỖI XƯỚC REAR A185", Inches(0.5), Inches(1), W - Inches(1), Inches(1), 28, True)
    _textbox(s, "Model: A185\nItem: Rear\nNgày báo cáo: 05/06/2026\nNgười lập: QA team", Inches(0.5), Inches(2.5), Inches(6), Inches(2), 16)
    # 2 QPN form as one large flat picture + root cause text under it
    s = prs.slides.add_slide(blank)
    pic = s.shapes.add_picture(_pic("#fff2cc", "QPN FORM  Date 15/05/2026  Xước 12ea", (1280, 560)),
                               Inches(0.3), Inches(0.2), width=Inches(12.5), height=Inches(3.6))
    if with_qpn_text_in_alt:
        pic._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("descr", "Quality Problem Notice")
    _textbox(s, "2. NGUYÊN NHÂN\n- Khay chứa Rear không có lớp lót nên sản phẩm va chạm gây xước (phát hiện 20/05/2026)\n"
                "- Jig ép có cạnh sắc", Inches(0.5), Inches(4.0), W - Inches(1), Inches(3), 14, True)
    # 3 temporary handling
    s = prs.slides.add_slide(blank)
    _textbox(s, "XỬ LÝ TẠM THỜI\n- Sorting 100% hàng tồn từ 21/05/2026: 1.200 pcs\n- Dán tem tạm thời", Inches(0.5), Inches(0.3), W - Inches(1), Inches(2), 16, True)
    s.shapes.add_picture(_pic("#ddd", "sorting"), Inches(1), Inches(3), width=Inches(5))
    # 4 improvement (heading + picture)
    s = prs.slides.add_slide(blank)
    _textbox(s, "3. CẢI TIẾN TRONG SẢN XUẤT\n" + improvement_vendor_text +
                "\n- Trước: khay nhựa cứng\n- Sau: khay có lót mút EVA (áp dụng từ 25/05/2026)", Inches(0.5), Inches(0.3), W - Inches(1), Inches(2.5), 16, True)
    s.shapes.add_picture(_pic("#cfc", "before/after"), Inches(1), Inches(3.5), width=Inches(5))
    # 5 continuation (no heading) with picture
    s = prs.slides.add_slide(blank)
    _textbox(s, "- Bổ sung vách ngăn từng ô trong khay\n- Cập nhật hướng dẫn thao tác", Inches(0.5), Inches(0.3), W - Inches(1), Inches(1.5), 16)
    s.shapes.add_picture(_pic("#ccf", "khay moi"), Inches(1), Inches(2.5), width=Inches(5))
    # 6 improvement in inspection (heading + picture)
    s = prs.slides.add_slide(blank)
    _textbox(s, "4. CẢI TIẾN TRONG KIỂM TRA\n- Bổ sung đèn 3 hướng tại OQC", Inches(0.5), Inches(0.3), W - Inches(1), Inches(1.5), 16, True)
    s.shapes.add_picture(_pic("#fcc", "den kiem tra"), Inches(1), Inches(2.5), width=Inches(5))
    # 7 long-term action (text only)
    if long_term:
        s = prs.slides.add_slide(blank)
        _textbox(s, "5. ĐỐI SÁCH LÂU DÀI\n- Đưa kiểm tra khay vào checklist đầu ca\n- Cập nhật SOP Rev.03", Inches(0.5), Inches(0.3), W - Inches(1), Inches(2), 16, True)
    prs.save(path)
    return path


REAL = {
    "260601030-VOC": dict(name="(CTMS)_11201_260601030-VOC_ Đối sách lỗi xước Rear A185 05.06.2026.pptx",
                          vendor_text="- Tại công đoạn lắp ráp Rear, thao tác đặt sản phẩm vào khay", vendors=[]),
    "260601037-VOC": dict(name="(CTMS)_11202_260601037-VOC_ Đối sách lỗi xước Rear A185 05.06.2026.pptx",
                          vendor_text="- Tại công đoạn Assy JT Tech Vina, thao tác lấy hàng", vendors=["JT Tech Vina"]),
    "260601038-VOC": dict(name="(CTMS)_11203_260601038-VOC_ Đối sách lỗi xước Rear A185 05.06.2026.pptx",
                          vendor_text="- Vendor Mtech: đổi khay\n- Tesung: kiểm tra lại bề mặt\n- Tại công đoạn Assy IT: thay khay",
                          vendors=["Mtech", "Tesung", "Assy IT"]),
    "260601039-VOC": dict(name="(CTMS)_11204_260601039-VOC_ Đối sách lỗi xước Rear A185 05.06.2026.pptx",
                          vendor_text="- Công đoạn Sơn IT: kiểm tra bề mặt sơn", vendors=["Sơn IT"]),
    "260601017": dict(name="(CTMS)_11205_260601017_ Đối sách lỗi xước Rear A185 05.06.2026.pptx",
                      vendor_text="- Tesung: bổ sung film bảo vệ", vendors=["Tesung"]),
}


@pytest.fixture(scope="module")
def real_like(tmp_path_factory):
    root = tmp_path_factory.mktemp("real5")
    return {k: build_real_like(root / v["name"], improvement_vendor_text=v["vendor_text"]) for k, v in REAL.items()}


# ---------------------------------------------------------------- 1. date only from Management Number
@pytest.mark.parametrize("key", list(REAL))
def test_occurrence_date_ignores_every_pptx_date(real_like, key):
    r = parse_pptx(real_like[key])
    rec = extract_record(r, heuristic_classify(r))
    assert rec.management_number == key
    assert rec.occurrence_date == dt.date(2026, 6, 1) and rec.occurrence_date_text == "01/06/2026"
    # the deck contains 05/06/2026, 15/05/2026, 20/05/2026, 21/05/2026, 25/05/2026 – none may leak
    assert "05/06/2026" not in rec.occurrence_date_text and "15/05/2026" not in rec.occurrence_date_text


# ---------------------------------------------------------------- 2/3. QPN precedence & warning
@pytest.mark.parametrize("key", list(REAL))
def test_heuristic_qpn_is_slide_2_by_structure(real_like, key):
    r = parse_pptx(real_like[key])
    c = heuristic_classify(r)
    assert c.qpn_slide == 2 and c.qpn_source == "structural"
    assert c.cause_slides == [2] and c.temporary_slides == [3]
    assert c.improvement_slides == [4, 5, 6, 7]
    assert c.improvement_image_slides == [4, 5, 6]          # 7 is text-only long-term action
    assert not any("QPN" in a for a in c.ambiguities)


@pytest.mark.parametrize("key", ["260601030-VOC", "260601037-VOC"])
def test_llm_cover_slide_qpn_is_overridden(real_like, key):
    r = parse_pptx(real_like[key])
    heur = heuristic_classify(r)
    raw = {"qpn_slide": 1, "defect_slide": 2, "cause_slides": [2], "temporary_slides": [3],
           "improvement_slides": [4, 5, 6, 7], "improvement_image_slides": [4], "confidence": 0.9}
    c = merge_llm_with_heuristic(raw, heur, r)
    assert c.qpn_slide == 2 and c.qpn_source == "structural"
    assert "LLM qpn_slide=1 rejected" in c.qpn_override
    assert not any("chỉ do AI" in a for a in c.ambiguities)
    rec = extract_record(r, c)
    assert not any("chỉ do AI" in x or "Không tìm thấy QPN" in x for x in rec.review_reasons)


def test_cover_never_becomes_qpn_even_without_structural_candidate(tmp_path):
    # same deck but slide 2 without any picture -> no deterministic QPN; LLM cover suggestion still rejected
    p = build_real_like(tmp_path / "(CTMS)_1_260601030-VOC_x.pptx", improvement_vendor_text="- x")
    prs = Presentation(p)
    s2 = prs.slides[1]
    for shp in list(s2.shapes):
        if shp.shape_type == 13:
            shp._element.getparent().remove(shp._element)
    prs.save(p)
    r = parse_pptx(p)
    heur = heuristic_classify(r)
    assert heur.qpn_slide is None
    c = merge_llm_with_heuristic({"qpn_slide": 1, "cause_slides": [2]}, heur, r)
    assert c.qpn_slide is None and "cover slide without QPN evidence" in c.qpn_override
    rec = extract_record(r, c)
    assert "Không tìm thấy QPN trong báo cáo" in rec.review_reasons


def test_explicit_evidence_outranks_structure(tmp_path):
    p = build_real_like(tmp_path / "(CTMS)_1_260601030-VOC_alt.pptx", improvement_vendor_text="- x", with_qpn_text_in_alt=True)
    c = heuristic_classify(parse_pptx(p))
    assert c.qpn_slide == 2 and c.qpn_source == "metadata"


# ---------------------------------------------------------------- 4. Nội dung lỗi vs QPN image
@pytest.mark.parametrize("key", ["260601030-VOC", "260601037-VOC"])
def test_defect_from_filename_when_qpn_is_only_an_image(real_like, key):
    # file name "... Đối sách lỗi xước Rear A185 05.06.2026" -> defect names from the name, not OCR / not Qwen
    r = parse_pptx(real_like[key])
    rec = extract_record(r, heuristic_classify(r))
    assert rec.defect_content == "xước"
    assert not any("Nội dung lỗi" in x for x in rec.review_reasons)
    # cause text under the QPN picture is cause, never defect text
    assert "Khay chứa Rear" in rec.root_cause and "Khay" not in rec.defect_content


def test_defect_only_in_image_and_no_marker_in_filename(tmp_path):
    p = build_real_like(tmp_path / "(CTMS)_1_260601030-VOC_ Bao cao A185 Rear 05.06.2026.pptx", improvement_vendor_text="- x")
    r = parse_pptx(p)
    rec = extract_record(r, heuristic_classify(r))
    assert rec.defect_content == ""
    assert "Nội dung lỗi chỉ có trong hình ảnh QPN, không có text để sao chép – cần bổ sung thủ công" in rec.review_reasons
    assert "Không tìm thấy Nội dung lỗi trong báo cáo" not in rec.review_reasons


def test_defect_text_copied_when_recoverable(tmp_path):
    p = build_real_like(tmp_path / "(CTMS)_1_260601030-VOC_ Bao cao A185 txt.pptx", improvement_vendor_text="- x")
    prs = Presentation(p)
    _textbox(prs.slides[1], "HIỆN TRẠNG\nXước mặt sau Rear: 12ea\nMẻ cạnh: 3ea", Inches(0.3), Inches(3.7), Inches(6), Inches(0.6), 12, True)
    prs.save(p)
    r = parse_pptx(p)
    rec = extract_record(r, heuristic_classify(r))
    assert rec.defect_content == "Xước mặt sau Rear: 12ea\nMẻ cạnh: 3ea"
    assert not any("Nội dung lỗi" in x for x in rec.review_reasons)


# ---------------------------------------------------------------- 5. images: union, never narrowed
@pytest.mark.parametrize("key,llm_imgs", [("260601038-VOC", [4]), ("260601017", [5]), ("260601039-VOC", [])])
def test_llm_cannot_narrow_structural_improvement_images(real_like, key, llm_imgs):
    r = parse_pptx(real_like[key])
    heur = heuristic_classify(r)
    raw = {"qpn_slide": 2, "cause_slides": [2], "temporary_slides": [3], "improvement_slides": [4, 5, 6, 7],
           "improvement_image_slides": llm_imgs, "confidence": 0.9}
    c = merge_llm_with_heuristic(raw, heur, r)
    assert c.improvement_image_slides == [4, 5, 6]
    assert c.image_slides_structural == [4, 5, 6] and c.image_slides_llm == llm_imgs
    assert c.improvement_slides == [4, 5, 6, 7]
    rec = extract_record(r, c)
    assert "Cập nhật SOP Rev.03" in rec.improvement and "ĐỐI SÁCH LÂU DÀI" not in rec.improvement and "Sorting 100%" not in rec.improvement
    assert rec.improvement_image_slides == [4, 5, 6]


def test_llm_can_add_only_safe_candidates(real_like):
    r = parse_pptx(real_like["260601038-VOC"])
    heur = heuristic_classify(r)
    raw = {"qpn_slide": 2, "cause_slides": [2], "temporary_slides": [3], "improvement_slides": [4, 5, 6, 7],
           "improvement_image_slides": [1, 2, 3, 4, 7]}
    c = merge_llm_with_heuristic(raw, heur, r)
    assert c.improvement_image_slides == [4, 5, 6]          # cover / QPN / temporary / text-only rejected
    assert any("rejected" in n for n in c.notes)


# ---------------------------------------------------------------- end-to-end with a fake qwen3:4b
class FakeQwen:
    """Returns what the real batch returned: cover as QPN and a narrowed image list."""
    def __init__(self):
        self.last_call = {"status": 200, "response_chars": 120}

    def generate_json(self, model, prompt, system="", **kw):
        return {"qpn_slide": 1, "defect_slide": 2, "cause_slides": [2], "temporary_slides": [3],
                "improvement_slides": [4, 5, 6, 7], "improvement_image_slides": [4], "confidence": 0.9}


def _template_with_rows(template: Path, rows):
    wb = load_workbook(template)
    ws = wb["Kiểm chứng"]
    for i, (mgmt, vendor) in enumerate(rows):
        ws.cell(row=4 + i, column=1, value=i + 1)
        ws.cell(row=4 + i, column=2, value=mgmt)
        if vendor is not None:
            ws.cell(row=4 + i, column=3, value=vendor)
        for k in range(12, 20):
            ws.cell(row=4 + i, column=k, value="OK")
    wb.save(template)


def test_five_report_batch_expectations(real_like, template, tmp_path, monkeypatch):
    import app.batch_processor as bp
    monkeypatch.setattr(bp, "OllamaClient", lambda *a, **k: FakeQwen())
    rows = [("260601030-VOC", None), ("260601037-VOC", "NCC JT"), ("260601038-VOC", None),
            ("260601039-VOC", "HK"), (260601017, "Mtech")]
    _template_with_rows(template, rows)
    out = tmp_path / "out" / "k.xlsx"
    opts = BatchOptions(files=[real_like[k] for k in REAL], template=template, output_file=out,
                        ollama_server="127.0.0.1:11434", model="qwen3:4b", use_ollama=True)
    proc = BatchProcessor(opts, on_file=lambda *a: None)
    s = proc.run()
    assert s.failed == 0 and s.not_written == 0 and s.completed + s.needs_review == 5
    ws = load_workbook(out)["Kiểm chứng"]
    by_key = {fr.management_number: fr for fr in proc.results}
    for i, key in enumerate(REAL):
        fr = by_key[key]
        row = 4 + i
        assert fr.excel_row == row and fr.classifier.startswith("qwen")
        assert fr.qpn_slide == 2 and fr.qpn_source == "structural" and "rejected" in fr.qpn_override
        assert fr.cause_slides == [2] and fr.temporary_slides == [3]
        assert fr.improvement_slides == [4, 5, 6, 7] and fr.improvement_image_slides == [4, 5, 6]
        assert fr.image_slides_structural == [4, 5, 6] and fr.image_slides_llm == [4]
        assert fr.occurrence_date == "01/06/2026"
        assert ws.cell(row=row, column=4).value.date() == dt.date(2026, 6, 1)
        assert ws.cell(row=row, column=5).value == "A185" and ws.cell(row=row, column=6).value == "Rear"
        assert "Cập nhật SOP Rev.03" in ws.cell(row=row, column=10).value
        assert "Sorting 100%" not in ws.cell(row=row, column=10).value
        assert all(ws.cell(row=row, column=k).value == "OK" for k in range(12, 20))
        assert ws.cell(row=row, column=7).value == "xước"                            # from the file name
        assert not any("Nội dung lỗi" in x for x in fr.review_reasons)
        assert not any("chỉ do AI" in x for x in fr.review_reasons)
    # vendor expectations
    r030, r037, r038, r039, r017 = (by_key[k] for k in REAL)
    assert r030.vendor == "" and "Không xác định được Vendor từ nội dung báo cáo" in r030.review_reasons
    assert ws.cell(row=4, column=3).value is None
    assert r037.vendor == "JT Tech Vina" and ws.cell(row=5, column=3).value == "NCC JT"
    assert not any("Vendor trong Excel" in x for x in r037.review_reasons)
    assert r038.vendor == "Mtech\nTesung\nAssy IT" and ws.cell(row=6, column=3).value == "Mtech\nTesung\nAssy IT"
    assert ws.cell(row=7, column=3).value == "HK" and any("Vendor trong Excel: HK" in x and "Sơn IT" in x for x in r039.review_reasons)
    assert ws.cell(row=8, column=3).value == "Mtech" and any("Vendor trong Excel: Mtech" in x and "Tesung" in x for x in r017.review_reasons)
    assert ws.cell(row=8, column=2).value == 260601017                              # bare key untouched
    # diagnostics persisted in batch_result.json
    import json
    res = json.loads((out.parent / "logs" / "batch_result.json").read_text(encoding="utf-8"))
    r0 = res["results"][0]
    assert r0["qpn_source"] == "structural" and r0["image_slides_structural"] == [4, 5, 6] and r0["image_slides_llm"] == [4]
    assert "rejected" in r0["qpn_override"]
