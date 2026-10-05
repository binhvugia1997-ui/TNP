"""PROMPT-001 (v1.0.1) regression tests.

BUG A – the "QPN" image is the Quality Problem Notice PANEL only (object/region based), never the whole
         "1. HIỆN TRẠNG" slide; fail-closed when the panel cannot be isolated.
BUG B – the improvement text stops before inspection/control sections and follow-up/effectiveness sections,
         fixed in the segmentation layer (a mid-slide heading terminates the previous section); inspection
         sections contribute zero text and zero pictures; the LLM cannot reintroduce them.
"""
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches

import app
from app.batch_processor import BatchOptions, BatchProcessor, format_file_diagnostics
from app.classifier import Classification, heuristic_classify, improvement_subkind, section_kind_of_heading
from app.extractor import _semantic_kind, extract_record, split_sections
from app.improvement_pictures import select_after_pictures
from app.pptx_parser import parse_pptx
from app.qpn_region import locate_qpn_region
from app.qpn_renderer import SlideRenderer, render_qpn, render_qpn_panel

from tests.test_content_region import _furniture, _images_at, _pic, _prefill, _shape, _tb, COL, SHEET

QPN_COLOR = (255, 0, 0)            # only inside the real QPN panel picture
MGMT = "260923045-VOC"

PROD_TEXT = ("Cải tiến lỗi mẻ xước (Áp dụng từ ngày 20/09/2026 - Công đoạn Assy):\n"
             "- Cải tiến jig nén tape sealing\n"
             "+ Trước: Jig nén bằng nhôm, cạnh sắc, không có lớp đệm\n"
             "+ Sau: Bọc silicon 2mm toàn bộ mặt tiếp xúc của jig, bo tròn cạnh R1.0")
INSPECTION_HEADING = "Cải tiến trong kiểm tra"
INSPECTION_TEXT = ("- Đào tạo người kiểm tra về tiêu chuẩn giới hạn lỗi mẻ xước\n"
                   "- Kiểm tra lại toàn bộ hàng tồn kho 1.250 pcs\n"
                   "- Bổ sung công đoạn kiểm tra nghiêng 45° dưới đèn 1000 lux")
FOLLOWUP_HEADING = "Theo dõi hiệu quả cải tiến"
FOLLOWUP_TEXT = ("- Theo dõi hiệu quả 4 tuần liên tiếp sau cải tiến\n"
                 "- Audit kiểm tra hàng tuần tại công đoạn Assy")
LONG_TERM_TEXT = ("- Duy trì audit kiểm tra hàng tuần, theo dõi hiệu quả đến hết Q4\n"
                  "- Theo dõi tỷ lệ lỗi hàng tuần trên dashboard")
FORBIDDEN = ["đào tạo người kiểm tra", "hàng tồn kho", "công đoạn kiểm tra nghiêng", "đối sách lâu dài",
             "theo dõi hiệu quả", "audit kiểm tra"]


def _qpn_slide(prs, mode: str):
    """'1. HIỆN TRẠNG' page: title, sidebar circle, dotted separator, note and footer – all physically close to
    the QPN panel.  mode: picture | named_picture | table | ambiguous."""
    W = prs.slide_width
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _tb(s, "1. HIỆN TRẠNG", 0.5, 0.25, 8, 0.7, 24, True)                                   # title band
    _shape(s, MSO_SHAPE.OVAL, "Hiện trạng", 0.15, 2.8, 1.3, 1.3)                             # sidebar circle
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.6), Inches(0.98), W - Inches(0.4), Inches(0.98))
    ln.line.dash_style = 4                                                                   # dotted separator
    if mode in ("picture", "named_picture"):
        pic = s.shapes.add_picture(_pic(QPN_COLOR, "qpn", (960, 540)), Inches(1.7), Inches(1.05), Inches(9.6), Inches(5.4))
        if mode == "named_picture":
            pic.name = "QPN_260923045"
    elif mode == "table":
        _tb(s, "Quality Problem Notice", 1.7, 1.05, 6, 0.5, 16, True)
        tbl = s.shapes.add_table(4, 4, Inches(1.7), Inches(1.6), Inches(9.6), Inches(4.4)).table
        for c, h in enumerate(["Model", "Defect Info", "Cause & Action", "PIC"]):
            tbl.cell(0, c).text = h
        tbl.cell(1, 0).text = "A185"
        tbl.cell(1, 1).text = "Mẻ xước: 13ea"
        tbl.cell(1, 2).text = "Xem báo cáo đối sách"
        tbl.cell(1, 3).text = "Nguyễn Văn A"
    else:                                                                                   # two similar pictures
        s.shapes.add_picture(_pic("#ffe0b2", "a", (480, 360)), Inches(1.7), Inches(1.2), Inches(4.5), Inches(3.4))
        s.shapes.add_picture(_pic("#c8e6c9", "b", (480, 360)), Inches(6.8), Inches(1.2), Inches(4.5), Inches(3.4))
    _tb(s, "* Ghi chú: QPN do khách hàng phát hành ngày 23/09/2026", 1.7, 6.55, 8, 0.3, 8)   # note
    _furniture(s, W)
    return s


def make_deck(path: Path, qpn_mode: str = "picture", mgmt: str = MGMT) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W = prs.slide_width
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)                                                          # 1 cover
    _tb(s, "BÁO CÁO ĐỐI SÁCH LỖI MẺ XƯỚC REAR A185", 0.5, 1, 12, 1, 28, True)
    _tb(s, f"Model: A185\nItem: Rear\nManagement No: {mgmt}", 0.5, 2.5, 6, 2, 16)

    _qpn_slide(prs, qpn_mode)                                                                # 2 QPN

    s = prs.slides.add_slide(blank)                                                          # 3 cause
    _tb(s, "2. NGUYÊN NHÂN", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Nguyên nhân", 0.15, 2.5, 1.3, 1.3)
    _tb(s, "- Jig nén bằng nhôm cạnh sắc gây mẻ xước khi nén", 1.7, 1.1, 11, 2.2, 13)
    _furniture(s, W)

    s = prs.slides.add_slide(blank)                                                          # 4 production -> inspection -> follow-up (one slide)
    _tb(s, "3. CẢI TIẾN TRONG SẢN XUẤT", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Cải tiến", 0.15, 3.0, 1.3, 1.3)
    _tb(s, PROD_TEXT, 1.7, 1.0, 5.6, 1.9, 12)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Trước cải tiến", 7.5, 1.0, 1.5, 0.35)
    s.shapes.add_picture(_pic("#ffe0b2", "before"), Inches(7.5), Inches(1.4), Inches(1.6), Inches(1.2))
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 9.4, 1.0, 1.5, 0.35)
    s.shapes.add_picture(_pic("#c8e6c9", "after"), Inches(9.4), Inches(1.4), Inches(1.6), Inches(1.2))
    _tb(s, INSPECTION_HEADING, 1.7, 3.0, 5.6, 0.4, 13, True)                                # mid-slide heading
    _tb(s, INSPECTION_TEXT, 1.7, 3.4, 5.6, 1.4, 12)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 3.0, 1.5, 0.35)
    s.shapes.add_picture(_pic("#b3e5fc", "inspection"), Inches(7.5), Inches(3.4), Inches(3.5), Inches(1.4))
    _tb(s, FOLLOWUP_HEADING, 1.7, 4.9, 5.6, 0.4, 13, True)                                  # mid-slide heading
    _tb(s, FOLLOWUP_TEXT, 1.7, 5.3, 5.6, 1.2, 12)
    _furniture(s, W)

    s = prs.slides.add_slide(blank)                                                          # 5 long-term = follow-up only
    _tb(s, "4. ĐỐI SÁCH LÂU DÀI", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Đối sách lâu dài", 0.15, 2.5, 1.3, 1.3)
    _tb(s, LONG_TERM_TEXT, 1.7, 1.1, 11, 2, 13)
    _furniture(s, W)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def _name(mode="picture"):
    return f"(CTMS)_1_{MGMT}_ Đối sách LỖI MẺ XƯỚC 24.9.2026 {mode}.pptx"


def _run(files, template, out, **kw):
    opts = BatchOptions(files=[Path(f) for f in files], template=template, output_file=out, use_ollama=False, **kw)
    proc = BatchProcessor(opts)
    return proc.run(), proc


def _norm(s: str) -> str:
    return " ".join((s or "").lower().split())


# ------------------------------------------------------------------ versioning
def test_single_canonical_version():
    assert app.__version__ == "1.2.0" and app.BUILD_NUMBER == 11 and app.BUILD_ID == "011"
    assert app.APP_TITLE == "Report Extractor v1.2.0" and app.BUILD_LABEL == "Build 011"
    assert app.VERSION_LINE == "version=1.2.0 build=011"
    src = Path(app.__file__).parent
    hits = [p for p in src.glob("*.py") if p.name != "__init__.py" and "1.0.2" in p.read_text(encoding="utf-8")]
    assert hits == [], f"version string duplicated in {hits}"


# ------------------------------------------------------------------ BUG A: QPN panel only
@pytest.mark.parametrize("mode", ["picture", "named_picture", "table"])
def test_qpn_region_excludes_title_sidebar_footer(tmp_path, mode):
    r = parse_pptx(make_deck(tmp_path / _name(mode), mode))
    s = r.slide(2)
    loc = locate_qpn_region(s, r.slide_width, r.slide_height)
    assert loc.region is not None, loc.reason
    reg = loc.region
    title = next(b for b in s.text_blocks if "HIỆN TRẠNG" in b.text)
    sidebar = next(b for b in s.text_blocks if b.text.strip() == "Hiện trạng")
    note = next(b for b in s.text_blocks if b.text.startswith("* Ghi chú"))
    footer = next(b for b in s.text_blocks if "Confidential" in b.text)
    assert reg.top >= title.bottom and reg.left >= sidebar.right          # physically adjacent but excluded
    assert reg.bottom <= note.top and reg.bottom < footer.top
    assert reg.width * reg.height < 0.8 * r.slide_width * r.slide_height   # never the whole slide
    expect = {"picture": ("dominant_picture",), "named_picture": ("picture_evidence",), "table": ("table_panel", "text_panel")}[mode]
    assert reg.method in expect


def test_qpn_image_is_panel_only(tmp_path):
    r = parse_pptx(make_deck(tmp_path / _name("picture"), "picture"))
    res = render_qpn_panel(r, 2, tmp_path / "qpn.png", SlideRenderer())
    assert res.ok and res.backend == "picture"
    with Image.open(res.path) as im:
        assert im.size == (960, 540)
        px = [im.getpixel((x, y)) for x in (2, im.width // 2, im.width - 3) for y in (2, im.height // 2, im.height - 3)]
    assert all(p == QPN_COLOR for p in px)                                   # no title / sidebar / white slide left


def test_qpn_table_panel_is_cropped_from_render(tmp_path):
    r = parse_pptx(make_deck(tmp_path / _name("table"), "table"))
    res = render_qpn_panel(r, 2, tmp_path / "qpn.png", SlideRenderer())
    assert res.ok and res.crop_px is not None
    l, t, rr, b = res.crop_px
    with Image.open(res.path) as im:
        assert im.size == (rr - l, b - t)
    # crop proportional to the EMU region (+ small pad), far from the full slide render
    reg = res.region
    assert abs((rr - l) / (b - t) - reg.width / reg.height) < 0.15
    assert t > 0 and l > 0


def test_qpn_fail_closed_no_whole_slide(tmp_path):
    r = parse_pptx(make_deck(tmp_path / _name("ambiguous"), "ambiguous"))
    res = render_qpn_panel(r, 2, tmp_path / "qpn.png", SlideRenderer())
    assert not res.ok and res.path is None and "nổi trội" in res.reason
    assert not (tmp_path / "qpn.png").exists()
    with pytest.raises(RuntimeError):
        render_qpn(r, 2, tmp_path / "qpn2.png", SlideRenderer())


def test_batch_qpn_panel_and_fail_closed(template, tmp_path):
    ok_deck = make_deck(tmp_path / "a" / _name("picture"), "picture")
    bad_deck = make_deck(tmp_path / "b" / _name("ambiguous").replace(MGMT, "260923046-VOC"), "ambiguous", "260923046-VOC")
    out = tmp_path / "out" / "k.xlsx"
    _prefill(template, [{"mgmt": MGMT}, {"mgmt": "260923046-VOC"}])        # PROMPT-004: rows must pre-exist
    summary, proc = _run([ok_deck, bad_deck], template, out)
    fr_ok, fr_bad = proc.results
    assert fr_ok.qpn_image and Path(fr_ok.qpn_image).exists()
    with Image.open(fr_ok.qpn_image) as im:
        assert im.size == (960, 540)
    assert fr_ok.qpn_region.startswith("dominant_picture")
    assert any("title" in e for e in fr_ok.qpn_excluded)
    assert not fr_bad.qpn_image
    assert fr_bad.status == "needs_review"
    assert any("Không tách được QPN khỏi slide 2" in m for m in fr_bad.review_reasons)
    ws = load_workbook(out)[SHEET]
    rows = {ws.cell(row=r, column=COL["mgmt"]).value: r for r in range(4, 8) if ws.cell(row=r, column=COL["mgmt"]).value}
    assert len(_images_at(ws, rows[MGMT], COL["qpn"])) == 1
    assert len(_images_at(ws, rows["260923046-VOC"], COL["qpn"])) == 0
    diag = format_file_diagnostics(fr_bad)
    assert "Vùng QPN" in diag and "không tách được" in diag


# ------------------------------------------------------------------ BUG B: improvement text boundaries
def test_improvement_subkind():
    assert improvement_subkind("Cải tiến trong kiểm tra") == "inspection"
    assert improvement_subkind("Cải tiến trong sản xuất") == "production"
    assert improvement_subkind("Theo dõi hiệu quả cải tiến") == "followup"
    assert improvement_subkind("Cải tiến jig tại công đoạn kiểm tra OQC") == "production"   # mixed -> production


def test_mid_slide_heading_terminates_section(tmp_path):
    r = parse_pptx(make_deck(tmp_path / _name(), "picture"))
    secs = [s for s in split_sections(r.slide(4), "improvement") if s.text.strip()]
    kinds = [s.kind for s in secs]
    assert kinds[0] == "improvement" and kinds[1] == "inspection" and kinds[2] in ("followup", "verify"), kinds
    assert "+ Sau: Bọc silicon" in secs[0].text and "đào tạo" not in secs[0].text.lower()


def test_improvement_text_stops_before_inspection_and_followup(tmp_path):
    r = parse_pptx(make_deck(tmp_path / _name(), "picture"))
    cls = heuristic_classify(r)
    rec = extract_record(r, cls, {"rear": "Rear"}, ["A185"])
    imp = _norm(rec.improvement)
    assert "+ sau: bọc silicon 2mm" in imp
    for bad in FORBIDDEN:
        assert bad not in imp, bad
    assert "jig nén bằng nhôm" in _norm(rec.root_cause)                       # root-cause extraction untouched
    assert any(e.startswith("S4 inspection") for e in rec.excluded_sections)
    assert any(e.startswith(("S4 followup", "S4 verify")) for e in rec.excluded_sections)
    assert any(e.startswith("S5 followup") for e in rec.excluded_sections)   # ĐỐI SÁCH LÂU DÀI = follow-up only
    sel = select_after_pictures(r, cls.improvement_image_slides or cls.improvement_slides)
    assert len(sel.after) == 1                                              # inspection picture never selected
    assert all(p.slide == 4 and p.block.top < Inches(3.0) for p in sel.after)


def test_qwen_selecting_mixed_slide_cannot_reintroduce_inspection(tmp_path):
    """The LLM may only point at slides; segmentation still drops inspection/follow-up parts."""
    r = parse_pptx(make_deck(tmp_path / _name(), "picture"))
    cls = Classification(management_number=MGMT, qpn_slide=2, cause_slides=[3], improvement_slides=[4, 5],
                         improvement_image_slides=[4, 5], source="qwen", confidence=0.9)
    rec = extract_record(r, cls, {"rear": "Rear"}, ["A185"])
    imp = _norm(rec.improvement)
    for bad in FORBIDDEN:
        assert bad not in imp, bad
    assert "+ sau: bọc silicon 2mm" in imp
    sel = select_after_pictures(r, [4, 5])
    assert len(sel.after) == 1 and sel.after[0].block.top < Inches(3.0)


def test_rerun_and_force_are_stable(template, tmp_path):
    deck = make_deck(tmp_path / _name(), "picture")
    _prefill(template, [{"mgmt": MGMT, "vendor": "Doaltech", "model": "A185", "item": "Rear"}])
    out = tmp_path / "out" / "k.xlsx"
    _, p1 = _run([deck], template, out)
    ws = load_workbook(out)[SHEET]
    imp1 = ws.cell(row=4, column=COL["improvement"]).value
    for bad in FORBIDDEN:
        assert bad not in _norm(imp1), bad
    assert len(_images_at(ws, 4, COL["qpn"])) == 1 and len(_images_at(ws, 4, COL["image"])) == 1
    _, p2 = _run([deck], template, out, force_reprocess=True)
    ws2 = load_workbook(out)[SHEET]
    assert ws2.cell(row=4, column=COL["improvement"]).value == imp1
    assert len(_images_at(ws2, 4, COL["qpn"])) == 1 and len(_images_at(ws2, 4, COL["image"])) == 1
    with Image.open(p2.results[0].qpn_image) as im:
        assert im.size == (960, 540)
    _, p3 = _run([deck], template, out)                                        # plain rerun: complete -> skipped
    assert p3.results[0].status == "skipped"


# ------------------------------------------------------------------ HOTFIX: real data 260918080-VOC
REAL_PROD_1 = "Lỗi xước rear ( Áp dụng cải tiến 17/9 – Công đoạn Assy Daoltech):"
REAL_PROD_2 = "Cải tiến tại công đoạn lắp ráp lỗi lệch ant (Áp dụng cải tiến từ ngày 22.9.2026):"
REAL_PROD_3 = "Cải tiến lỗi mẻ rear (Áp dụng cải tiến 17/9 – Công đoạn Assy Daoltech):"
REAL_BODY_1 = ("- Tại công đoạn nén tape sealing\n"
               "+ Trước: Jig nén cạnh sắc, không có đệm silicon\n"
               "+ Sau: Bọc silicon toàn bộ cạnh jig, bo tròn R1.0")
REAL_BODY_2 = ("- Tại công đoạn lắp ANT\n"
               "+ Trước: Lắp ANT bằng tay, không có cữ định vị\n"
               "+ Sau: Bổ sung jig định vị ANT, kiểm tra vị trí bằng gauge")


@pytest.mark.parametrize("heading", [REAL_PROD_1, REAL_PROD_2, REAL_PROD_3])
def test_ap_dung_cai_tien_date_is_production(heading):
    assert improvement_subkind(heading) == "production"
    assert section_kind_of_heading(heading) == "improvement"


@pytest.mark.parametrize("heading,kinds", [
    ("Áp dụng cải tiến và theo dõi hiệu quả cải tiến", ("followup", "verify")),
    ("Theo dõi hiệu quả cải tiến trên data kiểm tra OQC", ("followup", "verify")),
    ("Duy trì và áp dụng cải tiến", ("followup",)),
    ("Audit kiểm tra thường xuyên", ("followup",)),
    ("Theo dõi data OQC", ("followup", "verify")),
])
def test_positive_followup_evidence_still_followup(heading, kinds):
    assert improvement_subkind(heading) == "followup"
    assert _semantic_kind(heading) in kinds                       # the kind the segmentation actually uses


def make_real_production_deck(path: Path) -> Path:
    """S2 QPN, S3 cause, S4 '3. CẢI TIẾN TRONG SẢN XUẤT' with the two REAL child headings (+ Before/After
    pictures), S5 inspection slide, S6 'Hiệu quả cải tiến' / 'Duy trì và áp dụng cải tiến'."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W = prs.slide_width
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    _tb(s, "BÁO CÁO ĐỐI SÁCH LỖI XƯỚC, LỆCH ANT REAR A185", 0.5, 1, 12, 1, 28, True)
    _tb(s, "Model: A185\nItem: Rear\nManagement No: 260918080-VOC", 0.5, 2.5, 6, 2, 16)
    _qpn_slide(prs, "picture")
    s = prs.slides.add_slide(blank)
    _tb(s, "2. NGUYÊN NHÂN", 0.5, 0.3, 8, 0.7, 24, True)
    _tb(s, "- Jig nén cạnh sắc gây xước rear", 1.7, 1.1, 11, 2.2, 13)
    _furniture(s, W)
    s = prs.slides.add_slide(blank)                                                          # 4 production
    _tb(s, "3. CẢI TIẾN TRONG SẢN XUẤT", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Cải tiến", 0.15, 3.0, 1.3, 1.3)
    _tb(s, REAL_PROD_1, 1.7, 1.0, 5.6, 0.4, 12, True)                                       # bold child heading
    _tb(s, REAL_BODY_1, 1.7, 1.4, 5.6, 1.9, 12)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Trước cải tiến", 7.5, 1.0, 1.5, 0.35)
    s.shapes.add_picture(_pic("#ffe0b2", "b1"), Inches(7.5), Inches(1.4), Inches(1.6), Inches(1.2))
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 9.4, 1.0, 1.5, 0.35)
    s.shapes.add_picture(_pic("#c8e6c9", "a1"), Inches(9.4), Inches(1.4), Inches(1.6), Inches(1.2))
    _tb(s, REAL_PROD_2, 1.7, 3.6, 5.6, 0.4, 12, True)
    _tb(s, REAL_BODY_2, 1.7, 4.0, 5.6, 1.9, 12)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Trước cải tiến", 7.5, 3.6, 1.5, 0.35)
    s.shapes.add_picture(_pic("#ffe0b2", "b2"), Inches(7.5), Inches(4.0), Inches(1.6), Inches(1.2))
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 9.4, 3.6, 1.5, 0.35)
    s.shapes.add_picture(_pic("#c8e6c9", "a2"), Inches(9.4), Inches(4.0), Inches(1.6), Inches(1.2))
    _furniture(s, W)
    s = prs.slides.add_slide(blank)                                                          # 5 inspection
    _tb(s, "4. CẢI TIẾN TRONG KIỂM TRA", 0.5, 0.3, 8, 0.7, 24, True)
    _tb(s, INSPECTION_TEXT, 1.7, 1.0, 5.6, 2.5, 12)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 1.0, 1.5, 0.35)
    s.shapes.add_picture(_pic("#b3e5fc", "insp"), Inches(7.5), Inches(1.4), Inches(4), Inches(2.5))
    _furniture(s, W)
    s = prs.slides.add_slide(blank)                                                          # 6 verify / follow-up
    _tb(s, "5. HIỆU QUẢ CẢI TIẾN", 0.5, 0.3, 8, 0.7, 24, True)
    _tb(s, "- Tỷ lệ lỗi xước giảm từ 1.2% xuống 0%", 1.7, 1.0, 11, 1.5, 13)
    _tb(s, "Duy trì và áp dụng cải tiến", 1.7, 3.0, 5.6, 0.4, 13, True)
    _tb(s, "- Theo dõi data OQC hàng tuần", 1.7, 3.4, 11, 1.2, 12)
    _furniture(s, W)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


@pytest.fixture
def real_prod_deck(tmp_path):
    p = make_real_production_deck(tmp_path / "(CTMS)_1_260918080-VOC_ Đối sách LỖI XƯỚC 24.9.2026.pptx")
    return parse_pptx(p)


def test_parent_production_context_keeps_child_headings(real_prod_deck):
    r = real_prod_deck
    secs = split_sections(r.slide(4), "improvement")
    assert [s.kind for s in secs] == ["improvement", "improvement"], [(s.kind, s.heading) for s in secs]
    cls = Classification(management_number="260918080-VOC", qpn_slide=2, cause_slides=[3],
                         improvement_slides=[4, 5, 6], improvement_image_slides=[4, 5], verify_slides=[6],
                         source="qwen", confidence=0.9)
    rec = extract_record(r, cls, {"rear": "Rear"}, ["A185"])
    imp = rec.improvement
    for part in (REAL_PROD_1, REAL_BODY_1, REAL_PROD_2, REAL_BODY_2):
        for line in part.split("\n"):
            assert line.strip() in imp, line
    assert not any(e.startswith(("S4 followup", "S4 verify", "S4 inspection")) for e in rec.excluded_sections)
    assert any(e.startswith("S5 inspection") for e in rec.excluded_sections)
    assert any(e.startswith("S6 verify") for e in rec.excluded_sections)
    assert any(e.startswith("S6 followup") for e in rec.excluded_sections)
    assert "Duy trì" not in imp and "Tỷ lệ lỗi" not in imp and "Đào tạo" not in imp
    assert not any("Không tìm thấy Nội dung đối sách" in m for m in rec.review_reasons)
    assert not any("Không tìm thấy ảnh Sau cải tiến" in m for m in rec.review_reasons)


def test_production_pictures_follow_the_same_segmentation(real_prod_deck):
    sel = select_after_pictures(real_prod_deck, [4, 5, 6])
    after = sorted((p.slide, p.block.left) for p in sel.after)
    assert after == [(4, Inches(9.4)), (4, Inches(9.4))]                      # both After pictures of S4
    before = [p for p in sel.rejected if p.slide == 4 and p.kind == "before"]
    assert len(before) == 2 and all(p.block.left == Inches(7.5) for p in before)
    assert not any("followup section block" in (p.reason or "") for p in sel.rejected if p.slide == 4)
    insp = [p for p in sel.rejected if p.slide == 5]
    assert insp and all(p.reason == "inspection/control improvement slide" for p in insp)
    assert not sel.reasons
