"""PROMPT-024R: safe cancel-all ("Dừng tất cả") alongside the graceful file-atomic stop.

Covers §43–§50 at the batch level (real BatchProcessor + real workbook), §49 duplicate/idle semantics
at the service level, and the §59 narrowly scoped stale-pytest config migration. These are headless
tests: Windows UI behavior still requires the manual acceptance pass.
"""
from __future__ import annotations

import json
import threading
from pathlib import Path

from openpyxl import load_workbook

from app.batch_processor import BatchOptions, BatchProcessor, BatchSummary
from app.config import AppConfig
from app.excel_writer import ExcelWriter
from app.gui_controller import UiEvent
from tests.test_application_service import BridgeService, _make_service, _scan
from tests.test_prompt025_multi_item import NAME, _deck, _two_item_slide


def _opts(files, template, out, **kw):
    return BatchOptions(files=[Path(f) for f in files], template=template, output_file=out,
                        use_ollama=False, row_mode="append", **kw)


def _two_item_deck(tmp_path):
    return _deck(tmp_path / "in" / NAME, [_two_item_slide()])


# ================================================================== §43 stop-after-current stays file-atomic
def test_stop_after_current_completes_multi_item_report_and_queue_never_starts(sample_tree, tmp_path):
    deck = _two_item_deck(tmp_path)
    files = [deck, sample_tree["files"][1], sample_tree["files"][2]]
    out = tmp_path / "o" / "r.xlsx"
    events = []
    proc = BatchProcessor(_opts(files, sample_tree["template"], out),
                          on_file=lambda i, s, d: events.append((i, s, d)))

    def _stop_once(i, s, d):
        if i == 0 and s == "reading":
            proc.request_stop()

    proc.on_file = lambda i, s, d: (_stop_once(i, s, d), events.append((i, s, d)))
    summary = proc.run()

    assert summary.stopped is True and summary.cancel_requested is False
    assert summary.cancelled == 0 and summary.failed == 0
    # A finished COMPLETELY: both PROMPT-025 improvement items were processed before stopping
    assert summary.completed + summary.needs_review == 1
    fr_a = proc.results[0]
    assert len(fr_a.improvement_items) == 2                       # all items of the current file finished
    assert len(list((out.parent / "assets" / "0001_IMPROVEMENT_regions").glob("region_*.png"))) == 2
    # B and C never started
    assert len(proc.results) == 1
    assert not any(i in (1, 2) for i, s, d in events)
    assert out.exists() and load_workbook(out)["Kiểm chứng"].max_row >= 4


# ================================================================== §44 cancel before report commit
def test_cancel_all_before_commit_cancels_current_report_and_queue(sample_tree, tmp_path):
    deck = _two_item_deck(tmp_path)
    files = [deck, sample_tree["files"][1], sample_tree["files"][2]]
    out = tmp_path / "o" / "r.xlsx"
    events = []
    proc = BatchProcessor(_opts(files, sample_tree["template"], out),
                          on_file=lambda i, s, d: events.append((i, s, d)))

    def _cancel_once(i, s, d):
        if i == 0 and s == "reading":
            proc.request_cancel()

    proc.on_file = lambda i, s, d: (_cancel_once(i, s, d), events.append((i, s, d)))
    summary = proc.run()

    # §5/§50: CANCELLED != FAILED – a user choice is not an application failure
    assert summary.cancel_requested is True and summary.stopped is True
    assert summary.cancelled == 1 and summary.failed == 0 and summary.completed == 0
    fr_a = proc.results[0]
    assert fr_a.status == "cancelled" and fr_a.error == "" and fr_a.excel_row is None
    # B and C never started; only the cancelled report has a result entry
    assert len(proc.results) == 1 and not any(i in (1, 2) for i, s, d in events)
    # §10: the workbook is either untouched or fully valid; no partial row for the cancelled report
    if out.exists():
        ws = load_workbook(out)["Kiểm chứng"]
        assert all((ws.cell(row=r, column=5).value or "") != "SM-A185"
                   for r in range(4, ws.max_row + 1))
    res = json.loads((out.parent / "logs" / "batch_result.json").read_text(encoding="utf-8"))
    assert res["results"][0]["status"] == "cancelled"
    assert res["meta"]["summary"]["cancelled"] == 1 and res["meta"]["summary"]["failed"] == 0


# ================================================================== §45 cancel between improvement items
def test_cancel_between_improvement_items_never_commits_partial_report(sample_tree, tmp_path):
    deck = _two_item_deck(tmp_path)
    files = [deck, sample_tree["files"][1]]
    out = tmp_path / "o" / "r.xlsx"
    proc = BatchProcessor(_opts(files, sample_tree["template"], out),
                          on_file=lambda i, s, d: proc.request_cancel()
                          if (i == 0 and s == "extracting_images") else None)
    summary = proc.run()

    # cancellation acknowledged inside the item-level evidence stage: the CURRENT report is cancelled,
    # never committed as a partial success, and the next report never starts
    assert summary.cancelled == 1 and summary.failed == 0 and summary.completed == 0
    assert summary.cancel_requested is True and summary.stopped is True
    assert len(proc.results) == 1 and proc.results[0].status == "cancelled"
    assert proc.results[0].excel_row is None
    if out.exists():
        ws = load_workbook(out)["Kiểm chứng"]                    # workbook stays valid / empty of results
        assert ws.cell(row=4, column=5).value in (None, "")


# ================================================================== §46 cancel during the atomic commit
def test_cancel_during_atomic_commit_finishes_commit_then_stops(sample_tree, tmp_path):
    files = [sample_tree["files"][0], sample_tree["files"][1], sample_tree["files"][2]]
    out = tmp_path / "o" / "r.xlsx"
    proc = BatchProcessor(_opts(files, sample_tree["template"], out))
    real_save = ExcelWriter.save
    state = {"fired": False}

    def _save_then_cancel(self):
        if not state["fired"]:
            state["fired"] = True
            proc.request_cancel()          # cancel acknowledged while INSIDE the critical commit section
        return real_save(self)

    try:
        ExcelWriter.save = _save_then_cancel
        summary = proc.run()
    finally:
        ExcelWriter.save = real_save

    # §9/§11-B: the critical section completes; the committed report stays completed, queue stops
    assert state["fired"] is True
    assert summary.completed + summary.needs_review == 1
    assert summary.cancel_requested is True and summary.stopped is True
    assert len(proc.results) == 1
    assert proc.results[0].status in ("completed", "needs_review")
    assert proc.results[0].excel_row is not None
    ws = load_workbook(out)["Kiểm chứng"]                        # workbook valid + row really committed
    assert ws.cell(row=proc.results[0].excel_row, column=5).value == "A185"


# ================================================================== §47 cancel after a successful commit
def test_cancel_after_successful_commit_keeps_report_completed(sample_tree, tmp_path):
    files = [sample_tree["files"][0], sample_tree["files"][1]]
    out = tmp_path / "o" / "r.xlsx"
    proc = BatchProcessor(_opts(files, sample_tree["template"], out))
    real_save = ExcelWriter.save
    state = {"fired": False}

    def _save_completed_then_cancel(self):
        result = real_save(self)
        if not state["fired"]:
            state["fired"] = True
            proc.request_cancel()          # cancel acknowledged only AFTER the commit succeeded
        return result

    try:
        ExcelWriter.save = _save_completed_then_cancel
        summary = proc.run()
    finally:
        ExcelWriter.save = real_save

    assert summary.completed + summary.needs_review == 1         # committed result preserved (§11-B)
    assert summary.cancel_requested is True and summary.stopped is True
    assert summary.failed == 0 and summary.cancelled == 0        # nothing was cancelled: commit had finished
    assert len(proc.results) == 1 and proc.results[0].excel_row is not None


# ================================================================== §48 duplicate cancel requests
def test_duplicate_cancel_requests_are_idempotent(sample_tree, tmp_path):
    out = tmp_path / "o" / "r.xlsx"
    proc = BatchProcessor(_opts([sample_tree["files"][0], sample_tree["files"][1]],
                                sample_tree["template"], out))
    proc.request_cancel()
    proc.request_cancel()
    proc.request_cancel()                                        # many clicks, one deterministic effect
    summary = proc.run()
    assert summary.cancel_requested is True and summary.stopped is True
    assert summary.completed == 0 and summary.failed == 0 and summary.cancelled == 0  # nothing started
    assert proc.results == []
    assert summary.errors == []


# ================================================================== §49 cancel while idle
def test_cancel_before_batch_start_is_harmless_and_deterministic(sample_tree, tmp_path):
    out = tmp_path / "o" / "r.xlsx"
    proc = BatchProcessor(_opts([sample_tree["files"][0]], sample_tree["template"], out))
    proc.request_cancel()                                        # requested before any worker exists
    summary = proc.run()
    assert summary.cancel_requested is True and summary.stopped is True
    assert summary.total == 1 and summary.completed == 0 and summary.failed == 0
    assert proc.results == []


# ================================================================== GuiController job-state semantics
def test_controller_cancel_state_and_done_event_mark_cancelled_not_error(sample_tree, tmp_path):
    service, config = _make_service(sample_tree, tmp_path)
    c = service.controller
    assert c.request_cancel() is False                           # idle: harmless (§49)
    c.files = [Path(sample_tree["files"][0]), Path(sample_tree["files"][1])]
    from app.gui_controller import RowState
    c.rows = [RowState(index=0, path=c.files[0]), RowState(index=1, path=c.files[1])]
    c.state = "cancelling"
    # the worker was mid-pipeline with report A when cancel-all was acknowledged
    c.apply_event(UiEvent("row", (0, "extracting_images", "")))
    c.apply_event(UiEvent("done", BatchSummary(total=2, cancelled=1, stopped=True, cancel_requested=True)))
    assert c.state == "idle"
    assert c.rows[0].stage == "cancelled"                        # NOT "error" (§5/§50)
    assert c.rows[0].status_vi == "Đã hủy"
    assert c.rows[1].stage == "waiting"                          # untouched queue = Chưa xử lý (§12)
    lines = c.summary_lines()
    assert lines[0].startswith("Đã dừng toàn bộ xử lý")
    assert any("Đã hủy: 1" in line for line in lines)


# ================================================================== service/bridge cancel-all contract
class _CancellableGatedProcessor:
    """Async worker fixture: cancel is observed at one explicit SAFE boundary mid first report."""

    instances = []

    def __init__(self, options, on_prescan=None, on_file=None, on_batch=None, on_done=None, on_log=None):
        self.options = options
        self.on_file = on_file
        self.on_batch = on_batch
        self.on_done = on_done
        self.on_log = on_log
        self.results = []
        self.summary = None
        self._thread = None
        self.cancel_event = threading.Event()
        self.stop_event = threading.Event()
        self.first_started = threading.Event()
        self.safe_boundary = threading.Event()        # test opens it once the cancel decision should apply
        self.processed = []
        self.__class__.instances.append(self)

    def start(self):
        self._thread = threading.Thread(target=self.run, name="fake-cancel-worker", daemon=True)
        self._thread.start()

    def run(self):
        total = len(self.options.files)
        self.on_file(0, "reading", "")
        self.first_started.set()
        assert self.safe_boundary.wait(timeout=10)
        if self.cancel_event.is_set():                          # safe boundary -> cancel current report
            self.on_file(0, "cancelled", "Đã hủy theo yêu cầu")
            self.summary = BatchSummary(total=total, cancelled=1, stopped=True, cancel_requested=True)
            self.on_done(self.summary)
            return
        self.processed.append(self.options.files[0])            # file-atomic completion
        self.on_file(0, "completed", "")
        self.on_batch(1, total)
        for index in range(1, total):
            if self.cancel_event.is_set() or self.stop_event.is_set():
                break
            self.processed.append(self.options.files[index])
            self.on_file(index, "completed", "")
            self.on_batch(index + 1, total)
        self.summary = BatchSummary(total=total, completed=len(self.processed),
                                    stopped=self.stop_event.is_set(),
                                    cancel_requested=self.cancel_event.is_set())
        self.on_done(self.summary)

    def request_stop(self):
        self.stop_event.set()

    def request_cancel(self):
        self.stop_event.set()
        self.cancel_event.set()

    def is_running(self):
        return bool(self._thread and self._thread.is_alive())


def test_service_cancel_all_job_state_duplicate_and_idle(sample_tree, tmp_path):
    _CancellableGatedProcessor.instances.clear()
    service, config = _make_service(sample_tree, tmp_path, processor_factory=_CancellableGatedProcessor)
    _scan(service, config)
    bridge = BridgeService(service)

    assert bridge.cancel_all()["data"]["requested"] is False     # §49 idle: harmless, deterministic, ok

    started = service.start_processing(config, use_ollama=False)
    worker = _CancellableGatedProcessor.instances[-1]
    assert started["job"]["status"] in ("processing", "stopping", "cancelling")
    assert worker.first_started.wait(timeout=5)

    cancelled = bridge.cancel_all()                              # §6 bridge contract
    assert cancelled["ok"] is True
    payload = cancelled["data"]
    assert payload["requested"] is True and payload["mode"] == "cancel_all"
    assert payload["job"]["status"] == "cancelling"              # worker has not acknowledged yet
    assert payload["job"]["cancelRequested"] is True

    duplicate = bridge.cancel_all()                              # §48 duplicate clicks: idempotent
    assert duplicate["ok"] is True and duplicate["data"]["requested"] is False
    assert len(_CancellableGatedProcessor.instances) == 1        # no second worker

    worker.safe_boundary.set()
    worker._thread.join(timeout=10)
    assert not worker._thread.is_alive()

    final = service.dashboard_state()
    assert final["job"]["status"] == "done"                      # "Đã dừng" only after acknowledgement (§16)
    assert final["job"]["cancelRequested"] is True
    assert final["job"]["cancelledCount"] == 1
    assert final["job"]["error"] == ""                           # §18: not surfaced as a failure
    statuses = {row["fileName"]: row["status"] for row in final["reports"]}
    assert statuses[worker.options.files[0].name] == "cancelled"
    # §12: untouched queued reports keep their PRE-PROCESSING lifecycle state (Chưa xử lý / new row),
    # never completed/failed/cancelled.
    assert statuses[worker.options.files[1].name] in ("waiting", "new_row")
    assert service.controller.request_cancel() is False          # idle again


def test_service_stop_after_current_still_file_atomic(sample_tree, tmp_path):
    _CancellableGatedProcessor.instances.clear()
    service, config = _make_service(sample_tree, tmp_path, processor_factory=_CancellableGatedProcessor)
    _scan(service, config)

    service.start_processing(config, use_ollama=False)
    worker = _CancellableGatedProcessor.instances[-1]
    assert worker.first_started.wait(timeout=5)

    stopping = service.stop_after_current()
    assert stopping["job"]["status"] == "stopping"
    assert stopping["job"]["cancelRequested"] is False           # distinct from cancel-all (§16)
    worker.safe_boundary.set()
    worker._thread.join(timeout=10)

    final = service.dashboard_state()
    assert len(worker.processed) == 1                            # current file finished fully
    assert final["job"]["status"] == "done"
    assert final["job"]["stopped"] is True and final["job"]["cancelRequested"] is False
    assert final["job"]["cancelledCount"] == 0


# ================================================================== §59 stale pytest path migration
def test_stale_pytest_paths_are_reset_but_user_paths_are_kept(tmp_path, monkeypatch):
    # the migration is production-only; opt out of the disposable-test-runtime guard for this check
    import app.runtime_paths as rp
    monkeypatch.delenv(rp.TEST_RUNTIME_ROOT_ENV, raising=False)
    config_path = tmp_path / "config.json"
    stale_folder = r"C:\Users\ADMINPC\AppData\Local\Temp\pytest-of-ADMINPC\pytest-12\reports"
    stale_output = r"C:\Users\ADMINPC\AppData\Local\Temp\pytest-of-ADMINPC\pytest-12\out\K.xlsx"
    legit_temp = r"C:\Users\ADMINPC\AppData\Local\Temp\Bao cao thang 9"   # Temp, but NOT a pytest pattern
    config_path.write_text(json.dumps({
        "last_report_folder": stale_folder,
        "last_template": r"D:\Templates\Verification.xlsx",
        "last_output_folder": legit_temp,
        "last_output_file": stale_output,
        "update_path": r"C:\Temp\pytest-3\update",
    }, ensure_ascii=False), encoding="utf-8")

    cfg = AppConfig.load(config_path)
    assert cfg.last_report_folder == ""                          # obviously stale -> reset
    assert cfg.last_output_file == ""
    assert cfg.update_path == ""
    assert cfg.last_template == r"D:\Templates\Verification.xlsx"   # legitimate value untouched
    assert cfg.last_output_folder == legit_temp                  # legitimate Temp path kept (§59)

    saved = cfg.save(tmp_path / "config2.json")
    again = AppConfig.load(saved)
    assert again.last_report_folder == "" and again.last_output_file == ""
