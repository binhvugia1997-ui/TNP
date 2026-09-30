"""FINAL content-region rules, reproduced structurally on the real production layout:

* text fields receive ONLY the main business-content region (no slide title, sidebar circle, caption
  buttons, logos, footer);
* "Hình ảnh cải tiến" receives ONLY the "Sau cải tiến" pictures of production improvements;
* incremental master update keeps working with the new rules.
"""
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.batch_processor import BatchOptions, BatchProcessor
from app.classifier import heuristic_classify
from app.content_region import (ROLE_CAPTION, ROLE_CONTENT, ROLE_FURNITURE, ROLE_SIDEBAR, ROLE_TITLE, caption_kind,
                                classify_blocks, inline_anchor_kind, is_slide_level_heading)
from app.extractor import extract_record
from app.image_extractor import build_after_pictures_image
from app.improvement_pictures import AMBIGUOUS_REASON, select_after_pictures
from app.pptx_parser import parse_pptx

CAUSE_BLOCK_1 = ("Nguyên nhân trong kiểm tra:\n"
                 "- Mẻ xước: Sản phẩm bị va chạm với thành khay trong quá trình vận chuyển giữa các công đoạn\n"
                 "- Bong sơn: Lực miết urethane không đều, độ bám dính lớp sơn tại góc R không đạt")
CAUSE_BLOCK_2 = ("Nguyên nhân trong kiểm tra:\n"
                 "- Do nhân lực kiểm tra cuối chưa cố định, nhân viên mới chưa nắm rõ tiêu chuẩn phân biệt lỗi mẻ xước nhẹ")
IMP_ITEM_1 = ("Cải tiến lỗi Mẻ xước (Áp dụng cải tiến từ ngày 20/09/2026 - Công đoạn Assy Daoltech):\n"
              "- Cải tiến jig nén tape sealing\n"
              "+ Trước: Jig nén bằng nhôm, cạnh sắc, không có lớp đệm nên gây mẻ xước khi nén\n"
              "+ Sau: Bọc silicon 2mm toàn bộ mặt tiếp xúc của jig, bo tròn cạnh R1.0")
IMP_ITEM_2 = ("Cải tiến lỗi bong sơn (Áp dụng từ ngày 22/09/2026 - Công đoạn Sơn IT):\n"
              "- Thay đổi tool miết urethane\n"
              "+ Trước: Tool miết đầu nhọn, thao tác miết 1 lần, lực không đều\n"
              "+ Sau: Tool miết đầu tròn R3, thao tác miết 2 lần theo chiều cố định")
INSPECTION_TEXT = ("Cải tiến trong kiểm tra:\n"
                   "- Bổ sung tiêu chuẩn giới hạn lỗi mẻ xước kèm ảnh mẫu tại vị trí kiểm tra cuối\n"
                   "- Cố định 2 nhân viên kiểm tra cuối đã được đào tạo (ngày 21/09/2026)")
LONG_TERM_TEXT = ("- Đưa hạng mục kiểm tra jig nén và tool miết vào checklist đầu ca\n"
                  "- Cập nhật SOP công đoạn Assy Rev.04, triển khai ngang cho các model cùng cấu trúc")
TEMP_TEXT = "- Sorting 100% lô hàng tồn tại kho: 1.250 pcs\n- Tăng cường kiểm tra ngoại quan tại OQC trong 2 tuần"


def _pic(color, label, size=(480, 360)):
    im = Image.new("RGB", size, color)
    bio = BytesIO()
    im.save(bio, "PNG")
    bio.seek(0)
    return bio


def _tb(slide, text, left, top, width, height, size=12, bold=False):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
    return tb


def _shape(slide, kind, text, left, top, width, height, size=11):
    shp = slide.shapes.add_shape(kind, Inches(left), Inches(top), Inches(width), Inches(height))
    shp.text_frame.text = text
    shp.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    return shp


def _furniture(slide, W):
    slide.shapes.add_picture(_pic("#eeeeee", "logo", (200, 80)), W - Inches(1.4), Inches(0.1), Inches(1.2), Inches(0.45))
    _tb(slide, "CTMS – Confidential", 0.3, 7.05, 4, 0.35, 8)
    _tb(slide, "▶", 0.2, 0.35, 0.4, 0.4, 14)


def make_real_layout_report(path: Path, mgmt="260920045-VOC") -> Path:
    """Cover, 1. NGUYÊN NHÂN, 2. XỬ LÝ TẠM THỜI, 3. CẢI TIẾN TRONG SẢN XUẤT (2 items, before/after pictures),
    4. CẢI TIẾN TRONG KIỂM TRA (pictures), 5. ĐỐI SÁCH LÂU DÀI, 6. continuation slide with an unlabelled picture."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W = prs.slide_width
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)                                   # 1 cover
    _tb(s, "BÁO CÁO ĐỐI SÁCH LỖI MẺ XƯỚC, BONG SƠN REAR A185", 0.5, 1, 12, 1, 28, True)
    _tb(s, f"Model: A185\nItem: Rear\nManagement No: {mgmt}", 0.5, 2.5, 6, 2, 16)

    s = prs.slides.add_slide(blank)                                   # 2 root cause
    _tb(s, "1. NGUYÊN NHÂN", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Nguyên nhân", 0.15, 2.5, 1.3, 1.3)
    _tb(s, CAUSE_BLOCK_1, 1.7, 1.1, 11, 2.2, 13)
    _tb(s, CAUSE_BLOCK_2, 1.7, 3.6, 11, 1.5, 13)
    _furniture(s, W)

    s = prs.slides.add_slide(blank)                                   # 3 temporary
    _tb(s, "2. XỬ LÝ TẠM THỜI", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Xử lý tạm thời", 0.15, 2.5, 1.3, 1.3)
    _tb(s, TEMP_TEXT, 1.7, 1.1, 11, 2, 13)
    s.shapes.add_picture(_pic("#fff2cc", "sorting"), Inches(1.7), Inches(3.5), Inches(4), Inches(2.5))
    _furniture(s, W)

    s = prs.slides.add_slide(blank)                                   # 4 production improvement (2 items)
    _tb(s, "3. CẢI TIẾN TRONG SẢN XUẤT", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Cải tiến trong kiểm tra", 0.15, 3.0, 1.3, 1.3)      # sidebar mentions inspection!
    _tb(s, IMP_ITEM_1, 1.7, 1.0, 5.6, 2.8, 12)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Trước cải tiến", 7.5, 1.0, 1.5, 0.35)
    s.shapes.add_picture(_pic("#ffe0b2", "before1a"), Inches(7.5), Inches(1.4), Inches(1.3), Inches(1.2))
    s.shapes.add_picture(_pic("#ffe0b2", "before1b"), Inches(8.9), Inches(1.4), Inches(1.3), Inches(1.2))
    s.shapes.add_picture(_pic("#999999", "arrow", (400, 60)), Inches(7.5), Inches(2.62), Inches(2.0), Inches(0.3))
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 2.95, 1.5, 0.35)
    s.shapes.add_picture(_pic("#c8e6c9", "after1a"), Inches(7.5), Inches(3.35), Inches(1.3), Inches(1.2))
    s.shapes.add_picture(_pic("#c8e6c9", "after1b"), Inches(8.9), Inches(3.35), Inches(1.3), Inches(1.2))
    s.shapes.add_picture(_pic("#c8e6c9", "after1c"), Inches(10.3), Inches(3.35), Inches(1.3), Inches(1.2))
    _tb(s, IMP_ITEM_2, 1.7, 4.6, 5.6, 2.4, 12)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Trước cải tiến", 7.5, 4.6, 1.5, 0.35)
    s.shapes.add_picture(_pic("#ffe0b2", "before2a"), Inches(7.5), Inches(5.0), Inches(1.3), Inches(1.2))
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 9.2, 4.6, 2.7, 0.35)
    s.shapes.add_picture(_pic("#c8e6c9", "after2a"), Inches(9.2), Inches(5.0), Inches(1.3), Inches(1.2))
    s.shapes.add_picture(_pic("#c8e6c9", "after2b"), Inches(10.6), Inches(5.0), Inches(1.3), Inches(1.2))
    _furniture(s, W)

    s = prs.slides.add_slide(blank)                                   # 5 inspection improvement (pictures excluded)
    _tb(s, "4. CẢI TIẾN TRONG KIỂM TRA", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Cải tiến trong kiểm tra", 0.15, 3.0, 1.3, 1.3)
    _tb(s, INSPECTION_TEXT, 1.7, 1.0, 5.6, 2.5, 12)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 7.5, 1.0, 1.5, 0.35)
    s.shapes.add_picture(_pic("#b3e5fc", "standard"), Inches(7.5), Inches(1.4), Inches(4), Inches(2.5))
    _furniture(s, W)

    s = prs.slides.add_slide(blank)                                   # 6 long-term
    _tb(s, "5. ĐỐI SÁCH LÂU DÀI", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Đối sách lâu dài", 0.15, 2.5, 1.3, 1.3)
    _tb(s, LONG_TERM_TEXT, 1.7, 1.1, 11, 2, 13)
    _furniture(s, W)

    s = prs.slides.add_slide(blank)                                   # 7 continuation: picture without any anchor
    _tb(s, "Hình ảnh cải tiến công đoạn Assy (tiếp theo)", 1.7, 1.0, 8, 0.6, 12)
    s.shapes.add_picture(_pic("#d1c4e9", "unknown"), Inches(1.7), Inches(1.8), Inches(5), Inches(3.5))
    _furniture(s, W)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


@pytest.fixture(scope="module")
def real_deck(tmp_path_factory):
    p = tmp_path_factory.mktemp("real") / "(CTMS)_20601_260920045-VOC_ Đối sách LỖI MẺ XƯỚC, BONG SƠN 24.9.2026.pptx"
    make_real_layout_report(p)
    r = parse_pptx(p)
    return r, heuristic_classify(r), extract_record(r, heuristic_classify(r), {"rear": "Rear"}, ["A185"])


# ------------------------------------------------------------------ roles
def test_block_roles_on_real_layout(real_deck):
    r, _, _ = real_deck
    roles = {ro.text.split("\n")[0]: ro.role for ro in classify_blocks(r.slide(4))}
    assert roles["3. CẢI TIẾN TRONG SẢN XUẤT"] == ROLE_TITLE
    assert roles["Cải tiến trong kiểm tra"] == ROLE_SIDEBAR
    assert roles["Trước cải tiến"] == ROLE_CAPTION and roles["Sau cải tiến"] == ROLE_CAPTION
    assert roles["CTMS – Confidential"] == ROLE_FURNITURE and roles["▶"] == ROLE_FURNITURE
    assert roles[IMP_ITEM_1.split("\n")[0]] == ROLE_CONTENT and roles[IMP_ITEM_2.split("\n")[0]] == ROLE_CONTENT
    assert caption_kind("Trước cải tiến") == "before" and caption_kind("Sau cải tiến") == "after"
    assert caption_kind("After") == "after" and caption_kind("Hình ảnh trước") == "before"
    assert caption_kind("Sau khi lắp jig mới, tỷ lệ lỗi giảm") is None
    assert inline_anchor_kind("+ Trước: Jig nén bằng nhôm") == "before" and inline_anchor_kind("+ Sau: Bọc silicon") == "after"
    assert inline_anchor_kind("Sau đó kiểm tra lại") is None
    assert is_slide_level_heading("1. NGUYÊN NHÂN") and is_slide_level_heading("ĐỐI SÁCH LÂU DÀI")
    assert not is_slide_level_heading("Nguyên nhân trong kiểm tra:")


# ------------------------------------------------------------------ root cause text
def test_root_cause_is_main_content_only(real_deck):
    _, cls, rec = real_deck
    assert cls.cause_slides == [2]
    assert rec.root_cause == CAUSE_BLOCK_1 + "\n\n" + CAUSE_BLOCK_2          # verbatim, both blocks, in order
    assert "1. NGUYÊN NHÂN" not in rec.root_cause
    assert rec.root_cause.count("Nguyên nhân trong kiểm tra:") == 2          # repeated sub-heading kept
    assert not rec.root_cause.startswith("Nguyên nhân\n")                     # sidebar circle excluded
    for junk in ("CTMS", "Confidential", "▶", "logo"):
        assert junk not in rec.root_cause


# ------------------------------------------------------------------ improvement text
def test_improvement_text_is_main_content_only(real_deck):
    _, cls, rec = real_deck
    imp = rec.improvement
    for line in (IMP_ITEM_1 + "\n" + IMP_ITEM_2 + "\n" + INSPECTION_TEXT + "\n" + LONG_TERM_TEXT).split("\n"):
        assert line in imp, line
    assert imp.index("Cải tiến lỗi Mẻ xước") < imp.index("Cải tiến lỗi bong sơn") < imp.index("Cải tiến trong kiểm tra:") \
        < imp.index("checklist đầu ca")
    assert "+ Trước: Jig nén bằng nhôm" in imp and "+ Sau: Bọc silicon 2mm" in imp     # textual Trước/Sau retained
    assert "Trước cải tiến" not in imp and "Sau cải tiến" not in imp                     # caption buttons excluded
    for junk in ("3. CẢI TIẾN TRONG SẢN XUẤT", "4. CẢI TIẾN TRONG KIỂM TRA", "5. ĐỐI SÁCH LÂU DÀI", "CTMS", "▶"):
        assert junk not in imp
    assert imp.count("Cải tiến trong kiểm tra") == 1                     # content heading once, sidebar circles never
    assert "Sorting 100%" not in imp and "XỬ LÝ TẠM THỜI" not in imp     # temporary handling still excluded
    assert "Sorting 100%" in rec.temporary_excluded
    assert "BÁO CÁO ĐỐI SÁCH" not in imp


# ------------------------------------------------------------------ improvement pictures
def test_after_only_pictures(real_deck):
    r, cls, rec = real_deck
    assert 4 in cls.improvement_image_slides and 5 in cls.improvement_image_slides
    sel = select_after_pictures(r, cls.improvement_image_slides)
    names = [p.block.shape_name for p in sel.after]
    pics4 = {b.shape_id: b for b in r.slide(4).pictures}
    # exactly the 5 After pictures of slide 4 in visual order (row 1 left->right, then row 2)
    after_lefts = [(p.slide, p.block.top, p.block.left) for p in sel.after]
    assert [p.slide for p in sel.after] == [4, 4, 4, 4, 4]
    assert after_lefts == sorted(after_lefts)
    tops = sorted({p.block.top for p in sel.after})
    assert len(tops) == 2 and sum(1 for p in sel.after if p.block.top == tops[0]) == 3
    befores = [p.block.shape_id for p in sel.rejected if p.slide == 4 and p.kind == "before"]
    assert len(befores) == 3 and all(sid in pics4 for sid in befores)
    assert not any(p.slide == 4 and p.kind == "ambiguous" for p in sel.rejected)
    excluded = [p for p in sel.rejected if p.kind == "excluded"]
    assert any("arrow" in p.reason or "thin" in p.reason for p in excluded if p.slide == 4)
    assert any("logo" in p.reason for p in excluded if p.slide == 4)
    assert all(p.kind == "excluded" for p in sel.rejected if p.slide == 5)      # inspection pictures
    assert 5 not in sel.slides_with_after and sel.slides_with_after == [4]
    assert names and len(names) == 5
    assert rec.after_picture_slides == [4] and len(rec.after_pictures) == 5


def test_ambiguous_picture_is_omitted_and_flagged(real_deck):
    r, cls, rec = real_deck
    sel = select_after_pictures(r, [7])
    assert sel.after == [] and sel.reasons == [AMBIGUOUS_REASON.format(n=7)]
    assert any(p.kind == "ambiguous" for p in sel.rejected)


def test_after_pictures_image_is_vertical_stack(real_deck, tmp_path):
    r, _, rec = real_deck
    out, problems = build_after_pictures_image(r, rec.after_pictures, tmp_path / "imp.jpg", pictures_dir=tmp_path / "pics")
    assert out and out.exists() and problems == []
    with Image.open(out) as im:
        assert im.height > im.width * 3                                # 5 pictures stacked vertically
    assert len(list((tmp_path / "pics").glob("after*_slide04.png"))) == 5


def test_inline_anchor_layout_without_caption_buttons(tmp_path):
    """No caption buttons: '+ Trước:' / '+ Sau:' lines own the vertical range next to them."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _tb(s, "3. CẢI TIẾN TRONG SẢN XUẤT", 0.5, 0.3, 8, 0.7, 24, True)
    _tb(s, "Cải tiến khay chứa:\n- Thay khay\n+ Trước: khay cứng\n\n\n+ Sau: khay lót EVA\n\n", 1.7, 1.0, 5.6, 5.6, 12)
    s.shapes.add_picture(_pic("#ffe0b2", "b"), Inches(7.5), Inches(2.5), Inches(3), Inches(1.6))   # beside "+ Trước"
    s.shapes.add_picture(_pic("#c8e6c9", "a"), Inches(7.5), Inches(4.6), Inches(3), Inches(1.6))   # beside "+ Sau"
    p = tmp_path / "260920046-VOC_x.pptx"
    prs.save(str(p))
    r = parse_pptx(p)
    sel = select_after_pictures(r, [1])
    assert len(sel.after) == 1 and sel.after[0].block.top == Inches(4.6)
    assert [x.kind for x in sel.rejected] == ["before"] and sel.reasons == []


# ------------------------------------------------------------------ incremental update with the new rules
SHEET = "Kiểm chứng"
COL = {"stt": 1, "mgmt": 2, "vendor": 3, "date": 4, "model": 5, "item": 6, "defect": 7, "qpn": 8, "cause": 9,
       "improvement": 10, "image": 11}


def _prefill(template, rows):
    wb = load_workbook(template)
    ws = wb[SHEET]
    for i, values in enumerate(rows):
        r = 4 + i
        ws.cell(row=r, column=1, value=i + 1)
        for key, val in values.items():
            ws.cell(row=r, column=COL[key], value=val)
        for k in range(12, 20):
            ws.cell(row=r, column=k, value="OK")
    wb.save(template)


def _run(files, template, out, **kw):
    events = []
    opts = BatchOptions(files=[Path(f) for f in files], template=template, output_file=out, use_ollama=False, **kw)
    proc = BatchProcessor(opts, on_file=lambda i, s, d: events.append((i, s, d)))
    return proc.run(), events, proc


def _images_at(ws, row, col):
    return [im for im in ws._images if im.anchor._from.row + 1 == row and im.anchor._from.col + 1 == col]


def test_incremental_fill_uses_new_rules(template, tmp_path):
    deck = make_real_layout_report(tmp_path / "(CTMS)_1_260920045-VOC_ Đối sách LỖI MẺ XƯỚC 24.9.2026.pptx")
    _prefill(template, [{"mgmt": "260920045-VOC", "vendor": "Doaltech", "model": "A185", "item": "Rear",
                         "defect": "MẺ XƯỚC (nhập tay)", "qpn": "QPN nhập tay"}])
    out = tmp_path / "out" / "k.xlsx"
    summary, _, proc = _run([deck], template, out)
    fr = proc.results[0]
    assert "root_cause" in fr.filled_fields and "improvement" in fr.filled_fields and "improvement_image" in fr.filled_fields
    assert "defect_content" not in fr.filled_fields and "vendor" not in fr.filled_fields
    ws = load_workbook(out)[SHEET]
    assert ws.cell(row=4, column=COL["cause"]).value == CAUSE_BLOCK_1 + "\n\n" + CAUSE_BLOCK_2     # filtered content
    imp = ws.cell(row=4, column=COL["improvement"]).value
    assert "+ Sau: Tool miết đầu tròn R3" in imp and "3. CẢI TIẾN" not in imp and "Sau cải tiến" not in imp
    assert ws.cell(row=4, column=COL["defect"]).value == "MẺ XƯỚC (nhập tay)"                        # preserved
    assert ws.cell(row=4, column=COL["vendor"]).value == "Doaltech" and ws.cell(row=4, column=COL["qpn"]).value == "QPN nhập tay"
    assert len(_images_at(ws, 4, COL["image"])) == 1                                                  # After-only sheet
    assert len(fr.after_pictures) == 5 and fr.after_picture_slides == [4]
    assert all(ws.cell(row=4, column=k).value == "OK" for k in range(12, 20))
    # complete row now (qpn cell has text but no picture -> qpn image still missing -> processed again, nothing rewritten)
    before = [ws.cell(row=4, column=c).value for c in range(1, 20)]
    s2, _, proc2 = _run([deck], template, out)
    ws2 = load_workbook(out)[SHEET]
    assert [ws2.cell(row=4, column=c).value for c in range(1, 20)] == before
    assert "root_cause" not in proc2.results[0].filled_fields and "improvement" not in proc2.results[0].filled_fields
    assert len(_images_at(ws2, 4, COL["image"])) == 1                                                  # not duplicated


def test_complete_row_still_skipped_and_duplicates_marked(template, tmp_path):
    deck = make_real_layout_report(tmp_path / "(CTMS)_2_260920045-VOC_ Đối sách LỖI MẺ XƯỚC 24.9.2026.pptx")
    _prefill(template, [{"mgmt": "260920045-VOC"}, {"mgmt": "260920045-VOC"}])
    wb = load_workbook(template)                      # QPN picture entered by hand (deck has no QPN slide)
    from openpyxl.drawing.image import Image as XLImage
    png = tmp_path / "manual.png"
    Image.new("RGB", (40, 30), "red").save(png)
    img = XLImage(str(png))
    img.anchor = "H4"
    wb[SHEET].add_image(img)
    wb.save(template)
    out = tmp_path / "out" / "k.xlsx"
    s1, _, proc = _run([deck], template, out)
    assert proc.results[0].excel_row == 4 and s1.failed == 0
    ws = load_workbook(out)[SHEET]
    assert ws.cell(row=5, column=COL["mgmt"]).fill.fgColor.rgb.endswith("FFC7CE")
    assert ws.cell(row=5, column=COL["cause"]).value is None
    s2, ev, _ = _run([deck], template, out)
    assert s2.skipped == 1 and ev[-1][1] == "skipped"
