"""PROMPT-011B BUG B – WinError 32 after content review: labels and Excel are separate transactions, prepared rows
are never reported as committed, locked target -> friendly message with the real path + [Thử lại]/[Để sau].

The lock is simulated at the exact boundary that fails on Windows (``os.replace`` of ``.saving.xlsx`` onto the
master that Excel holds open) – no Microsoft Excel needed."""
import json
import logging
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import excel_writer as ew
from app.config import AppConfig
from app.excel_writer import ExcelLockedError, ExcelWriter, is_sharing_violation, locked_file_message, probe_writable
from app.gui_controller import ContentReviewOutcome, GuiController
from app.image_learning import ImageLearning
from app.image_review import reapply_content_labels
from tests.test_content_group_candidates import BODY, DECK, HEAD, TABLE_TEXT, _by_text, _pipeline, make_group_deck
from tests.test_content_region import COL, SHEET, _prefill
from tests.test_gui_redesign import _make_app
from tests.test_prompt004 import MGMT

WIN32 = "[WinError 32] The process cannot access the file because it is being used by another process"


def _win_lock(path):
    e = PermissionError(13, "The process cannot access the file because it is being used by another process")
    e.filename = str(path)
    try:
        e.winerror = 32                                  # attribute exists natively on Windows only
    except Exception:
        pass
    return e


class _Lock:
    """Monkeypatch helper: ``os.replace`` onto ``target`` fails like Excel holding the file; everything else works."""
    def __init__(self, monkeypatch, target: Path, locked=True):
        self.target, self.locked, self.attempts = Path(target), locked, 0
        real = os.replace

        def fake(src, dst, *a, **k):
            if Path(dst) == self.target and self.locked:
                self.attempts += 1
                raise _win_lock(dst)
            return real(src, dst, *a, **k)
        monkeypatch.setattr(ew.os, "replace", fake)


def _setup(template, tmp_path):
    deck = make_group_deck(tmp_path / DECK)
    _prefill(template, [{"mgmt": MGMT, "vendor": "Mtech"}])
    out = tmp_path / "out" / "master.xlsx"
    out.parent.mkdir(parents=True)
    out.write_bytes(template.read_bytes())
    r, cls, sections, cands = _pipeline(deck)
    return deck, out, cands


# ------------------------------------------------------------------ classification
def test_sharing_violation_classification():
    assert is_sharing_violation(_win_lock("x.xlsx"))
    assert is_sharing_violation(OSError(WIN32))
    assert is_sharing_violation(PermissionError(13, "sharing violation"))
    assert not is_sharing_violation(PermissionError(13, "Permission denied"))     # read-only folder / ACL
    assert not is_sharing_violation(FileNotFoundError(2, "missing"))
    assert not is_sharing_violation(RuntimeError(WIN32))
    msg = locked_file_message(Path("D:/x/master.xlsx"))
    assert msg.startswith("Không thể cập nhật file Excel vì file đang được sử dụng.") and "D:/x/master.xlsx" in msg
    assert "Hãy đóng file Excel hoặc chương trình đang sử dụng file, sau đó thử lại." in msg


def test_probe_writable(tmp_path):
    f = tmp_path / "a.xlsx"
    assert probe_writable(f) is None                      # missing file: nothing to lock
    f.write_bytes(b"x")
    assert probe_writable(f) is None


# ------------------------------------------------------------------ writer commit boundary
def test_writer_locked_target_keeps_original_and_cleans_tmp(template, tmp_path, monkeypatch):
    deck, out, cands = _setup(template, tmp_path)
    before = out.read_bytes()
    lock = _Lock(monkeypatch, out)
    w = ExcelWriter(template, out)
    assert w.replace_improvement_text(4, "X") is True
    with pytest.raises(ExcelLockedError) as ei:
        w.save()
    w.close()
    assert ei.value.path == out and str(out) in str(ei.value) and "WinError" not in str(ei.value)
    assert lock.attempts == 1
    assert out.read_bytes() == before and load_workbook(out)[SHEET].cell(row=4, column=COL["improvement"]).value is None
    assert not out.with_name("master.saving.xlsx").exists()                     # temp cleaned
    backups = list((out.parent / "backup").glob("*_backup_*.xlsx"))
    assert len(backups) == 1 and load_workbook(backups[0])[SHEET].cell(row=4, column=COL["mgmt"]).value == MGMT


def test_writer_unrelated_permission_error_not_misclassified(template, tmp_path, monkeypatch):
    deck, out, cands = _setup(template, tmp_path)

    def denied(src, dst, *a, **k):
        raise PermissionError(13, "Permission denied")
    monkeypatch.setattr(ew.os, "replace", denied)
    w = ExcelWriter(template, out)
    w.replace_improvement_text(4, "X")
    with pytest.raises(PermissionError) as ei:
        w.save()
    w.close()
    assert not isinstance(ei.value, ExcelLockedError)
    assert not out.with_name("master.saving.xlsx").exists()


def test_writer_atomic_replace_after_handles_closed(template, tmp_path, monkeypatch):
    deck, out, cands = _setup(template, tmp_path)
    order = []
    real_save = ew.load_workbook(out).__class__.save
    monkeypatch.setattr(ew.load_workbook(out).__class__, "save", lambda self, p: (order.append(("save", Path(p).name)),
                                                                                    real_save(self, p)))
    real_replace = os.replace
    monkeypatch.setattr(ew.os, "replace", lambda s, d: (order.append(("replace", Path(d).name)), real_replace(s, d)))
    w = ExcelWriter(template, out)
    w.replace_improvement_text(4, "X")
    w.save()
    w.close()
    assert order == [("save", "master.saving.xlsx"), ("replace", "master.xlsx")]
    assert load_workbook(out)[SHEET].cell(row=4, column=COL["improvement"]).value == "X"


# ------------------------------------------------------------------ reapply transaction
def test_reapply_success_reports_exact_committed_rows(template, tmp_path, caplog):
    deck, out, cands = _setup(template, tmp_path)
    lrn = ImageLearning(tmp_path / "lrn")
    nested = _by_text(cands, "- Cập nhật checksheet")
    lrn.content.store.label(nested, "EXCLUDE_CONTENT")
    with caplog.at_level(logging.INFO, logger="report_extractor.image_review"):
        res = reapply_content_labels(template, out, [nested], lrn)
    assert res.ok and res.prepared_rows == [4] and res.updated_rows == [4] and res.committed_rows == [4]
    assert not res.locked and res.locked_path is None
    assert load_workbook(out)[SHEET].cell(row=4, column=COL["improvement"]).value == f"{HEAD}\n\n{BODY}\n\n{TABLE_TEXT}"
    text = caplog.text
    assert "CONTENT_REAPPLY_START" in text and "CONTENT_REAPPLY_PREPARED rows=1" in text and "CONTENT_REAPPLY_COMMIT_OK rows=1" in text


def test_reapply_locked_target(template, tmp_path, monkeypatch, caplog):
    deck, out, cands = _setup(template, tmp_path)
    before = out.read_bytes()
    lock = _Lock(monkeypatch, out)
    lrn = ImageLearning(tmp_path / "lrn")
    nested = _by_text(cands, "- Cập nhật checksheet")
    lrn.content.store.label(nested, "EXCLUDE_CONTENT")
    with caplog.at_level(logging.INFO):
        res = reapply_content_labels(template, out, [nested], lrn)
    assert lock.attempts == 1
    assert res.prepared_rows == [4] and res.updated_rows == [] and res.committed_rows == []       # nothing committed
    assert res.locked and res.locked_path == str(out)
    assert len(res.errors) == 1 and str(out) in res.errors[0] and "đang được sử dụng" in res.errors[0]
    assert "WinError" not in res.errors[0] and "WinError" not in " ".join(res.messages)
    assert "being used by another process" in res.technical_error                     # technical detail kept apart
    assert "CONTENT_REAPPLY_COMMIT_FAILED" in caplog.text and "being used by another process" in caplog.text
    assert "prepared_rows=1" in caplog.text
    # original master intact + loadable, backup valid, temp gone, labels on disk
    assert out.read_bytes() == before and load_workbook(out)[SHEET].cell(row=4, column=COL["mgmt"]).value == MGMT
    assert not out.with_name("master.saving.xlsx").exists()
    backups = list((out.parent / "backup").glob("*_backup_*.xlsx"))
    assert len(backups) == 1 and load_workbook(backups[0]) is not None
    labels = (tmp_path / "lrn" / "content_labels.jsonl").read_text(encoding="utf-8").splitlines()
    assert len(labels) == 1 and json.loads(labels[0])["candidate_id"] == nested.candidate_id
    # retry after unlock: same saved labels, commit succeeds
    lock.locked = False
    res2 = reapply_content_labels(template, out, [nested], lrn)
    assert res2.ok and res2.updated_rows == [4]
    assert load_workbook(out)[SHEET].cell(row=4, column=COL["improvement"]).value == f"{HEAD}\n\n{BODY}\n\n{TABLE_TEXT}"
    assert len((tmp_path / "lrn" / "content_labels.jsonl").read_text(encoding="utf-8").splitlines()) == 1


# ------------------------------------------------------------------ controller: two transactions + retry
def _ctl(tmp_path, template, out):
    ctl = GuiController(AppConfig(period_mode="all"), config_path=tmp_path / "cfg" / "config.json",
                        learning_dir_override=tmp_path / "lrn")
    ctl.template, ctl.output = str(template), str(out)
    return ctl


def test_controller_success_message(template, tmp_path):
    deck, out, cands = _setup(template, tmp_path)
    ctl = _ctl(tmp_path, template, out)
    nested = _by_text(cands, "- Cập nhật checksheet")
    ctl.set_pending_content_label(nested, "EXCLUDE_CONTENT")
    ok, msg = ctl.save_content_confirmations(cands)
    assert ok and msg == "Đã lưu 1 nhãn nội dung. Đã cập nhật nội dung cải tiến cho 1 dòng Excel."
    o = ctl.last_content_review
    assert (o.labels_saved, o.excel_rows_prepared, o.excel_rows_committed) == (1, 1, 1) and not o.excel_pending
    assert ctl.content_reapply_pending == 0


def test_controller_locked_then_retry(template, tmp_path, monkeypatch, caplog):
    deck, out, cands = _setup(template, tmp_path)
    lock = _Lock(monkeypatch, out)
    ctl = _ctl(tmp_path, template, out)
    nested = _by_text(cands, "- Cập nhật checksheet")
    ctl.set_pending_content_label(nested, "EXCLUDE_CONTENT")
    with caplog.at_level(logging.INFO):
        ok, msg = ctl.save_content_confirmations(cands)
    assert ok                                                   # the review transaction succeeded
    assert msg.startswith("Đã lưu 1 nhãn nội dung.\n\nChưa thể cập nhật 1 dòng Excel vì file đang được sử dụng:\n\n")
    assert str(out) in msg and msg.endswith("Hãy đóng file rồi thử lại.")
    assert "Đã cập nhật nội dung cải tiến" not in msg and "WinError" not in msg
    assert "CONTENT_REVIEW_LABELS_SAVED count=1" in caplog.text
    o = ctl.last_content_review
    assert (o.labels_saved, o.excel_rows_prepared, o.excel_rows_committed) == (1, 1, 0)
    assert o.excel_pending and o.locked_path == str(out)
    assert ctl.content_reapply_pending == 1 and not ctl.pending_content_labels      # labels NOT pending again
    labels_file = tmp_path / "lrn" / "content_labels.jsonl"
    assert len(labels_file.read_text(encoding="utf-8").splitlines()) == 1
    # retry while still locked: still pending, still no commit, no duplicate label
    ok2, msg2 = ctl.retry_content_reapply()
    assert not ok2 and ctl.last_content_review.excel_pending and ctl.content_reapply_pending == 1
    assert lock.attempts == 2 and len(labels_file.read_text(encoding="utf-8").splitlines()) == 1
    # user closes Excel -> retry commits
    lock.locked = False
    ok3, msg3 = ctl.retry_content_reapply()
    assert ok3 and "Đã cập nhật nội dung cải tiến cho 1 dòng Excel." in msg3
    assert msg3.startswith("Nhãn nội dung đã lưu trước đó (1 nhãn) được dùng lại.")
    assert ctl.last_content_review.excel_rows_committed == 1 and ctl.content_reapply_pending == 0
    assert len(labels_file.read_text(encoding="utf-8").splitlines()) == 1         # no duplicate labels
    assert load_workbook(out)[SHEET].cell(row=4, column=COL["improvement"]).value == f"{HEAD}\n\n{BODY}\n\n{TABLE_TEXT}"
    assert ctl.retry_content_reapply() == (False, "Không có dòng Excel nào đang chờ cập nhật.")


def test_controller_later_then_next_save_merges_pending(template, tmp_path, monkeypatch):
    deck, out, cands = _setup(template, tmp_path)
    lock = _Lock(monkeypatch, out)
    ctl = _ctl(tmp_path, template, out)
    nested = _by_text(cands, "- Cập nhật checksheet")
    ctl.set_pending_content_label(nested, "EXCLUDE_CONTENT")
    ctl.save_content_confirmations(cands)
    assert ctl.content_reapply_pending == 1
    lock.locked = False                                          # "Để sau", later a new label is saved
    tbl = _by_text(cands, "Hạng mục")
    ctl.set_pending_content_label(tbl, "EXCLUDE_CONTENT")
    ok, msg = ctl.save_content_confirmations(cands)
    assert ok and "Đã cập nhật nội dung cải tiến cho 1 dòng Excel." in msg and ctl.content_reapply_pending == 0
    assert load_workbook(out)[SHEET].cell(row=4, column=COL["improvement"]).value == f"{HEAD}\n\n{BODY}"


def test_outcome_dataclass():
    o = ContentReviewOutcome(labels_saved=61, excel_rows_prepared=7, excel_rows_committed=0, excel_error="x")
    assert o.excel_pending
    assert not ContentReviewOutcome(labels_saved=61, excel_rows_prepared=7, excel_rows_committed=7).excel_pending
    assert not ContentReviewOutcome(labels_saved=61).excel_pending


# ------------------------------------------------------------------ GUI: retry dialog, no re-review
def test_gui_retry_dialog(monkeypatch, tmp_path):
    gui, a, reg, _ = _make_app(monkeypatch, tmp_path)
    calls = []
    a.ctl.last_content_review = ContentReviewOutcome(labels_saved=61, excel_rows_prepared=7, excel_error="locked",
                                                     locked_path="D:/x/master.xlsx")
    a.ctl.pending_content_reapply = [object()] * 7
    monkeypatch.setattr(a.ctl, "save_content_confirmations",
                        lambda c: (True, "Đã lưu 61 nhãn nội dung.\n\nChưa thể cập nhật 7 dòng Excel vì file đang "
                                         "được sử dụng:\n\nD:/x/master.xlsx\n\nHãy đóng file rồi thử lại."))
    monkeypatch.setattr(a, "open_content_review", lambda: calls.append("review"))
    n_tops = len(reg["toplevels"])
    a._creview_save()
    assert len(reg["toplevels"]) == n_tops + 1 and a.excel_retry_win is reg["toplevels"][-1]
    assert a.btn_excel_retry.cfg["text"] == "Thử lại" and a.btn_excel_later.cfg["text"] == "Để sau"
    assert "D:/x/master.xlsx" in a.lbl_excel_retry.cfg["text"] and "Đã lưu 61 nhãn nội dung." in a.lbl_excel_retry.cfg["text"]
    assert a.btn_reapply_saved.cfg.get("state") == "normal"
    # retry still locked -> dialog stays, text refreshed, review never reopened
    def retry_locked():
        calls.append("retry")
        return False, "vẫn khoá"
    monkeypatch.setattr(a.ctl, "retry_content_reapply", retry_locked)
    a._excel_retry()
    assert a.excel_retry_win is not None and a.lbl_excel_retry.cfg["text"] == "vẫn khoá" and calls == ["retry"]
    # retry succeeds -> dialog closed
    def retry_ok():
        calls.append("retry")
        a.ctl.last_content_review = ContentReviewOutcome(labels_saved=61, excel_rows_prepared=7, excel_rows_committed=7)
        a.ctl.pending_content_reapply = []
        return True, "Đã cập nhật nội dung cải tiến cho 7 dòng Excel."
    monkeypatch.setattr(a.ctl, "retry_content_reapply", retry_ok)
    a._excel_retry()
    assert a.excel_retry_win is None and calls == ["retry", "retry"] and a.btn_reapply_saved.cfg.get("state") == "disabled"


def test_gui_later_keeps_labels_and_button(monkeypatch, tmp_path):
    gui, a, reg, _ = _make_app(monkeypatch, tmp_path)
    a.ctl.last_content_review = ContentReviewOutcome(labels_saved=3, excel_rows_prepared=2, excel_error="locked")
    a.ctl.pending_content_reapply = [object(), object()]
    a._open_excel_retry("msg")
    a._excel_retry_later()
    assert a.excel_retry_win is None and a.ctl.content_reapply_pending == 2
    assert a.btn_reapply_saved.cfg.get("state") == "normal"
    assert any("Để sau" in ln and "đã được lưu" in ln for ln in a.txt_log.lines)
