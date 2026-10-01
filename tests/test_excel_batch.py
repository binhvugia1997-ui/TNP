"""Excel writing, rendering, batch processing, error recovery, duplicates."""
import json
import threading
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.batch_processor import BatchOptions, BatchProcessor
from app.excel_writer import ExcelWriter, TemplateError, validate_template
from app.extractor import ExtractedRecord
from app.history import fingerprint
from app.image_extractor import build_improvement_image, combine_vertically
from app.pptx_parser import parse_pptx
from app.qpn_renderer import SlideRenderer, render_qpn


def _record(**kw):
    rec = ExtractedRecord(management_number="260918080-VOC", model="A185", item="Rear",
                          defect_content="Xước: 15ea", root_cause="NGUYÊN NHÂN\n- a",
                          improvement="3. CẢI TIẾN\n" + "dòng dài " * 300, qpn_slide=2)
    for k, v in kw.items():
        setattr(rec, k, v)
    return rec


def test_template_detection_and_mapping(template, tmp_path):
    w = ExcelWriter(template, tmp_path / "out.xlsx")
    assert w.ws.title == "Kiểm chứng"
    assert w.header_row == 2 and w.data_start == 4
    for f in ("management_number", "vendor", "occurrence_date", "model", "item", "defect_content",
              "qpn", "root_cause", "improvement", "improvement_image", "week_1", "week_8"):
        assert f in w.columns, f
    assert w.item_mapping["nắp lưng"] == "Rear" and w.item_mapping["Rear"] == "Rear"
    assert "A185" in w.known_models
    ok, msg = validate_template(template)
    assert ok, msg


def test_never_overwrites_template_and_rejects_same_path(template, tmp_path):
    before = template.read_bytes()
    with pytest.raises(TemplateError):
        ExcelWriter(template, template)
    w = ExcelWriter(template, tmp_path / "o.xlsx")
    w.append_record(_record())
    w.save()
    assert template.read_bytes() == before


def test_excel_row_insertion_rules(template, tmp_path):
    out = tmp_path / "out.xlsx"
    w = ExcelWriter(template, out)
    row = w.append_record(_record())
    row2 = w.append_record(_record(model="A175"))
    w.save()
    assert (row, row2) == (4, 5)
    ws = load_workbook(out)["Kiểm chứng"]
    c = w.columns
    assert ws.cell(row=4, column=c["management_number"]).value == "260918080-VOC"
    assert ws.cell(row=4, column=c["model"]).value == "A185"
    assert ws.cell(row=5, column=c["model"]).value == "A175"
    assert ws.cell(row=4, column=c["stt"]).value == 1 and ws.cell(row=5, column=c["stt"]).value == 2
    # blank Vendor / Occurrence date / WEEK +1..+8
    assert ws.cell(row=4, column=c["vendor"]).value is None
    assert ws.cell(row=4, column=c["occurrence_date"]).value is None
    for i in range(1, 9):
        assert ws.cell(row=4, column=c[f"week_{i}"]).value is None
    # full text, wrap + top
    cell = ws.cell(row=4, column=c["improvement"])
    assert cell.value.startswith("3. CẢI TIẾN") and len(cell.value) > 2000
    assert cell.alignment.wrap_text is True and cell.alignment.vertical == "top"
    assert ws.row_dimensions[4].height > 100
    # template preserved
    assert ws.cell(row=1, column=1).value == "BẢNG KIỂM CHỨNG ĐỐI SÁCH CẢI TIẾN"
    assert ws.cell(row=3, column=c["week_1"]).value == "WEEK +1"
    assert ws.column_dimensions["J"].width == 70
    assert ws.cell(row=5, column=c["improvement"]).border.left.style == "thin"   # style copied
    assert "Phân loại" in load_workbook(out).sheetnames


def test_images_embedded_and_scaled(template, tmp_path, a185_report):
    r = parse_pptx(a185_report)
    qpn, backend = render_qpn(r, 2, tmp_path / "qpn.png", SlideRenderer(prefer=("builtin",)))
    assert backend == "builtin" and 1500 <= Image.open(qpn).width <= 1920    # white slide margins trimmed
    imp, _ = build_improvement_image(r, [5, 6], tmp_path / "imp.jpg", SlideRenderer(prefer=("builtin",)))
    with Image.open(imp) as im:
        assert im.height > im.width          # vertically combined
    out = tmp_path / "out.xlsx"
    w = ExcelWriter(template, out)
    w.append_record(_record(), qpn_png=qpn, improvement_jpg=imp)
    w.save()
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert any(n.startswith("xl/media/") for n in names)          # embedded, not linked
        drawing = z.read("xl/drawings/drawing1.xml").decode()
    assert drawing.count("<pic>") == 2
    ws = load_workbook(out)["Kiểm chứng"]
    assert len(ws._images) == 2
    assert ws.row_dimensions[4].height <= 409


def test_reopen_output_and_append(template, tmp_path, a185_report):
    out = tmp_path / "out.xlsx"
    r = parse_pptx(a185_report)
    qpn, _ = render_qpn(r, 2, tmp_path / "q.png", SlideRenderer(prefer=("builtin",)))
    w = ExcelWriter(template, out)
    w.append_record(_record(), qpn_png=qpn)
    w.save()
    w.close()
    w2 = ExcelWriter(template, out)
    assert w2.data_start == 4 and w2.next_row() == 5
    w2.append_record(_record(model="A166"), qpn_png=qpn)
    w2.save()
    w2.append_record(_record(model="A175"))
    w2.save()                                    # second save after reload must not fail
    ws = load_workbook(out)["Kiểm chứng"]
    assert [ws.cell(row=r, column=5).value for r in (4, 5, 6)] == ["A185", "A166", "A175"]
    assert len(ws._images) == 2


def test_combine_vertically():
    a = Image.new("RGB", (400, 200), "red")
    b = Image.new("RGB", (800, 200), "blue")
    sheet = combine_vertically([a, b], width=400, gap=10)
    assert sheet.size == (400, 200 + 100 + 10)


# ---------------------------------------------------------------- batch
def _run_batch(files, template, out, **kw):
    events = []
    opts = BatchOptions(files=[Path(f) for f in files], template=template, output_file=out, use_ollama=False, row_mode="append", **kw)
    proc = BatchProcessor(opts, on_file=lambda i, s, d: events.append((i, s, d)))
    summary = proc.run()
    return summary, events, proc


def test_batch_continues_after_failed_report(sample_tree, tmp_path):
    bad = tmp_path / "in" / "A165_broken.pptx"
    bad.parent.mkdir()
    bad.write_bytes(b"this is not a pptx")
    files = [sample_tree["files"][0], bad, sample_tree["files"][2]]
    out = tmp_path / "Output" / "Kiem_chung.xlsx"
    summary, events, _ = _run_batch(files, sample_tree["template"], out)
    assert summary.total == 3 and summary.failed == 1 and summary.completed == 2
    stages = [s for i, s, d in events if i == 1]
    assert stages[-1] == "error"
    assert [s for i, s, d in events if i == 2][-1] == "completed"
    # logs
    err = (out.parent / "logs" / "errors.log").read_text(encoding="utf-8")
    assert "A165_broken.pptx" in err
    res = json.loads((out.parent / "logs" / "batch_result.json").read_text(encoding="utf-8"))
    assert len(res["results"]) == 3
    r0 = res["results"][0]
    assert r0["status"] == "completed" and r0["qpn_slide"] == 2 and r0["model"] == "A185"
    assert r0["cause_slides"] == [3] and r0["improvement_slides"] == [5, 6, 7]
    assert res["results"][1]["status"] == "error" and res["results"][1]["error"]
    assert res["meta"]["summary"]["failed"] == 1
    # assets
    assert (out.parent / "assets" / "0001_QPN.png").exists()
    assert list((out.parent / "assets" / "0001_IMPROVEMENT_pics").glob("after*.png"))   # independent After pictures
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=4, column=5).value == "A185" and ws.cell(row=5, column=5).value == "A175"
    assert ws.cell(row=6, column=5).value is None


def test_batch_needs_review_status(sample_tree, tmp_path):
    out = tmp_path / "o" / "r.xlsx"
    summary, events, _ = _run_batch([sample_tree["files"][3]], sample_tree["template"], out)   # no QPN
    assert summary.needs_review == 1 and summary.completed == 0
    assert [s for i, s, d in events][-1] == "needs_review"


def test_duplicate_protection_and_force(sample_tree, tmp_path):
    out = tmp_path / "o" / "r.xlsx"
    f = [sample_tree["files"][0]]
    s1, _, _ = _run_batch(f, sample_tree["template"], out)
    s2, ev, _ = _run_batch(f, sample_tree["template"], out)
    assert s1.completed == 1 and s2.skipped == 1 and ev[-1][1] == "skipped"
    s3, _, _ = _run_batch(f, sample_tree["template"], out, force_reprocess=True)
    assert s3.completed == 1
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=5, column=5).value == "A185"
    assert fingerprint(f[0]) == fingerprint(f[0])


def test_stop_after_current_file(sample_tree, tmp_path):
    out = tmp_path / "o" / "r.xlsx"
    opts = BatchOptions(files=list(sample_tree["files"]), template=sample_tree["template"], output_file=out, use_ollama=False, row_mode="append")
    proc = BatchProcessor(opts)
    proc.on_file = lambda i, s, d: proc.request_stop() if (i == 0 and s == "reading") else None
    summary = proc.run()
    assert summary.stopped is True
    assert summary.completed + summary.needs_review == 1
    assert out.exists()


def test_batch_runs_in_worker_thread(sample_tree, tmp_path):
    out = tmp_path / "o" / "r.xlsx"
    done = threading.Event()
    opts = BatchOptions(files=[sample_tree["files"][0]], template=sample_tree["template"], output_file=out, use_ollama=False, row_mode="append")
    proc = BatchProcessor(opts, on_done=lambda s: done.set())
    t = proc.start()
    assert t is not threading.current_thread()
    assert done.wait(60)
    assert proc.summary.completed == 1


def test_source_pptx_untouched(sample_tree, tmp_path):
    f = sample_tree["files"][0]
    before = f.read_bytes()
    _run_batch([f], sample_tree["template"], tmp_path / "o" / "r.xlsx")
    assert f.read_bytes() == before
