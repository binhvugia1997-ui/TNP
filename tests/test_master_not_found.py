"""PROMPT-004 §1/§11: a Management Number absent from the master is reported – NEVER written as a new row."""
import shutil
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

import app.batch_processor as bp
from app.batch_processor import BatchOptions, BatchProcessor
from app.config import AppConfig
from app.gui_controller import GuiController
from app.prescan import (ACTION_MASTER_NOT_FOUND, ACTION_OUTSIDE_PERIOD, ACTION_PROCESS, ACTION_SOURCE_DUPLICATE,
                         CACHE_FILE_NAME, PROCESS_ACTIONS, FastScanCache, MasterLookup, month_period, prescan)

SHEET = "Kiểm chứng"
SEP = month_period(2026, 9)
RUN_TODAY = date(2026, 9, 20)


def _fake(folder: Path, name: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    p = folder / name
    p.write_bytes(b"not a pptx")
    return p


def _run(files, template, out, **kw):
    events = []
    opts = BatchOptions(files=[Path(f) for f in files], template=template, output_file=out, use_ollama=False,
                        **{"today": RUN_TODAY, "period": SEP, **kw})
    proc = BatchProcessor(opts, on_file=lambda i, s, d: events.append((i, s, d)))
    return proc.run(), events, proc


def _keys(out):
    ws = load_workbook(out)[SHEET]
    return [ws.cell(row=r, column=2).value for r in range(4, 10)]


def test_absent_key_is_not_written_no_parse_no_row(sample_tree, template, tmp_path, monkeypatch):
    calls = {"parse": 0}
    real = bp.parse_pptx

    def spy(p):
        calls["parse"] += 1
        return real(p)
    monkeypatch.setattr(bp, "parse_pptx", spy)
    out = tmp_path / "Output" / "k.xlsx"
    s, ev, proc = _run([sample_tree["files"][0]], template, out)
    it = proc.prescan_result.items[0]
    assert it.action == ACTION_MASTER_NOT_FOUND and it.action not in PROCESS_ACTIONS and s.candidates == 0
    assert s.not_written == 1 and s.completed == 0 and s.new_rows == 0 and s.master_not_found == 1
    assert calls["parse"] == 0                                                   # cheap decision, no PPTX, no Qwen
    fr = proc.results[0]
    assert fr.status == "not_written" and not fr.new_row and fr.excel_row is None
    assert "Không tìm thấy Management Number 260918080-VOC trong Excel" in fr.error and "không tạo dòng mới" in fr.error
    assert _keys(out) == [None] * 6                                              # master untouched
    assert ev[-1][1] == "not_written"
    assert s.backup_file == "" and not (out.parent / "backup").exists()         # nothing modified -> no backup


def test_bare_and_suffixed_forms_match_existing_rows(sample_tree, template, tmp_path):
    wb = load_workbook(template)
    ws = wb[SHEET]
    ws.cell(row=4, column=2, value=260918080)             # numeric bare form stored by Excel
    ws.cell(row=5, column=2, value=" 260918081-voc ")     # spacing / case
    wb.save(template)
    a = _fake(tmp_path / "in", "(CTMS)_11108_260918080_ Đối sách LỖI X 18.9.2026.pptx")
    b = _fake(tmp_path / "in", "260918081-VOC_b.pptx")
    c = _fake(tmp_path / "in", "260918080-VOC_c.pptx")   # suffixed form of a bare key: NOT the same key
    s, ev, proc = _run([a, b, c], template, tmp_path / "Output" / "k.xlsx")
    acts = [it.action for it in proc.prescan_result.items]
    assert acts[0] == ACTION_PROCESS and acts[1] == ACTION_PROCESS and acts[2] == ACTION_MASTER_NOT_FOUND
    assert proc.prescan_result.items[0].excel_row == 4 and proc.prescan_result.items[1].excel_row == 5


def test_invalid_key_outside_period_duplicates_and_absent_key_all_leave_master_untouched(template, tmp_path):
    nokey = _fake(tmp_path / "in", "bao_cao_khong_ma.pptx")
    baddate = _fake(tmp_path / "in", "(CTMS)_11107_261345001_ Đối sách.pptx")
    octo = _fake(tmp_path / "in", "261001002-VOC_october.pptx")
    d1 = _fake(tmp_path / "in" / "A", "260915003-VOC_a.pptx")
    d2 = _fake(tmp_path / "in" / "B", "260915003-VOC_b.pptx")
    import os
    os.utime(d1, (1_700_000_000, 1_700_000_000))
    os.utime(d2, (1_700_000_900, 1_700_000_900))
    out = tmp_path / "Output" / "k.xlsx"
    s, ev, proc = _run([nokey, baddate, octo, d1, d2], template, out)
    acts = [it.action for it in proc.prescan_result.items]
    assert acts[0] == "INVALID_MANAGEMENT_NUMBER" and acts[1] == "INVALID_MANAGEMENT_NUMBER"
    assert acts[2] == ACTION_OUTSIDE_PERIOD and acts[3] == ACTION_SOURCE_DUPLICATE
    assert acts[4] == ACTION_MASTER_NOT_FOUND and s.new_rows == 0 and s.not_written == 3
    assert _keys(out) == [None] * 6


def test_existing_partial_rows_still_filled_and_duplicates_marked(sample_tree, template, tmp_path):
    wb = load_workbook(template)
    ws = wb[SHEET]
    ws.cell(row=4, column=2, value="260918080-VOC")
    ws.cell(row=5, column=2, value="260918081-VOC")
    ws.cell(row=6, column=2, value="260918081-VOC")
    wb.save(template)
    out = tmp_path / "Output" / "k.xlsx"
    s, ev, proc = _run([sample_tree["files"][0], sample_tree["files"][1]], template, out)
    assert [it.action for it in proc.prescan_result.items] == [ACTION_PROCESS, ACTION_PROCESS]
    assert s.new_rows == 0 and all(not r.new_row for r in proc.results)
    ws2 = load_workbook(out)[SHEET]
    assert ws2.cell(row=4, column=5).value == "A185" and ws2.cell(row=5, column=5).value == "A185"
    assert ws2.cell(row=6, column=5).value is None and ws2.cell(row=7, column=2).value is None
    assert ws2.cell(row=6, column=2).fill.fgColor.rgb.endswith("FFC7CE")
    s2, ev2, proc2 = _run([sample_tree["files"][0]], template, out)
    assert s2.skipped == 1 and s2.new_rows == 0 and load_workbook(out)[SHEET].cell(row=7, column=2).value is None


def test_cache_hit_never_hides_a_missing_row(tmp_path):
    f = _fake(tmp_path / "in", "260915001-VOC_r.pptx")
    cache = FastScanCache(tmp_path / CACHE_FILE_NAME)
    cache.record("260915001-VOC", f, "completed")
    res = prescan([f], SEP, cache, MasterLookup(lambda m: [], lambda r: []), today=RUN_TODAY)
    it = res.items[0]
    assert it.action == ACTION_MASTER_NOT_FOUND and it.cache_decision == "miss:master_row_missing"
    assert res.candidates == [] and res.counts()["fast_skipped"] == 0 and res.counts()["master_not_found"] == 1
    assert "Không tìm thấy Management Number trong Excel: 1" in res.summary_lines_vi()


def test_new_monthly_workbook_without_the_key_reports_not_written(sample_tree, template, tmp_path):
    wb = load_workbook(template)
    wb[SHEET].cell(row=4, column=2, value="260918080-VOC")
    wb.save(template)
    out1 = tmp_path / "Output" / "Kiem_chung_08_2026.xlsx"
    s1, _, _ = _run([sample_tree["files"][0]], template, out1)
    assert s1.completed == 1
    # a different workbook created from a template WITHOUT the key -> reported, not created
    tpl2 = template.with_name("tpl2.xlsx")
    wb = load_workbook(template)
    wb[SHEET].cell(row=4, column=2).value = None
    wb.save(tpl2)
    template = tpl2
    out2 = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    s2, ev2, proc2 = _run([sample_tree["files"][0]], template, out2)
    assert s2.fast_skipped == 0 and s2.new_rows == 0 and s2.not_written == 1
    assert load_workbook(out2)[SHEET].cell(row=4, column=2).value is None


def test_gui_prescan_counts_not_found_outside_the_denominator(sample_tree, template, tmp_path, monkeypatch):
    class _D(date):
        @classmethod
        def today(cls):
            return RUN_TODAY
    monkeypatch.setattr(bp, "date", _D)
    reports = tmp_path / "Bao_cao" / "2026-09"
    reports.mkdir(parents=True)
    shutil.copy(sample_tree["files"][0], reports / "260918080-VOC_a.pptx")
    shutil.copy(sample_tree["files"][1], reports / "260918081-VOC_b.pptx")
    wb = load_workbook(template)
    wb[SHEET].cell(row=4, column=2, value="260918080-VOC")
    wb.save(template)
    ctl = GuiController(AppConfig(), config_path=tmp_path / "c.json")
    ctl.set_report_folder(str(reports))
    ctl.set_template(str(template))
    ctl.set_output(str(tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"))
    assert ctl.start(use_ollama=False)
    ctl.processor._thread.join(120)
    ctl.pump()
    lines = ctl.summary_lines()
    assert "Không tìm thấy Management Number trong Excel: 1" in lines and "Cần xử lý thực tế: 1" in lines
    assert ctl.progress.total == 1 and len(ctl.files) == 1                        # not in the denominator
    scan = {r.path.name: r for r in ctl.scan_rows()}
    row = scan["260918081-VOC_b.pptx"]
    assert row.status_vi == "Không tìm thấy Management Number trong Excel" and not row.will_process
    done = ctl.result_for(0)
    assert done.management_number == "260918080-VOC" and done.status == "completed"
    assert ctl.diagnostics_for(0)["Dòng sử dụng"] == "4" and "Dòng Excel" not in ctl.diagnostics_for(0)


def test_excel_writer_create_row_is_never_called_by_the_batch():
    src = Path(bp.__file__).read_text(encoding="utf-8")
    assert "create_row(" not in src and "append_record(" in src          # append mode (non-master) untouched
