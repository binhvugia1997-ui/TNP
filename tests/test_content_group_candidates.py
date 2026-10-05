"""BUGFIX – content-learning review showed only a small child text ("Massage") for an improvement candidate.

Root cause (structural, in :mod:`app.pptx_parser`):

* children of a ``p:grpSp`` are stored in the group's CHILD coordinate space (``a:chOff``/``a:chExt``) – python-pptx
  returns those raw values and ``_iter_shapes`` used them as slide coordinates.  Once the author had moved/resized
  the group (the normal case) every text box inside it was mis-placed (typically into the title band / sidebar
  column), so the heading/body of the improvement item got hard-excluded or re-ordered while an unrelated small
  label kept the region evidence;
* copy-pasted grouped shapes may share one ``cNvPr id``: sections, candidates, labels and the rebuilt text are
  keyed by ``(slide, shape_id)``, so one shape's evidence was attached to another shape's text.

The fixtures below rebuild that PPTX structure (group → "Massage" label, heading, body, nested group, table).
"""
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.classifier import heuristic_classify
from app.content_learning import build_content_candidates, decide_content, rebuild_improvement_text
from app.content_region import ROLE_CONTENT, classify_blocks
from app.extractor import collect_sections, extract_record
from app.image_learning import ImageLearning
from app.image_review import reapply_content_labels
from app.pptx_parser import _NS_A, _NS_P, GroupXform, _dedupe_shape_ids, parse_pptx, Block
from tests.test_after_evidence import MGMT, _pics, _prod_head, _tb
from tests.test_content_region import COL, SHEET, _prefill, _shape

HEAD = "Cải tiến lỗi Bẩn, chấm đen (Áp dụng từ ngày 24/09/2026 - Công đoạn Assy Mtech):"
BODY = ("- Bổ sung máy massage rung làm sạch bụi trước khi dán tape\n"
        "+ Trước: Vệ sinh bằng khăn khô, bụi còn bám trên bề mặt gây chấm đen\n"
        "+ Sau: Dùng máy massage rung + khí ion 3s/pcs, kiểm tra bụi dưới đèn 1000lux")
NESTED = "- Cập nhật checksheet vệ sinh đầu ca, ký xác nhận 2 lần/ngày"
TABLE = [["Hạng mục", "Trước", "Sau"], ["Tỷ lệ chấm đen", "1.8%", "0.2%"]]
TABLE_TEXT = "Hạng mục | Trước | Sau\nTỷ lệ chấm đen | 1.8% | 0.2%"
TEMP = "- Sàng lọc 100% tồn kho, cách ly 120 pcs lỗi"
FOOTER = "CTMS – Confidential"
DECK = f"(CTMS)_{MGMT}_ĐỐI SÁCH LỖI BẨN CHẤM ĐEN 24.09.2026.pptx"


# ------------------------------------------------------------------ fixture: improvement item authored as a GROUP
def _grp_tb(container, text, left, top, width, height, size=12, bold=False):
    tb = container.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    tf.paragraphs[0].text = lines[0]
    for ln in lines[1:]:
        tf.add_paragraph().text = ln
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
    return tb


def _move_group(g, left, top, scale=1.0):
    """What PowerPoint does when a group is dragged/resized: ``a:off``/``a:ext`` change, the children keep their
    child-space coordinates (``a:chOff``/``a:chExt`` unchanged)."""
    xfrm = g._element.find("{%s}grpSpPr/{%s}xfrm" % (_NS_P, _NS_A))
    off, ext = xfrm.find("{%s}off" % _NS_A), xfrm.find("{%s}ext" % _NS_A)
    off.set("x", str(int(Inches(left))))
    off.set("y", str(int(Inches(top))))
    ext.set("cx", str(int(int(ext.get("cx")) * scale)))
    ext.set("cy", str(int(int(ext.get("cy")) * scale)))


def make_group_deck(path: Path, nested=True, table=True, dup_ids=False, scale=1.0) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W = prs.slide_width
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    _tb(s, "BÁO CÁO ĐỐI SÁCH LỖI BẨN CHẤM ĐEN", 0.5, 1, 12, 1, 28, True)
    _tb(s, f"Model: A253\nItem: Front\nManagement No: {MGMT}", 0.5, 2.5, 6, 2, 16)
    s = prs.slides.add_slide(blank)
    _tb(s, "1. NGUYÊN NHÂN", 0.5, 0.3, 8, 0.7, 24, True)
    _tb(s, "- Bụi bám trên bề mặt trước khi dán tape", 1.7, 1.1, 11, 1.5, 13)
    s = prs.slides.add_slide(blank)
    _tb(s, "2. XỬ LÝ TẠM THỜI", 0.5, 0.3, 8, 0.7, 24, True)
    _shape(s, MSO_SHAPE.OVAL, "Xử lý tạm thời", 0.15, 2.5, 1.3, 1.3)
    _tb(s, TEMP, 1.7, 1.1, 11, 1.5, 13)
    # slide 4 – the real-world structure: one logical improvement item = one group, moved after authoring
    s = prs.slides.add_slide(blank)
    _prod_head(s, W)                                     # title band, sidebar label, footer
    g = s.shapes.add_group_shape()
    _grp_tb(g, "Massage", 0.0, 0.0, 1.0, 0.3, 10)        # small equipment label (first child in XML order)
    _grp_tb(g, HEAD, 0.0, 0.35, 8.0, 0.4, 12, True)      # improvement heading
    _grp_tb(g, BODY, 0.0, 0.8, 8.0, 1.4, 12)             # actual improvement text
    if nested:
        inner = g.shapes.add_group_shape()
        _grp_tb(inner, NESTED, 0.0, 2.3, 8.0, 0.4, 12)
    if table:
        gf = s.shapes.add_table(2, 3, Inches(0.0), Inches(2.8), Inches(6.0), Inches(0.8))
        for ri, row in enumerate(TABLE):
            for ci, v in enumerate(row):
                gf.table.cell(ri, ci).text = v
        g._element.append(gf._element)                   # python-pptx has no GroupShapes.add_table: move the XML
    _move_group(g, 1.7, 1.0, scale)
    _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 10.0, 1.0, 1.5, 0.35)
    _pics(s, ["#c8e6c9"], 10.0, 1.4)
    if dup_ids:   # copy-paste artefact: the label carries the same cNvPr id as the body textbox
        sps = g._element.findall("{%s}sp" % _NS_P)
        sps[0].find(".//{%s}cNvPr" % _NS_P).set("id", sps[2].find(".//{%s}cNvPr" % _NS_P).get("id"))
    prs.save(str(path))
    return path


def _pipeline(path: Path):
    r = parse_pptx(path)
    cls = heuristic_classify(r)
    sections = collect_sections(r, cls)
    cands = decide_content(build_content_candidates(r, sections, [4], MGMT, str(path)))
    return r, cls, sections, cands


def _by_text(cands, start):
    return next(c for c in cands if c.text.startswith(start))


# ------------------------------------------------------------------ parser: group geometry
def test_group_children_get_absolute_slide_geometry(tmp_path):
    r = parse_pptx(make_group_deck(tmp_path / DECK))
    blocks = {b.text.split("\n")[0]: b for b in r.slides[3].blocks if b.kind != "picture"}
    for key in ("Massage", HEAD, BODY.split("\n")[0], NESTED):
        assert abs(blocks[key].left - Inches(1.7)) <= 2, key            # group offset applied, not raw child x=0
    assert abs(blocks["Massage"].top - Inches(1.0)) <= 2
    assert abs(blocks[HEAD].top - Inches(1.35)) <= 2
    assert abs(blocks[BODY.split("\n")[0]].top - Inches(1.8)) <= 2
    assert abs(blocks[NESTED].top - Inches(3.3)) <= 2                    # nested group composes with the parent
    tbl = next(b for b in r.slides[3].blocks if b.kind == "table")
    assert abs(tbl.top - Inches(3.8)) <= 2 and abs(tbl.left - Inches(1.7)) <= 2


def test_resized_group_scales_children(tmp_path):
    r = parse_pptx(make_group_deck(tmp_path / DECK, nested=False, table=False, scale=0.5))
    body = next(b for b in r.slides[3].blocks if b.text.startswith("- Bổ sung"))
    assert abs(body.left - Inches(1.7)) <= 2 and abs(body.top - Inches(1.4)) <= 2       # 1.0 + 0.8 * 0.5
    assert abs(body.width - Inches(4.0)) <= 2                                           # 8.0 * 0.5


def test_group_xform_math():
    xf = GroupXform(ox=1000, oy=2000, sx=2.0, sy=0.5, chx=100, chy=100)
    assert xf.apply(100, 100, 50, 40) == (1000, 2000, 100, 20)
    assert xf.apply(150, 300, 10, 10) == (1100, 2100, 20, 5)
    assert GroupXform.IDENTITY.apply(7, 8, 9, 10) == (7, 8, 9, 10)


def test_before_fix_structure_misplaces_group_children(tmp_path):
    """Document the failure mode: raw child coordinates put the heading into the title band -> hard excluded."""
    r = parse_pptx(make_group_deck(tmp_path / DECK))
    for b in r.slides[3].blocks:                        # simulate the old parser: raw child-space coordinates
        if b.text.startswith((HEAD, "Massage")):
            b.left, b.top = b.left - Inches(1.7), b.top - Inches(1.0)
    roles = {br.block.text.split("\n")[0]: br.role for br in classify_blocks(r.slides[3])}
    assert roles[HEAD] != ROLE_CONTENT                  # the heading of the improvement item is lost
    r2 = parse_pptx(make_group_deck(tmp_path / DECK))   # fixed parser: all item text is CONTENT
    roles2 = {br.block.text.split("\n")[0]: br.role for br in classify_blocks(r2.slides[3])}
    assert roles2[HEAD] == ROLE_CONTENT and roles2[BODY.split("\n")[0]] == ROLE_CONTENT and roles2[NESTED] == ROLE_CONTENT


# ------------------------------------------------------------------ candidates: source text == evidence region
def test_candidate_is_not_collapsed_to_massage(tmp_path):
    r, cls, sections, cands = _pipeline(make_group_deck(tmp_path / DECK))
    included = [c for c in cands if c.decision == "include"]
    assert [c.text for c in included] == [HEAD, BODY, NESTED, TABLE_TEXT]            # source order, verbatim
    body = _by_text(cands, "- Bổ sung")
    assert body.text == BODY and body.confidence >= 0.9
    assert any("CẢI TIẾN TRONG SẢN XUẤT" in e for e in body.evidence)
    assert any(e.startswith("thuộc mục Cải tiến lỗi Bẩn, chấm đen") for e in body.evidence)
    assert body.nearest_heading.startswith("Cải tiến lỗi Bẩn, chấm đen")
    massage = _by_text(cands, "Massage")
    assert massage.decision == "exclude" and massage.candidate_id != body.candidate_id
    assert not any("Cải tiến lỗi Bẩn" in e for e in massage.evidence)               # no borrowed region evidence
    assert len({c.candidate_id for c in cands}) == len(cands)


def test_nested_group_and_table_text_are_candidates(tmp_path):
    r, cls, sections, cands = _pipeline(make_group_deck(tmp_path / DECK))
    nested = _by_text(cands, "- Cập nhật checksheet")
    tbl = _by_text(cands, "Hạng mục")
    assert nested.decision == "include" and nested.text == NESTED
    assert tbl.decision == "include" and tbl.shape_kind == "table" and tbl.text == TABLE_TEXT
    assert tbl.text.splitlines() == ["Hạng mục | Trước | Sau", "Tỷ lệ chấm đen | 1.8% | 0.2%"]   # cell order


def test_sidebar_title_footer_caption_temporary_excluded(tmp_path):
    r, cls, sections, cands = _pipeline(make_group_deck(tmp_path / DECK))
    excluded = {c.text.split("\n")[0]: c for c in cands if c.decision == "exclude"}
    for key in ("3. CẢI TIẾN TRONG SẢN XUẤT", "Cải tiến trong sản xuất", FOOTER, "Sau cải tiến", "Massage"):
        assert key in excluded and excluded[key].hard_excluded, key
    rec = extract_record(r, cls)
    text = rec.improvement
    assert text == f"{HEAD}\n\n{BODY}\n\n{NESTED}\n\n{TABLE_TEXT}"                     # exact, verbatim, no dupes
    for bad in (TEMP, "XỬ LÝ TẠM THỜI", FOOTER, "Massage", "Sau cải tiến", "3. CẢI TIẾN TRONG SẢN XUẤT"):
        assert bad not in text
    assert text.count(BODY.split("\n")[0]) == 1


def test_source_order_follows_slide_layout(tmp_path):
    r, cls, sections, cands = _pipeline(make_group_deck(tmp_path / DECK))
    order = [c.order for c in cands if c.decision == "include"]
    assert order == sorted(order)
    tops = [c.bounds[1] for c in cands if c.decision == "include"]
    assert tops == sorted(tops)


# ------------------------------------------------------------------ duplicate shape ids
def test_duplicate_shape_ids_are_separated(tmp_path):
    r = parse_pptx(make_group_deck(tmp_path / DECK, dup_ids=True))
    ids = [b.shape_id for b in r.slides[3].blocks if b.kind != "arrow"]
    assert len(ids) == len(set(ids))
    _, _, _, cands = _pipeline(tmp_path / DECK)
    massage, body = _by_text(cands, "Massage"), _by_text(cands, "- Bổ sung")
    assert massage.shape_id != body.shape_id and massage.candidate_id != body.candidate_id
    assert massage.decision == "exclude" and body.decision == "include" and body.text == BODY
    assert any(e.startswith("thuộc mục Cải tiến lỗi Bẩn") for e in body.evidence)
    assert not any("Cải tiến lỗi Bẩn" in e for e in massage.evidence)


def test_dedupe_keeps_first_id_and_is_noop_when_unique():
    blocks = [Block(kind="paragraph", text="a", shape_id=5), Block(kind="paragraph", text="b", shape_id=7),
              Block(kind="picture", shape_id=9)]
    _dedupe_shape_ids(blocks)
    assert [b.shape_id for b in blocks] == [5, 7, 9]
    blocks = [Block(kind="paragraph", text="Massage", shape_id=10), Block(kind="paragraph", text="body", shape_id=10),
              Block(kind="arrow", shape_id=10), Block(kind="table", text="t", shape_id=3)]
    _dedupe_shape_ids(blocks)
    assert [b.shape_id for b in blocks] == [10, 11, 10, 3]            # arrows are not part of the text/picture key


# ------------------------------------------------------------------ reapply uses exactly the reviewed text
def test_reapply_uses_exact_candidate_text(template, tmp_path):
    deck = make_group_deck(tmp_path / DECK)
    _prefill(template, [{"mgmt": MGMT, "vendor": "Mtech"}])
    out = tmp_path / "out" / "k.xlsx"
    out.parent.mkdir(parents=True)
    out.write_bytes(template.read_bytes())
    r, cls, sections, cands = _pipeline(deck)
    lrn = ImageLearning(tmp_path / "lrn")
    nested = _by_text(cands, "- Cập nhật checksheet")
    lrn.content.store.label(nested, "EXCLUDE_CONTENT")
    res = reapply_content_labels(template, out, [nested], lrn)
    assert not res.errors and res.updated_rows == [4]
    cell = load_workbook(out)[SHEET].cell(row=4, column=COL["improvement"]).value
    assert cell == f"{HEAD}\n\n{BODY}\n\n{TABLE_TEXT}"                 # exactly the remaining candidates, verbatim
    # rebuild from the reviewed candidates gives the same string (candidate text == written text)
    rebuilt = rebuild_improvement_text(r, sections, decide_content(cands, None, {nested.candidate_id: "EXCLUDE_CONTENT"}))
    assert rebuilt == cell


# ------------------------------------------------------------------ label compatibility (ids recorded before the fix)
def test_old_labels_survive_reading_order_shift(tmp_path):
    from app.image_learning import lookup_override
    ov = {f"{MGMT}|S4|SH10|5": "EXCLUDE_CONTENT", f"{MGMT}|S4|SH12|6": "IMPROVEMENT_CONTENT"}
    assert lookup_override(ov, f"{MGMT}|S4|SH10|5") == "EXCLUDE_CONTENT"           # exact
    assert lookup_override(ov, f"{MGMT}|S4|SH10|7") == "EXCLUDE_CONTENT"           # same shape, new reading order
    assert lookup_override(ov, f"{MGMT}|S4|SH1|7") is None                         # unknown shape
    assert lookup_override(ov, f"{MGMT}|S5|SH10|5") is None                        # other slide
    ov2 = {f"{MGMT}|S4|SH10|5": "EXCLUDE_CONTENT", f"{MGMT}|S4|SH10|9": "IMPROVEMENT_CONTENT"}
    assert lookup_override(ov2, f"{MGMT}|S4|SH10|7") is None                       # ambiguous -> no guess
    _, _, _, cands = _pipeline(make_group_deck(tmp_path / DECK))
    nested = _by_text(cands, "- Cập nhật checksheet")
    old_id = nested.candidate_id.rsplit("|", 1)[0] + "|99"
    out = decide_content(cands, None, {old_id: "EXCLUDE_CONTENT"})
    c = next(x for x in out if x.candidate_id == nested.candidate_id)
    assert c.decision == "exclude" and c.decision_source == "user"
