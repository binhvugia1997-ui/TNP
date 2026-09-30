"""GUI controller tests – no display required (the Tk view is a thin layer over this)."""
import json
import sys
import threading
import time
import types
from pathlib import Path

from openpyxl import load_workbook

from app.config import AppConfig
from app.gui_controller import (DEFAULT_OUTPUT_NAME, STATUS_VI, GuiController, UiEvent, default_output_path,
                                format_elapsed, status_label)
from app.ollama_client import OllamaError


def _ctl(tmp_path, sample_tree=None, **cfg):
    c = AppConfig(**cfg)
    ctl = GuiController(c, config_path=tmp_path / "config.json")
    if sample_tree is not None:
        ctl.set_report_folder(str(sample_tree["reports"]))
        ctl.set_template(str(sample_tree["template"]))
    return ctl


def _prefill(template: Path, keys):
    wb = load_workbook(template)
    ws = wb["Kiểm chứng"]
    for i, k in enumerate(keys):
        ws.cell(row=4 + i, column=2, value=k)
    wb.save(template)


def _run_to_end(ctl, use_ollama=False, timeout=120):
    assert ctl.start(use_ollama=use_ollama)
    assert ctl.state == "running" and not ctl.controls_enabled()
    t = ctl.processor._thread
    t.join(timeout)
    events = ctl.pump()
    return events


# ------------------------------------------------------------------ discovery / defaults
def test_discovery_count_and_default_output(sample_tree, tmp_path):
    ctl = _ctl(tmp_path)
    n = ctl.set_report_folder(str(sample_tree["reports"]))
    assert n == len(sample_tree["files"]) == len(ctl.rows) == ctl.progress.total
    assert ctl.output == str(Path(sample_tree["reports"]) / DEFAULT_OUTPUT_NAME)
    assert default_output_path("") == ""
    # management number pre-filled from the file name before processing
    assert {r.management_number for r in ctl.rows} >= {"260918080-VOC", "260918081-VOC"}
    assert all(r.stage == "waiting" and r.status_vi == "Đang chờ" for r in ctl.rows)
    # a user-chosen output name is kept when the folder changes
    ctl.set_output(str(tmp_path / "my.xlsx"))
    ctl.set_report_folder(str(sample_tree["reports"]))
    assert ctl.output == str(tmp_path / "my.xlsx")


def test_discovery_lists_legacy_ppt(tmp_path):
    (tmp_path / "old.ppt").write_bytes(b"x")
    (tmp_path / "new.pptx").write_bytes(b"x")
    ctl = _ctl(tmp_path)
    assert ctl.set_report_folder(str(tmp_path)) == 2


# ------------------------------------------------------------------ validation
def test_validation_messages(sample_tree, tmp_path):
    ctl = _ctl(tmp_path)
    errs = ctl.validate()
    assert any("Thư mục báo cáo" in e for e in errs) and any("Kiểm chứng" in e for e in errs)
    empty = tmp_path / "empty"
    empty.mkdir()
    ctl.set_report_folder(str(empty))
    assert any("Không tìm thấy file .ppt/.pptx" in e for e in ctl.validate())
    ctl.set_report_folder(str(sample_tree["reports"]))
    ctl.set_template(str(sample_tree["template"]))
    ctl.set_output(str(sample_tree["template"]))                       # same file -> refused
    assert any("không được trùng" in e for e in ctl.validate())
    ctl.set_output(str(tmp_path / "out" / "ket_qua.txt"))
    assert any(".xlsx" in e for e in ctl.validate())
    ctl.set_output(str(tmp_path / "out" / "ket_qua.xlsx"))
    assert ctl.validate() == []                                          # Ollama is not required
    # template without the required headers
    from openpyxl import Workbook
    bad = tmp_path / "bad.xlsx"
    wb = Workbook()
    wb.active.title = "Kiểm chứng"
    wb.active.append(["STT", "Model"])
    wb.save(bad)
    ctl.set_template(str(bad))
    assert any("không hợp lệ" in e for e in ctl.validate())


# ------------------------------------------------------------------ status mapping
def test_status_mapping_vietnamese():
    assert STATUS_VI["completed"] == "Hoàn thành"
    assert STATUS_VI["needs_review"] == "Cần kiểm tra"
    assert STATUS_VI["not_written"] == "Không tìm thấy Management Number"
    assert STATUS_VI["error"] == "Lỗi" and STATUS_VI["skipped"] == "Bỏ qua"
    assert status_label("analyzing") == "Đang phân tích Qwen"
    assert status_label("extracting") == "Đang trích xuất nguyên nhân / đối sách cải tiến"
    assert status_label("extracting_qpn") == "Đang trích xuất QPN"
    assert status_label("writing_excel") == "Đang ghi Excel"
    assert format_elapsed(65) == "1:05" and format_elapsed(3725) == "1:02:05"


# ------------------------------------------------------------------ worker events / rows / summary
def test_worker_events_update_rows_and_summary(sample_tree, tmp_path):
    ctl = _ctl(tmp_path, sample_tree)
    ctl.cfg.row_mode = "append"
    ctl.set_output(str(tmp_path / "out" / "k.xlsx"))
    events = _run_to_end(ctl)
    kinds = [e.kind for e in events]
    assert "row" in kinds and "progress" in kinds and "done" in kinds
    assert ctl.state == "idle" and ctl.controls_enabled()
    stages_seen = {e.payload[1] for e in events if e.kind == "row"}
    assert {"reading", "analyzing_heuristic", "extracting", "writing_excel"} <= stages_seen   # real pipeline events
    assert "analyzing" not in stages_seen                                                    # no fake Qwen stage
    r0 = next(r for r in ctl.rows if r.management_number == "260918080-VOC")
    assert r0.stage == "completed" and r0.status_vi == "Hoàn thành"
    assert r0.model == "A185" and r0.item == "Rear" and r0.vendor == "Doaltech"
    assert r0.as_values()[:3] == (r0.index + 1, "260918080-VOC", ctl.files[r0.index].name)
    assert r0.as_values()[6] == "Hoàn thành"
    s = ctl.summary
    assert s.total == len(ctl.files) and s.completed + s.needs_review + s.failed == s.total
    lines = ctl.summary_lines()
    assert lines[0] == f"Tổng: {s.total}" and f"Hoàn thành: {s.completed}" in lines
    assert any(ln.startswith("Không tìm thấy Management Number:") for ln in lines)
    assert any(ln.startswith("Thời gian:") for ln in lines)
    assert ctl.output_file() and ctl.output_folder() and ctl.log_file()
    assert ctl.progress.percent == 100.0


def test_missing_management_number_row_is_separate_status(sample_tree, tmp_path, template):
    _prefill(template, ["260918080-VOC"])                        # only the first key exists
    ctl = _ctl(tmp_path)
    ctl.set_report_folder(str(sample_tree["reports"]))
    ctl.set_template(str(template))
    ctl.set_output(str(tmp_path / "out" / "k.xlsx"))
    _run_to_end(ctl)
    by_key = {r.management_number: r for r in ctl.rows}
    assert by_key["260918080-VOC"].stage == "completed"
    other = by_key["260918081-VOC"]
    assert other.stage == "not_written" and other.status_vi == "Không tìm thấy Management Number"
    assert "Không tìm thấy Management Number 260918081-VOC" in other.note
    s = ctl.summary
    assert s.not_written >= 1 and s.failed == 0                  # not a generic program error
    assert f"Không tìm thấy Management Number: {s.not_written}" in ctl.summary_lines()
    assert "Lỗi: 0" in ctl.summary_lines()


def test_start_stop_state_transitions(sample_tree, tmp_path):
    ctl = _ctl(tmp_path, sample_tree)
    ctl.cfg.row_mode = "append"
    ctl.set_output(str(tmp_path / "out" / "k.xlsx"))
    assert ctl.can_start() and not ctl.request_stop()             # nothing to stop yet
    assert ctl.start(use_ollama=False)
    assert ctl.state == "running" and not ctl.can_start()
    assert not ctl.start(use_ollama=False)                        # no double start
    # wait until the first report is actually being processed, then ask to stop
    for _ in range(200):
        ctl.pump()
        if any(r.stage != "waiting" for r in ctl.rows):
            break
        time.sleep(0.02)
    assert ctl.request_stop() and ctl.state == "stopping"
    assert not ctl.request_stop()
    ctl.processor._thread.join(120)
    ctl.pump()
    assert ctl.state == "idle" and ctl.summary is not None and ctl.summary.stopped
    finished = [r for r in ctl.rows if r.is_final]
    assert finished and finished[0].stage in ("completed", "needs_review")   # current report finished, never killed
    assert any(r.stage == "waiting" for r in ctl.rows)                        # remaining reports untouched
    assert ctl.summary_lines()[0] == "Đã dừng theo yêu cầu."


def test_start_refused_when_invalid(tmp_path):
    ctl = _ctl(tmp_path)
    assert not ctl.start()
    assert ctl.state == "idle" and ctl.processor is None


def test_qwen_unavailable_falls_back_to_heuristic(sample_tree, tmp_path):
    ctl = _ctl(tmp_path, sample_tree, ollama_server="127.0.0.1:1", model="qwen3:4b", request_timeout=1)
    ctl.cfg.row_mode = "append"
    ctl.set_output(str(tmp_path / "out" / "k.xlsx"))
    ok, msg = ctl.check_ollama()
    assert not ok and msg.startswith("Ollama: không kết nối được")
    _run_to_end(ctl, use_ollama=True)                              # user chose to continue anyway
    assert ctl.summary.failed == 0
    d = ctl.diagnostics_for(0)
    assert d["Bộ phân loại"] == "heuristic" and "heuristic fallback" in d["Thời gian Ollama"]


def test_check_ollama_messages(tmp_path, monkeypatch):
    ctl = _ctl(tmp_path, ollama_server="http://127.0.0.1:11434", model="qwen3:4b")

    class Good:
        def __init__(self, *a, **k): pass
        def test_connection(self): return {"ok": True, "server": "http://127.0.0.1:11434", "models": ["qwen3:4b", "llama3"]}

    class NoModel(Good):
        def test_connection(self): return {"ok": True, "server": "x", "models": ["llama3"]}

    class Down(Good):
        def test_connection(self): raise OllamaError("Không kết nối được Ollama tại http://127.0.0.1:11434")

    ctl._client_factory = Good
    assert ctl.check_ollama() == (True, "Ollama: Sẵn sàng — qwen3:4b") and ctl.available_models == ["qwen3:4b", "llama3"]
    ctl._client_factory = NoModel
    ok, msg = ctl.check_ollama()
    assert not ok and "không có trên máy chủ" in msg and "llama3" in msg
    ctl._client_factory = Down
    ok, msg = ctl.check_ollama()
    assert not ok and "không kết nối được" in msg
    ctl.model = ""
    ctl._client_factory = Good
    assert "chưa chọn model" in ctl.check_ollama()[1]
    # async variant delivers the same result through the event queue
    ctl.model = "qwen3:4b"
    ctl.check_ollama_async()
    for _ in range(50):
        evs = ctl.pump()
        if evs:
            break
        threading.Event().wait(0.05)
    assert ctl.ollama_ok and ctl.ollama_status == "Ollama: Sẵn sàng — qwen3:4b"


def test_diagnostics_mapping(sample_tree, tmp_path):
    ctl = _ctl(tmp_path, sample_tree)
    ctl.cfg.row_mode = "append"
    ctl.set_output(str(tmp_path / "out" / "k.xlsx"))
    idx = next(r.index for r in ctl.rows if r.management_number == "260918080-VOC")
    assert ctl.diagnostics_for(idx) is None                        # not processed yet
    _run_to_end(ctl)
    d = ctl.diagnostics_for(idx)
    for key in ("Management Number", "Ngày phát sinh", "Vendor", "Model", "Item", "Slide QPN", "Slide nguyên nhân",
                "Slide xử lý tạm thời", "Slide đối sách cải tiến", "Slide ảnh cải tiến", "Bộ phân loại",
                "Độ tin cậy AI", "Lý do cần kiểm tra", "Lỗi", "Thời gian Ollama"):
        assert key in d
    assert d["Management Number"] == "260918080-VOC" and d["Ngày phát sinh"] == "18/09/2026"
    assert d["Vendor"] == "Doaltech" and d["Model"] == "A185" and d["Item"] == "Rear"
    assert d["Slide QPN"].startswith("2") and d["Slide nguyên nhân"] == "[3]" and d["Slide xử lý tạm thời"] == "[4]"
    assert d["Slide ảnh cải tiến"] == "[5, 6]" and d["Slide đối sách cải tiến"] == "[5, 6, 7]" and d["Bộ phân loại"] == "heuristic"
    assert "<think>" not in d["_text"] and "thinking" not in d["Ghi chú phân loại"].lower()


def test_settings_persistence(sample_tree, tmp_path):
    cfg_path = tmp_path / "config.json"
    ctl = GuiController(AppConfig(), config_path=cfg_path)
    ctl.set_report_folder(str(sample_tree["reports"]))
    ctl.set_template(str(sample_tree["template"]))
    ctl.set_output(str(tmp_path / "o" / "k.xlsx"))
    ctl.set_ollama("192.168.1.50:11434", "qwen3:4b")
    ctl.save_settings()
    data = json.loads(cfg_path.read_text(encoding="utf-8"))
    assert data["last_report_folder"] == str(sample_tree["reports"])
    assert data["last_template"] == str(sample_tree["template"])
    assert data["last_output_folder"] == str(tmp_path / "o") and data["last_output_file"].endswith("k.xlsx")
    assert data["ollama_server"] == "http://192.168.1.50:11434" and data["model"] == "qwen3:4b"
    assert not any(k for k in data if "report" in k and "content" in k)   # no report contents
    ctl2 = GuiController(AppConfig.load(cfg_path), config_path=cfg_path)
    assert ctl2.report_folder == str(sample_tree["reports"]) and ctl2.template == str(sample_tree["template"])
    assert ctl2.server == "http://192.168.1.50:11434" and ctl2.model == "qwen3:4b"
    # defaults for a fresh install
    fresh = GuiController(AppConfig(), config_path=tmp_path / "none.json")
    assert fresh.server == "http://127.0.0.1:11434" and fresh.model == "qwen3:4b"


def test_apply_event_marks_unfinished_rows_after_done(sample_tree, tmp_path):
    ctl = _ctl(tmp_path, sample_tree)
    ctl.rows[0].stage = "writing_excel"
    from app.batch_processor import BatchSummary
    ctl.apply_event(UiEvent("done", BatchSummary(total=len(ctl.rows), stopped=True)))
    assert ctl.rows[0].stage == "error" and ctl.rows[1].stage == "waiting"
    assert ctl.summary_lines()[0] == "Đã dừng theo yêu cầu."


# ------------------------------------------------------------------ the Tk view imports and builds with a fake tkinter
def test_gui_view_builds_with_mocked_tkinter(monkeypatch, sample_tree, tmp_path):
    class Widget:
        def __init__(self, *a, **k): self.k = k; self.children = []
        def __getattr__(self, name): return lambda *a, **k: None
        def __setitem__(self, k, v): pass
        def __getitem__(self, k): return ""
        def get_children(self): return []
        def insert(self, *a, **k): return "I001"
        def selection(self): return ()
        def get(self): return self.k.get("value", "")
    class Var(Widget):
        def __init__(self, value="", **k): self.v = value
        def set(self, v): self.v = v
        def get(self): return self.v
    tkmod = types.ModuleType("tkinter")
    for n in ("Tk", "Text", "Toplevel", "StringVar", "BooleanVar", "TclError"):
        setattr(tkmod, n, Var if n.endswith("Var") else (Exception if n == "TclError" else Widget))
    ttkmod = types.ModuleType("tkinter.ttk")
    for n in ("Style", "Label", "LabelFrame", "Frame", "Entry", "Button", "Combobox", "Checkbutton", "Treeview",
              "Scrollbar", "Progressbar"):
        setattr(ttkmod, n, Widget)
    fd = types.ModuleType("tkinter.filedialog"); mb = types.ModuleType("tkinter.messagebox")
    for m in (fd, mb):
        m.__getattr__ = lambda name: (lambda *a, **k: None)
    tkmod.ttk, tkmod.filedialog, tkmod.messagebox = ttkmod, fd, mb
    monkeypatch.setitem(sys.modules, "tkinter", tkmod)
    monkeypatch.setitem(sys.modules, "tkinter.ttk", ttkmod)
    monkeypatch.setitem(sys.modules, "tkinter.filedialog", fd)
    monkeypatch.setitem(sys.modules, "tkinter.messagebox", mb)
    monkeypatch.delitem(sys.modules, "app.gui", raising=False)
    import importlib
    gui = importlib.import_module("app.gui")
    cfg = AppConfig(last_report_folder=str(sample_tree["reports"]), last_template=str(sample_tree["template"]))
    app = gui.ReportExtractorApp(cfg)
    assert len(app.ctl.rows) == len(sample_tree["files"])
    assert [c[1] for c in gui.COLUMNS] == ["STT", "Management Number", "Tên file", "Vendor", "Model", "Item", "Trạng thái", "Ghi chú"]
    app._set_running(True)
    app._set_running(False)
    app._poll()
