"""PROMPT-020: report-scoped improvement-image ownership, export paths, and Excel writes."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE

from app.batch_processor import BatchOptions, BatchProcessor
from app.classifier import heuristic_classify
from app.excel_writer import ExcelWriter
from app.image_learning import ImageLearning, build_candidates
from app.image_review import reapply_labels
from app.improvement_pictures import group_refs, select_after_pictures
from app.pptx_parser import parse_pptx
from tests.test_after_evidence import _deck, _pics, _prod_head, _shape, _tb
from tests.test_content_region import COL, SHEET
from tools.make_samples import make_template

MN_A = "260920101-VOC"
MN_B = "260920102-VOC"
MN_C = "260920103-VOC"
MNS = (MN_A, MN_B, MN_C)


def _same_defect_slide(after_colors, before_color="#ffe0b2", heading="Cải tiến lỗi Hằn"):
    def build(slide, width):
        _prod_head(slide, width)
        _tb(slide, heading, 1.7, 0.9, 5.6, 0.4, 12)
        _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Trước cải tiến", 7.5, 0.9, 1.5, 0.35)
        _pics(slide, [before_color], 7.5, 1.3, w=1.3, h=0.8)
        _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Sau cải tiến", 9.1, 0.9, 1.5, 0.35)
        _pics(slide, list(after_colors), 9.1, 1.3, w=1.3, h=0.8)
    return build


def _master_template(path: Path, management_numbers):
    template = make_template(path)
    wb = load_workbook(template)
    ws = wb[SHEET]
    for offset, mn in enumerate(management_numbers):
        ws.cell(row=4 + offset, column=COL["mgmt"], value=mn)
    wb.save(template)
    return template


def _build_report(path: Path, after_colors, before_color="#ffe0b2", heading="Cải tiến lỗi Hằn"):
    return _deck(path, [_same_defect_slide(after_colors, before_color, heading)])


def _report_selection(path: Path):
    report = parse_pptx(path)
    cls = heuristic_classify(report)
    slides = cls.improvement_image_slides or cls.improvement_slides
    selection = select_after_pictures(report, slides)
    return report, slides, selection


def _row_images(ws, row: int):
    images = [image for image in ws._images
              if image.anchor._from.row == row - 1 and image.anchor._from.col == COL["image"] - 1]
    return sorted(images, key=lambda image: (image.anchor._from.row, image.anchor._from.rowOff,
                                             image.anchor._from.col, image.anchor._from.colOff))


def _image_bytes(image) -> bytes:
    cached = getattr(image, "_test_bytes", None)
    if cached is None:
        cached = image._data()
        image._test_bytes = cached
    image.ref = BytesIO(cached)
    return cached


def _image_color(image):
    with Image.open(BytesIO(_image_bytes(image))) as opened:
        return opened.convert("RGB").getpixel((5, 5))


def _row_colors(ws, row: int):
    return [_image_color(image) for image in _row_images(ws, row)]


def _image_has_rgb(image, rgb, tolerance=4, minimum=15):
    with Image.open(BytesIO(_image_bytes(image))) as opened:
        pixels = opened.convert("RGB").tobytes()
    matches = sum(1 for i in range(0, len(pixels), 3)
                  if all(abs(pixels[i + channel] - rgb[channel]) <= tolerance for channel in range(3)))
    return matches >= minimum


def _row_contains_colors(ws, row, colors):
    images = _row_images(ws, row)
    return all(any(_image_has_rgb(image, color) for image in images) for color in colors)


def _row_image_bytes(ws, row):
    return [_image_bytes(image) for image in _row_images(ws, row)]


def _run_batch(files, template: Path, output: Path):
    processor = BatchProcessor(BatchOptions(
        files=[Path(path) for path in files], template=template, output_file=output,
        use_ollama=False, use_fast_cache=False, use_image_learning=False,
    ))
    summary = processor.run()
    assert summary.failed == 0
    return processor


def _write_rows_for_batch(root: Path, order, paths_by_mn):
    template = _master_template(root / "template.xlsx", MNS)
    output = root / "result.xlsx"
    processor = _run_batch([paths_by_mn[mn] for mn in order], template, output)
    saved = load_workbook(output)[SHEET]
    results = {result.management_number: result for result in processor.results}
    per_mn = {}
    for index, mn in enumerate(MNS):
        row = 4 + index
        assert saved.cell(row, COL["mgmt"]).value == mn
        assert results[mn].excel_row == row
        per_mn[mn] = _row_image_bytes(saved, row)
    return per_mn


def test_same_defect_across_mn_rows_is_ordered_report_scoped_and_before_stays_excluded(tmp_path):
    colors = {
        MN_A: ["#c8e6c9", "#a5d6a7"],
        MN_B: ["#81c784", "#66bb6a"],
        MN_C: ["#4caf50", "#388e3c"],
    }
    paths = {mn: _build_report(tmp_path / f"{mn}_same-defect.pptx", colors[mn]) for mn in MNS}

    # Even if a caller aggregates refs for diagnostics/layout, identical item headings are distinct
    # report/item groups; the owner heading is intentionally identical in every report.
    refs = []
    expected_owner_ids = set()
    for mn in MNS:
        _report, _slides, selected = _report_selection(paths[mn])
        expected_owner_ids.update(ref.owner_id for ref in selected.after)
        assert len(selected.after) == 2
        assert all(ref.temporal_role == "AFTER" and ref.excel_output_eligible for ref in selected.after)
        assert all(ref.temporal_role != "AFTER" for ref in selected.rejected if ref.block.image_blob)
        refs.extend(selected.after)
    assert len(expected_owner_ids) == 1
    grouped = group_refs(refs)
    assert sorted(len(group) for group in grouped.values()) == [2, 2, 2]

    forward = _write_rows_for_batch(tmp_path / "forward", MNS, paths)
    reverse = _write_rows_for_batch(tmp_path / "reverse", tuple(reversed(MNS)), paths)
    for index, mn in enumerate(MNS):
        row = 4 + index
        expected_colors = [tuple(int(color[i:i + 2], 16) for i in (1, 3, 5)) for color in colors[mn]]
        saved = load_workbook(tmp_path / "forward" / "result.xlsx")[SHEET]
        assert len(_row_images(saved, row)) == 1  # two authored After pictures are one rendered crop
        assert _row_contains_colors(saved, row, expected_colors)
        assert not _row_contains_colors(saved, row, [(255, 224, 178)])
    assert forward == reverse


def test_identical_image_bytes_in_different_reports_are_kept_in_both_rows(tmp_path):
    identical = "#c8e6c9"
    path_a = _build_report(tmp_path / f"{MN_A}_same-bytes.pptx", [identical])
    path_b = _build_report(tmp_path / f"{MN_B}_same-bytes.pptx", [identical])
    report_a, slides_a, selection_a = _report_selection(path_a)
    report_b, slides_b, selection_b = _report_selection(path_b)
    after_a = next(ref for ref in selection_a.after if ref.temporal_role == "AFTER")
    after_b = next(ref for ref in selection_b.after if ref.temporal_role == "AFTER")
    assert after_a.block.image_blob == after_b.block.image_blob

    template = _master_template(tmp_path / "template.xlsx", [MN_A, MN_B])
    output = tmp_path / "out" / "result.xlsx"
    _run_batch([path_a, path_b], template, output)
    saved = load_workbook(output)[SHEET]
    expected_color = (200, 230, 201)
    assert _row_contains_colors(saved, 4, [expected_color])
    assert _row_contains_colors(saved, 5, [expected_color])
    assert len(_row_images(saved, 4)) == len(_row_images(saved, 5)) == 1


def _candidates_for(path: Path, management_number: str):
    report, slides, selection = _report_selection(path)
    return build_candidates(report, slides, selection, management_number, str(path))


def test_image_label_reapply_uses_report_scoped_asset_paths_for_duplicate_basenames(tmp_path):
    # The basename intentionally collides while the source paths and Management Numbers differ.
    colors = {MN_A: "#c8e6c9", MN_B: "#ffcc80", MN_C: "#d1c4e9"}
    paths = {mn: _build_report(tmp_path / f"folder-{mn}" / "same-name.pptx", [colors[mn]]) for mn in MNS}
    template = _master_template(tmp_path / "template.xlsx", MNS)
    learning = ImageLearning(tmp_path / "learning")
    candidates = {mn: _candidates_for(paths[mn], mn) for mn in MNS}
    for mn in MNS:
        for candidate in candidates[mn]:
            if candidate.excel_output_eligible and candidate.temporal_role == "AFTER":
                learning.store.label(candidate, "AFTER")

    expected = {mn: tuple(int(colors[mn][i:i + 2], 16) for i in (1, 3, 5)) for mn in MNS}

    def reapply(order, target_name):
        output = tmp_path / target_name / "result.xlsx"
        writer = ExcelWriter(template, output)
        result = reapply_labels(template, output, [c for mn in order for c in candidates[mn]], learning,
                                assets_dir=tmp_path / target_name / "assets", writer=writer)
        writer.close()
        assert result.errors == []
        asset_dirs = list((tmp_path / target_name / "assets").glob("*_IMPROVEMENT_regions"))
        assert len(asset_dirs) == len(order)      # same basenames still get distinct per-report export folders
        saved = load_workbook(output)[SHEET]
        return result, {mn: _row_image_bytes(saved, 4 + index) for index, mn in enumerate(MNS)}

    forward_result, forward = reapply(MNS, "reapply-forward")
    reverse_order = tuple(reversed(MNS))
    reverse_result, reverse = reapply(reverse_order, "reapply-reverse")
    assert forward_result.updated_rows == [4, 5, 6]
    assert reverse_result.updated_rows == [6, 5, 4]
    for order_name in ("reapply-forward", "reapply-reverse"):
        saved = load_workbook(tmp_path / order_name / "result.xlsx")[SHEET]
        for index, mn in enumerate(MNS):
            row = 4 + index
            assert len(_row_images(saved, row)) == 1
            assert _row_contains_colors(saved, row, [expected[mn]])
            assert not _row_contains_colors(saved, row, [(255, 224, 178)])
    assert forward == reverse
    assert learning.store.counts()["after"] == 3


def test_force_reprocess_replaces_only_the_matching_management_number_images(tmp_path):
    path_a = _build_report(tmp_path / f"{MN_A}_force.pptx", ["#c8e6c9", "#a5d6a7"])
    path_b = _build_report(tmp_path / f"{MN_B}_force.pptx", ["#81c784", "#66bb6a"])
    template = _master_template(tmp_path / "template.xlsx", [MN_A, MN_B])
    output = tmp_path / "out" / "result.xlsx"

    first = _run_batch([path_a, path_b], template, output)
    assert {fr.management_number: fr.excel_row for fr in first.results} == {MN_A: 4, MN_B: 5}
    saved = load_workbook(output)[SHEET]
    original_b = _row_image_bytes(saved, 5)
    assert len(original_b) == 1
    assert _row_contains_colors(saved, 5, [(129, 199, 132), (102, 187, 106)])

    # A full reprocess of A produces different evidence. B's row must remain byte-for-byte intact.
    _build_report(path_a, ["#d1c4e9", "#b39ddb"], before_color="#fff9c4")
    forced = BatchProcessor(BatchOptions(
        files=[path_a], template=template, output_file=output, use_ollama=False,
        use_fast_cache=False, use_image_learning=False, force_reprocess=True,
    ))
    summary = forced.run()
    assert summary.failed == 0
    assert forced.results[0].management_number == MN_A and forced.results[0].excel_row == 4
    saved = load_workbook(output)[SHEET]
    assert len(_row_images(saved, 4)) == 1
    assert _row_contains_colors(saved, 4, [(209, 196, 233), (179, 157, 219)])
    assert _row_image_bytes(saved, 5) == original_b


def test_replace_improvement_images_clears_and_rewrites_only_the_requested_row(tmp_path):
    template = _master_template(tmp_path / "template.xlsx", [MN_A, MN_B])
    shared_path = tmp_path / "shared-temporary-image.png"
    Image.new("RGB", (240, 160), (200, 230, 201)).save(shared_path)
    output = tmp_path / "out.xlsx"

    writer = ExcelWriter(template, output)
    assert writer.replace_improvement_images(4, [shared_path]) == 1
    # Deliberately reuse and overwrite the exact same temporary path for another report before the save.
    Image.new("RGB", (240, 160), (129, 199, 132)).save(shared_path)
    assert writer.replace_improvement_images(5, [shared_path]) == 1
    writer.save()
    Image.new("RGB", (240, 160), (209, 196, 233)).save(shared_path)
    assert writer.replace_improvement_images(4, [shared_path]) == 1
    writer.save()
    writer.close()

    saved = load_workbook(output)[SHEET]
    assert _row_colors(saved, 4) == [(209, 196, 233)]
    assert _row_colors(saved, 5) == [(129, 199, 132)]
