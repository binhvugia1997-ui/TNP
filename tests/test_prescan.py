"""Fast pre-scan layer: processing period, 7-day success cache, master pre-check, source duplicates, GUI counters.

Every test that claims "PPTX never opened" patches ``app.batch_processor.parse_pptx`` / ``classify`` /
``export_after_pictures`` with spies that raise – so a single accidental open fails the test loudly.
"""
import json
import os
import shutil
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

import app.batch_processor as bp
from app.batch_processor import BatchOptions, BatchProcessor
from app.config import AppConfig
from app.excel_writer import MANAGED_FIELDS
from app.extractor import derive_occurrence_date, management_number_from_filename
from app.gui_controller import FINAL_STATUSES, GuiController, UiEvent
from app.prescan import (ACTION_FAST_SKIP, ACTION_OUTSIDE_PERIOD, ACTION_PROCESS_NEW_ROW,
                         ACTION_PROCESS, ACTION_SOURCE_DUPLICATE, AUTO_PERIOD_FAIL_VI, CACHE_FILE_NAME,
                         PERIOD_DIFFERS_VI, FastScanCache, MasterLookup, ProcessingPeriod, auto_period_from_excel,
                         cache_cutoff, detect_month_from_excel_name, month_period, prescan, range_period)

TODAY = date(2026, 10, 14)          # fixed "today" of the exact 7-day expiry rule (#40)
RUN_TODAY = date(2026, 9, 20)       # "today" for September batches (cache entries of 2026-09-18 still live)
SEP = month_period(2026, 9)


# ---------------------------------------------------------------- helpers
def _fake_pptx(folder: Path, name: str, mtime: float = None, size: int = 100) -> Path:
    """A file that is NOT a valid PPTX – opening it would raise, which is exactly what we want to detect."""
    folder.mkdir(parents=True, exist_ok=True)
    p = folder / name
    p.write_bytes(b"x" * size)
    if mtime is not None:
        os.utime(p, (mtime, mtime))
    return p


def _prefill(template: Path, keys, complete=()):
    """Put Management Numbers into the master; rows in ``complete`` get every managed field filled
    (text cells + a real picture in the QPN / Hình ảnh cải tiến cells)."""
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image
    wb = load_workbook(template)
    ws = wb["Kiểm chứng"]
    pic = template.with_name("pic.png")
    Image.new("RGB", (20, 20), "red").save(pic)
    for i, k in enumerate(keys):
        r = 4 + i
        ws.cell(row=r, column=2, value=k)
        if k in complete:
            for col in (3, 4, 5, 6, 7, 9, 10):
                ws.cell(row=r, column=col, value="x")
            for col in (8, 11):
                img = XLImage(str(pic))
                img.anchor = f"{'H' if col == 8 else 'K'}{r}"
                ws.add_image(img)
    wb.save(template)


class Spies:
    """Installed into app.batch_processor: any call = PPTX opened / Qwen called / images extracted."""

    def __init__(self, monkeypatch, allow=False):
        self.parse = self.classify = self.images = 0
        real_parse, real_classify, real_images = bp.parse_pptx, bp.classify, bp.export_after_pictures

        def parse(path):
            self.parse += 1
            if not allow:
                raise AssertionError(f"PPTX opened: {path}")
            return real_parse(path)

        def classify(report, client, model):
            self.classify += 1
            if client is not None:
                raise AssertionError("Qwen called")
            return real_classify(report, client, model)

        def images(*a, **k):
            self.images += 1
            return real_images(*a, **k)

        monkeypatch.setattr(bp, "parse_pptx", parse)
        monkeypatch.setattr(bp, "classify", classify)
        monkeypatch.setattr(bp, "export_after_pictures", images)


def _run(files, template, out, **kw):
    events = []
    opts = BatchOptions(files=[Path(f) for f in files], template=template, output_file=out, use_ollama=False,
                        **{"today": RUN_TODAY, **kw})
    proc = BatchProcessor(opts, on_file=lambda i, s, d: events.append((i, s, d)))
    summary = proc.run()
    return summary, events, proc


def _final(events, idx):
    return [s for i, s, d in events if i == idx][-1]


# ================================================================ 37. period detection
@pytest.mark.parametrize("name", ["Kiem_chung_09_2026.xlsx", "Kiem_chung_2026_09.xlsx", "Kiem_chung_09-2026.xlsx",
                                  "Kiem_chung_09.2026.xlsx", "Kiem_chung_Tháng_09_2026.xlsx", "Kiem_chung_T09_2026.xlsx",
                                  "Kiem_chung_Thang_09_2026.xlsx", "Kiem_chung_2026-09.xlsx", "Thang 9 2026.xlsx",
                                  "Tháng 9-2026.xlsx"])
def test_detect_month_from_excel_name(name):                                       # 1-6
    assert detect_month_from_excel_name(name) == (2026, 9)
    per = auto_period_from_excel(f"C:/x/{name}")
    assert (per.start, per.end, per.mode) == (date(2026, 9, 1), date(2026, 9, 30), "auto")


@pytest.mark.parametrize("name", ["Kiem_chung_v2_09.xlsx", "Kiem_chung.xlsx", "Kiem_chung_Ket_qua.xlsx",
                                  "Kiem_chung_13_2026.xlsx", "Kiem_chung_09_2026_10_2026.xlsx", "Kiem_chung_202609.xlsx"])
def test_ambiguous_excel_name_is_not_guessed(name):                                 # 7
    assert detect_month_from_excel_name(name) is None


def test_auto_mode_failure_blocks_start_with_vietnamese_message(sample_tree, tmp_path):   # 8
    ctl = GuiController(AppConfig(), config_path=tmp_path / "c.json")
    ctl.set_report_folder(str(sample_tree["reports"]))
    ctl.set_template(str(sample_tree["template"]))
    ctl.set_output(str(tmp_path / "Kiem_chung_v2_09.xlsx"))
    assert ctl.period_mode == "auto"
    assert AUTO_PERIOD_FAIL_VI in ctl.validate()
    assert ctl.period_label() == "Thời gian xử lý: " + AUTO_PERIOD_FAIL_VI
    assert not ctl.start(use_ollama=False)
    ctl.set_period("all")
    assert AUTO_PERIOD_FAIL_VI not in ctl.validate()
    assert ctl.period_label() == "Thời gian xử lý: Tất cả"


def test_manual_month_and_range_override_excel_name(sample_tree, tmp_path):        # 9, 10
    ctl = GuiController(AppConfig(), config_path=tmp_path / "c.json")
    ctl.set_output(str(tmp_path / "Kiem_chung_09_2026.xlsx"))
    assert ctl.effective_period()[0].start == date(2026, 9, 1)
    assert ctl.set_period("month", 10, 2026) == ""
    per, _ = ctl.effective_period()
    assert (per.start, per.end) == (date(2026, 10, 1), date(2026, 10, 31))
    assert ctl.period_label() == "Thời gian xử lý: Tháng 10/2026"
    assert ctl.period_warning() == PERIOD_DIFFERS_VI
    assert ctl.set_period("range", start="15/09/2026", end="30/09/2026") == ""
    per, _ = ctl.effective_period()
    assert (per.start, per.end) == (date(2026, 9, 15), date(2026, 9, 30))
    assert ctl.period_label() == "Thời gian xử lý: 15/09/2026 → 30/09/2026"
    assert ctl.period_warning() == PERIOD_DIFFERS_VI
    assert ctl.set_period("month", 9, 2026) == "" and ctl.period_warning() == ""
    assert ctl.build_options(False).period.start == date(2026, 9, 1)


def test_invalid_range_rejected_not_swapped():                                       # 11
    per, err = range_period("30/09/2026", "01/09/2026")
    assert per is None and "Từ ngày" in err and "Đến ngày" in err
    assert range_period("31/02/2026", "01/03/2026")[1].startswith("Từ ngày: Ngày không hợp lệ")
    assert range_period("01/09/2026", "x")[1].startswith("Đến ngày: Ngày không hợp lệ")
    per, err = range_period("01/09/2026", "30/09/2026")
    assert err == "" and per.contains(date(2026, 9, 1)) and per.contains(date(2026, 9, 30))
    with pytest.raises(ValueError):
        month_period(2026, 13)
    assert month_period(2026, 2).end == date(2026, 2, 28)                          # real calendar length


def test_period_boundaries_inclusive_and_from_management_number_only(tmp_path):     # 12-16
    files = [_fake_pptx(tmp_path / "in", f"(CTMS)_11107_{m}_ Đối sách LỖI X 05.10.2026.pptx")
             for m in ("260831001", "260901002", "260915003", "260930004", "261001005")]
    res = prescan(files, SEP, None, None, today=TODAY)
    actions = [it.action for it in res.items]
    assert actions == [ACTION_OUTSIDE_PERIOD, ACTION_PROCESS, ACTION_PROCESS, ACTION_PROCESS, ACTION_OUTSIDE_PERIOD]
    # the human-readable 05.10.2026 in the file name and the folder name are ignored – only YYMMDD counts
    assert [it.occurrence_date for it in res.items][1] == derive_occurrence_date("260901002")[0]
    assert res.items[0].reason.startswith("Bỏ qua ngoài thời gian xử lý")
    assert res.items[1].management_number == management_number_from_filename(files[1].name)


# ================================================================ 38. early filtering
def test_outside_period_never_opens_pptx_no_qwen_no_images(sample_tree, tmp_path, monkeypatch):   # 17-19
    spies = Spies(monkeypatch)
    aug = _fake_pptx(tmp_path / "in" / "2026-08" / "Vendor A", "260831001-VOC_report.pptx")
    oct_ = _fake_pptx(tmp_path / "in" / "2026-10" / "Vendor B", "261001002-VOC_report.pptx")
    tpl = sample_tree["template"]
    summary, events, proc = _run([aug, oct_], tpl, tmp_path / "Kiem_chung_09_2026.xlsx", period=SEP,
                                 ollama_server="http://127.0.0.1:1", model="qwen3:4b")
    assert spies.parse == spies.classify == spies.images == 0
    assert _final(events, 0) == "outside_period" and _final(events, 1) == "outside_period"
    assert summary.outside_period == 2 and summary.candidates == 0 and summary.failed == 0
    assert [r.status for r in proc.results] == ["outside_period", "outside_period"]


def test_period_filter_runs_before_cache_and_master(tmp_path, monkeypatch):          # 20
    cache = FastScanCache(tmp_path / CACHE_FILE_NAME)
    f = _fake_pptx(tmp_path / "in", "261001002-VOC_r.pptx")
    cache.record("261001002-VOC", f, "completed")
    calls = []
    master = MasterLookup(lambda m: calls.append(m) or [4], lambda r: [])
    res = prescan([f], SEP, cache, master, today=TODAY)
    assert res.items[0].action == ACTION_OUTSIDE_PERIOD
    assert res.items[0].cache_decision == "" and calls == []                        # neither layer consulted


def test_auto_period_recalculates_when_excel_changes(tmp_path):                      # 21, 65
    ctl = GuiController(AppConfig(), config_path=tmp_path / "c.json")
    ctl.set_output(str(tmp_path / "Kiem_chung_09_2026.xlsx"))
    assert ctl.effective_period()[0].start == date(2026, 9, 1)
    ctl.set_output(str(tmp_path / "Kiem_chung_10_2026.xlsx"))
    assert ctl.effective_period()[0].start == date(2026, 10, 1)
    assert ctl.period_label() == "Thời gian xử lý: 01/10/2026 → 31/10/2026 (từ tên file Excel)"
    ctl.set_output(str(tmp_path / "Ket_qua.xlsx"))
    ctl.set_template(str(tmp_path / "Kiem_chung_T08_2026.xlsx"))                     # falls back to master name
    assert ctl.effective_period()[0].start == date(2026, 8, 1)


def test_recursive_nested_discovery(tmp_path):                                       # 22
    from app.scanner import scan_folder
    a = _fake_pptx(tmp_path / "Bao_cao" / "2026-08" / "Vendor A", "260815001-VOC_a.pptx")
    b = _fake_pptx(tmp_path / "Bao_cao" / "2026-09" / "Vendor B" / "deep" / "deeper", "260915002-VOC_b.pptx")
    _fake_pptx(tmp_path / "Bao_cao" / "2026-09" / "Vendor B", "~$260915002-VOC_b.pptx")
    found = scan_folder(tmp_path / "Bao_cao")
    assert found == [a, b]


# ================================================================ 39. cache
def _cached_setup(tmp_path, sample_tree, status="completed"):
    tpl = tmp_path / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl)
    _prefill(tpl, ["260915001-VOC"], complete=["260915001-VOC"])
    f = _fake_pptx(tmp_path / "in", "260915001-VOC_report.pptx")
    out = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    cache = FastScanCache(out.parent / "logs" / CACHE_FILE_NAME)
    cache.record("260915001-VOC", f, status)
    return tpl, f, out


def test_recent_success_fast_skip_before_parser(sample_tree, tmp_path, monkeypatch):   # 23-26
    spies = Spies(monkeypatch)
    tpl, f, out = _cached_setup(tmp_path, sample_tree)
    summary, events, proc = _run([f], tpl, out, period=SEP, ollama_server="http://127.0.0.1:1", model="qwen3:4b")
    assert _final(events, 0) == "fast_skip" and events[-1][2] == "Bỏ qua nhanh — đã xử lý gần đây"
    assert spies.parse == spies.classify == spies.images == 0
    assert summary.fast_skipped == 1 and summary.candidates == 0
    assert proc.prescan_result.items[0].action == ACTION_FAST_SKIP


def test_changed_mtime_or_size_is_cache_miss(tmp_path):                              # 27, 28
    cache = FastScanCache(tmp_path / CACHE_FILE_NAME)
    f = _fake_pptx(tmp_path / "in", "261010001-VOC_r.pptx", mtime=1_700_000_000)
    assert cache.record("261010001-VOC", f, "completed")
    assert cache.lookup("261010001-VOC", f, TODAY) == (True, "recent_success_cache")
    os.utime(f, (1_700_000_100, 1_700_000_100))
    assert cache.lookup("261010001-VOC", f, TODAY) == (False, "source_changed")
    f2 = _fake_pptx(tmp_path / "in2", "261010002-VOC_r.pptx", mtime=1_700_000_000, size=50)
    cache.record("261010002-VOC", f2, "completed")
    _fake_pptx(tmp_path / "in2", "261010002-VOC_r.pptx", mtime=1_700_000_000, size=51)   # same mtime, new size
    assert cache.lookup("261010002-VOC", f2, TODAY) == (False, "source_changed")
    assert cache.lookup("261010003-VOC", f2, TODAY) == (False, "not_cached")
    res = prescan([f], ProcessingPeriod("all"), cache, None, today=TODAY)
    assert res.items[0].action == ACTION_PROCESS and res.items[0].cache_decision == "miss:source_changed"


@pytest.mark.parametrize("status", ["needs_review", "error", "not_written", "waiting", "reading"])
def test_unsafe_outcomes_are_never_cached(tmp_path, status):                         # 29-32
    cache = FastScanCache(tmp_path / CACHE_FILE_NAME)
    f = _fake_pptx(tmp_path / "in", "261010001-VOC_r.pptx")
    assert cache.record("261010001-VOC", f, status) is False
    assert cache.lookup("261010001-VOC", f, TODAY) == (False, "not_cached")
    assert cache.record("", f, "completed") is False                                 # missing Management Number


def test_batch_only_caches_safe_statuses(sample_tree, tmp_path, monkeypatch):          # 29-32 end-to-end
    Spies(monkeypatch, allow=True)
    tpl = tmp_path / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl)
    # 260918080 real report (completes), 260918083 (no QPN -> needs review), 260918099 fake (error)
    _prefill(tpl, ["260918080-VOC", "260918083-VOC", "260918099-VOC"])
    bad = _fake_pptx(tmp_path / "in", "260918099-VOC_bad.pptx")
    out = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    summary, events, proc = _run([sample_tree["files"][0], sample_tree["files"][3], bad], tpl, out, period=SEP)
    data = json.loads((out.parent / "logs" / CACHE_FILE_NAME).read_text(encoding="utf-8"))
    statuses = {r.management_number: r.status for r in proc.results}
    assert statuses["260918099-VOC"] == "error" and statuses["260918083-VOC"] == "needs_review"
    assert set(data["entries"]) == {m for m, s in statuses.items() if s == "completed"} == {"260918080-VOC"}
    ent = data["entries"]["260918080-VOC"]
    assert set(ent) == {"path", "size", "mtime", "occurrence_date", "processed_at", "status"}
    assert ent["occurrence_date"] == "2026-09-18"
    # interrupted run: stop before the first report -> nothing new cached
    opts = BatchOptions(files=[sample_tree["files"][1]], template=tpl, output_file=out, use_ollama=False,
                        today=RUN_TODAY, period=SEP)
    proc2 = BatchProcessor(opts)
    proc2.request_stop()
    proc2.run()
    data2 = json.loads((out.parent / "logs" / CACHE_FILE_NAME).read_text(encoding="utf-8"))
    assert set(data2["entries"]) == {"260918080-VOC"}


def test_corrupt_or_missing_cache_does_not_block(sample_tree, tmp_path, monkeypatch):   # 33, 34
    Spies(monkeypatch, allow=True)
    tpl = tmp_path / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl)
    _prefill(tpl, ["260918080-VOC"])
    out = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    cache_file = out.parent / "logs" / CACHE_FILE_NAME
    cache_file.parent.mkdir(parents=True)
    cache_file.write_text("{not json", encoding="utf-8")
    summary, events, proc = _run([sample_tree["files"][0]], tpl, out, period=SEP)
    assert summary.completed == 1 and proc.cache.problem
    assert json.loads(cache_file.read_text(encoding="utf-8"))["entries"]["260918080-VOC"]["status"] == "completed"
    cache_file.unlink()
    summary2, _, _ = _run([sample_tree["files"][0]], tpl, out, period=SEP, force_reprocess=True)
    assert summary2.completed == 1 and cache_file.exists()


def test_cache_write_failure_keeps_business_output(sample_tree, tmp_path, monkeypatch):   # 35
    Spies(monkeypatch, allow=True)
    tpl = tmp_path / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl)
    _prefill(tpl, ["260918080-VOC"])
    out = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    monkeypatch.setattr(FastScanCache, "flush", lambda self: (_ for _ in ()).throw(OSError("disk full")))
    summary, _, _ = _run([sample_tree["files"][0]], tpl, out, period=SEP)
    assert summary.completed == 1 and summary.failed == 0
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=4, column=9).value                                            # Nguyên nhân written


# ================================================================ 40. exact 7-day expiry
def test_seven_day_expiry_boundaries(tmp_path):                                      # 36-40
    assert cache_cutoff(TODAY) == date(2026, 10, 7)
    cache = FastScanCache(tmp_path / CACHE_FILE_NAME)
    old_mtime = 1_600_000_000                      # 2020 – far older than 7 days: must NOT matter
    for m in ("261006001", "261007002", "261008003", "261014004"):
        f = _fake_pptx(tmp_path / "in", f"{m}-VOC.pptx", mtime=old_mtime)
        cache.record(f"{m}-VOC", f, "completed")
    expired = cache.cleanup(TODAY)
    assert sorted(expired) == ["261006001-VOC", "261007002-VOC"]
    assert set(cache.entries) == {"261008003-VOC", "261014004-VOC"}
    data = json.loads((tmp_path / CACHE_FILE_NAME).read_text(encoding="utf-8"))
    assert set(data["entries"]) == {"261008003-VOC", "261014004-VOC"}
    # lookup also honours the boundary without cleanup
    f = _fake_pptx(tmp_path / "in", "261007002-VOC.pptx", mtime=old_mtime)
    cache.record("261007002-VOC", f, "completed")
    assert cache.lookup("261007002-VOC", f, TODAY) == (False, "expired")
    f8 = tmp_path / "in" / "261008003-VOC.pptx"
    assert cache.lookup("261008003-VOC", f8, TODAY) == (True, "recent_success_cache")


def test_cleanup_only_touches_cache_file(sample_tree, tmp_path):                      # 41
    out = tmp_path / "Output"
    (out / "logs").mkdir(parents=True)
    xlsx = out / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], xlsx)
    src = _fake_pptx(tmp_path / "in", "261001001-VOC.pptx")
    log = out / "logs" / "errors.log"
    log.write_text("x", encoding="utf-8")
    before = {p: p.stat().st_mtime_ns for p in (xlsx, src, log)}
    cache = FastScanCache(out / "logs" / CACHE_FILE_NAME)
    cache.record("261001001-VOC", src, "completed")
    assert cache.cleanup(TODAY) == ["261001001-VOC"]
    assert {p: p.stat().st_mtime_ns for p in (xlsx, src, log)} == before
    assert xlsx.exists() and src.exists() and log.exists()


# ================================================================ 41. master Excel pre-check
def test_master_complete_no_pptx_open_and_absent_key_is_not_found(sample_tree, tmp_path, monkeypatch):   # 42-44
    spies = Spies(monkeypatch)
    tpl = tmp_path / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl)
    _prefill(tpl, ["260915001-VOC"], complete=["260915001-VOC"])
    complete = _fake_pptx(tmp_path / "in", "260915001-VOC_done.pptx")
    out = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    summary, events, proc = _run([complete], tpl, out, period=SEP,
                                 ollama_server="http://127.0.0.1:1", model="qwen3:4b")
    assert spies.parse == spies.classify == 0
    assert _final(events, 0) == "skipped" and "Bỏ qua — đã cập nhật" in proc.results[0].error
    assert summary.skipped == 1 and summary.candidates == 0 and summary.prescan["master_complete"] == 1
    data = json.loads((out.parent / "logs" / CACHE_FILE_NAME).read_text(encoding="utf-8"))
    assert data["entries"]["260915001-VOC"]["status"] == "skipped"
    # absent key: pre-scan plans a new row (PROMPT-004D) – it IS a candidate
    absent = _fake_pptx(tmp_path / "in", "260916002-VOC_new.pptx")
    res = prescan([absent], SEP, None, MasterLookup(lambda m: [], lambda r: []), today=RUN_TODAY)
    assert res.items[0].action == ACTION_PROCESS_NEW_ROW and len(res.candidates) == 1
    assert res.counts()["new_rows"] == 1 and "Management Number mới (sẽ thêm dòng): 1" in res.summary_lines_vi()


def test_incomplete_master_row_proceeds(sample_tree, tmp_path, monkeypatch):          # 45
    spies = Spies(monkeypatch, allow=True)
    tpl = tmp_path / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl)
    _prefill(tpl, ["260918080-VOC"])
    out = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    summary, events, proc = _run([sample_tree["files"][0]], tpl, out, period=SEP)
    assert spies.parse == 1 and summary.completed == 1 and summary.candidates == 1
    assert proc.prescan_result.items[0].master_decision == "found_incomplete"


def test_cache_never_hides_manually_cleared_field(sample_tree, tmp_path, monkeypatch):   # 46
    spies = Spies(monkeypatch, allow=True)
    tpl = tmp_path / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl)
    _prefill(tpl, ["260918080-VOC"])
    out = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    s1, _, _ = _run([sample_tree["files"][0]], tpl, out, period=SEP)
    assert s1.completed == 1 and spies.parse == 1
    s2, ev2, p2 = _run([sample_tree["files"][0]], tpl, out, period=SEP)
    assert _final(ev2, 0) == "fast_skip" and spies.parse == 1                        # yesterday's success
    # the user clears "Nguyên nhân" by hand
    wb = load_workbook(out)
    wb["Kiểm chứng"].cell(row=4, column=9).value = None
    wb.save(out)
    s3, ev3, p3 = _run([sample_tree["files"][0]], tpl, out, period=SEP)
    it = p3.prescan_result.items[0]
    assert it.cache_decision == "miss:master_row_incomplete" and it.action == ACTION_PROCESS
    assert spies.parse == 2 and s3.completed == 1
    assert load_workbook(out)["Kiểm chứng"].cell(row=4, column=9).value               # refilled


def test_force_bypasses_cache_and_complete_rows_but_respects_period(sample_tree, tmp_path, monkeypatch):   # 47, 48
    spies = Spies(monkeypatch, allow=True)
    tpl = tmp_path / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl)
    _prefill(tpl, ["260918080-VOC", "260918083-VOC"])
    out = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    _run([sample_tree["files"][0]], tpl, out, period=SEP)
    assert spies.parse == 1
    oct_ = _fake_pptx(tmp_path / "in", "261001009-VOC_october.pptx")
    s, ev, p = _run([sample_tree["files"][0], oct_], tpl, out, period=SEP, force_reprocess=True)
    assert spies.parse == 2 and _final(ev, 0) == "completed"                         # reprocessed despite cache
    assert _final(ev, 1) == "outside_period" and s.candidates == 1                   # --force ≠ all months
    assert p.prescan_result.items[0].cache_decision == "bypassed_force"


# ================================================================ 42. source duplicates
def test_source_duplicates_newest_wins_and_tie_is_deterministic(sample_tree, tmp_path, monkeypatch):   # 49-52
    spies = Spies(monkeypatch)
    old = _fake_pptx(tmp_path / "in" / "A", "report_260923130.pptx", mtime=1_700_000_000)
    new = _fake_pptx(tmp_path / "in" / "B", "copy_260923130.pptx", mtime=1_700_000_500)
    res = prescan([old, new], SEP, None, None, today=TODAY)
    assert [it.action for it in res.items] == [ACTION_SOURCE_DUPLICATE, ACTION_PROCESS]
    assert res.duplicates["260923130"] == {"selected": str(new), "ignored": [str(old)]}
    assert len(res.candidates) == 1
    # tie on mtime -> path order (A before B) regardless of discovery order
    os.utime(old, (1_700_000_500, 1_700_000_500))
    for order in ([old, new], [new, old]):
        r2 = prescan(order, SEP, None, None, today=TODAY)
        sel = [it for it in r2.items if it.action == ACTION_PROCESS]
        assert [it.path for it in sel] == [old]
    # ignored duplicate is never opened in a real batch; sources untouched
    tpl = tmp_path / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl)
    _prefill(tpl, ["260923130"], complete=["260923130"])
    before = {p: (p.stat().st_size, p.stat().st_mtime_ns) for p in (old, new)}
    s, ev, p = _run([new, old], tpl, tmp_path / "Output" / "k.xlsx", period=SEP)
    assert spies.parse == 0 and _final(ev, 0) == "source_duplicate" and _final(ev, 1) == "skipped"
    assert s.source_duplicates == 1 and {q: (q.stat().st_size, q.stat().st_mtime_ns) for q in (old, new)} == before


def test_excel_duplicate_rows_rule_unchanged(sample_tree, tmp_path, monkeypatch):      # 53
    Spies(monkeypatch, allow=True)
    tpl = tmp_path / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl)
    _prefill(tpl, ["260918080-VOC", "260918080-VOC"])
    out = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    s, ev, p = _run([sample_tree["files"][0]], tpl, out, period=SEP)
    ws = load_workbook(out)["Kiểm chứng"]
    assert ws.cell(row=4, column=9).value and not ws.cell(row=5, column=9).value       # topmost updated only
    assert ws.cell(row=5, column=2).fill.fgColor.rgb.endswith("FFC7CE")               # extra row red
    assert p.results[0].duplicate_rows == [5] and p.results[0].status == "completed"     # PROMPT-004 §15: diagnostics, not review
    # complete-row skip through the pre-scan also keeps the red marking
    tpl2 = tmp_path / "Kiem_chung_T09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl2)
    _prefill(tpl2, ["260915001-VOC", "260915001-VOC"], complete=["260915001-VOC"])
    f = _fake_pptx(tmp_path / "in", "260915001-VOC_done.pptx")
    s2, ev2, _ = _run([f], tpl2, tmp_path / "Output2" / "k.xlsx", period=SEP)
    ws2 = load_workbook(tmp_path / "Output2" / "k.xlsx")["Kiểm chứng"]
    assert _final(ev2, 0) == "skipped" and ws2.cell(row=5, column=2).fill.fgColor.rgb.endswith("FFC7CE")


# ================================================================ 43. GUI / counters
class _FixedDate(date):
    @classmethod
    def today(cls):
        return RUN_TODAY


def test_gui_counters_progress_denominator_and_eta(sample_tree, tmp_path, monkeypatch):   # 54-62
    Spies(monkeypatch, allow=True)
    reports = tmp_path / "Bao_cao"
    shutil.copy(sample_tree["files"][0], _fake_pptx(reports / "2026-09" / "Vendor A", "x.pptx").parent / "260918080-VOC_real.pptx")
    (reports / "2026-09" / "Vendor A" / "x.pptx").unlink()
    _fake_pptx(reports / "2026-08" / "Vendor A", "260820001-VOC_aug.pptx")
    _fake_pptx(reports / "2026-10" / "Vendor B", "261002002-VOC_oct.pptx")
    _fake_pptx(reports / "2026-09" / "Vendor A", "260915003-VOC_done.pptx")           # complete row
    shutil.copy(sample_tree["files"][0], _fake_pptx(reports / "2026-09" / "Vendor B", "260916004-VOC_absent.pptx"))  # not in master
    _fake_pptx(reports / "2026-09" / "Vendor B", "260917005-VOC_a.pptx", mtime=1_700_000_000)
    _fake_pptx(reports / "2026-09" / "Vendor B" / "dup", "260917005-VOC_b.pptx", mtime=1_700_000_900)
    tpl = tmp_path / "Kiem_chung_09_2026.xlsx"
    shutil.copy(sample_tree["template"], tpl)
    _prefill(tpl, ["260918080-VOC", "260915003-VOC", "260917005-VOC"], complete=["260915003-VOC", "260917005-VOC"])
    out = tmp_path / "Output" / "Kiem_chung_09_2026.xlsx"
    # a previous run cached the newer duplicate as recently successful
    dup_new = reports / "2026-09" / "Vendor B" / "dup" / "260917005-VOC_b.pptx"
    FastScanCache(out.parent / "logs" / CACHE_FILE_NAME).record("260917005-VOC", dup_new, "completed")

    ctl = GuiController(AppConfig(), config_path=tmp_path / "c.json")
    monkeypatch.setattr(bp, "date", _FixedDate)                                      # today = 20/09/2026
    ctl._today = lambda: RUN_TODAY
    ctl.set_report_folder(str(reports))
    ctl.set_template(str(tpl))
    ctl.set_output(str(out))
    assert ctl.period_label() == "Thời gian xử lý: 01/09/2026 → 30/09/2026 (từ tên file Excel)"   # 63
    assert ctl.validate() == []
    assert len(ctl.all_files) == 7
    assert ctl.start(use_ollama=False)                      # no reviewed list yet -> the real pre-scan runs first
    ctl.processor._thread.join(120)
    events = ctl.pump()
    kinds = [e.kind for e in events]
    assert kinds.index("prescan") < kinds.index("row")                                 # counters first
    c = ctl.scan_result.counts()                            # reviewed list = all 7 discovered files
    assert c == {"discovered": 7, "outside_period": 2, "source_duplicates": 1, "fast_skipped": 1,
                 "master_complete": 1, "new_rows": 1, "incomplete": 1, "invalid_management_number": 0, "candidates": 2}
    assert ctl.progress.total == 2 and ctl.progress.done == 2 and ctl.progress.percent == 100.0   # 61 (new row included)
    assert ctl.progress.text == "Đã xử lý: 2 / 2 — 100%"
    assert len(ctl.report_durations) == 2                                              # 62: skips not measured
    assert all(r.is_final for r in ctl.rows) and len(ctl.files) == 2                  # queue = candidates only
    scan = {r.path.name: r.status_vi for r in ctl.scan_rows()}
    assert scan["260820001-VOC_aug.pptx"] == scan["261002002-VOC_oct.pptx"] == "Ngoài thời gian xử lý"
    assert scan["260917005-VOC_a.pptx"] == "Trùng Management Number trong folder"
    assert scan["260917005-VOC_b.pptx"] == "Bỏ qua — đã xử lý gần đây"
    assert scan["260915003-VOC_done.pptx"] == "Bỏ qua — Excel đã đầy đủ"
    assert scan["260916004-VOC_absent.pptx"] == "Sẽ xử lý — Management Number mới"
    stages = {r.path.name: r.stage for r in ctl.rows}
    assert stages["260916004-VOC_absent.pptx"] == "completed_new"                      # PROMPT-004D: new row
    assert stages["260918080-VOC_real.pptx"] == "completed"
    lines = ctl.summary_lines()
    assert "Tổng file phát hiện: 7" in lines and "Ngoài thời gian xử lý: 2" in lines
    assert "Trùng Management Number trong folder: 1" in lines and "Bỏ qua nhanh — đã xử lý gần đây: 1" in lines
    assert "Bỏ qua — Excel đã đầy đủ: 1" in lines and "Management Number mới (sẽ thêm dòng): 1" in lines
    assert "Cần bổ sung dữ liệu: 1" in lines and "Cần xử lý thực tế: 2" in lines
    ws = load_workbook(out)["Kiểm chứng"]
    keys = [ws.cell(row=r, column=2).value for r in range(4, 9)]
    assert keys == ["260918080-VOC", "260915003-VOC", "260917005-VOC", "260916004-VOC", None]  # exactly one new row
    assert set(FINAL_STATUSES) >= {"outside_period", "source_duplicate", "fast_skip"}


def test_prescan_event_resets_eta_and_denominator(tmp_path):                           # 61, 62 (unit)
    from app.prescan import PreScanItem, PreScanResult
    ctl = GuiController(AppConfig(period_mode="all"), config_path=tmp_path / "c.json")
    ctl.files = [tmp_path / f"{i}.pptx" for i in range(4)]
    from app.gui_controller import RowState
    ctl.rows = [RowState(index=i, path=p) for i, p in enumerate(ctl.files)]
    t = [0.0]
    ctl._clock = lambda: t[0]
    ctl.state, ctl.started_at = "running", 0.0
    t[0] = 30.0                                                                          # 30 s of pre-scan
    res = PreScanResult(period=ProcessingPeriod("all"))
    res.items = [PreScanItem(index=i, path=p) for i, p in enumerate(ctl.files)]
    for it in res.items[:3]:
        it.action = ACTION_OUTSIDE_PERIOD
    ctl.apply_event(UiEvent("prescan", res))
    for i in range(3):
        ctl.apply_event(UiEvent("row", (i, "outside_period", "x")))
    assert ctl.progress.total == 1 and ctl.progress.done == 0 and ctl.report_durations == []
    assert ctl.progress.text == "Đang xử lý: 1 / 1 — 0%"
    ctl.apply_event(UiEvent("row", (3, "reading", "")))
    t[0] = 40.0
    ctl.apply_event(UiEvent("row", (3, "completed", "")))
    assert ctl.report_durations == [10.0]                                                # not 40 s


def test_period_settings_persist(tmp_path):                                              # 64
    cfg_path = tmp_path / "c.json"
    ctl = GuiController(AppConfig(), config_path=cfg_path)
    ctl.set_output(str(tmp_path / "Kiem_chung_09_2026.xlsx"))
    ctl.set_period("range", start="15/09/2026", end="30/09/2026")
    ctl.save_settings()
    ctl2 = GuiController(AppConfig.load(cfg_path), config_path=cfg_path)
    assert ctl2.period_mode == "range" and (ctl2.period_from, ctl2.period_to) == ("15/09/2026", "30/09/2026")
    ctl2.set_period("month", 10, 2026)
    ctl2.save_settings()
    ctl3 = GuiController(AppConfig.load(cfg_path), config_path=cfg_path)
    assert ctl3.period_mode == "month" and (ctl3.period_month, ctl3.period_year) == ("10", "2026")
    ctl3.set_period("auto")
    ctl3.save_settings()
    assert json.loads(cfg_path.read_text(encoding="utf-8"))["period_mode"] == "auto"
    # auto mode after reload follows the NEW Excel, not a stored month
    ctl4 = GuiController(AppConfig.load(cfg_path), config_path=cfg_path)
    ctl4.set_output(str(tmp_path / "Kiem_chung_11_2026.xlsx"))
    assert ctl4.effective_period()[0].start == date(2026, 11, 1)


def test_managed_fields_constant_still_drives_completeness():
    assert "root_cause" in MANAGED_FIELDS and "improvement_image" in MANAGED_FIELDS
